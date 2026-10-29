"""The yogurt pilot samples: shipped data, provenance and install/restore."""
import json, re, importlib.util
from pathlib import Path
import pytest
import sample_projects as sp
import storage as storage_backend
from food_bo import FoodOptimizer

# A timestamp, or a bare date (pending_batch_created is date-only).
TS = re.compile(r'"20\d\d-\d\d-\d\d(T[^"]*)?"')


def _norm(text):
    return TS.sub('"<ts>"', text)


def _builder():
    spec = importlib.util.spec_from_file_location("yogurt_build", sp.SAMPLE_DIR / "build.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)          # LocalStorage writes into the cwd
    return storage_backend.LocalStorage()


def test_every_shipped_file_exists():
    assert sp.available()
    for name in sp.YOGURT_NAMES:
        assert (sp.SAMPLE_DIR / f"{name}.json").is_file()
    assert sp.ROUND6_RESULTS.is_file()
    assert sp.ROUND6_RESULTS.name == sp.ROUND6_RESULTS_FILE_NAME == "Round 6 results (pilot).xlsx"


@pytest.mark.parametrize("name,scored,open_round", [
    (sp.ROUND6, 5, 6), (sp.ALL10, 10, None)])
def test_shipped_project_equals_a_rebuild_from_the_csvs(tmp_path, name, scored, open_round):
    _builder().build(name, scored, open_round, out_dir=str(tmp_path))
    rebuilt = (tmp_path / f"{name}.pkl").read_text(encoding="utf-8")
    shipped = (sp.SAMPLE_DIR / f"{name}.json").read_text(encoding="utf-8")
    assert _norm(rebuilt) == _norm(shipped)



@pytest.mark.parametrize("name", [sp.ROUND6, sp.ALL10])
def test_shipped_projects_carry_no_date(name):
    """The pilot's CSVs date no Batch, so the shipped files carry no date at
    all: not the day build.py ran as the day a pilot row was recorded, nor
    as the day the open Round 6 was opened."""
    text = (sp.SAMPLE_DIR / f"{name}.json").read_text(encoding="utf-8")
    assert not TS.search(text)
    state = json.loads(text)
    assert state["timestamps_history"] == [None] * len(state["results_history"])
    assert state["pending_batch_created"] is None



def test_the_shipped_round6_results_file_carries_no_print_date():
    """Its title says which Round of which project, not the day the file
    was made: the pilot's sheets were never printed on that day."""
    from openpyxl import load_workbook
    book = load_workbook(sp.ROUND6_RESULTS)
    assert book["Round overview"]["A1"].value == f"Round 6 · {sp.ROUND6}"
    for sheet in book:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                assert not (isinstance(value, str) and "Printed" in value)

def _today_range(before, after):
    return {before.astimezone().strftime("%Y-%m-%d"),
            after.astimezone().strftime("%Y-%m-%d")}


@pytest.mark.parametrize("name,n", [(sp.ROUND6, 15), (sp.ALL10, 30)])
def test_pilot_rows_show_a_blank_date_recorded(store, name, n):
    """3 · Results (history_frame, as ui_results draws it), the workbook and
    the CSV (history_export_frame) show every pilot row's Date recorded as
    a blank cell, never the build date, None or NaT."""
    import wording
    sp.ensure(store)
    opt = FoodOptimizer(name, storage=store)
    shown = opt.history_frame(include_amounts=True)
    assert len(shown) == n
    assert list(shown[wording.DATE_RECORDED_COLUMN]) == [""] * n
    exported = opt.history_export_frame()
    assert list(exported[wording.DATE_RECORDED_COLUMN]) == [""] * n
    csv = opt.history_csv()
    for bad in ("None", "NaT", "nan"):
        assert bad not in csv
    assert not re.search(r"20\d\d-\d\d-\d\d", csv)
    assert opt.all_formulations_workbook()[:2] == b"PK"


def test_a_result_recorded_now_gets_todays_date(store):
    """Results the user records in the app are stamped as before: Round 6
    scored now shows today on 3 · Results, beside 15 blank pilot dates.
    Round 6's sheets, opened by the pilot on no known day, print today's
    date, as does a formulation of the user's own added to it."""
    from datetime import datetime, timezone
    import wording
    sp.ensure(store)
    r6 = FoodOptimizer(sp.ROUND6, storage=store)
    before = datetime.now(timezone.utc)
    assert r6._sheet_date() in _today_range(before, datetime.now(timezone.utc))
    r6.add_to_pending_batch(dict(r6.pending_batch[0]["recipe"]), "My own")
    assert r6.pending_batch_created in _today_range(before, datetime.now(timezone.utc))
    b = _builder()
    by_note = {b.note_of(x): res for x, _, res, _ in b.pilot()[6]}
    by_note["My own"] = by_note[r6.pending_batch[0]["note"]]
    for row in list(r6.pending_batch):
        r6.tell(row["recipe"], dict(by_note[row["note"]]),
                formulation_no=row["formulation"], batch_no=6, note=row["note"])
    r6.set_pending_batch(None)
    today = _today_range(before, datetime.now(timezone.utc))
    frame = r6.history_frame(order=wording.SORT_BATCH_ORDER)
    dates = dict(zip(frame["Formulation"], frame[wording.DATE_RECORDED_COLUMN]))
    assert all(dates[no] == "" for no in range(1, 16))
    assert len(dates) == 19 and all(dates[no] in today for no in range(16, 20))
    reloaded = FoodOptimizer(sp.ROUND6, storage=store)
    exported = reloaded.history_export_frame()
    assert sorted(no for no, d in zip(exported["Formulation"],
                                      exported[wording.DATE_RECORDED_COLUMN])
                  if d in today) == [16, 17, 18, 19]


@pytest.mark.parametrize("name,n,open_rows", [(sp.ROUND6, 15, 3), (sp.ALL10, 30, 0)])
def test_shipped_projects_load_at_this_class_version(store, name, n, open_rows):
    FoodOptimizer.validate_state(sp.shipped_state(name))
    sp.ensure(store)
    opt = FoodOptimizer(name, storage=store)
    assert not opt.load_error
    assert len(opt.Y_history) == n
    assert len(opt.pending_batch or []) == open_rows
    assert opt.tasting_noise is None


def test_is_sample_is_an_exact_match():
    assert sp.is_sample(sp.ROUND6) and sp.is_sample(sp.ALL10)
    for other in ("Sample project", sp.ROUND6 + " copy", sp.ROUND6 + "_pre_restore",
                  sp.ROUND6 + "_sample_reset", "yogurt pilot - round 6 to score", ""):
        assert not sp.is_sample(other)


def test_ensure_writes_only_what_is_missing(store):
    assert sorted(sp.ensure(store)) == sorted(sp.YOGURT_NAMES)
    Path(f"{sp.ALL10}.pkl").unlink()
    before = Path(f"{sp.ROUND6}.pkl").stat().st_mtime_ns
    assert sp.ensure(store) == [sp.ALL10]
    assert Path(f"{sp.ROUND6}.pkl").stat().st_mtime_ns == before
    assert sp.ensure(store) == []


def test_ensure_never_overwrites_user_work(store):
    sp.ensure(store)
    opt = FoodOptimizer(sp.ROUND6, storage=store)
    opt.tasting_noise = 2.0
    opt.save()
    sp.ensure(store)
    assert FoodOptimizer(sp.ROUND6, storage=store).tasting_noise == 2.0


def test_restore_archives_then_writes_the_shipped_copy(store):
    sp.ensure(store)
    opt = FoodOptimizer(sp.ROUND6, storage=store)
    opt.tasting_noise = 3.0
    opt.save()
    Path("My own project.pkl").write_text('{"keep": true}', encoding="utf-8")
    first = sp.restore(sp.ROUND6, store)
    second = sp.restore(sp.ROUND6, store)
    assert first == f"{sp.ROUND6}_sample_reset"
    assert second == f"{sp.ROUND6}_sample_reset_1"
    assert FoodOptimizer(sp.ROUND6, storage=store).tasting_noise is None
    assert json.loads(Path(f"{first}.pkl").read_text(encoding="utf-8"))["tasting_noise"] == 3.0
    assert Path("My own project.pkl").read_text(encoding="utf-8") == '{"keep": true}'
    assert Path(f"{sp.ALL10}.pkl").exists()


def test_the_restored_copy_is_a_saved_copy_not_a_project(store):
    """The copy Start this sample over keeps is listed with the project's
    Saved copies, under its own reason, and never as a project of its own."""
    import wording
    sp.ensure(store)
    archived = sp.restore(sp.ROUND6, store)
    assert archived in store.list_archives()
    assert archived not in store.list_projects()
    assert sorted(store.list_projects()) == sorted(sp.YOGURT_NAMES)
    match = storage_backend.ARCHIVE_SUFFIX_RE.search(archived)
    assert storage_backend.ARCHIVE_SUFFIX_RE.sub("", archived) == sp.ROUND6
    assert wording.SAFETY_COPY_REASONS[match.group(1)] == "Before Start this example over"


def test_restore_without_a_file_just_writes_it(store):
    assert sp.restore(sp.ALL10, store) is None
    assert Path(f"{sp.ALL10}.pkl").exists()


def test_restore_refuses_a_non_sample(store):
    with pytest.raises(ValueError):
        sp.restore("My own project", store)


def test_round6_download_offer(store):
    sp.ensure(store)
    r6 = FoodOptimizer(sp.ROUND6, storage=store)
    assert sp.offers_round6_download(r6)
    assert not sp.offers_what_next(r6)
    assert not sp.offers_round6_download(FoodOptimizer(sp.ALL10, storage=store))
    assert sp.round6_results_bytes()[:2] == b"PK"      # an xlsx is a zip


def test_after_round6_is_recorded(store):
    sp.ensure(store)
    r6 = FoodOptimizer(sp.ROUND6, storage=store)
    b = _builder()
    b.score_round(r6, b.pilot()[6])
    assert not sp.offers_round6_download(r6)
    assert sp.offers_what_next(r6)
    b.add_round(r6, b.pilot()[7])           # the user goes on to Round 7
    assert not sp.offers_round6_download(r6)
    assert sp.offers_what_next(r6)


def test_an_own_formulation_added_to_round6_keeps_the_download(store):
    sp.ensure(store)
    r6 = FoodOptimizer(sp.ROUND6, storage=store)
    r6.add_to_pending_batch(dict(r6.pending_batch[0]["recipe"]), "My own")
    assert r6.pending_batch_no == 6 and len(r6.pending_batch) == 4
    assert sp.offers_round6_download(r6)
    assert not sp.offers_what_next(r6)


def test_a_round6_row_not_scored_still_offers_what_next(store):
    sp.ensure(store)
    r6 = FoodOptimizer(sp.ROUND6, storage=store)
    b = _builder()
    by_note = {b.note_of(x): res for x, _, res, _ in b.pilot()[6]}
    *scored, last = list(r6.pending_batch)
    for row in scored:
        r6.tell(row["recipe"], dict(by_note[row["note"]]),
                formulation_no=row["formulation"], batch_no=6, note=row["note"])
    assert not sp.offers_what_next(r6)          # Round 6 still open
    r6.record_skipped(last["formulation"], 6, last["recipe"])
    r6.set_pending_batch(None)
    assert len(r6.X_history) == 17
    assert not sp.offers_round6_download(r6)
    assert sp.offers_what_next(r6)


def test_round7_open_hides_the_download(store):
    sp.ensure(store)
    r6 = FoodOptimizer(sp.ROUND6, storage=store)
    b = _builder()
    b.score_round(r6, b.pilot()[6])
    b.add_round(r6, b.pilot()[7])
    assert r6.pending_batch_no == 7 and r6.pending_batch
    assert not sp.offers_round6_download(r6)
    assert sp.offers_what_next(r6)


def test_the_all10_sample_offers_neither(store):
    sp.ensure(store)
    a10 = FoodOptimizer(sp.ALL10, storage=store)
    assert not sp.offers_round6_download(a10)
    assert not sp.offers_what_next(a10)


def test_uploading_the_shipped_round6_file_records_the_pilots_results(store):
    """Spec 4.4: the bundled xlsx, read and saved the way 2 · Make a round's
    Or upload results from a file > Save uploaded results does
    (ui_batch._read_results_file, parse_batch_results, tell)."""
    from food_bo import uploaded_parts
    sp.ensure(store)
    opt = FoodOptimizer(sp.ROUND6, storage=store)
    with open(sp.ROUND6_RESULTS, "rb") as fh:
        upload = opt.results_from_workbook(fh)
    sheet, weighed, lots = uploaded_parts(upload)
    parsed, left_out = opt.parse_batch_results(
        sheet, opt.pending_batch, with_skipped=True, weighed=weighed)
    assert [no for no, _, _ in parsed] == [16, 17, 18] and left_out == []
    batch_no = opt.pending_batch_no
    by_number = {r["formulation"]: r["recipe"] for r in opt.pending_batch}
    for number, results, note in parsed:
        opt.tell(opt.amounts_as_weighed(by_number[number], weighed.get(number)),
                 results, formulation_no=number, batch_no=batch_no, note=note)
    opt.store_lots(batch_no, lots)
    opt.set_pending_batch(None)

    names = ("Consistency", "Creaminess", "Tanginess", "Uniformity")
    opt = FoodOptimizer(sp.ROUND6, storage=store)
    got = {no: tuple(opt.results_history[opt.index_of_formulation(no)][n] for n in names)
           for no in (16, 17, 18)}
    assert got == {16: (12, 5, 4, 13), 17: (12, 5, 4, 13), 18: (13, 7, 4, 12)}
    assert [opt.batch_history[opt.index_of_formulation(no)] for no in (16, 17, 18)] == [6, 6, 6]
    assert not opt.pending_batch                # Round 6 closed
    assert sp.offers_what_next(opt)
    assert not sp.offers_round6_download(opt)


# The Round 6 results file beside the projects, so Browse files lands on it.
_R6_COPY = "Round 6 results (pilot).xlsx"


def test_ensure_writes_the_round6_results_file_into_the_projects_folder(store, tmp_path):
    assert store.folder() == str(tmp_path.resolve())
    sp.ensure(store)
    copy = Path(store.folder()) / _R6_COPY
    assert copy.read_bytes() == sp.ROUND6_RESULTS.read_bytes()


def test_ensure_never_overwrites_a_changed_results_file(store):
    sp.ensure(store)
    Path(_R6_COPY).write_bytes(b"changed")
    Path(f"{sp.ALL10}.pkl").unlink()              # ensure has work to do
    assert sp.ensure(store) == [sp.ALL10]
    assert Path(_R6_COPY).read_bytes() == b"changed"
    sp.ensure(store)
    assert Path(_R6_COPY).read_bytes() == b"changed"


def test_ensure_writes_the_results_file_when_only_it_is_missing(store):
    sp.ensure(store)
    Path(_R6_COPY).unlink()
    assert sp.ensure(store) == []
    assert Path(_R6_COPY).read_bytes() == sp.ROUND6_RESULTS.read_bytes()


def test_restore_round6_rewrites_the_results_file(store):
    sp.ensure(store)
    Path(_R6_COPY).write_bytes(b"changed")
    sp.restore(sp.ROUND6, store)
    assert Path(_R6_COPY).read_bytes() == sp.ROUND6_RESULTS.read_bytes()
    Path(_R6_COPY).unlink()
    sp.restore(sp.ROUND6, store)
    assert Path(_R6_COPY).read_bytes() == sp.ROUND6_RESULTS.read_bytes()


def test_restore_all10_leaves_the_results_file(store):
    sp.ensure(store)
    Path(_R6_COPY).write_bytes(b"changed")
    sp.restore(sp.ALL10, store)
    assert Path(_R6_COPY).read_bytes() == b"changed"
    Path(_R6_COPY).unlink()
    sp.restore(sp.ALL10, store)
    assert not Path(_R6_COPY).exists()


def test_the_results_file_is_never_a_project_or_a_copy(store):
    sp.ensure(store)
    sp.restore(sp.ROUND6, store)
    assert Path(_R6_COPY).exists()
    listed = store.list_projects() + store.list_archives()
    assert not [n for n in listed if "results" in n.lower()]
    assert sorted(store.list_projects()) == sorted(sp.YOGURT_NAMES)


@pytest.mark.parametrize("doc", ["README.md", "desktop/start_here.txt", "desktop/shared/manual_walkthrough.md",
                                 "desktop/windows/start_here.txt"])
def test_the_docs_say_what_the_screen_says(doc):
    """The walkthrough names the results file by the name it has in the
    FoodOptimizer folder, and (v3) no stop: the Review has no stop rule."""
    import wording
    text = " ".join((Path(sp._ROOT) / doc).read_text(encoding="utf-8").split())
    assert 'Stop: score within noise of 100' not in text and "within noise" not in text
    assert sp.ROUND6_RESULTS_FILE_NAME in text
    assert "Browse files" in text and "FoodOptimizer" in text
    assert "it says to stop, or change the target" not in text
    # The welcome screen's two samples, by their own labels (spec §9).
    assert wording.OPEN_YOGURT_LABEL in text and wording.SEE_EXAMPLE_SETUP_LABEL in text
    assert sp.BURGER in text and "Try the sample project" not in text
    # v3: What next?'s two buttons, by their labels on screen.
    assert wording.start_round_label(7) in text and wording.SEE_WHAT_HAPPENED_NEXT in text
    assert "See what happened next" not in text


@pytest.mark.parametrize("doc", ["README.md", "build.py"])
def test_no_culture_says_where_it_comes_from(doc):
    """'No culture or fermentation' is a claim about the pilot; where it is
    made, its source is said with it: the pilot's method (mix, cook,
    chill) and its ingredients, none of them a culture."""
    import re
    text = " ".join((sp.SAMPLE_DIR / doc).read_text(encoding="utf-8").split())
    claims = [m.start() for m in re.finditer("no culture or fermentation", text, re.I)]
    assert claims
    for at in claims:
        said = text[at:at + 120]
        assert "the pilot's method is mix, cook, chill" in said, said
        assert "no culture among its ingredients" in said, said


def _xlsx_write_fails(monkeypatch):
    """Every write of the results file fails (a full or read-only disk);
    the project files are written by LocalStorage.save, not this."""
    def boom(self, data):
        raise PermissionError("read-only")
    monkeypatch.setattr(Path, "write_bytes", boom)


def test_a_failed_results_file_write_never_blocks_the_samples(store, monkeypatch):
    _xlsx_write_fails(monkeypatch)
    assert sorted(sp.ensure(store)) == sorted(sp.YOGURT_NAMES)
    assert all(Path(f"{n}.pkl").exists() for n in sp.YOGURT_NAMES)
    assert not Path(_R6_COPY).exists()


def test_a_failed_results_file_write_never_fails_a_restore(store, monkeypatch):
    sp.ensure(store)
    opt = FoodOptimizer(sp.ROUND6, storage=store)
    opt.tasting_noise = 3.0
    opt.save()
    Path(_R6_COPY).unlink()
    Path(_R6_COPY).mkdir()                      # a folder in its place
    assert sp.restore(sp.ROUND6, store) == f"{sp.ROUND6}_sample_reset"
    assert FoodOptimizer(sp.ROUND6, storage=store).tasting_noise is None
    Path(_R6_COPY).rmdir()
    _xlsx_write_fails(monkeypatch)              # or a disk that refuses it
    assert sp.restore(sp.ROUND6, store) == f"{sp.ROUND6}_sample_reset_1"
    assert FoodOptimizer(sp.ROUND6, storage=store).tasting_noise is None


def test_a_project_write_failure_still_raises(store, monkeypatch):
    def boom(name, state):
        raise OSError("disk full")
    monkeypatch.setattr(store, "save", boom)
    with pytest.raises(OSError):
        sp.ensure(store)


# ------------------------------------------------------------------ #
#  The campaign's own limits (paper, §5 and Appendix G)
# ------------------------------------------------------------------ #

CAMPAIGN_BOUNDS = {
    "Coconut cream": (46.0, 61.0), "Tap water": (25.0, 42.0),
    "Lemon juice": (0.0, 3.2), "Salt": (0.0, 1.8),
    "Nutritional yeast": (0.0, 2.5), "Tapioca starch": (0.9, 2.6),
    "Xanthan gum": (0.15, 0.65)}
CAMPAIGN_TOTAL = (97.69051, 98.69051)    # 100 g less the two fixed ingredients


@pytest.mark.parametrize("name", [sp.ROUND6, sp.ALL10])
def test_the_limits_are_the_campaigns_own(name):
    b = _builder()
    assert {n: (lo, hi) for n, lo, hi in b.INGREDIENTS.values()} == CAMPAIGN_BOUNDS
    assert (b.BATCH_TOTAL_MIN, b.BATCH_TOTAL_MAX) == CAMPAIGN_TOTAL
    state = sp.shipped_state(name)
    assert {v["name"]: tuple(v["bounds"]) for v in state["variables"]} == CAMPAIGN_BOUNDS
    [limit] = state["quantity_constraints"]
    assert (limit["name"], limit["min"], limit["max"]) == ("Ingredient total", *CAMPAIGN_TOTAL)
    assert sorted(limit["ingredients"]) == sorted(CAMPAIGN_BOUNDS)


def test_every_batch_egbo_chose_sits_inside_every_limit(store):
    """Rounds 3-10 (the engine's) sit inside every limit, as the app checks a
    limit; Rounds 1-2 (the random starting set) inside every ingredient's
    Lowest and Highest, but not all inside the total."""
    b = _builder()
    sp.ensure(store)
    opt = FoodOptimizer(sp.ALL10, storage=store)
    outside_total = []
    for it, batches in b.pilot().items():
        for batch, recipe, _, _ in batches:
            for ing, (lo, hi) in CAMPAIGN_BOUNDS.items():
                assert lo - 1e-9 <= recipe[ing] <= hi + 1e-9, (batch, ing, recipe[ing])
            if it > b.STARTING_ROUNDS:
                assert opt._check_constraints(recipe), batch
                total = sum(recipe.values())
                assert CAMPAIGN_TOTAL[0] - 1e-5 <= total <= CAMPAIGN_TOTAL[1] + 1e-5, batch
            elif not opt._check_constraints(recipe):
                outside_total.append(batch)
    assert outside_total == ["1.1", "1.2", "1.3", "2.1", "2.2", "2.3"]


def test_the_notes_say_where_each_batch_came_from():
    b = _builder()
    assert b.STARTING_ROUNDS == 2
    assert b.note_of("1.1") == "Random starting set 1.1"
    assert b.note_of("2.3") == "Random starting set 2.3"
    assert b.note_of("3.1") == "Engine's pick 3.1"
    assert all(len(b.note_of(x).split()) <= 5 for x in ("1.1", "10.3"))
    notes = sp.shipped_state(sp.ALL10)["notes_history"]
    assert notes[:6] == [f"Random starting set {r}.{k}" for r in (1, 2) for k in (1, 2, 3)]


def test_the_limit_reads_as_the_check_reads_it(store):
    """97.69 g would look inside 'at least 97.69 g' and still break the
    pilot's 97.69051 g limit (0.0005 g short, past the check's slack), so
    the limit reads to the digits the check needs, and an amount at either
    number shown meets it."""
    sp.ensure(store)
    opt = FoodOptimizer(sp.ROUND6, storage=store)
    [limit] = opt.quantity_constraints
    assert opt.limit_text(limit) == \
        "Ingredient total: at least 97.6905 g and at most 98.6905 g"
    names = limit["ingredients"]
    for total in (97.6905, 98.6905):
        recipe = {n: total / len(names) for n in names}
        assert opt._check_constraints(recipe), total
    # The sidebar "?" names the total in the same numbers.
    import wording
    assert "their total at 97.6905–98.6905 g per 100 g" in wording.SAMPLE_PROVENANCE_HELP


# ------------------------------------------------------------------ #
#  A copy an earlier build wrote is brought up to date
# ------------------------------------------------------------------ #

_OLD_BOUNDS = {"Coconut cream": (48.0, 61.0), "Lemon juice": (0.0, 3.0),
               "Tapioca starch": (1.0, 2.6), "Xanthan gum": (0.1, 0.65)}


def _named_build(name):
    """The state 5dc62ff shipped: the same pilot, its Rounds 3-10 notes
    naming the method ('Pilot batch 3.1 · EGBO')."""
    import re
    state = sp.shipped_state(name)
    pick = re.compile(r"^Engine's pick (\d+\.\d+)$")
    def old(note):
        m = pick.match(note) if isinstance(note, str) else None
        return f"Pilot batch {m.group(1)} \u00b7 EGBO" if m else note
    state["notes_history"] = [old(n) for n in state["notes_history"]]
    for row in state.get("pending_batch") or []:
        if "note" in row:
            row["note"] = old(row["note"])
    return state


def _v3_build(name):
    """The state db0d6b7 (v3) shipped: as _named_build, its attributes still
    named Tangy and Thickness uniformity (the pilot's names are Tanginess
    and Uniformity)."""
    import json
    text = (json.dumps(_named_build(name))
            .replace('"Tanginess"', '"Tangy"').replace('"Uniformity"', '"Thickness uniformity"'))
    return json.loads(text)


@pytest.mark.parametrize("name", [sp.ROUND6, sp.ALL10])
def test_the_build_that_named_the_method_is_known(name):
    assert sp.fingerprint(_named_build(name)) in sp._SUPERSEDED[name]


def _earlier_build(name):
    """The state the previous build (29b1b24) shipped: limits inferred from
    the data, 'Pilot starting set' notes and the old attribute names."""
    state = _v3_build(name)
    for v in state["variables"]:
        if v["name"] in _OLD_BOUNDS:
            v["bounds"] = list(_OLD_BOUNDS[v["name"]])
    state["quantity_constraints"][0].update({"min": 97.6, "max": 98.7})
    state["notes_history"] = [n.replace("Random starting set", "Pilot starting set")
                              for n in state["notes_history"]]
    return state


@pytest.mark.parametrize("name", [sp.ROUND6, sp.ALL10])
def test_the_earlier_build_is_known_and_this_one_is_not(name):
    assert sp.fingerprint(_earlier_build(name)) in sp._SUPERSEDED[name]
    assert sp.fingerprint(sp.shipped_state(name)) not in sp._SUPERSEDED[name]


@pytest.mark.parametrize("name", [sp.ROUND6, sp.ALL10])
def test_the_v3_build_with_the_old_attribute_names_is_known(name):
    assert sp.fingerprint(_v3_build(name)) in sp._SUPERSEDED[name]
    assert sp.fingerprint(_v3_build(name)) != sp.fingerprint(sp.shipped_state(name))


def test_refresh_gives_an_untouched_v3_copy_the_paper_names(store):
    for name in sp.YOGURT_NAMES:
        store.save(name, _v3_build(name))
    assert sorted(sp.refresh(store)) == sorted(sp.YOGURT_NAMES)
    for name in sp.YOGURT_NAMES:
        assert store.load(name) == sp.shipped_state(name)
        assert [o["name"] for o in store.load(name)["objectives"]] == [
            "Consistency", "Creaminess", "Tanginess", "Uniformity"]
    assert sp.refresh(store) == []                  # once only


def test_a_changed_v3_copy_keeps_its_names(store):
    state = _v3_build(sp.ROUND6)
    state["notes_history"][0] = "mine"
    store.save(sp.ROUND6, state)
    assert sp.refresh(store) == []
    assert store.load(sp.ROUND6)["notes_history"][0] == "mine"


def test_refresh_replaces_an_untouched_copy_from_an_earlier_build(store):
    for name in sp.YOGURT_NAMES:
        store.save(name, _earlier_build(name))
    assert sorted(sp.refresh(store)) == sorted(sp.YOGURT_NAMES)
    for name in sp.YOGURT_NAMES:
        assert store.load(name) == sp.shipped_state(name)
    assert sp.refresh(store) == []                  # once only


def test_ensure_replaces_an_untouched_copy_from_an_earlier_build(store):
    store.save(sp.ROUND6, _earlier_build(sp.ROUND6))
    assert sorted(sp.ensure(store)) == sorted(sp.YOGURT_NAMES)
    assert store.load(sp.ROUND6) == sp.shipped_state(sp.ROUND6)


def test_a_changed_copy_from_an_earlier_build_is_kept(store):
    state = _earlier_build(sp.ROUND6)
    state["tasting_noise"] = 2.0
    store.save(sp.ROUND6, state)
    before = Path(f"{sp.ROUND6}.pkl").read_text(encoding="utf-8")
    assert sp.refresh(store) == []
    assert sp.ensure(store) == [sp.ALL10]
    assert Path(f"{sp.ROUND6}.pkl").read_text(encoding="utf-8") == before
    # Start this sample over still brings it up to date, keeping a copy.
    archived = sp.restore(sp.ROUND6, store)
    assert store.load(sp.ROUND6) == sp.shipped_state(sp.ROUND6)
    assert store.load(archived)["tasting_noise"] == 2.0
