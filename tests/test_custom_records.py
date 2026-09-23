"""Custom unscored records persist independently and survive workbook round trips."""
import copy
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

import custom_records as cr
from food_bo import FoodOptimizer


@pytest.fixture
def project(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer('my_project', robust=False)
    opt.add_ingredient('Protein', 5, 20)
    opt.add_ingredient('Water', 0, 100)
    opt.set_formulation_total(100)
    opt.set_formula('Water', '= rest')
    opt.add_objective('Taste', 1, goal='max', min_val=0, max_val=10)
    opt.set_pending_batch([{'Protein': 10, 'Water': 90}, {'Protein': 15, 'Water': 85}], batch_no=1)
    return opt


def data(book):
    from workbook_flow import hide_metadata, restore_metadata
    restore_metadata(book)
    hide_metadata(book)
    out = io.BytesIO(); book.save(out); out.seek(0)
    return out


def marked(book):
    from workbook_flow import restore_metadata
    restore_metadata(book)
    return [c for s in book for row in s for c in row if c.comment and c.comment.text.startswith(cr.MARKER)]


def test_persistence_scopes_visibility_and_rename(project):
    a = cr.add_field(project, 'Operator', 'formulation')
    b = cr.add_field(project, 'Expiry date', 'ingredient')
    before = copy.deepcopy(project.pending_batch)
    cr.save_values(project, 1, [('formulation', '1', a, '0012'), ('ingredient', 'Protein', b, '2027-01-31')])
    cr.set_enabled(project, [])
    loaded = FoodOptimizer('my_project')
    assert cr.values(loaded, 1, 'formulation', 1)[a] == '0012'
    assert cr.fields(loaded) == []
    assert len(cr.export_frame(loaded)) == 2
    assert loaded.pending_batch == before
    assert not loaded.formulation_ids
    cr.set_enabled(loaded, [a, b])
    loaded.rename_variable('Protein', 'Pea protein')
    assert cr.values(FoodOptimizer('my_project'), 1, 'ingredient', 'Pea protein')[b] == '2027-01-31'
    assert cr.values(loaded, 2, 'ingredient', 'Pea protein') == {}


@pytest.mark.parametrize('name,scope', [('', 'ingredient'), ('x'*81, 'ingredient'), ('Operator', 'bad'), ('Lot', 'ingredient')])
def test_invalid_definition_atomic(project, name, scope):
    before = project.export_json()
    with pytest.raises(ValueError): cr.add_field(project, name, scope)
    assert project.export_json() == before


def test_duplicate_and_invalid_values_atomic(project):
    a = cr.add_field(project, 'Operator', 'formulation')
    with pytest.raises(ValueError): cr.add_field(project, ' OPERATOR ', 'formulation')
    cr.add_field(project, 'Operator', 'ingredient')
    before = project.export_json()
    with pytest.raises(ValueError): cr.save_values(project, 1, [('formulation', '1', a, 'ok'), ('ingredient', 'Water', a, 'wrong scope')])
    assert project.export_json() == before


@pytest.mark.parametrize('print_pack', [False, True])
def test_workbook_both_scopes_custom_only_and_hidden(project, print_pack):
    a = cr.add_field(project, 'Operator', 'formulation')
    b = cr.add_field(project, 'Expiry date', 'ingredient')
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=print_pack)))
    if not print_pack: assert [s.title for s in book if s.sheet_state == 'visible'] == ['Round overview', 'Preparation', 'Results']
    cells = marked(book)
    assert len(cells) == 4
    for i, c in enumerate(cells):
        assert not c.protection.locked
        c.value = '0012' if i == 0 else str(i)
    cr.set_enabled(project, [])
    uploaded = project.results_from_workbook(data(book))
    assert uploaded.frame.empty
    assert len(uploaded.custom_records) == 4
    cr.save_values(project, 1, uploaded.custom_records)
    assert cr.values(project, 1, 'formulation', 1)[a] == '0012'
    assert cr.values(project, 1, 'ingredient', 'Water')[b] == '3'
    cells[0].value = None
    assert len(cr.read_workbook(data(book), project, 1)) == 3
    assert cr.values(project, 1, 'formulation', 1)[a] == '0012'


@pytest.mark.parametrize('bad', ['formula', 'scope', 'conflict', 'metadata'])
def test_invalid_workbook_refused(project, bad):
    cr.add_field(project, 'Operator', 'formulation')
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    cells = marked(book)
    if bad == 'formula': cells[0].value = '=1+1'
    elif bad == 'scope': cells[0].comment.text = cr.MARKER + json.dumps({'round':1,'scope':'bad','field':'bad','subject':'1'})
    elif bad == 'metadata': cells[0].comment.text = cr.MARKER + '[]'
    else:
        cells[0].value = 'one'; cells[1].value = 'two'; cells[1].comment = copy.copy(cells[0].comment)
    with pytest.raises(ValueError): project.results_from_workbook(data(book))


def test_text_formula_export_is_literal(project):
    a = cr.add_field(project, '=Operator', 'formulation')
    cr.save_values(project, 1, [('formulation', '1', a, '=1+1')])
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    assert marked(book)[0].data_type == 's'
    assert cr.read_workbook(data(book), project, 1)[0][-1] == '=1+1'


@pytest.mark.parametrize('value', [None, [], {}, {'fields': [], 'values': []}, {'fields':[{'id':'x','name':'X','scope':[],'enabled':True}], 'values':{}}])
def test_malformed_state(value):
    assert not cr.valid(value)


def test_app_add_scopes_edit_twice_and_reload(project):
    at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'), default_timeout=180).run()
    assert not at.exception
    for name, scope in [('Operator', 'formulation'), ('Expiry date', 'ingredient')]:
        at.text_input(key='custom_field_name').set_value(name)
        at.selectbox(key='custom_field_scope').select(scope)
        at.button(key='custom_field_add').click().run()
        assert not at.exception
    loaded = FoodOptimizer('my_project')
    defs = cr.fields(loaded)
    assert {f['scope'] for f in defs} == {'formulation', 'ingredient'}
    for value in ['First entry', 'Corrected entry']:
        key = next(k for k in at.session_state.filtered_state if k.startswith('custom_values_') and '_round_formulation_' in k)
        field = next(f['id'] for f in defs if f['scope'] == 'formulation')
        at.session_state[key] = {'edited_rows': {0: {field: value}}, 'added_rows': [], 'deleted_rows': []}
        at.run()
        assert not at.exception
        assert cr.values(FoodOptimizer('my_project'), 1, 'formulation', 1)[field] == value
    at.multiselect(key='record_fields').set_value([]).run()
    assert not at.exception
    assert cr.values(FoodOptimizer('my_project'), 1, 'formulation', 1)[field] == 'Corrected entry'


def test_custom_only_upload_saves_without_completing_round(project):
    import ui_batch
    import wording
    field = cr.add_field(project, 'Operator', 'formulation')
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    marked(book)[0].value = 'A. Scientist'
    filled = data(book); filled.name = 'round.xlsx'
    at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'), default_timeout=180).run()
    at.session_state['_results_upload'] = ui_batch._read_results_file(at.session_state['optimizer'], filled)
    at.run()
    assert not at.exception
    assert not at.error
    at.button(key='save_uploaded').click().run()
    assert not at.exception
    assert not at.error
    loaded = FoodOptimizer('my_project')
    assert len(loaded.pending_batch) == 2
    assert not loaded.results_history
    assert cr.values(loaded, 1, 'formulation', 1)[field] == 'A. Scientist'


def test_json_validation_legacy_and_record_export(project):
    field = cr.add_field(project, 'Operator', 'formulation')
    cr.save_values(project, 1, [('formulation', '1', field, '=literal text')])
    state = json.loads(json.dumps(project.export_json()))
    assert FoodOptimizer.validate_state(state)['version'] == 15
    state.pop('custom_records')
    assert FoodOptimizer.validate_state(state)['version'] == 15
    state['custom_records'] = {'fields': [], 'values': {'1': {'ingredient': {'Water': {'missing':'bad'}}}}}
    with pytest.raises(ValueError): FoodOptimizer.validate_state(state)
    book = load_workbook(io.BytesIO(project.all_formulations_workbook()))
    assert 'Additional records' in book.sheetnames
    cell = next(c for row in book['Additional records'] for c in row if c.value == '=literal text')
    assert cell.data_type == 's'


def test_history_editor_does_not_revert_current_round_entries(project):
    field = cr.add_field(project, 'Operator', 'formulation')
    project.tell({'Protein': 10, 'Water': 90}, {'Taste':8}, formulation_no=1, batch_no=1)
    cr.save_values(project, 1, [('formulation', '1', field, 'First')])
    at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'), default_timeout=180).run()
    key = next(k for k in at.session_state.filtered_state if k.startswith('custom_values_') and '_history_formulation_' in k)
    at.session_state[key] = {'edited_rows': {0: {field: 'Edited in history'}}, 'added_rows': [], 'deleted_rows': []}
    at.run(); at.run()
    assert not at.exception
    assert cr.values(FoodOptimizer('my_project'), 1, 'formulation', 1)[field] == 'Edited in history'


def _portioned(opt):
    """A project whose Dry blend is a pre-mix kept at fixed proportions, so
    its parts are NOT rows of the flat list."""
    opt.add_ingredient('Dry blend', 10, 20)
    opt.add_premix('Dry blend', 'portioned', over_row=True)
    opt.set_premix_parts('Dry blend', [{'name': 'Pea protein isolate', 'share': 70},
                                       {'name': 'Methylcellulose', 'share': 30}])
    opt.set_records('lot', True)
    return opt


def test_a_lot_against_a_premix_part_never_locks_the_saved_copy(project):
    """0.7.1 offered the parts of a portioned pre-mix in the Lot table and
    then refused to open the copy the user saved: the part is not a row, and
    validate_state took that for damage. Those copies exist, so the loader
    keeps them."""
    opt = _portioned(project)
    opt.store_lots(1, {'Pea protein isolate': 'L-1', 'Protein': 'L-2'})
    state = json.loads(json.dumps(opt.export_json()))
    assert FoodOptimizer.validate_state(state)['version'] == 15
    reopened = FoodOptimizer('reopened', robust=False)
    reopened.import_json(state)
    assert reopened.lots[1]['Pea protein isolate'] == 'L-1'
    # ...and a name that is neither a row nor a part is still damage.
    state['lots']['1']['Nothing here'] = 'L-3'
    with pytest.raises(ValueError):
        FoodOptimizer.validate_state(state)


def test_the_lot_table_offers_rows_and_never_a_portioned_premixs_parts(project):
    """A part of a portioned pre-mix is weighed on the pre-mix's own
    preparation sheet, and its lot is written there. Offering it twice gave
    two places to write one fact — and the second one broke the copy."""
    opt = _portioned(project)
    offered = cr.lot_names(opt, 1)
    assert 'Dry blend' not in offered, offered
    assert 'Pea protein isolate' not in offered and 'Methylcellulose' not in offered
    # A row whose amount is calculated has no lot either: water has no lot
    # number at a bench.
    assert 'Water' not in offered, offered
    assert offered == ['Protein']
