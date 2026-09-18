"""Bench-facing pre-mix pages, totals, protection, and Excel round trips."""
import io

import pytest
from openpyxl import load_workbook

from food_bo import FoodOptimizer
import wording


@pytest.fixture
def premix_project(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer('pre-mix pages', robust=False)
    for field in ('vendor', 'sku', 'lot', 'actual'):
        opt.set_records(field, True)
    opt.set_amount_unit('g')
    opt.add_ingredient('Water', 0, 100)
    opt.add_objective('Taste', 1, goal='max', min_val=0, max_val=10)
    opt.add_premix('Dry blend', 'portioned')
    opt.set_premix_parts('Dry blend', [
        {'name': 'Flour', 'share': 60, 'unit': 'g', 'vendor': 'Mill'},
        {'name': 'Salt', 'share': 40, 'unit': 'g'}])
    opt.add_ingredient('Dry blend', 10, 60)
    opt.add_premix('Fat phase', 'weighed')
    opt.set_premix_parts('Fat phase', [
        {'name': 'Coconut oil', 'share': 50, 'unit': 'g'},
        {'name': 'Sunflower oil', 'share': 50, 'unit': 'g'}])
    opt.add_ingredient('Coconut oil', 0, 20)
    opt.add_ingredient('Sunflower oil', 0, 20)
    opt.set_pending_batch([
        {'Water': 60, 'Dry blend': 30, 'Coconut oil': 4, 'Sunflower oil': 6},
        {'Water': 40, 'Dry blend': 50, 'Coconut oil': 7, 'Sunflower oil': 3}], batch_no=1)
    opt._snapshot_premixes(1)
    return opt


def workbook(opt, total=100):
    return load_workbook(io.BytesIO(opt.workbook_bytes(opt.pending_batch, total)))


def row_for(sheet, label, column=1):
    return next(row for row in range(1, sheet.max_row + 1)
                if sheet.cell(row, column).value == label)


def test_portioned_page_is_first_and_names_its_round_quantity(premix_project):
    """The page asks for a MAKEABLE quantity - the greater of the need plus
    a tenth and 100 g, rounded up to the next 5 g - and says what the round
    takes out of it. `make 160.00 g` asked for a blend dispensed with
    nothing left in the bowl, on the paddle or on the scoop."""
    book = workbook(premix_project, 200)
    assert book.sheetnames[:2] == ['Pre-mix · Dry blend', 'Round 1']
    sheet = book.active
    assert sheet['A1'].value == \
        'Dry blend · Prepare 180 g (this round needs 160.00 g)'
    # The parts are scaled to what is MADE, and the Total adds the printed
    # numbers rather than the ones behind them.
    assert sheet['B4'].value == 108
    assert sheet['B5'].value == 72
    assert sheet['B6'].value == 180
    assert sheet['A6'].value == wording.TOTAL_LABEL
    assert sheet['F4'].value == 'Mill'
    # The pre-mix's own lot, which nothing anywhere used to give it.
    labels = [sheet.cell(r, 1).value for r in range(1, sheet.max_row + 1)]
    assert wording.PREMIX_LOT_LABEL in labels
    assert wording.PREMIX_BLEND_TIME_LABEL in labels


def test_prep_page_uses_the_rounds_original_makeup(premix_project):
    premix_project.set_premix_parts('Dry blend', [
        {'name': 'Flour', 'share': 90, 'unit': 'g'},
        {'name': 'Salt', 'share': 10, 'unit': 'g'}])
    sheet = workbook(premix_project).active
    # 80 g needed, so 100 g made: 60 / 40 of the round's own make-up.
    assert sheet['A1'].value == \
        'Dry blend · Prepare 100 g (this round needs 80.00 g)'
    assert sheet['B4'].value == 60
    assert sheet['B5'].value == 40


def test_portioned_blend_is_one_line_and_weighed_parts_are_indented(premix_project):
    book = workbook(premix_project)
    sheet = book['Formulation 1']
    labels = [sheet.cell(r, 2).value for r in range(1, sheet.max_row + 1)]
    assert labels.count('Dry blend') == 1
    assert 'Flour' not in labels
    group = row_for(sheet, wording.premix_group_line('Fat phase'), 2)
    assert sheet.cell(group, 2).font.bold
    for name in ('Coconut oil', 'Sunflower oil'):
        r = row_for(sheet, name, 2)
        assert sheet.cell(r, 2).alignment.indent == 1
        assert not sheet.cell(r, 4).protection.locked
    total = row_for(sheet, wording.premix_group_total_line('Fat phase'), 2)
    assert sheet.cell(total, 3).value == 10
    assert sheet.cell(total, 4).protection.locked


def test_the_round_sheet_separates_what_is_made_from_what_is_needed(
        premix_project):
    """One heading totalled two kinds of number. A pre-mix's parts ARE
    weighed together at the number printed; `Water 200` is two formulations
    added up and is weighed at that number nowhere."""
    sheet = workbook(premix_project, 200)['Round 1']
    made = row_for(sheet, wording.MAKE_FOR_ROUND_HEADING)
    assert sheet.cell(made + 1, 2).value == 'Amount (g)'
    assert sheet.cell(made + 2, 1).value == \
        'Dry blend · Prepare 180 g (this round needs 160.00 g)'
    assert sheet.cell(made + 2, 2).value == 180
    on_hand = row_for(sheet, wording.HAVE_ON_HAND_HEADING)
    assert sheet.cell(on_hand + 1, 2).value == 'Total for this round (g)'
    rows = {}
    for i in range(on_hand + 2, sheet.max_row + 1):
        name = sheet.cell(i, 1).value
        if name == wording.HAVE_ON_HAND_CAPTION or name is None:
            break
        rows[name] = sheet.cell(i, 2).value
    assert rows == {'Water': 200, 'Coconut oil': 22, 'Sunflower oil': 18}
    assert sheet.cell(on_hand + 5, 1).value == wording.HAVE_ON_HAND_CAPTION


def test_prep_page_only_unlocks_lot_cells(premix_project):
    sheet = workbook(premix_project).active
    assert sheet.protection.sheet
    # The parts' Lot cells, and the pre-mix's own four write-in lines.
    assert {c.coordinate for row in sheet for c in row
            if not c.protection.locked} == {
        'D4', 'D5', 'E4', 'E5', 'B8', 'B9', 'B10', 'B11'}
    assert sheet['D4'].fill.fgColor.rgb == '00FFF2CC'


def test_long_or_invalid_names_get_unique_legal_tabs(premix_project):
    batch = premix_project.pending_batch
    for name in ('A/' + 'x' * 40, 'A?' + 'x' * 40, 'Round 1', 'x' * 20 + "'long"):
        premix_project.add_premix(name, 'portioned')
        premix_project.set_premix_parts(name, [{'name': 'Flour', 'share': 100, 'unit': 'g'}])
    premix_project.set_pending_batch(batch, batch_no=1)
    book = workbook(premix_project)
    assert len({n.lower() for n in book.sheetnames}) == len(book.sheetnames)
    assert all(len(n) <= 31 and not any(c in n for c in '[]:*?/\\') for n in book.sheetnames)
    assert all(not n.startswith("'") and not n.endswith("'") for n in book.sheetnames)
    assert 'Round 1' in book.sheetnames and 'Formulation 1' in book.sheetnames


def test_setup_sheet_names_groups_parts_and_modes(premix_project):
    sheet = load_workbook(io.BytesIO(premix_project.all_formulations_workbook()))[wording.SET_UP_SHEET]
    r = row_for(sheet, wording.PREMIXES_HEADING)
    values = {c.value for row in sheet.iter_rows(min_row=r) for c in row if c.value is not None}
    assert {'Dry blend', 'Fat phase', 'Flour', 'Salt', 'Coconut oil', 'Sunflower oil',
            wording.PREMIX_MADE_AS_PORTIONED, wording.PREMIX_MADE_AS_WEIGHED} <= values
    first_group = row_for(sheet, 'Fat phase')
    assert sheet.cell(first_group, 5).value is None


@pytest.mark.parametrize('legacy_pointer', [False, True])
def test_lots_and_actual_part_amounts_survive_workbook_upload(premix_project, legacy_pointer):
    book = workbook(premix_project)
    if legacy_pointer:
        for row in book['Round 1']:
            for cell in row:
                if cell.value == wording.PREMIX_LOT_ON_ITS_PAGE:
                    cell.value = 'see its page'
    book.active['E4'] = 'F-123'
    page = book['Formulation 1']
    page.cell(row_for(page, 'Coconut oil', 2), 4, 4.5)
    page.cell(row_for(page, 'Taste', 2), 4, 8)
    source = io.BytesIO()
    book.save(source)
    source.seek(0)
    upload = premix_project.results_from_workbook(source)
    assert upload.lots == {'Flour': 'F-123'}
    assert upload.actual == {1: {'Coconut oil': 4.5}}
    assert upload.frame['Taste'].tolist() == [8]


def test_an_empty_premix_cannot_generate_a_round(premix_project):
    premix_project.add_premix('New pre-mix', 'portioned')
    from ui_helpers import readiness
    ready, message = readiness(premix_project)
    assert not ready and 'New pre-mix' in message
    with pytest.raises(ValueError, match='New pre-mix'):
        premix_project.ask(1)


def test_regenerating_a_recorded_round_preserves_its_makeup(premix_project, monkeypatch):
    opt = premix_project
    recipe = dict(opt.pending_batch[0])
    opt.tell(recipe, {'Taste': 8}, formulation_no=1, batch_no=1)
    opt.set_premix_parts('Dry blend', [
        {'name': 'Flour', 'share': 90, 'unit': 'g'},
        {'name': 'Salt', 'share': 10, 'unit': 'g'}])
    observed = []
    def generate(*args):
        observed.append(opt._property_parts('Dry blend')[0]['share'])
        return [recipe]
    monkeypatch.setattr(opt, '_ask_cold_start', generate)
    opt.set_pending_batch(None)
    opt.ask(1, batch_no=1)
    assert observed == [60]
    assert opt.premix_version_of('Dry blend', 0)[0]['share'] == 60
    opt.set_pending_batch(None)
    opt.ask(1)
    assert observed == [60, 90]
    assert opt.premix_version_of('Dry blend', 0)[0]['share'] == 60


def test_switching_a_weighed_group_does_not_count_old_parts_in_rule_preview(
        premix_project):
    opt = premix_project
    opt.set_pending_batch(None)
    opt.add_ingredient('Coconut oil', 3, 10)
    opt.add_ingredient('Sunflower oil', 3, 10)
    opt.add_ingredient('Dry blend', 1, 2)
    opt.set_formulation_total(10)
    opt.set_formula('Water', '= rest')
    frame = opt.ingredient_grid_frame()
    row = frame.index[frame[wording.NAME_LABEL] == 'Fat phase'][0]
    frame.loc[row, wording.MADE_AS_LABEL] = wording.PREMIX_MADE_AS_PORTIONED
    errors, _ = opt.apply_ingredient_grid(frame)
    assert errors == []
    assert set(opt._by_name()) == {'Water', 'Dry blend', 'Fat phase'}
