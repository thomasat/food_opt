"""0.8.0 review fixes, Task 3: a recorded formulation whose result will never
come is moved to Not scored in place (spec 2026-09-29 §1.1, §9 B2/I5).
It keeps its row, Round and partial values; it is left out exactly like a
waiting one and is not counted as waiting; entering the value, or Wait for
its results again, undoes it. Score without is gone."""
import json

import pytest

import insights
import wording
from food_bo import FoodOptimizer
from tests.test_ask_unchanged import exempt_and_waiting, fixture, run, same
from tests.test_complete_results import _full, _patty, _tell


@pytest.fixture(autouse=True)
def _isolate_cwd(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)


def _with_one_waiting(rows=6):
    """`rows` complete formulations, then Formulation 7 waiting for Cook loss."""
    opt = _patty(rows=rows)
    _tell(opt, 7, 3, drop=("Cook loss",))
    return opt


def test_moving_to_not_scored_keeps_the_row_and_its_partial_values():
    opt = _with_one_waiting()
    opt.set_aside_formulation(7)
    i = opt.index_of_formulation(7)
    assert opt.set_aside == [7]
    assert opt.is_set_aside(i)
    assert opt.batch_history[i] == 3
    assert opt.results_history[i] == {k: v for k, v in _full(7).items() if k != "Cook loss"}
    assert opt.waiting_rows() == [i]             # still out of the model
    assert opt.awaited_rows() == []              # but no longer waiting
    assert opt.complete_count() == 6
    assert i not in opt.best_candidates()


def test_it_reads_not_scored_with_what_was_not_measured():
    opt = _with_one_waiting()
    opt.set_aside_formulation(7)
    frame = opt.history_frame()
    row = frame[frame["Formulation"] == 7].iloc[0]
    assert row[wording.OVERALL_SCORE_COLUMN] == "Not scored · Cook loss not measured"
    assert frame["Formulation"].tolist()[-1] == 7          # the Not scored tier
    assert wording.not_scored_missing("Cook loss") == "Not scored · Cook loss not measured"


def test_entering_the_missing_value_scores_it_again():
    opt = _with_one_waiting()
    opt.set_aside_formulation(7)
    i = opt.index_of_formulation(7)
    opt.edit_result(i, dict(opt.results_history[i], **{"Cook loss": 12.0}))
    assert opt.set_aside == []
    assert opt.fully_measured(i) and not opt.is_set_aside(i)


def test_wait_again_undoes_it():
    opt = _with_one_waiting()
    opt.set_aside_formulation(7)
    opt.wait_again(7)
    assert opt.set_aside == []
    assert opt.awaited_rows() == [opt.index_of_formulation(7)]


def test_only_a_formulation_that_waits_can_be_moved():
    opt = _with_one_waiting()
    with pytest.raises(ValueError, match="already has all its results"):
        opt.set_aside_formulation(1)
    with pytest.raises(IndexError):
        opt.set_aside_formulation(99)


def test_deleting_the_measurement_it_lacked_makes_it_an_ordinary_formulation():
    opt = _with_one_waiting()
    opt.set_aside_formulation(7)
    opt.remove_objective("Cook loss")
    i = opt.index_of_formulation(7)
    assert opt.fully_measured(i) and not opt.is_set_aside(i)
    # Pruned from the file too: a later edit that left it short again would
    # otherwise bring the flag back (Task 3 review, M-1).
    assert opt.set_aside == []
    assert FoodOptimizer(opt.project_name).set_aside == []


def test_deleting_the_formulation_drops_its_number():
    opt = _with_one_waiting()
    opt.set_aside_formulation(7)
    opt.delete_formulation(7)
    assert opt.set_aside == []


def test_undoing_its_round_drops_its_number():
    opt = _with_one_waiting()
    opt.set_aside_formulation(7)
    opt.undo_last_batch()
    assert opt.index_of_formulation(7) is None
    assert opt.set_aside == []
    assert FoodOptimizer(opt.project_name).set_aside == []


def test_rewinding_past_it_drops_its_number():
    opt = _with_one_waiting()
    opt.set_aside_formulation(7)
    opt.rewind_to(opt.index_of_formulation(6))
    assert opt.set_aside == []


def test_a_saved_copy_never_carries_a_stale_number():
    opt = _with_one_waiting()
    opt.set_aside_formulation(7)
    opt.set_aside.append(99)                     # no such formulation
    assert opt.export_json()['set_aside'] == [7]


def test_with_nothing_named_it_reads_plain_not_scored():
    assert wording.not_scored_missing("") == wording.NOT_SCORED


def test_set_aside_survives_a_saved_copy_and_a_reopen():
    opt = _with_one_waiting()
    opt.set_aside_formulation(7)
    state = json.loads(json.dumps(opt.export_json()))
    assert state['set_aside'] == [7]
    assert state['CLASS_VERSION'] == FoodOptimizer.CLASS_VERSION == 17
    assert FoodOptimizer.validate_state(state)['version'] == 17
    other = FoodOptimizer("copy_of_set_aside")
    other.import_json(state)
    assert other.set_aside == [7]
    assert FoodOptimizer(opt.project_name).set_aside == [7]


def test_a_file_without_set_aside_opens_with_none():
    opt = _with_one_waiting()
    state = json.loads(json.dumps(opt.export_json()))
    state.pop('set_aside')
    state['CLASS_VERSION'] = 16
    assert FoodOptimizer.validate_state(state)['version'] == 16
    other = FoodOptimizer("older_copy")
    other.import_json(state)
    assert other.set_aside == []


@pytest.mark.parametrize("bad", ["7", [0], ["seven"], [7.5], None])
def test_the_validator_refuses_a_damaged_set_aside(bad):
    opt = _with_one_waiting()
    state = json.loads(json.dumps(opt.export_json()))
    state['set_aside'] = bad
    with pytest.raises(ValueError):
        FoodOptimizer.validate_state(state)


def test_the_workbook_ticks_it_not_scored():
    import io
    from openpyxl import load_workbook
    opt = _with_one_waiting()
    opt.set_aside_formulation(7)
    book = load_workbook(io.BytesIO(opt.all_formulations_workbook()))
    sheet = book.worksheets[0]
    header = [c.value for c in sheet[2]]         # row 1 is the sheet's title
    rows = [[c.value for c in r] for r in sheet.iter_rows(min_row=3)]
    row = next(r for r in rows if r[header.index("Formulation")] == 7)
    assert row[header.index(wording.NOT_SCORED)] == wording.TICKED_BOX
    assert row[header.index(wording.OVERALL_SCORE_COLUMN)] in ("", None)


def test_insights_counts_nothing_from_it():
    opt = _with_one_waiting()
    before_key, before_rounds = insights.data_key(opt), insights.round_keys(opt)
    assert insights.measured(opt, "Firmness", 6) is not None
    opt.set_aside_formulation(7)
    for name in ("Firmness", "Juiciness", "Cook loss", insights.SCORE):
        assert insights.measured(opt, name, 6) is None
    assert insights.fit(opt, "Firmness", list(range(7))).n_rows == 6
    assert insights.data_key(opt) != before_key
    assert insights.round_keys(opt)[3] != before_rounds[3]


def test_suggestions_are_bit_identical_whether_it_waits_or_is_set_aside():
    """The owner's rule (§9 I5): Move to Not scored changes no suggestion.
    Formulation 7 of the fixture project waits for Cook loss; moved to Not
    scored, ask() hands the model the same rows and suggests the same."""
    opt = exempt_and_waiting()
    opt.set_aside_formulation(7)
    same(run(opt), fixture()["exempt_and_waiting"])


def test_score_without_is_gone():
    assert not hasattr(FoodOptimizer, "score_without")
    assert not hasattr(wording, "score_without_button")
    assert not hasattr(wording, "scored_without")


def test_the_new_strings_pass_the_vocabulary_guard():
    from tests.test_food_bo import _BANNED
    for text in (wording.not_scored_missing("Cook loss"), wording.save_missing("Cook loss"),
                 wording.MOVE_TO_NOT_SCORED, wording.WAIT_AGAIN,
                 wording.never_coming_warning(33, "Cook loss"), wording.moved_to_not_scored(33),
                 wording.waiting_again(33, "Cook loss"),
                 wording.saved_not_scored(33, "Cook loss"),
                 wording.formulation_has_all_results(33)):
        assert not any(p.search(text) for p in _BANNED), text
    assert wording.never_coming_warning(33, "Cook loss") == (
        "Cook loss for Formulation 33 will never come? It stays in your records, unscored, "
        "and stops counting in the review.")  # 469a1d1: Insights is now 4 · Review
