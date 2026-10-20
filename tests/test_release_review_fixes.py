"""Review fixes before the public release (0.8.0): the goal words on
Results and Make a round, the last-opened file on a shared host, the score
axis near 0, the record chips, the workbook's sheet list and the All
formulations column widths."""
import io
import json
import os

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

import ui_batch
import ui_insights
import ui_results
import ui_setup
import wording
from food_bo import FoodOptimizer
from tests.test_app_ui import APP_PATH, _submit_button
from ui_helpers import goal_shown


@pytest.fixture
def goals(tmp_path, monkeypatch):
    """One measurement of each goal: higher, lower, a target."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("goals")
    opt.set_amount_unit("g")
    opt.add_ingredient("Oil", 0, 20)
    opt.add_ingredient("Starch", 0, 5)
    opt.add_objective("Juiciness", 3.0, goal="max", min_val=0, max_val=10)
    opt.add_objective("Cook loss", 2.0, goal="min", min_val=0, max_val=40, unit="%")
    opt.add_objective("Firmness", 1.0, goal="target", target=6, min_val=0,
                      max_val=10, unit="N")
    opt.tell({"Oil": 5.0, "Starch": 1.0},
             {"Juiciness": 6.0, "Cook loss": 12.0, "Firmness": 5.0},
             formulation_no=1, batch_no=1)
    return opt


# ------------------------------------------------------------------ #
#  1 · one wording for which way is better
# ------------------------------------------------------------------ #

def test_goal_shown_says_what_the_headers_say():
    assert goal_shown({"goal": "max"}) == wording.SCORE_GOAL_HIGHER == "Higher is better"
    assert goal_shown({"goal": "min"}) == wording.SCORE_GOAL_LOWER == "Lower is better"
    assert goal_shown({"goal": "target", "target": 6, "unit": "N"}) == "Target 6 N"
    # food_bo's own text, as closeness_details hands it over.
    assert goal_shown(wording.GOAL_LABELS["max"]) == wording.SCORE_GOAL_HIGHER
    assert goal_shown(wording.GOAL_LABELS["min"]) == wording.SCORE_GOAL_LOWER
    assert goal_shown("Target 6 N") == "Target 6 N"


def test_measurements_against_goals_say_higher_or_lower_is_better(goals):
    frame = ui_results._off_by_table(goals, 0).data
    shown = dict(zip(frame[wording.MEASUREMENT_COLUMN], frame[wording.GOAL_LABEL]))
    assert shown["Juiciness"] == "Higher is better", shown
    assert shown["Cook loss"] == "Lower is better", shown
    assert shown["Firmness"] == "Target 6 N", shown
    assert not any(str(v).startswith("Prefer") for v in shown.values()), shown


def test_make_a_round_entry_boxes_say_higher_or_lower_is_better(goals):
    goals.set_pending_batch([{"Oil": 6.0, "Starch": 2.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "goals"
    at.run()
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    assert not at.exception
    labels = [n.label for n in at.number_input if n.key and n.key.startswith("f2_")]
    assert any(label.endswith("· Higher is better") for label in labels), labels
    assert any(label.endswith("· Lower is better") for label in labels), labels
    assert not any("Prefer" in label for label in labels), labels


# ------------------------------------------------------------------ #
#  4 · the edit-in-place status line, ten words at most
# ------------------------------------------------------------------ #

def test_formulations_saved_line_is_ten_words_at_most():
    line = wording.FORMULATIONS_SAVED_NEW_WORKBOOK
    assert line == "Formulations saved. Download a new workbook."
    assert len(line.split()) <= 10


# ------------------------------------------------------------------ #
#  5 · no last-opened file on a shared host
# ------------------------------------------------------------------ #

@pytest.mark.parametrize("shared", [False, True])
def test_opening_a_project_writes_last_opened_only_off_a_shared_host(
        tmp_path, monkeypatch, shared):
    monkeypatch.chdir(tmp_path)
    for name in ("first", "second"):
        FoodOptimizer(name).add_ingredient("Water", 0, 100)
    real_isdir = os.path.isdir
    monkeypatch.setattr(os.path, "isdir",
                        lambda p: (shared if p == "/mount/src" else real_isdir(p)))
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "first"
    at.run()
    at.sidebar.selectbox(key="project_select").select("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert at.session_state["optimizer"].project_name == "second"
    assert (tmp_path / ".last_opened").exists() is not shared


# ------------------------------------------------------------------ #
#  6 · scores near 0: the axis a pad under the lowest dot
# ------------------------------------------------------------------ #

def _y_domain(chart):
    found = []

    def walk(node):
        if isinstance(node, dict):
            y = node.get("y")
            if isinstance(y, dict) and "domain" in y.get("scale", {}):
                found.append(y["scale"]["domain"])
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for value in node:
                walk(value)
    walk(chart.to_dict())
    assert found
    return found[0]


@pytest.mark.parametrize("ratings", [(0.0, 0.3, 1.0, 2.0), (0.05, 0.3, 1.0, 2.0),
                                     (0.2, 0.5, 1.0, 3.0)])
def test_scores_near_zero_sit_clear_of_the_axis_bottom(tmp_path, monkeypatch, ratings):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("near_zero")
    opt.set_amount_unit("g")
    opt.add_ingredient("Oil", 0, 20)
    opt.add_ingredient("Starch", 0, 5)
    opt.add_objective("Taste", 100, goal="max", min_val=0, max_val=10)
    for k, rating in enumerate(ratings):
        opt.tell({"Oil": 2.0 + k, "Starch": 1.0}, {"Taste": rating},
                 formulation_no=k + 1, batch_no=1 + k // 2)
    scores = [float(y) for y in opt.Y_history]
    assert min(scores) < 5 and max(scores) <= 30
    bottom, top = _y_domain(ui_insights.improving_chart(opt))
    assert bottom < min(scores), (bottom, scores)        # no dot on the axis
    # A score is never below 0: at most a small pad under 0.
    pad = 0.04 * max(max(scores) - min(scores), 4.0)
    assert bottom >= -pad - 1e-9, (bottom, pad)
    assert top > max(scores)


# ------------------------------------------------------------------ #
#  7 · a long field name wraps inside its chip
# ------------------------------------------------------------------ #

def test_record_chips_wrap_inside_the_box():
    css = ui_setup._RECORD_CHIP_STYLE
    rules = [r for r in css.replace("<style>", "").replace("</style>", "").split("}") if r]
    assert all(r.lstrip().startswith(".st-key-record_fields") for r in rules), rules
    span = next(r for r in rules if r.split("{")[0].strip().endswith("span"))
    assert "max-width: 100%" in span and "white-space: normal" in span, span
    assert "max-width: none" not in css


# ------------------------------------------------------------------ #
#  8 · the workbook's sheet list, read and closed
# ------------------------------------------------------------------ #

def _book(*names):
    wb = openpyxl.Workbook()
    wb.active.title = names[0]
    for name in names[1:]:
        wb.create_sheet(name)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_preparation_sheet_is_read_from_the_file():
    assert ui_batch._has_preparation_sheet(
        _book(wording.OVERVIEW_SHEET, wording.PREPARATION_SHEET, wording.RESULTS_SHEET))
    assert not ui_batch._has_preparation_sheet(
        _book(wording.OVERVIEW_SHEET, wording.RESULTS_SHEET))
    assert not ui_batch._has_preparation_sheet(b"not a workbook")


def test_preparation_sheet_check_opens_no_openpyxl_workbook(monkeypatch):
    def refuse(*_a, **_k):
        raise AssertionError("a workbook was opened")
    monkeypatch.setattr(openpyxl, "load_workbook", refuse)
    assert ui_batch._has_preparation_sheet(
        _book(wording.OVERVIEW_SHEET, wording.PREPARATION_SHEET))


# ------------------------------------------------------------------ #
#  9 · All formulations: the last measurement on screen at 1440 px
# ------------------------------------------------------------------ #

def test_all_formulations_cap_fits_four_measurements_at_1440():
    cap = ui_results.all_formulations_cap
    assert cap(1) == cap(2) == ui_results.ALL_FORMULATIONS_MAX_WIDTH == 200
    assert cap(4) == 145                  # 400 + 4 × 145 inside 980 px
    assert cap(9) == ui_results.ALL_FORMULATIONS_MIN_WIDTH == 140


def test_all_formulations_columns_are_capped_with_the_header_in_the_tooltip(goals):
    goals.add_objective("Mouth-coating after swallowing", 1.0, goal="target",
                        target=3, min_val=0, max_val=15)
    goals.tell({"Oil": 7.0, "Starch": 2.0},
               {"Juiciness": 7.0, "Cook loss": 10.0, "Firmness": 6.0,
                "Mouth-coating after swallowing": 4.0},
               formulation_no=2, batch_no=2)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "goals"
    at.run()
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception
    config = next(c for c in (json.loads(d.proto.columns or "{}") for d in at.dataframe)
                  if wording.OVERALL_SCORE_COLUMN in c and wording.FORMULATION_CAP in c)
    cap = ui_results.all_formulations_cap(4)
    widths = {name: c.get("width") for name, c in config.items() if not c.get("hidden")}
    assert widths and all(w is None or w <= cap for w in widths.values()), widths
    # A full header that would not fit is never cut mid-word: the column
    # shows its short name and the full header is its tooltip.
    with_tip = [c for c in config.values() if c.get("help")]
    assert len(with_tip) == 4, config
    for column in with_tip:
        assert column["help"].startswith(column["label"].split(" (")[0]), column
        assert column["help"] != column["label"], column
        assert 8 * len(column["label"]) + 24 <= cap or column.get("width") == cap, column
    cook = config["Cook loss (%)"]
    assert cook["label"] == "Cook loss (%)"
    assert cook["help"] == "Cook loss (0–40, %) · Lower is better"
    long = next(c for c in config.values()
                if str(c.get("label", "")).startswith("Mouth-coating after swallowing"))
    assert long["help"] == "Mouth-coating after swallowing (0–15) · Target 3"
    # A header that fits is not given a tooltip that repeats it.
    assert "help" not in config[wording.OVERALL_SCORE_COLUMN], config
