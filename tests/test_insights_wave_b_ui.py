"""Insights wave B on screen: the predictions on Make a round, the answer and
compare chart of question 2, and question 4. AppTest, like
tests/test_insights_ui.py."""
import io

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

import ui_insights
import wording
from food_bo import FoodOptimizer
from tests.test_app_ui import APP_PATH, _submit_button, burger, scored  # noqa: F401
from tests.test_insights_ui import _open_insights, insights_scored  # noqa: F401


def _app():
    return AppTest.from_file(APP_PATH, default_timeout=180)




def _batch_table(at):
    """v3: the Round table (ingredients down, the Round across)."""
    return next(t for t in at.table if "Pea protein" in t.value.index)


# ------------------------------------------------------------------ #
#  Task 2: predictions on Make a round (v3: gone from that screen; what
#  is left checks it fits nothing and leaves the workbook alone)
# ------------------------------------------------------------------ #

@pytest.fixture
def warm_round(insights_scored):
    """Nine formulations with all their results (the spread-out start is
    over) and Round 4 open with three formulations."""
    insights_scored.set_pending_batch([
        {"Pea protein": 12.0, "Methylcellulose": 1.0},
        {"Pea protein": 18.0, "Methylcellulose": 2.0},
        {"Pea protein": 22.0, "Methylcellulose": 2.5}])
    return insights_scored


def test_editing_the_round_hides_the_column(warm_round):
    at = _app()
    at.run()
    at.button(key=f"edit_round_burger_{warm_round.pending_batch_no}_start").click().run()
    assert not at.exception
    assert any(i.value == wording.EDIT_FORMULATIONS_ACTIVE for i in at.info)
    [grid] = [d for d in at.dataframe if "Pea protein (g)" in d.value.columns]   # the editor
    assert 'Predicted Overall score (95\xa0% range)' not in grid.value.columns
    assert 'Exploring suggestions can predict below Best so far.' not in [c.value for c in at.caption]
    assert 'Predicted measurements (95\xa0% range)' not in [e.label for e in at.expander]


def test_the_round_tab_closed_fits_nothing(warm_round):
    at = _app()
    at.session_state["_loaded_project"] = "burger"   # no landing on Make a round
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception
    assert ui_insights.FITTED_KEY not in at.session_state


def _cells(book):
    return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in book.worksheets]


def test_the_workbook_is_unchanged(warm_round):
    opt = FoodOptimizer("burger")
    before = _cells(openpyxl.load_workbook(io.BytesIO(
        opt.workbook_bytes(opt.pending_batch, opt.open_round_size()))))
    at = _app()
    at.run()
    assert _batch_table(at) is not None          # v3: the Round table, no forecast
    opt = at.session_state["optimizer"]
    after = _cells(openpyxl.load_workbook(io.BytesIO(
        opt.workbook_bytes(opt.pending_batch, opt.open_round_size()))))
    assert after == before
    assert 'Predicted Overall score (95\xa0% range)' not in opt.batch_frame(opt.pending_batch).columns


def test_a_project_switch_drops_the_shared_models(warm_round):
    at = _open_insights(_app())
    # v3: the Review fits nothing, so the shared models are put in place
    # here, as Make a round's fit leaves them; the switch must drop them.
    at.session_state[ui_insights.FITTED_KEY] = {"stale": True}
    at.session_state[ui_insights.CACHE_KEY] = {"stale": True}
    FoodOptimizer("second").add_ingredient("Flour", 0, 100)
    at.run()
    at.selectbox(key="project_select").set_value("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert ui_insights.FITTED_KEY not in at.session_state
    assert ui_insights.CACHE_KEY not in at.session_state


# ------------------------------------------------------------------ #
#  Task 3: which ingredients matter most
# ------------------------------------------------------------------ #



# ------------------------------------------------------------------ #
#  Task 5: which formulation is predicted to score highest
# ------------------------------------------------------------------ #





def test_a_project_switch_drops_the_best(warm_round):
    at = _open_insights(_app())
    # Review v3 fits nothing: the key as Make a round's work leaves it.
    at.session_state[ui_insights.BEST_KEY] = {"stale": True}
    FoodOptimizer("second").add_ingredient("Flour", 0, 100)
    at.run()
    at.selectbox(key="project_select").set_value("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert ui_insights.BEST_KEY not in at.session_state






# ------------------------------------------------------------------ #
#  One vocabulary across Insights: trials, process settings, values
# ------------------------------------------------------------------ #





# ------------------------------------------------------------------ #
#  Final fix wave (waves B + C): carried from wave B's final review
# ------------------------------------------------------------------ #





