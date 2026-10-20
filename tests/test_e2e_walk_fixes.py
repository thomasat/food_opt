"""The end-to-end walk of 0.8.0 before its public release: each finding,
triggered where the reader meets it (AppTest, or the method the screen
calls) and pinned in its new words."""
import io
import os

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

import sample_projects as sp
import storage as storage_backend
import wording
from food_bo import FoodOptimizer
from tests.test_food_bo import _add, _ing_row, _meas_row

APP_PATH = os.path.join(os.path.dirname(__file__), "..", "app.py")


def _errors(at):
    return [e.value for e in at.error]


def _app(name, tab=None):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = name
    if tab:
        at.session_state["main_tab"] = tab
    return at


@pytest.fixture
def oat(tmp_path, monkeypatch):
    """The walk's Oat cookie: 30-50 g oat flour, 10-30 g butter, 5-25 g
    sugar, scored on Crunch and Sweetness out of 10."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("Oat cookie")
    opt.set_amount_unit("g")
    opt.add_ingredient("Oat flour", 30, 50)
    opt.add_ingredient("Butter", 10, 30)
    opt.add_ingredient("Sugar", 5, 25)
    opt.add_objective("Crunch", 1.0, goal="target", target=7,
                      min_val=0, max_val=10, unit="score")
    opt.add_objective("Sweetness", 1.0, goal="max", min_val=0, max_val=10,
                      unit="score")
    return opt


# ---- 1 · two new rows of one name ---------------------------------- #

def test_two_unsaved_rows_of_one_name_say_so(oat):
    frame = _add(oat.ingredient_grid_frame(), **_ing_row("Cocoa", high=5.0))
    frame = _add(frame, **_ing_row("Cocoa", high=6.0))
    errors, _ = oat.apply_ingredient_grid(frame)
    assert [text for _, text in errors] == [
        "Cocoa is on two rows; rename or delete one."]


def test_a_new_row_wearing_a_saved_rows_name_keeps_its_words(oat):
    errors, _ = oat.apply_ingredient_grid(
        _add(oat.ingredient_grid_frame(), **_ing_row("Butter", high=5.0)))
    assert [text for _, text in errors] == [
        "Butter is already the name of an ingredient. Choose another name."]


# ---- 2 · shares that do not add up to 100 -------------------------- #

def test_shares_scaled_to_100_are_named(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("shares")
    opt.add_ingredient("Oat flour", 30, 50)
    frame = _add(opt.measurement_grid_frame(),
                 **_meas_row("Crunch", share=60.0))
    frame = _add(frame, **_meas_row("Sweetness", share=30.0))
    errors, messages = opt.apply_measurement_grid(frame)
    assert errors == []
    assert ("info", "Shares scaled to add up to 100 %: Crunch 66.7, "
                    "Sweetness 33.3.") in messages, messages


# ---- 3 · a result outside its scale -------------------------------- #

def test_a_typed_result_outside_its_scale_reads_naturally(oat):
    oat.set_pending_batch([{"Oat flour": 40.0, "Butter": 20.0, "Sugar": 10.0}])
    df = __import__("pandas").DataFrame(
        {"Formulation": [oat.pending_batch[0]["formulation"]],
         "Crunch": [12.0], "Sweetness": [5.0]})
    with pytest.raises(ValueError) as refused:
        oat.parse_batch_results(df, oat.pending_batch)
    number = oat.pending_batch[0]["formulation"]
    assert str(refused.value) == (
        f"Formulation {number}: Crunch 12 is outside 0–10. Fix it, or raise "
        "Scale maximum in Set up.")
    assert "12 score" not in str(refused.value)


# ---- 4 · a workbook with nothing filled in ------------------------- #

def test_an_empty_workbook_names_the_results_sheet(oat):
    oat.set_pending_batch([{"Oat flour": 40.0, "Butter": 20.0, "Sugar": 10.0}])
    book = oat.workbook_bytes(oat.pending_batch, None, print_pack=False)
    assert load_workbook(io.BytesIO(book)).sheetnames[:2] == [
        wording.OVERVIEW_SHEET, wording.RESULTS_SHEET]
    with pytest.raises(ValueError) as refused:
        oat.results_from_workbook(io.BytesIO(book))
    assert str(refused.value) == (
        "Nothing is filled in on the Results sheet. Enter a result, or tick "
        "Not scored.")


def test_an_empty_printed_pack_names_its_round_sheet(oat):
    oat.set_pending_batch([{"Oat flour": 40.0, "Butter": 20.0, "Sugar": 10.0}])
    book = oat.workbook_bytes(oat.pending_batch, None, print_pack=True)
    sheet = wording.batch_sheet_name(oat.pending_batch_no)
    assert sheet in load_workbook(io.BytesIO(book)).sheetnames
    with pytest.raises(ValueError) as refused:
        oat.results_from_workbook(io.BytesIO(book))
    assert str(refused.value) == (
        f"Nothing is filled in on the {sheet} sheet. Enter a result, or tick "
        "Not scored.")


def test_the_upload_caption_names_the_results_sheet():
    assert wording.UPLOAD_HELP_CAPTION == (
        "Upload the workbook with Results filled in.")
    assert "Measured cell" not in wording.UPLOAD_HELP_DETAIL
    assert wording.RESULTS_SHEET in wording.UPLOAD_HELP_DETAIL


# ---- 5 · Add a formulation of your own with a blank amount --------- #

def test_one_blank_amount_says_leave_it_out(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("own")
    opt.set_amount_unit("g")
    opt.add_ingredient("Pea protein", 0, 25)
    opt.add_ingredient("Salt", 0, 3)
    opt.add_objective("Taste", 1.0, goal="max", min_val=0, max_val=10)
    at = _app("own", wording.TAB_BATCH)
    at.run()
    at.number_input(key="own_Pea protein").set_value(10.0)
    next(b for b in at.button if b.label == wording.ADD_TO_THIS_BATCH).click()
    at.run()
    assert not at.exception
    assert "Enter an amount for Salt; type 0 to leave it out." in \
        _errors(at), _errors(at)


def test_a_blank_whose_lowest_is_above_0_is_not_offered_0(oat):
    at = _app("Oat cookie", wording.TAB_BATCH)
    at.run()
    at.number_input(key="own_Oat flour").set_value(40.0)
    at.number_input(key="own_Sugar").set_value(10.0)
    next(b for b in at.button if b.label == wording.ADD_TO_THIS_BATCH).click()
    at.run()
    assert not at.exception
    assert "Enter an amount for Butter." in _errors(at), _errors(at)


def test_several_blanks_keep_leave_one_out():
    assert wording.enter_amounts_for("Salt and Water", n=2) == (
        "Enter an amount for Salt and Water; type 0 to leave one out.")


# ---- 7 · Open a saved copy: the answer under the control ----------- #

def _flatten(node):
    out = [node]
    for child in getattr(node, "children", {}).values():
        out += _flatten(child)
    return out


def test_the_copy_question_is_drawn_under_open_a_saved_copy(oat):
    at = _app("Oat cookie", wording.TAB_SETUP)
    at.session_state["_restore_candidate"] = FoodOptimizer(
        "Oat cookie").export_json()
    at.run()
    assert not at.exception
    fold = next(e for e in at.sidebar.expander
                if e.label == wording.MANAGE_PROJECT)
    inside = _flatten(fold)
    kinds = [(getattr(n, "type", ""), getattr(n, "label", None) or
              getattr(n, "value", None)) for n in inside]
    warning = next(i for i, (t, v) in enumerate(kinds)
                   if t == "alert" or (isinstance(v, str)
                                       and v.startswith("This copy holds")))
    replace = next(i for i, (_, v) in enumerate(kinds)
                   if v == wording.YES_REPLACE)
    delete = next(i for i, (_, v) in enumerate(kinds)
                  if v == wording.DELETE_PROJECT_LABEL)
    uploader = next(i for i, (_, v) in enumerate(kinds)
                    if v == wording.OPEN_SAVED_COPY_CAPTION)
    assert uploader < warning < replace < delete, kinds


def test_a_file_that_is_not_a_copy_is_refused_under_the_control(oat):
    at = _app("Oat cookie", wording.TAB_SETUP)
    at.session_state["_restore_candidate"] = {"hello": "world"}
    at.run()
    assert not at.exception
    fold = next(e for e in at.sidebar.expander
                if e.label == wording.MANAGE_PROJECT)
    assert any(e.value.startswith("This file") or "copy" in e.value
               for e in fold.error), [e.value for e in fold.error]


# ---- 8 · clearing the Batch size box ------------------------------- #

def test_unscale_round_goes_back_to_the_generated_amounts(oat):
    made = [{"Oat flour": 40.0, "Butter": 20.0, "Sugar": 10.0},
            {"Oat flour": 35.0, "Butter": 15.0, "Sugar": 12.0}]
    oat.set_pending_batch(made)
    oat.scale_round(90.0)
    assert oat.pending_batch_total == 90.0
    assert oat.pending_batch[0]["recipe"]["Oat flour"] == pytest.approx(51.43,
                                                                        abs=0.01)
    oat.unscale_round()
    assert oat.pending_batch_total is None
    for row, recipe in zip(oat.pending_batch, made):
        assert row["recipe"] == pytest.approx(recipe)


def test_clearing_the_box_returns_the_round_to_its_generated_amounts(oat):
    made = [{"Oat flour": 40.0, "Butter": 20.0, "Sugar": 10.0}]
    oat.set_pending_batch(made)
    at = _app("Oat cookie", wording.TAB_BATCH)
    at.run()
    at.number_input(key="scale_total").set_value(90.0)
    at.run()
    assert not at.exception
    assert FoodOptimizer("Oat cookie").pending_batch_total == 90.0
    at.number_input(key="scale_total").set_value(None)
    at.run()
    assert not at.exception
    back = FoodOptimizer("Oat cookie")
    assert back.pending_batch_total is None
    assert back.pending_batch[0]["recipe"] == pytest.approx(made[0])
    # Empty, not the 0.00 a cleared box posts back.
    assert at.session_state["scale_total"] is None
    at.run()                    # and it stays empty
    assert at.session_state["scale_total"] is None
    assert FoodOptimizer("Oat cookie").pending_batch_total is None


# ---- 9 · the most a Round can be made to, with limits -------------- #

def test_the_reach_respects_a_group_limit(oat):
    assert oat.total_reach() == (45.0, 105.0)
    assert oat.limited_total_reach() is None
    oat.add_quantity_constraint(["Butter", "Sugar"], max_val=45)
    assert oat.limited_total_reach() == (45.0, 95.0)


def test_a_batch_size_past_the_limits_is_refused(oat):
    oat.add_quantity_constraint(["Butter", "Sugar"], max_val=45)
    oat.set_pending_batch([{"Oat flour": 40.0, "Butter": 20.0, "Sugar": 10.0}])
    at = _app("Oat cookie", wording.TAB_BATCH)
    at.run()
    at.number_input(key="scale_total").set_value(100.0)
    at.run()
    assert not at.exception
    said = [c.value for c in at.caption]
    assert any(c.startswith("Batch size: enter 95 g or less, the most your "
                            "limits allow.") for c in said), said
    assert FoodOptimizer("Oat cookie").pending_batch_total is None
    at.number_input(key="scale_total").set_value(95.0)
    at.run()
    assert FoodOptimizer("Oat cookie").pending_batch_total == 95.0


def test_a_default_batch_size_past_the_limits_is_refused(oat):
    oat.add_quantity_constraint(["Butter", "Sugar"], max_val=45)
    with pytest.raises(ValueError) as refused:
        oat.set_formulation_total(100)
    assert str(refused.value) == (
        "Your limits allow at most 95 g; enter 95 g or less, or change a "
        "limit.")
    oat.set_formulation_total(95)
    assert oat.formulation_total == 95.0


# ---- 11 · the yogurt pilot's count --------------------------------- #

def test_the_yogurt_caption_counts_what_the_project_holds(tmp_path,
                                                          monkeypatch):
    monkeypatch.chdir(tmp_path)
    store = storage_backend.LocalStorage()
    sp.ensure(store)
    opt = FoodOptimizer(sp.ROUND6, storage=store)
    assert len(opt.Y_history) == 15
    from tests.test_v3_results import upload_round6
    upload_round6(opt)
    at = _app(sp.ROUND6)
    at.run()
    assert not at.exception
    captions = [c.value for c in at.sidebar.caption]
    assert "Real tasting data; 18 formulations scored." in captions, captions


# ---- 12 · a limit on a project with nothing made ------------------- #

def test_a_limit_on_a_project_with_no_formulations_says_nothing_of_them(
        monkeypatch):
    import streamlit as st_module
    import ui_setup
    calls = []

    class Opt:
        recipe_history, skipped, pending_batch = [], [], None

    monkeypatch.setattr(ui_setup, "flash",
                        lambda kind, text, **k: calls.append(text))
    monkeypatch.setattr(ui_setup, "saved_ok", lambda opt: True)
    monkeypatch.setattr(ui_setup, "_nothing_made_fits", lambda opt: False)
    monkeypatch.setattr(st_module, "rerun", lambda: None)
    ui_setup._report_limit(Opt(), "Limit added.")
    made = Opt()
    made.recipe_history = [{"Oat flour": 40.0}]
    ui_setup._report_limit(made, "Limit added.")
    assert calls == ["Limit added.", f"Limit added. {wording.LIMIT_KEPT}"]


# ---- 13 · a copy's file name is never said ------------------------- #

def test_the_restored_line_names_the_copy_by_its_reason():
    line = wording.restored_flash("18 formulations", "Yogurt pilot",
                                  "Yogurt pilot_pre_restore")
    assert line == ("Restored 18 formulations into Yogurt pilot. A copy is in "
                    "Saved copies: Before opening a saved copy.")
    assert "_pre_restore" not in wording.restored_flash(
        "1 formulation", "X", "X_pre_restore_3")
    assert wording.copy_saved_as("X_archived") == (
        "A copy is in Saved copies: Before clearing project contents.")
