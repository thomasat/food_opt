"""v3 lane C (final-v3, 3 · Results): "Round N results" — ratings against
their targets, no forecasts — and "How the score is worked out", whose
points are food_bo's own and add up to the stored Overall score; the
workbook's Preparation sheet only when a Round has pre-mixes; the shipped
Round 6 results file rebuilt from the current workbook."""
import importlib.util
import io

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

import sample_projects as sp
import storage as storage_backend
import ui_results
import wording
from food_bo import FoodOptimizer, uploaded_parts
from tests.test_app_ui import APP_PATH

NAMES = ("Consistency", "Creaminess", "Tanginess", "Uniformity")


def _builder():
    spec = importlib.util.spec_from_file_location("yogurt_build", sp.SAMPLE_DIR / "build.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return storage_backend.LocalStorage()


def upload_round6(opt, source=None):
    """2 · Make a round > Or upload results from a file > Save uploaded
    results, with the shipped file (or `source`)."""
    with open(sp.ROUND6_RESULTS, "rb") if source is None else source as fh:
        upload = opt.results_from_workbook(fh)
    sheet, weighed, lots = uploaded_parts(upload)
    parsed, left_out = opt.parse_batch_results(
        sheet, opt.pending_batch, with_skipped=True, weighed=weighed)
    by_number = {r["formulation"]: r["recipe"] for r in opt.pending_batch}
    for number, results, note in parsed:
        opt.tell(opt.amounts_as_weighed(by_number[number], weighed.get(number)),
                 results, formulation_no=number, batch_no=opt.pending_batch_no, note=note)
    opt.store_lots(opt.pending_batch_no, lots)
    opt.set_pending_batch(None)
    return parsed, left_out


@pytest.fixture
def round6(store):
    sp.ensure(store)
    opt = FoodOptimizer(sp.ROUND6, storage=store)
    upload_round6(opt)
    return FoodOptimizer(sp.ROUND6, storage=store)


# ------------------------------------------------------------------ #
#  Round N results
# ------------------------------------------------------------------ #

def test_round6_results_table(round6):
    table = ui_results.round_results(round6)
    assert table["round"] == 6
    assert table["columns"] == [
        "Formulation", "Overall score", "Consistency (target 12)",
        "Creaminess (target 5)", "Tanginess (target 5)", "Uniformity (target 13)"]
    rows = table["frame"].values.tolist()
    assert rows == [
        ["F16 · new best (tied)", "98.3", "12", "5", "4", "13"],
        ["F17 · new best (tied)", "98.3", "12", "5", "4", "13"],
        ["F18", "91.7", "13", "7", "4", "12"]]
    assert table["best"] == [0, 1]
    tangy, cons, cream, thick = (table["columns"][k] for k in (4, 2, 3, 5))
    # Off target shows its direction (owner, v3 final cleanup): above
    # orange, below purple; one number a cell, no gap text.
    assert table["off"] == {(0, tangy): "below", (1, tangy): "below", (2, tangy): "below",
                            (2, cons): "above", (2, cream): "above", (2, thick): "below"}
    assert table["caption"] == "Panel median (n = 5), 0–15."
    assert table["key"] is True


def test_the_new_best_is_measured_against_every_earlier_formulation(round6):
    """F16 and F17 beat F10's 95.0 (Round 4) by more than the tolerance;
    Round 5 raised nothing; Round 1's top is the first best."""
    assert [int(round6.formulation_ids[i]) for i in ui_results.new_best_rows(round6, 6)] \
        == [16, 17]
    assert ui_results.new_best_rows(round6, 5) == []
    assert len(ui_results.new_best_rows(round6, 1)) == 1


def test_a_single_new_best_is_not_called_tied(round6):
    i = round6.index_of_formulation(17)
    round6.edit_result(i, {**round6.results_history[i], "Consistency": 11.0})
    table = ui_results.round_results(round6)
    assert table["frame"].iloc[0, 0] == "F16 · new best"
    assert table["frame"].iloc[1, 0] == "F17"
    assert table["best"] == [0]


def test_a_tie_with_an_earlier_best_is_no_new_best(store):
    """Equal within 1e-6 to the best before the Round: not new."""
    sp.ensure(store)
    opt = FoodOptimizer(sp.ROUND6, storage=store)
    f10 = opt.results_history[opt.index_of_formulation(10)]
    for row in list(opt.pending_batch):
        opt.tell(row["recipe"], dict(f10), formulation_no=row["formulation"], batch_no=6)
    opt.set_pending_batch(None)
    table = ui_results.round_results(opt)
    assert table["best"] == []
    assert not any("new best" in v for v in table["frame"].iloc[:, 0])


def test_a_project_without_a_recorded_panel_claims_none(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("Mine", robust=False)
    opt.add_ingredient("Oil", 0, 10)
    opt.add_objective("Taste", 60, goal="max", min_val=0, max_val=10)
    opt.add_objective("Cook loss", 40, goal="target", target=20, min_val=0, max_val=100, unit="%")
    opt.tell({"Oil": 5}, {"Taste": 7, "Cook loss": 25}, batch_no=1)
    table = ui_results.round_results(opt)
    # Taste has no target: plain. Cook loss 25 against 20 is above.
    assert table["caption"] == "" and table["key"] is True
    assert table["off"] == {(0, table["columns"][3]): "above"}
    assert table["columns"][2:] == ["Taste (0–10, higher is better)",
                                    "Cook loss (0–100, %, target 20)"]
    assert sp.panel_size(opt) is None


def test_waiting_and_not_scored_rows_say_so(store):
    sp.ensure(store)
    opt = FoodOptimizer(sp.ROUND6, storage=store)
    a, b, c = list(opt.pending_batch)
    opt.tell(a["recipe"], {"Consistency": 12, "Creaminess": 5, "Tanginess": 5,
                           "Uniformity": 13}, formulation_no=a["formulation"], batch_no=6)
    opt.tell(b["recipe"], {"Consistency": 12}, formulation_no=b["formulation"], batch_no=6)
    opt.record_skipped(c["formulation"], 6, c["recipe"])
    opt.set_pending_batch(None)
    frame = ui_results.round_results(opt)["frame"]
    assert frame["Overall score"].tolist() == ["100.0", wording.MEASURED_WAITING,
                                               wording.NOT_SCORED]
    assert frame.iloc[0, 0] == "F16 · new best"
    assert frame.iloc[1, 2:].tolist() == ["12", "", "", ""]


def test_results_says_no_forecast(round6):
    """The Results tab after Round 6: the table and the fold, and nothing
    predicted, no range, no In/Above/Below, no stop."""
    at = AppTest.from_file(APP_PATH, default_timeout=300)
    at.session_state["_loaded_project"] = sp.ROUND6
    at.run()
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception
    tab = at.tabs[2]
    said = " ".join([m.value for m in tab.markdown] + [c.value for c in tab.caption]
                    + [e.label for e in tab.expander])
    assert "Round 6 results" in said
    assert wording.SCORE_WORKED_OUT_FOLD in said
    # "Above target" is the colour key's (v3 final cleanup), not a forecast's
    # Above / In / Below.
    said = said.replace(wording.ABOVE_TARGET, "")
    for word in ("redict", "orecast", "95 %", "Above", " range", "measured against",
                 'Review says stop.', "noise", "Stop"):
        assert word not in said, word
    tables = [t.value for t in tab.table]
    assert any("F16 · new best (tied)" in str(t.index.tolist()) for t in tables)
    assert not [k for k in (b.key for b in tab.button) if k == "link_round_checked"]


# ------------------------------------------------------------------ #
#  How the score is worked out
# ------------------------------------------------------------------ #

def test_breakdown_equals_the_stored_score_for_every_pilot_formulation(store):
    sp.ensure(store)
    opt = FoodOptimizer(sp.ALL10, storage=store)
    assert len(opt.Y_history) == 30
    for i in range(30):
        parts = opt.score_breakdown(i)
        assert abs(sum(p["points"] for p in parts) - float(opt.Y_history[i])) < 1e-9
        shown = ui_results.score_breakdown_frame(opt, i)
        assert abs(shown["total"] - float(opt.Y_history[i])) < 1e-9
        assert shown["frame"].iloc[-1].tolist() == [
            "Overall score", "", "", "100", f"{float(opt.Y_history[i]):.1f}"]


def test_breakdown_of_f16(round6):
    shown = ui_results.score_breakdown_frame(round6, round6.index_of_formulation(16))
    assert shown["line"] == "Each point off costs 1/15 of its share."
    assert shown["frame"].values.tolist() == [
        ["Consistency", "12", "12", "25", "25.0"],
        ["Creaminess", "5", "5", "25", "25.0"],
        ["Tanginess", "4", "5", "25", "23.3"],
        ["Uniformity", "13", "13", "25", "25.0"],
        ["Overall score", "", "", "100", "98.3"]]
    # Only the rating carries the direction; Points earned stays neutral.
    assert shown["off"] == {(2, "Rating"): "below"}


@pytest.mark.parametrize("goal", ["max", "min", "target"])
def test_breakdown_equals_the_score_for_every_goal_and_a_partial_row(tmp_path, monkeypatch, goal):
    """Higher, lower and target goals, clipped ratings, and a formulation
    recorded before a measurement was added (its shares rescaled)."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("Goals", robust=False)
    opt.add_ingredient("Oil", 0, 10)
    opt.add_objective("Taste", 50, goal=goal, target=6, min_val=0, max_val=10)
    opt.add_objective("Bite", 30, goal="target", target=40, min_val=0, max_val=100)
    opt.tell({"Oil": 1}, {"Taste": 7.5, "Bite": 55}, batch_no=1)
    opt.tell({"Oil": 2}, {"Taste": 12, "Bite": -3}, batch_no=1)
    opt.add_objective("Crunch", 20, goal="min", min_val=0, max_val=5)
    opt.tell({"Oil": 3}, {"Taste": 2, "Bite": 40, "Crunch": 1.5}, batch_no=2)
    for i in range(len(opt.Y_history)):
        parts = opt.score_breakdown(i)
        assert abs(sum(p["points"] or 0.0 for p in parts) - float(opt.Y_history[i])) < 1e-9
        assert opt._compute_utility(opt.results_history[i], opt._number_at(i)) \
            == opt.Y_history[i]
    first = {p["name"]: p for p in opt.score_breakdown(0)}
    assert first["Crunch"]["points"] is None and not first["Crunch"]["applies"]
    assert abs(first["Taste"]["share_of_score"] + first["Bite"]["share_of_score"] - 100.0) \
        < 1e-9
    frame = ui_results.score_breakdown_frame(opt, 0)["frame"]
    assert frame.set_index("Attribute").loc["Crunch", "Rating"] == wording.DOES_NOT_APPLY


# ------------------------------------------------------------------ #
#  The workbook: Preparation only with pre-mixes; uploads either way
# ------------------------------------------------------------------ #

@pytest.fixture
def plain(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("plain", robust=False)
    opt.add_ingredient("Protein", 5, 20)
    opt.add_ingredient("Water", 0, 100)
    opt.add_objective("Taste", 1, goal="max", min_val=0, max_val=10)
    opt.set_pending_batch([{"Protein": 10, "Water": 90}], batch_no=1)
    return opt


def _with_premix(opt):
    opt.add_premix("Seasoning", "portioned")
    opt.set_premix_parts("Seasoning", [{"name": "Salt", "share": 100, "unit": "g"}])
    opt.add_ingredient("Seasoning", 1, 3)
    opt.set_pending_batch([{"Protein": 10, "Seasoning": 2, "Water": 88}], batch_no=1)
    return opt


def _book(opt):
    return load_workbook(io.BytesIO(opt.workbook_bytes(opt.pending_batch, 100, print_pack=False)))


def _save(book):
    out = io.BytesIO()
    book.save(out)
    out.seek(0)
    return out


def _texts(book):
    return [str(c.value) for s in book if s.sheet_state == "visible"
            for row in s for c in row if c.value is not None]


def _enter_taste(book, value=8):
    sheet = book["Results"]
    row = next(c.row for r in sheet for c in r
               if c.column == 1 and str(c.value or "").startswith("Taste"))
    sheet.cell(row, 2, value)


def test_without_premixes_no_preparation_sheet_or_mention(plain):
    book = _book(plain)
    assert [s.title for s in book if s.sheet_state == "visible"] == ["Round overview", "Results"]
    assert not [t for t in _texts(book) if "Preparation" in t]
    links = {c.hyperlink.location for s in book for r in s for c in r if c.hyperlink}
    assert links == {"'Round overview'!A1", "'Results'!A1"}
    _enter_taste(book)
    assert plain.results_from_workbook(_save(book)).frame["Taste"].tolist() == [8]


def test_with_premixes_the_preparation_sheet_is_there_and_named(plain):
    book = _book(_with_premix(plain))
    assert [s.title for s in book if s.sheet_state == "visible"] == [
        "Round overview", "Preparation", "Results"]
    assert book["Results"]["A2"].value.startswith(wording.sheet_navigation(True))
    assert "Round overview → Preparation → Results" in wording.sheet_navigation(True)
    _enter_taste(book)
    assert plain.results_from_workbook(_save(book)).frame["Taste"].tolist() == [8]


def test_a_premix_workbook_missing_its_preparation_sheet_is_refused_plainly(plain):
    """A pre-mix Round's workbook carries its Preparation sheet; one deleted
    by hand is refused with the Round-information message, never read
    without the pre-mix pages."""
    book = _book(_with_premix(plain))
    book.remove(book["Preparation"])
    _enter_taste(book, 6)
    with pytest.raises(ValueError, match="Round information"):
        plain.results_from_workbook(_save(book))


def test_an_older_workbook_with_an_empty_preparation_sheet_still_uploads(plain):
    """Downloaded before v3: an empty Preparation sheet beside the two."""
    book = _book(plain)
    sheet = book.create_sheet("Preparation", 1)
    sheet["A1"] = "Preparation — Pre-mixes for this round"
    sheet["A4"] = "No pre-mixes are required. Follow the method on Round overview."
    _enter_taste(book, 7)
    assert plain.results_from_workbook(_save(book)).frame["Taste"].tolist() == [7]


# ------------------------------------------------------------------ #
#  The shipped Round 6 results file
# ------------------------------------------------------------------ #

def _cells(book):
    return {s.title: [[c.value for c in row] for row in s.iter_rows()]
            for s in book if s.sheet_state == "visible"}


def test_the_shipped_round6_file_equals_a_rebuild(tmp_path):
    """build.py's round6_results() writes it from the current workbook:
    same sheets and cells, the pilot's values and notes."""
    out = tmp_path / "rebuilt.xlsx"
    _builder().write_round6_results(str(out), work_dir=str(tmp_path / "work"))
    shipped, rebuilt = load_workbook(sp.ROUND6_RESULTS), load_workbook(out)
    assert _cells(shipped) == _cells(rebuilt)
    assert [s.title for s in shipped if s.sheet_state == "visible"] == [
        "Round overview", "Results"]


def test_the_shipped_round6_file_records_the_pilots_values(store):
    sp.ensure(store)
    opt = FoodOptimizer(sp.ROUND6, storage=store)
    parsed, left_out = upload_round6(opt)
    assert left_out == []
    by_no = {no: (tuple(r[n] for n in NAMES), note) for no, r, note in parsed}
    assert by_no == {16: ((12, 5, 4, 13), "Engine's pick 6.1"),
                     17: ((12, 5, 4, 13), "Engine's pick 6.2"),
                     18: ((13, 7, 4, 12), "Engine's pick 6.3")}
    scores = {int(n): round(float(y), 1) for n, y in zip(opt.formulation_ids, opt.Y_history)}
    assert (scores[16], scores[17], scores[18], scores[10]) == (98.3, 98.3, 91.7, 95.0)


# ------------------------------------------------------------------ #
#  The Overall score on screen: one decimal everywhere (owner, v3 final
#  fixes); full precision kept in the data and the exports
# ------------------------------------------------------------------ #

def test_the_score_lines_show_one_decimal():
    assert wording.score_shown(98.33333333) == "98.3"
    assert wording.score_shown(95.0) == "95.0"
    assert wording.overall_score_line(98.33333333, 100.0) == "Overall score 98.3 of 100."
    assert wording.batch_recorded_progress(2, 86.6049, 88.1) == (
        "Round 2 recorded · Best overall score 86.6 → 88.1")


def test_view_formulation_names_the_tie_at_one_decimal(round6):
    assert ui_results._progress_line(round6) == "Formulations 16 and 17 are tied at 98.3."


def test_a_tie_is_found_at_full_precision_with_a_tolerance(round6):
    """98.33 and 98.33 + 1e-9 tie; 98.33 and 98.34 (both 98.3 on screen)
    do not read as tied at one value."""
    f16, f17 = round6.index_of_formulation(16), round6.index_of_formulation(17)
    round6.Y_history[f17] = float(round6.Y_history[f16]) + 1e-9
    assert "tied at 98.3" in ui_results._tie(round6, f16, f17)
    round6.Y_history[f17] = float(round6.Y_history[f16]) + 0.01
    assert "tied at" not in ui_results._tie(round6, f16, f17)


# ------------------------------------------------------------------ #
#  No tasting noise on Results (owner, v3 final cleanup): the line under
#  Best so far says the next best as a measurement, never "within noise"
#  or what a rise "needs"
# ------------------------------------------------------------------ #

_NOISE_WORDS = ("noise", "needed", "need ", "within", "apart", "tell", "stop", "Stop")


@pytest.fixture
def round5(store):
    """The Round 6 sample as shipped: Rounds 1–5 scored, Best so far F10
    (95.0), the next best different formulation F13 (93.33)."""
    sp.ensure(store)
    return FoodOptimizer(sp.ROUND6, storage=store)


def test_the_runner_up_line_is_a_plain_measurement():
    assert wording.runner_up_line(13, "93.3", "1.7") == (
        "Formulation 13: 93.3 (1.7 lower than the best so far).")


def test_view_formulation_on_round5_names_the_next_best_without_noise(round5):
    line = ui_results._runner_up_line(round5)
    assert line == "Formulation 13: 93.3 (1.7 lower than the best so far)."
    said = " ".join([ui_results._progress_line(round5), line])
    for word in _NOISE_WORDS:
        assert word not in said, word


def test_a_close_lead_is_no_tie_and_no_noise(round6):
    """F17 0.5 above F16: not tied (that is equal scores only); the flash
    says Best so far moved and the line names F16 0.5 lower."""
    f16, f17 = round6.index_of_formulation(16), round6.index_of_formulation(17)
    round6.Y_history[f17] = float(round6.Y_history[f16]) + 0.5
    assert ui_results._tie(round6, f16, f17) == ""
    assert ui_results._move_sentence(round6, 16, 17) == wording.best_moved(16, 17)
    assert ui_results._runner_up_line(round6) == (
        "Formulation 16: 98.3 (0.5 lower than the best so far).")


def test_a_lead_under_a_tenth_reads_at_two_decimals(round6):
    """0.03 lower never reads '0.0 lower'."""
    f16, f17 = round6.index_of_formulation(16), round6.index_of_formulation(17)
    round6.Y_history[f17] = float(round6.Y_history[f16]) + 0.03
    assert ui_results._runner_up_line(round6) == (
        "Formulation 16: 98.3 (0.03 lower than the best so far).")


def test_no_runner_up_line_on_an_exact_tie(round6):
    assert ui_results._runner_up_line(round6) == ""


def test_results_on_round5_says_no_noise(round5):
    at = AppTest.from_file(APP_PATH, default_timeout=300)
    at.session_state["_loaded_project"] = sp.ROUND6
    at.run()
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception
    tab = at.tabs[2]
    said = " ".join([m.value for m in tab.markdown] + [c.value for c in tab.caption]
                    + [str(c.help or "") for c in tab.caption])
    assert "Formulation 13: 93.3 (1.7 lower than the best so far)." in said
    for word in ("noise", "needed", "within noise", "tell apart"):
        assert word not in said, word


def test_all_formulations_draws_the_score_at_one_decimal(round6):
    """The grid's own formatter left the score's text alone: 98.3, never
    88.30 (it read every column as a two-decimal number)."""
    frame = round6.history_frame()
    cells = ui_results._table_style(round6, frame)._translate(False, False)["body"]
    col = list(frame.columns).index(wording.OVERALL_SCORE_COLUMN) + 1
    shown = [row[col]["display_value"] for row in cells]
    assert "98.3" in shown and "91.7" in shown and "95.0" in shown
    assert not any(v.count(".") == 1 and len(v.split(".")[1]) == 2 for v in shown), shown


def test_all_formulations_ratings_read_as_the_round_table_does(round6):
    """A rating reads 12 or 4.5, as in Round results — never 12.00; the
    score keeps its one decimal and an amount its two."""
    frame = round6.history_frame(include_amounts=True)
    cells = ui_results._table_style(round6, frame)._translate(False, False)["body"]
    headers = list(frame.columns)
    table = ui_results.round_results(round6)
    for o in round6.objectives:
        col = headers.index(round6._measurement_column(o)) + 1
        shown = [row[col]["display_value"] for row in cells]
        assert not any(v.endswith(".00") or (v.count(".") == 1 and v.endswith("0"))
                       for v in shown), (o["name"], shown)
        nums = [v for v in shown if v]
        assert nums and all(v == ui_results._rating_text(float(v)) for v in nums
                            if v != wording.DOES_NOT_APPLY), shown
    score = [row[headers.index(wording.OVERALL_SCORE_COLUMN) + 1]["display_value"]
             for row in cells]
    assert "95.0" in score
    amount = next(c for c in headers if c.startswith("Coconut cream"))
    assert all(len(row[headers.index(amount) + 1]["display_value"].split(".")[1]) == 2
               for row in cells if row[headers.index(amount) + 1]["display_value"])
    # The same text as the Round table, row for row (Round 6's ratings).
    f16 = next(i for i, row in enumerate(cells)
               if row[headers.index(wording.FORMULATION_CAP) + 1]["display_value"] == "16")
    first = round6.measurements_by_importance()[0]
    head = next(c for c in table["columns"] if c.startswith(first["name"]))
    r16 = next(r for _i, r in table["frame"].iterrows()
               if str(r[table["columns"][0]]).startswith(("F16", "Formulation 16", "16")))
    assert cells[f16][headers.index(round6._measurement_column(first)) + 1][
        "display_value"] == r16[head]


@pytest.mark.parametrize("which", ["light", "dark"])
def test_off_by_text_takes_the_direction_colour(round6, monkeypatch, which):
    """View formulation's Off by: purple below target, orange above, as in
    Round results and the Review strip; plain on target."""
    import ui_insights
    monkeypatch.setattr(ui_insights, "theme", lambda: which)
    seen = set()
    for index in range(len(round6.Y_history)):
        styler = ui_results._off_by_table(round6, index)
        styler._compute()
        cells = styler._translate(False, False)["body"]
        col = list(styler.data.columns).index(wording.OFF_BY_COLUMN) + 1
        parts = round6.score_breakdown(index)
        for row, part in zip(cells, parts):
            text = row[col]["display_value"]
            css = dict(styler.ctx.get((cells.index(row), col - 1), []))
            kind = ui_results._direction(part)
            if kind:
                seen.add(kind)
                assert text.endswith("too high" if kind == "above" else "too low"), text
                assert css.get("color") == ui_insights.direction_colour(kind, which), (text, css)
            else:
                assert "color" not in css, (text, css)
    assert seen == {"above", "below"}


def test_all_formulations_shows_one_decimal_and_the_export_keeps_two(round6):
    frame = round6.history_frame()
    shown = dict(zip(frame["Formulation"].astype(int), frame[wording.OVERALL_SCORE_COLUMN]))
    assert shown[16] == "98.3" and shown[18] == "91.7" and shown[10] == "95.0"
    out = round6.history_export_frame()
    kept = dict(zip(out["Formulation"].astype(int), out[wording.OVERALL_SCORE_COLUMN]))
    assert kept[16] == 98.33


def test_results_tab_says_the_score_at_one_decimal(round6):
    at = AppTest.from_file(APP_PATH, default_timeout=300)
    at.session_state["_loaded_project"] = sp.ROUND6
    at.run()
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception
    tab = at.tabs[2]
    said = " ".join([m.value for m in tab.markdown] + [c.value for c in tab.caption])
    assert "Overall score 98.3 of 100." in said
    assert "Formulations 16 and 17 are tied at 98.3." in said
    assert "98.33" not in said


@pytest.mark.parametrize("which", ["light", "dark"])
def test_round_results_colour_and_bold_the_direction(round6, monkeypatch, which):
    import ui_insights
    monkeypatch.setattr(ui_insights, "theme", lambda: which)
    table = ui_results.round_results(round6)
    styled = ui_results._styled(table["frame"], table["off"], tinted=set(table["best"]))
    css = {}
    styled._compute()
    for style in styled._translate(False, False)["cellstyle"]:
        for sel in style["selectors"]:
            r, c = map(int, sel.split("row")[1].split("_col"))
            css[(r, c)] = "; ".join(f"{k}: {v}" for k, v in style["props"])
    above, below = (ui_insights.direction_colour(k, which) for k in ("above", "below"))
    cols = list(table["frame"].columns)[1:]
    tangy = cols.index(table["columns"][4])
    cons = cols.index(table["columns"][2])
    assert f"color: {below}" in css[(0, tangy)] and f"color: {above}" in css[(2, cons)]
    shown = styled.data
    assert shown.iloc[2, cons] == "**13**" and shown.iloc[0, tangy] == "**4**"
    assert shown.iloc[0, cons] == "12"


def test_the_results_key_is_coloured_under_the_table(round6):
    import ui_insights
    html = ui_results.results_key_markup(ui_results.round_results(round6))
    assert html.startswith("Panel median (n = 5), 0–15.")
    assert f"color: {ui_insights.direction_colour('above')}" in html
    assert ">Above target<" in html and ">Below target<" in html
