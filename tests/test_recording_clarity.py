"""Recording scopes and unknown property data must be explicit."""
import io
from pathlib import Path
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from food_bo import FoodOptimizer
import wording

@pytest.fixture
def opt(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    o=FoodOptimizer('my_project',robust=False)
    o.add_ingredient('Protein',5,20)
    o.add_ingredient('Water',0,100)
    o.add_objective('Taste',1,goal='max',min_val=0,max_val=10)
    o.set_formulation_total(100)
    o.set_formula('Water','= rest')
    return o

def test_unknown_property_blocks_generation_until_explicit_zero(opt):
    opt.add_property('Fat')
    opt.set_property_value('Protein','Fat',10)
    opt.add_constraint('Fat',max_val=15)
    before=opt.export_json()
    with pytest.raises(ValueError,match='Complete Fat for: Water'):
        opt.ask(1)
    assert opt.export_json()==before
    assert not opt._check_constraints({'Protein':10,'Water':90})
    opt.set_property_value('Water','Fat',0)
    rows=opt.ask(1)
    assert rows and opt._check_constraints(rows[0])
    opt.set_property_value('Protein','Fat',None)
    with pytest.raises(ValueError,match='Protein'): opt.ask(1)

def test_empty_property_bounds_do_not_restrict_or_require_data(opt):
    opt.add_constraint('Fat')
    assert opt.ask(1)

def test_old_example_note_is_shortened_but_custom_note_kept(opt):
    opt.set_targets_source(wording.LEGACY_SAMPLE_TARGETS_SOURCE)
    assert FoodOptimizer('my_project').targets_source==wording.SAMPLE_TARGETS_SOURCE
    opt.set_targets_source('My own panel protocol')
    assert FoodOptimizer('my_project').targets_source=='My own panel protocol'

def test_lot_editor_autosaves_and_prefills_workbook(opt):
    opt.set_records('lot',True)
    opt.set_records('actual',True)
    opt.set_pending_batch([{'Protein':10,'Water':90}],batch_no=1)
    at=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'), default_timeout=120).run()
    assert not at.exception
    assert 'actual' not in at.multiselect(key='record_fields').value
    assert at.checkbox(key='record_actual').value
    assert any(m.value == f'**{wording.LOT_ENTRY_HEADING}**' for m in at.tabs[1].markdown)
    assert not any(e.label == wording.CUSTOM_RECORDS_HEADING for e in at.tabs[1].expander)
    for value in ['00123','00124','']:
        key=next(k for k in at.session_state.filtered_state if k.startswith('custom_lots_'))
        at.session_state[key]={'edited_rows':{0:{'Lot':value}},'added_rows':[],'deleted_rows':[]}
        at.run()
        assert not at.exception
        saved=FoodOptimizer('my_project')
        assert saved.lots[1]['Protein']==value
        book=load_workbook(io.BytesIO(saved.workbook_bytes(saved.pending_batch,100,print_pack=False)))
        if value:
            assert any(c.value==value for sheet in book for row in sheet for c in row)
    at.multiselect(key='record_fields').set_value([]).run()
    assert FoodOptimizer('my_project').records('actual')
    assert not any(m.value == f'**{wording.LOT_ENTRY_HEADING}**' for m in at.tabs[1].markdown)
    at.checkbox(key='record_actual').uncheck().run()
    assert not FoodOptimizer('my_project').records('actual')
    assert any(b.proto.popover.label=='How the score is calculated' for b in at.get('popover'))

def test_incomplete_lower_limit_can_be_saved_but_not_used(opt):
    opt.add_property('Protein content')
    opt.add_constraint('Protein content',min_val=5)
    with pytest.raises(ValueError,match='Complete Protein content'):
        opt.ask(1)
    assert opt.constraints[0]['min']==5
    assert opt.constraints[0]['max'] is None

def test_workbook_with_previous_goal_label_still_imports(opt):
    opt.set_pending_batch([{'Protein':10,'Water':90}],batch_no=1)
    book=load_workbook(io.BytesIO(opt.workbook_bytes(opt.pending_batch,100,print_pack=True)))
    sheet=book[wording.batch_sheet_name(1)]
    cell=next(c for row in sheet for c in row if c.value=='Taste · Prefer higher values')
    cell.value='Taste · Higher is better'
    sheet.cell(cell.row,2,7)
    source=io.BytesIO();book.save(source);source.seek(0)
    uploaded=opt.results_from_workbook(source)
    assert uploaded.frame['Taste'].tolist()==[7]
