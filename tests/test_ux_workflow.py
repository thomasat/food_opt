"""User-facing editing and workbook interoperability regressions."""
import io
import math

import pytest
from openpyxl import load_workbook

from food_bo import FoodOptimizer
import wording


@pytest.fixture
def project(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer('editable', robust=False)
    opt.add_ingredient('Protein', 5, 20)
    opt.add_ingredient('Water', 0, 100)
    opt.set_formulation_total(100)
    opt.set_formula('Water', '= rest')
    opt.add_objective('Taste', 1, goal='max', min_val=0, max_val=10)
    opt.set_pending_batch([{'Protein': 10, 'Water': 90}, {'Protein': 15, 'Water': 85}], batch_no=1)
    return opt


def data(book):
    out = io.BytesIO(); book.save(out); out.seek(0)
    return out


def result_row(sheet, label='Taste', col=1):
    return next(c.row for row in sheet for c in row if c.column == col and (c.value == label or str(c.value).startswith(label + " (") or str(c.value).startswith(label + " ·")))


def test_edit_recalculates_preserves_original_and_reloads(project):
    assert project.edit_pending_formulations({1: {'Protein': 12}})
    row = project.pending_batch[0]
    assert row['recipe'] == {'Protein': 12, 'Water': 88}
    assert row['original_recipe'] == {'Protein': 10, 'Water': 90}
    assert FoodOptimizer('editable').pending_batch[0] == row
    assert project.storage.list_archives()
    assert not project.edit_pending_formulations({1: {'Protein': 12}})


@pytest.mark.parametrize('changes', [{1: {'Protein': math.nan}}, {1: {'Protein': 21}},
                                    {1: {'Water': 80}}, {3: {'Protein': 10}},
                                    {1: {'Protein': 12}, 2: {'Protein': -1}}])
def test_invalid_edits_are_atomic(project, changes):
    before = project.export_json()
    with pytest.raises(ValueError):
        project.edit_pending_formulations(changes)
    assert project.export_json() == before


def test_compact_workbook_results_and_actual_import(project):
    project.set_records('actual', True)
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    assert [s.title for s in book if s.sheet_state == 'visible'] == ['Round overview', 'Preparation', 'Results']
    assert book.active.title == 'Round overview'
    results = book['Results']
    results.cell(result_row(results), 2, 8)
    actual = result_row(results, 'Actual amounts and settings — optional')
    results.cell(actual + 3, 2, 11)
    upload = project.results_from_workbook(data(book))
    assert upload.frame['Taste'].tolist() == [8]
    assert upload.actual == {1: {'Protein': 11}}


def test_old_workbook_after_edit_is_refused(project):
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    project.edit_pending_formulations({1: {'Protein': 12}})
    with pytest.raises(ValueError, match='changed after'):
        project.results_from_workbook(data(book))


def test_legacy_results_on_both_surfaces_merge_and_conflicts_fail(project):
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100)))
    book['Round 1'].cell(result_row(book['Round 1']), 2, 8)
    page = book['Formulation 2']
    page.cell(result_row(page, col=2), 4, 7)
    upload = project.results_from_workbook(data(book))
    assert upload.frame['Taste'].tolist() == [8, 7]
    page = book['Formulation 1']
    page.cell(result_row(page, col=2), 4, 9)
    with pytest.raises(ValueError, match='conflicting Taste'):
        project.results_from_workbook(data(book))


def test_legacy_without_metadata_still_imports(project):
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100)))
    from workbook_flow import METADATA_SHEET
    book.remove(book[METADATA_SHEET])
    book['Round 1']['A1'].comment = None
    book['Round 1'].cell(result_row(book['Round 1']), 2, 8)
    assert project.results_from_workbook(data(book)).frame['Taste'].tolist() == [8]


def test_compact_preparation_lots_and_records(project):
    project.add_premix('Seasoning', 'portioned')
    project.set_premix_parts('Seasoning', [{'name': 'Salt', 'share': 100, 'unit': 'g'}])
    project.add_ingredient('Seasoning', 1, 3)
    project.set_records('lot', True)
    project.set_records('actual', True)
    project.set_pending_batch([{'Protein': 10, 'Seasoning': 2, 'Water': 88}], batch_no=2)
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    sheet = book['Preparation']
    sheet.cell(result_row(sheet, wording.PREMIX_LOT_LABEL), 2, 'Blend-42')
    sheet.cell(result_row(sheet, wording.PREMIX_BLENDED_BY_LABEL), 2, 'Alex')
    book['Results'].cell(result_row(book['Results']), 2, 7)
    upload = project.results_from_workbook(data(book))
    assert upload.lots['Seasoning'] == 'Blend-42'
    assert any(r['value'] == 'Alex' for r in upload.bench_records)


def test_the_compact_sheets_pointer_never_comes_back_as_a_lot(project):
    """The compact workbook rewrites a portioned pre-mix's Lot cell as a
    live link to the Preparation sheet. The reader skipped only the printed
    wording, so every upload of the app's default download filed the
    substitute as the pre-mix's lot number."""
    project.add_premix('Seasoning', 'portioned')
    project.set_premix_parts('Seasoning', [{'name': 'Salt', 'share': 100, 'unit': 'g'}])
    project.add_ingredient('Seasoning', 1, 3)
    project.set_records('lot', True)
    project.set_pending_batch([{'Protein': 10, 'Seasoning': 2, 'Water': 88}], batch_no=2)
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    printed = [c.value for row in book['Round overview'] for c in row]
    assert wording.PREMIX_LOT_ON_PREPARATION in printed
    book['Results'].cell(result_row(book['Results']), 2, 7)
    assert project.results_from_workbook(data(book)).lots == {}


def test_guided_example_generates_feasible_rounds(tmp_path, monkeypatch):
    kind = 'Burger formulation'
    import sample_projects
    from storage import LocalStorage
    monkeypatch.chdir(tmp_path)
    assert sample_projects.OPTIONS == [kind]
    opt = sample_projects.build(kind, sample_projects.NAMES[kind], LocalStorage(), LocalStorage())
    assert not opt.X_history
    # Where firmness 6 and juiciness 7 came from: the control, the cook,
    # the panel and the direction of failure, with the allergens.
    assert '80/20 beef control cooked to 71 °C core' in opt.targets_source
    assert 'Contains soy and wheat (gluten).' in opt.targets_source
    suggestions = opt.ask(3)
    assert len(suggestions) == 3
    for recipe in suggestions:
        assert opt._check_constraints(recipe)
        assert recipe['Hydration water'] == pytest.approx(2.2 * (recipe['Textured pea protein'] + recipe['Textured soy protein']))
        assert recipe['Remaining water'] >= 0
        assert sum(v for k, v in recipe.items() if k != 'Mixing time after fat') == pytest.approx(100)
    opt.set_pending_batch(suggestions)
    book = load_workbook(io.BytesIO(opt.workbook_bytes(opt.pending_batch, opt.open_round_size(), print_pack=False)))
    assert [s.title for s in book if s.sheet_state == 'visible'] == ['Round overview', 'Preparation', 'Results']


def test_a_250_g_round_of_the_sample_raises_no_caution(tmp_path, monkeypatch):
    """250 g is the first size anybody tries, and the sample's Seasoning
    blend is fixed at 2.2 g per 100 g. 2.2 * 2.5 is 5.5; the stored amount
    is 5.499999999999999, and one unit in the last place turned ordinary
    bench work into "Seasoning blend goes past the amounts you allowed" on
    screen, on tab 3 and on the printed Round sheet."""
    import sample_projects
    from storage import LocalStorage
    monkeypatch.chdir(tmp_path)
    opt = sample_projects.build('Burger formulation', 'Sample project',
                                LocalStorage(), LocalStorage())
    opt.set_pending_batch(opt.ask(3))
    opt.scale_round(250.0)
    assert opt.open_round_size() == 250.0
    assert opt.scaled_caution(opt.pending_batch, 250.0, sized=True) == ""
    assert opt.scaled_limit_caution(opt.pending_batch, 250.0, sized=True) == ""
    assert opt.scaled_cautions(opt.pending_batch, 250.0, sized=True) == [
        opt.scaled_amounts_note(opt.pending_batch, 250.0, sized=True)]


def test_recorded_row_in_partial_round_cannot_be_edited(project):
    project.tell(project.pending_batch[0]['recipe'], {'Taste': 8}, formulation_no=1, batch_no=1)
    assert project.pending_batch
    with pytest.raises(ValueError, match='Only unrecorded'):
        project.edit_pending_formulations({1: {'Protein': 12}})
    assert project.edit_pending_formulations({2: {'Protein': 12}})
    assert project.recipe_history[0] == {'Protein': 10, 'Water': 90}


def test_editor_save_and_cancel_preserve_result_drafts(project):
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    app = str(Path(__file__).resolve().parents[1] / 'app.py')
    at = AppTest.from_file(app, default_timeout=120).run()
    at.number_input(key='f1_Taste').set_value(8).run()
    next(b for b in at.button if b.label == wording.EDIT_FORMULATIONS).click().run()
    assert not at.exception
    assert any(i.value == wording.EDIT_FORMULATIONS_ACTIVE for i in at.info)
    assert any(c.value == wording.EDIT_FORMULATIONS_CAPTION for c in at.caption)
    key = 'edit_round_editable_1'
    at.session_state[key + '_grid'] = {'edited_rows': {0: {'Protein (g)': 12}}, 'added_rows': [], 'deleted_rows': []}
    at.run()
    at.session_state[key + '_grid'] = {'edited_rows': {0: {'Protein (g)': 12}}, 'added_rows': [], 'deleted_rows': []}
    at.button(key=key + '_save').click().run()
    assert not at.exception
    assert FoodOptimizer('editable').pending_batch[0]['recipe']['Protein'] == 12
    assert at.number_input(key='f1_Taste').value == 8
    at.button(key=key + '_start').click().run()
    at.button(key=key + '_cancel').click().run()
    assert not at.exception
    assert at.number_input(key='f1_Taste').value == 8


def test_actual_and_measurement_columns_align_for_every_formulation(project):
    project.set_records('actual', True)
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    sheet = book['Results']
    actual = result_row(sheet, 'Actual amounts and settings — optional')
    for col in range(2, sheet.max_column + 1):
        number = project._formulation_column_number(sheet.cell(actual + 2, col).value)
        if number:
            sheet.cell(result_row(sheet), col, 6 + number)
            sheet.cell(actual + 3, col, 10 + number)
    upload = project.results_from_workbook(data(book))
    assert upload.frame['Taste'].tolist() == [7, 8]
    assert upload.actual == {1: {'Protein': 11}, 2: {'Protein': 12}}


def test_compact_measurement_units_are_visible_before_entry(project):
    project.add_objective('Cook loss', 1, goal='min', min_val=0, max_val=40, unit='%')
    project.set_pending_batch([{'Protein': 10, 'Water': 90}], batch_no=2)
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    sheet = book['Results']
    row = result_row(sheet, 'Cook loss (%)')
    sheet.cell(row, 2, 12.5)
    sheet.cell(result_row(sheet), 2, 7)
    assert project.results_from_workbook(data(book)).frame['Cook loss'].tolist() == [12.5]


def test_excel_cached_measurement_formula_is_preserved(project):
    from zipfile import ZipFile, ZIP_DEFLATED
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    sheet = book['Results']
    sheet.cell(result_row(sheet), 2, '=4+4')
    saved = data(book)
    with pytest.raises(ValueError, match='no saved result'):
        project.results_from_workbook(saved)
    # Excel supplies the cached number on recalculation; openpyxl cannot
    # calculate, so emulate the standard OOXML value Excel would write.
    original = ZipFile(data(book))
    out = io.BytesIO()
    with ZipFile(out, 'w', ZIP_DEFLATED) as rewritten:
        for name in original.namelist():
            payload = original.read(name)
            if name == 'xl/worksheets/sheet3.xml':
                payload = payload.replace(b'<f>4+4</f><v></v>', b'<f>4+4</f><v>8</v>')
            rewritten.writestr(name, payload)
    out.seek(0)
    assert project.results_from_workbook(out).frame['Taste'].tolist() == [8]


def test_workbook_separates_formulations_and_has_internal_navigation(project):
    project.set_records('actual', True)
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    for title in ('Round overview', 'Results'):
        sheet = book[title]
        assert 'top in Numbers' in sheet['A2'].value
        headers = [c for row in sheet for c in row if str(c.value or '').startswith('Formulation ')]
        assert len(headers) >= 2
        for header in headers:
            assert header.border.left.style == 'medium'
            assert header.fill.fgColor.rgb in ('00DDEBF7', '00E2EFDA')
        links = [c for row in sheet for c in row if c.hyperlink]
        assert {c.hyperlink.location for c in links} >= {"'Results'!A1", "'Preparation'!A1", "'Round overview'!A1"}
        assert all(c.hyperlink.target is None for c in links)
        assert all(not sheet.column_dimensions[c.column_letter].hidden for c in links)
    sheet = book['Results']
    sheet.cell(result_row(sheet), 2, 8)
    assert project.results_from_workbook(data(book)).frame['Taste'].tolist() == [8]


def test_calculated_range_marker_can_be_saved_without_changing_rule(project):
    frame = project.ingredient_grid_frame()
    row = frame.loc[frame[wording.NAME_LABEL].eq('Water')].iloc[0]
    assert row[wording.HIGHEST_LABEL] == ""
    from food_bo import _range_from_cells
    for marker in (wording.CALCULATED_RANGE, wording.WORKED_OUT, wording.OLD_CALCULATED_RANGE):
        legacy = row.copy()
        legacy[wording.HIGHEST_LABEL] = marker
        assert _range_from_cells(row) == _range_from_cells(legacy)
    errors, _ = project.apply_ingredient_grid(frame)
    assert not errors
    assert project._var_by_name('Water')['balance']


@pytest.mark.parametrize('label', ['Ingredients varied separately', 'Variable-ratio blend', wording.PREMIX_MADE_AS_WEIGHED])
def test_previous_variable_blend_label_still_imports(tmp_path, monkeypatch, label):
    import pandas as pd
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer('legacy-blend')
    opt.load_ingredients_from_csv(pd.DataFrame([
        {'Name': 'Oils', 'Made as': label},
        {'Name': 'Coconut oil', 'Part of': 'Oils', 'Lowest': 2, 'Highest': 5, 'Unit': 'g'},
        {'Name': 'Sunflower oil', 'Part of': 'Oils', 'Lowest': 1, 'Highest': 4, 'Unit': 'g'},
    ]))
    assert opt.premix_mode('Oils') == wording.PREMIX_MADE_AS_WEIGHED
    assert opt.premixes['Oils']['mode'] == 'weighed'


def test_previous_workbook_calculated_labels_keep_actual_amounts(project):
    project.set_records('actual', True)
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    sheet = book['Results']
    sheet.cell(result_row(sheet), 2, 8)
    actual = result_row(sheet, 'Actual amounts and settings — optional')
    water_row = next(r for r in range(actual + 3, sheet.max_row + 1)
                     if str(sheet.cell(r, 1).value).startswith('Water'))
    sheet.cell(water_row, 2, 89)
    for page in book:
        for row in page:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = cell.value.replace('· calculated', '· worked out')
    uploaded = project.results_from_workbook(data(book))
    assert uploaded.actual == {1: {'Water': 89}}
    assert uploaded.frame['Taste'].tolist() == [8]


@pytest.mark.parametrize('edit', ['one', 'all', 'rename', 'delete', 'add', 'invalid'])
def test_share_preview_matches_save_without_mutation(project, monkeypatch, edit):
    project.add_objective('Texture', 1, goal='max', min_val=0, max_val=10)
    project.add_objective('Cook loss', 1, goal='min', min_val=0, max_val=40)
    frame = project.measurement_grid_frame().reset_index(drop=True)
    if edit == 'one':
        frame.loc[0, wording.SHARE_COLUMN] = 60
    elif edit == 'all':
        frame[wording.SHARE_COLUMN] = [17, 21, 29]
    elif edit == 'rename':
        frame.loc[0, wording.MEASUREMENT_COLUMN] = 'Overall liking'
        frame.loc[0, wording.SHARE_COLUMN] = 52
    elif edit == 'delete':
        frame = frame.iloc[1:].copy()
    elif edit == 'add':
        row = frame.iloc[0].copy()
        from food_bo import GRID_ID
        row[GRID_ID] = ''
        row[wording.MEASUREMENT_COLUMN] = 'Appearance'
        row[wording.SHARE_COLUMN] = 25
        frame.loc[len(frame)] = row
    else:
        frame.loc[0, wording.SHARE_COLUMN] = -1
    before = project.export_json()
    with monkeypatch.context() as guard:
        guard.setattr(project, 'save', lambda: pytest.fail('Preview wrote to storage'))
        preview = project.preview_measurement_shares(frame)
    assert project.export_json() == before
    if edit == 'invalid':
        assert preview is None
        return
    errors, _ = project.apply_measurement_grid(frame)
    assert not errors
    assert preview == project.share_percents()
    assert sum(preview.values()) == 100


def test_process_only_workbook_uses_trial_headers_and_imports(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer('fermentation', robust=False)
    opt.add_process_parameter('Temperature', 25, 40, unit='°C')
    opt.add_process_parameter('Time', 4, 24, unit='h')
    opt.add_objective('Acidity', 1, goal='target', target=4.5, min_val=3, max_val=7, unit='pH')
    opt.set_records('actual', True)
    opt.set_pending_batch([{'Temperature': 30, 'Time': 12}, {'Temperature': 35, 'Time': 18}])
    book = load_workbook(io.BytesIO(opt.workbook_bytes(opt.pending_batch, None, print_pack=False)))
    for title in ('Round overview', 'Results'):
        cells = [c.value for row in book[title] for c in row]
        assert 'Trial 1' in cells
        assert 'Trial 2' in cells
        assert 'Formulation 1' not in cells
    sheet = book['Results']
    sheet.cell(result_row(sheet, 'Acidity'), 2, 4.6)
    actual = result_row(sheet, 'Actual amounts and settings — optional')
    sheet.cell(actual + 3, 2, 31)
    uploaded = opt.results_from_workbook(data(book))
    assert uploaded.frame['Acidity'].tolist() == [4.6]
    assert uploaded.actual == {1: {'Temperature': 31}}


def test_record_selector_keeps_existing_values_when_hidden(project):
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    project.set_records('vendor', True)
    project._var_by_name('Protein')['vendor'] = 'Example supplier'
    project.save()
    app = str(Path(__file__).resolve().parents[1] / 'app.py')
    at = AppTest.from_file(app, default_timeout=120).run()
    at.multiselect(key='record_fields').set_value([]).run()
    assert not at.exception
    saved = FoodOptimizer('editable')
    assert not saved.records('vendor')
    assert saved._var_by_name('Protein')['vendor'] == 'Example supplier'
    at.multiselect(key='record_fields').set_value(['vendor', 'lot']).run()
    assert not at.exception
    saved = FoodOptimizer('editable')
    assert saved.records('vendor') and saved.records('lot')
    assert saved._var_by_name('Protein')['vendor'] == 'Example supplier'


@pytest.mark.parametrize('print_pack', [False, True])
def test_metadata_is_hidden_without_machine_comments_and_legacy_still_imports(project, print_pack):
    import custom_records
    import workbook_flow
    custom_records.add_field(project, 'Operator', 'formulation')
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=print_pack)))
    assert book[workbook_flow.METADATA_SHEET].sheet_state == 'veryHidden'
    assert all(c.comment is None for s in book for row in s for c in row)
    # An ordinary save keeps metadata and both measurements and custom records.
    workbook_flow.restore_metadata(book)
    custom = next(c for s in book for row in s for c in row
                  if c.comment and c.comment.text.startswith(custom_records.MARKER))
    custom.value = '0012'
    sheet = book['Round 1' if print_pack else 'Results']
    sheet.cell(result_row(sheet), 2, 8)
    legacy_bytes = data(book)
    legacy = project.results_from_workbook(legacy_bytes)
    assert legacy.frame['Taste'].tolist() == [8]
    assert legacy.custom_records[0][-1] == '0012'
    workbook_flow.hide_metadata(book)
    saved = load_workbook(data(book))
    assert all(c.comment is None for s in saved for row in s for c in row)
    modern = project.results_from_workbook(data(saved))
    assert modern.frame.equals(legacy.frame)
    assert modern.custom_records == legacy.custom_records


@pytest.mark.parametrize('damage', ['missing', 'version', 'json', 'shape', 'sheet', 'address'])
def test_missing_or_damaged_hidden_metadata_refused(project, damage):
    import json
    import workbook_flow as wf
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, 100, print_pack=False)))
    metadata = book[wf.METADATA_SHEET]
    if damage == 'missing':
        book.remove(metadata)
    elif damage == 'version':
        metadata['A1'] = 'unknown-version'
    elif damage == 'json':
        metadata['A2'] = '{'
    elif damage == 'shape':
        metadata['A2'] = '{}'
    else:
        records = json.loads(metadata['A2'].value)
        records[0][0 if damage == 'sheet' else 1] = 'missing' if damage == 'sheet' else 'A1:A2'
        metadata['A2'] = json.dumps(records)
    with pytest.raises(ValueError, match='round information'):
        project.results_from_workbook(data(book))


def test_large_metadata_survives_excel_cell_limit():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    import workbook_flow as wf
    book = Workbook()
    text = wf.MARKER + 'Protein & water — ' * 10000
    book.active['A1'].comment = Comment(text, 'Food Opt')
    book.active['A2'].comment = Comment('Keep this user note.', 'Scientist')
    wf.hide_metadata(book)
    assert book[wf.METADATA_SHEET].max_row > 2
    assert all(len(c.value) <= 20000 for row in book[wf.METADATA_SHEET] for c in row)
    saved = load_workbook(data(book))
    assert saved.active['A1'].comment is None
    assert saved.active['A2'].comment.text == 'Keep this user note.'
    wf.restore_metadata(saved)
    assert saved.active['A1'].comment.text == text


@pytest.mark.parametrize('label', ['Fixed-ratio pre-mix', wording.PREMIX_MADE_AS_PORTIONED])
def test_previous_and_current_fixed_premix_labels_keep_composition(tmp_path, monkeypatch, label):
    import pandas as pd
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer('fixed-blend')
    opt.load_ingredients_from_csv(pd.DataFrame([
        {'Name': 'Dry blend', 'Preparation': label, 'Lowest': 2, 'Highest': 10, 'Unit': 'g'},
        {'Name': 'Flour', 'Part of': 'Dry blend', 'Composition (%)': 60, 'Unit': 'g'},
        {'Name': 'Starch', 'Part of': 'Dry blend', 'Composition (%)': 40, 'Unit': 'g'},
    ]))
    assert opt.premix_mode('Dry blend') == wording.PREMIX_MADE_AS_PORTIONED
    assert opt.premixes['Dry blend']['mode'] == 'portioned'
    assert [p['share'] for p in opt.premix_parts('Dry blend')] == [60, 40]
    loaded = FoodOptimizer('fixed-blend')
    assert loaded.premixes == opt.premixes


def _sample(tmp_path, monkeypatch):
    import sample_projects
    from storage import LocalStorage
    monkeypatch.chdir(tmp_path)
    return sample_projects.build('Burger formulation', 'Sample project',
                                 LocalStorage(), LocalStorage())


def test_every_formulation_the_sample_suggests_is_a_patty_that_forms(
        tmp_path, monkeypatch):
    """A formed plant-based patty is 55-65 % moisture; beef 80/20 is ~60 %.
    The shipped space ran 42-60 % and the three formulations it generated
    came out at 47, 50 and 52 % — three crumbles and a wasted day. Six bands
    moving independently add 25 g of swing on the dry side and the row that
    fills to the total hands every gram of it to the water, so the solids
    are held between 36 and 43 g."""
    opt = _sample(tmp_path, monkeypatch)
    limit = next(qc for qc in opt.quantity_constraints
                 if qc.get('source') is None)
    assert sorted(limit['ingredients']) == sorted([
        'Textured pea protein', 'Textured soy protein', 'Dry blend',
        'Wheat gluten', 'Coconut oil', 'Sunflower oil'])
    assert (limit['min'], limit['max']) == (36.0, 43.0)
    for recipe in opt.ask(3):
        water = recipe['Hydration water'] + recipe['Remaining water']
        assert 52.0 <= water <= 66.0, recipe
        assert 36.0 <= sum(recipe[n] for n in limit['ingredients']) <= 43.0


def test_the_sample_method_says_how_hot_how_long_and_how_to_cook_it(
        tmp_path, monkeypatch):
    """What shipped had no temperature, no duration, no geometry and no
    cook in it, and its first step was a disclaimer. 'Use a fixed hydration
    protocol' tells three operators to be consistent without saying what
    about, and the round's variance then swamps the formulation effect it
    was built to measure."""
    opt = _sample(tmp_path, monkeypatch)
    method = "\n".join(opt.method_lines())
    for fact in ("45 °C", "≤ 5 °C", "180 °C", "74 °C core", "60 s on low",
                 "3 min per side", "100 mm across, 12 mm thick",
                 "Contains soy and wheat (gluten).",
                 "Cook loss = (raw − cooked)/raw × 100"):
        assert fact in method, fact
    # And it never states the formulation total, which is a lie on every
    # sheet printed at any other batch size. The Total row says it.
    assert "totals 100 g" not in method
    assert "they total 100 g" not in method
    assert "made to the batch size" in method
    # The disclaimer is kept, at the foot, not in the position a method's
    # first step belongs in.
    assert method.splitlines()[0].startswith("1. Hydrate")
    assert "demonstrate the app" in method.splitlines()[-1]


def test_a_resized_round_says_so_on_every_sheet_it_prints(tmp_path,
                                                          monkeypatch):
    """The screen said 'Made to 250 g — …' the moment the box moved; every
    cell of the workbook printed from those numbers said it nowhere, so the
    bench got a sheet of unfamiliar amounts with nothing reconciling them."""
    opt = _sample(tmp_path, monkeypatch)
    opt.set_pending_batch(opt.ask(3))
    opt.scale_round(250.0)
    note = opt.scaled_amounts_note(opt.pending_batch, 250.0, sized=True)
    # The shipped sentence, unchanged (H5): the sheets say what the screen
    # says, in the screen's own words.
    assert note == ("Made to 250 g — every amount is 2.5 × the amounts you "
                    "set per 100 g.")
    book = load_workbook(io.BytesIO(opt.workbook_bytes(
        opt.pending_batch, 250.0, sized=True)))
    pages = [s.title for s in book if s.title.startswith('Formulation ')]
    assert len(pages) == 3
    for title in ['Round 1'] + pages:
        assert note in str(book[title]['A2'].value), title
    compact = load_workbook(io.BytesIO(opt.workbook_bytes(
        opt.pending_batch, 250.0, sized=True, print_pack=False)))
    assert note in str(compact['Round overview']['A2'].value)


def test_a_setting_is_printed_on_a_step_its_dial_can_be_set_to(tmp_path,
                                                               monkeypatch):
    """`109.04` s asks the bench to round it, and three benches round three
    ways. It also sat directly under `Total (g) 100.00` with no number
    format, where it read for a moment as another mass."""
    opt = _sample(tmp_path, monkeypatch)
    for recipe in opt.ask(3):
        assert recipe['Mixing time after fat'] % 5 == 0, recipe
        assert 45 <= recipe['Mixing time after fat'] <= 150
    opt.set_pending_batch(opt.ask(3))
    sheet = load_workbook(io.BytesIO(opt.workbook_bytes(
        opt.pending_batch, 100.0)))['Round 1']
    flat = [c.value for row in sheet.iter_rows() for c in row]
    heading = flat.index(wording.SETTINGS_SHEET_HEADING)
    total = flat.index(wording.ROUND_TOTAL_COLUMN + ' (g)') if (
        wording.ROUND_TOTAL_COLUMN + ' (g)') in flat else 0
    assert heading > flat.index(wording.METHOD_SHEET_HEADING)
    assert opt.settings_step_note() == "Settings are set to the nearest 5 s."
    assert opt.settings_step_note() in flat


def test_cook_loss_asks_for_a_per_cent_on_every_sheet(tmp_path, monkeypatch):
    """A blank cell cannot display a unit after its value, so the label is
    the only place it can be said — and a bench handed "Cook loss" writes
    18, 0.18 or 18.4 g. The compact sheet said it; the print pack did not."""
    opt = _sample(tmp_path, monkeypatch)
    opt.set_pending_batch(opt.ask(3))
    for print_pack in (True, False):
        book = load_workbook(io.BytesIO(opt.workbook_bytes(
            opt.pending_batch, 100.0, print_pack=print_pack)))
        for sheet in book:
            labels = [str(c.value) for row in sheet.iter_rows() for c in row
                      if c.value and 'Cook loss' in str(c.value)]
            assert all('Cook loss (%)' in label or 'raw − cooked' in label
                       for label in labels), (sheet.title, labels)
    # ...and the whole interesting range of the measurement is the scale.
    assert [(o['min_val'], o['max_val']) for o in opt.objectives
            if o['name'] == 'Cook loss'] == [(0, 40)]


def test_the_bench_is_given_somewhere_to_write_the_numbers_cook_loss_needs(
        tmp_path, monkeypatch):
    """Twenty per cent of the score rests on (raw − cooked) / raw × 100, and
    the workbook asked for the answer with nowhere to write either weight.
    The two temperatures are what a methylcellulose system stands on."""
    import custom_records
    opt = _sample(tmp_path, monkeypatch)
    assert [f['name'] for f in custom_records.fields(opt, 'formulation')] == [
        'Raw weight (g)', 'Cooked weight (g)',
        'Water temperature at addition (°C)',
        'Mass temperature out of the bowl (°C)']


def test_a_premix_is_not_made_at_a_quantity_that_prints_one_column_twice(
        tmp_path, monkeypatch):
    """At exactly 100 g, `Amount (g)` and `Composition (%)` are the same
    column printed twice. The floor is the smallest quantity that blends
    evenly, and one pre-mix can be told it needs more than the usual 100 g."""
    opt = _sample(tmp_path, monkeypatch)
    assert opt.premix_smallest_quantity('Dry blend') == 150.0
    assert opt.premix_smallest_quantity('Seasoning blend') == 100.0
    opt.set_pending_batch(opt.ask(3))
    made = dict((name, make) for name, make, _ in opt.round_make_quantities())
    assert made['Dry blend'] == 150.0


def test_a_bench_record_cites_the_sheet_and_cell_it_was_really_read_from(
        project):
    """The provenance table read `Round 1 | B54` off the layout the app
    expands an uploaded file into. The file the bench uploaded has no sheet
    called Round 1: the number was written in Results, and that is where a
    reader has to go to check it."""
    book = load_workbook(io.BytesIO(project.workbook_bytes(
        project.pending_batch, 100, print_pack=False)))
    sheet = book['Results']
    row = result_row(sheet)
    sheet.cell(row, 2, 7)
    upload = project.results_from_workbook(data(book))
    cited = [r for r in upload.bench_records if str(r['value']) == '7']
    assert cited, upload.bench_records
    assert cited[0]['sheet'] == 'Results'
    assert cited[0]['cell'] == f'B{row}'
    assert {r['sheet'] for r in upload.bench_records} <= set(book.sheetnames)


def test_the_download_caption_describes_the_file_the_box_produces(project):
    """Ticking `Include individual formulation pages for printing` renames
    every sheet; the caption above the box went on promising a sheet called
    Results, which that file does not have."""
    from streamlit.testing.v1 import AppTest
    from pathlib import Path
    app = str(Path(__file__).resolve().parents[1] / 'app.py')
    at = AppTest.from_file(app, default_timeout=180).run()
    said = [c.value for c in at.caption]
    assert wording.COMPACT_WORKBOOK_CAPTION in said
    assert wording.PRINT_PACK_WORKBOOK_CAPTION not in said
    at.checkbox(key='workbook_print_pack').set_value(True).run()
    assert not at.exception
    said = [c.value for c in at.caption]
    assert wording.PRINT_PACK_WORKBOOK_CAPTION in said
    assert wording.COMPACT_WORKBOOK_CAPTION not in said
    # ...and the sheet the standing caption names is in the file it names it
    # for.
    book = load_workbook(io.BytesIO(project.workbook_bytes(
        project.pending_batch, 100, print_pack=False)))
    assert 'Results' in book.sheetnames
