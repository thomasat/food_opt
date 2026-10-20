"""Make a round and Results: every rewritten refusal names the row or cell
and says the fix (error-message review, 0.8.0).

Each test triggers the refusal where the reader meets it — the screen
through AppTest, or the method the screen calls — and pins the sentence.
"""
import io
import os

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

import wording
from food_bo import FoodOptimizer

APP_PATH = os.path.join(os.path.dirname(__file__), "..", "app.py")


class _FakeUpload(io.BytesIO):
    """What st.file_uploader hands back; AppTest cannot drive the widget."""

    def __init__(self, name, data):
        super().__init__(data)
        self.name = name
        self.size = len(data)
        self.file_id = name


def _uploads(monkeypatch, prefix, name, data):
    """Stand in for the one uploader whose key starts with `prefix`."""
    import streamlit as st_module
    real = st_module.file_uploader

    def _uploader(label, *a, **k):
        if str(k.get("key") or "").startswith(prefix):
            return _FakeUpload(name, data)
        return real(label, *a, **k)

    monkeypatch.setattr(st_module, "file_uploader", _uploader)


def _errors(at):
    return [e.value for e in at.error]


@pytest.fixture
def burger(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("burger")
    opt.set_amount_unit("g")
    opt.add_ingredient("Pea protein", 0, 25)
    opt.add_ingredient("Methylcellulose", 0, 3)
    opt.add_objective("Juiciness", 1.0, goal="target", target=7,
                      min_val=0, max_val=10, unit="/10")
    opt.add_objective("Firmness", 1.5, goal="target", target=6,
                      min_val=0, max_val=10, unit="N")
    return opt


@pytest.fixture
def open_batch(burger):
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0},
                              {"Pea protein": 20.0, "Methylcellulose": 2.0}])
    return burger


def _app(tab=None):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    if tab:
        at.session_state["main_tab"] = tab
    return at


# ------------------------------------------------------------------ #
#  Generate
# ------------------------------------------------------------------ #

def test_generate_with_a_property_gap_names_the_property_and_the_rows(burger):
    burger.add_property("Fat")
    burger.set_property_value("Pea protein", "Fat", 5)
    burger.add_constraint("Fat", max_val=10)
    at = _app(wording.TAB_BATCH)
    at.run()
    next(b for b in at.button if b.key == "generate").click()
    at.run()
    assert not at.exception
    assert ("Enter Fat for Methylcellulose under Ingredient properties, or "
            "delete the Fat limit.") in _errors(at), _errors(at)


def test_generate_with_everything_fixed_says_which_cells(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("fixed")
    opt.add_ingredient("Salt", 2, 2)
    with pytest.raises(ValueError) as refused:
        opt.ask(1)
    assert str(refused.value) == (
        "Every row is fixed. In Set up, give one a Highest above its Lowest.")


def test_a_failed_generate_says_what_to_try(open_batch, monkeypatch):
    open_batch.set_pending_batch(None)

    def broken(self, *a, **k):
        raise RuntimeError("solver blew up")

    monkeypatch.setattr(FoodOptimizer, "ask", broken)
    at = _app(wording.TAB_BATCH)
    at.run()
    next(b for b in at.button if b.key == "generate").click()
    at.run()
    assert not at.exception
    assert ("Could not make formulations. Ask for fewer, or widen a limit "
            "in Set up.") in _errors(at), _errors(at)
    assert "solver" not in " ".join(_errors(at))


# ------------------------------------------------------------------ #
#  Batch size box
# ------------------------------------------------------------------ #

def test_batch_size_refusals_name_the_box_and_the_number_to_type():
    assert wording.batch_size_too_big("28 g", "25 g") == (
        "Batch size: enter 28 g or less, the most these ingredients make. "
        "The Round stays at 25 g.")
    assert wording.batch_size_too_small("40 g", "100 g") == (
        "Batch size: enter 40 g or more, the least these ingredients make. "
        "The Round stays at 100 g.")
    assert wording.batch_size_leaves_rest_below_zero(
        "Water", "11 g", "g", "100 g") == (
        "Batch size: enter 11 g or more, or Water goes below 0 g. "
        "The Round stays at 100 g.")
    # A Round made as generated has no size to stay at.
    assert wording.batch_size_too_big("28 g", "") == (
        "Batch size: enter 28 g or less, the most these ingredients make.")


def test_a_too_small_batch_size_is_refused_on_screen(open_batch):
    open_batch._var_by_name("Pea protein")['bounds'] = (5.0, 25.0)
    open_batch.save()
    at = _app(wording.TAB_BATCH)
    at.run()
    at.number_input(key="scale_total").set_value(2.0)
    at.run()
    assert not at.exception
    said = [c.value for c in at.caption]
    assert any(c.startswith("Batch size: enter 5 g or more, the least these "
                            "ingredients make.") for c in said), said


# ------------------------------------------------------------------ #
#  Add a formulation of your own · Save results
# ------------------------------------------------------------------ #

def test_add_your_own_names_every_box_left_blank(burger):
    at = _app(wording.TAB_BATCH)
    at.run()
    next(b for b in at.button if b.label == wording.ADD_TO_THIS_BATCH).click()
    at.run()
    assert not at.exception
    assert ("Enter an amount for Pea protein and Methylcellulose; type 0 to "
            "leave one out.") in _errors(at), _errors(at)


def test_an_out_of_range_result_on_save_names_the_formulation(open_batch):
    """Once a row of the Round is recorded the grid calls rows by their
    numbers, and the refusal does too."""
    open_batch.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                    {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1)
    at = _app(wording.TAB_BATCH)
    at.run()
    at.number_input(key="f2_Firmness").set_value(12.0)
    at.number_input(key="f2_Juiciness").set_value(5.0)
    at.run()
    next(b for b in at.button if b.label == wording.SAVE_RESULTS).click()
    at.run()
    assert not at.exception
    assert any(e.startswith("Formulation 2: Firmness 12 N is outside")
               for e in _errors(at)), _errors(at)


def test_every_row_not_scored_says_how_to_save():
    assert wording.nothing_to_save(1) == (
        "This formulation is ticked Not scored. Untick it and enter results "
        "to save.")
    assert wording.nothing_to_save(3).startswith(
        "All 3 formulations are ticked Not scored.")


# ------------------------------------------------------------------ #
#  Edit formulations of the open Round
# ------------------------------------------------------------------ #

@pytest.fixture
def editable(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer('editable', robust=False)
    opt.add_ingredient('Protein', 5, 20)
    opt.add_ingredient('Water', 0, 100)
    opt.set_formulation_total(100)
    opt.set_formula('Water', '= rest')
    opt.add_objective('Taste', 1, goal='max', min_val=0, max_val=10)
    opt.set_pending_batch([{'Protein': 10, 'Water': 90},
                           {'Protein': 15, 'Water': 85}], batch_no=1)
    return opt


def test_a_cleared_cell_in_edit_formulations_names_row_and_column(editable):
    at = AppTest.from_file(APP_PATH, default_timeout=180).run()
    next(b for b in at.button
         if b.label == wording.EDIT_FORMULATIONS).click().run()
    key = 'edit_round_editable_1'
    edits = {'edited_rows': {1: {'Protein (g)': None}}, 'added_rows': [],
             'deleted_rows': []}
    at.session_state[key + '_grid'] = edits
    at.run()
    at.session_state[key + '_grid'] = edits
    at.button(key=key + '_save').click().run()
    assert not at.exception
    assert "Formulation 2: enter an amount for Protein." in _errors(at), \
        _errors(at)


def test_edit_refusals_name_the_formulation(editable):
    import math
    for bad in (None, math.nan, math.inf, "abc"):
        with pytest.raises(ValueError) as refused:
            editable.edit_pending_formulations({1: {'Protein': bad}})
        assert str(refused.value) == "Formulation 1: enter an amount for Protein."


def test_an_edit_that_misses_the_size_says_what_it_adds_up_to(tmp_path,
                                                              monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer('nobalance', robust=False)
    opt.set_amount_unit('g')
    opt.add_ingredient('Protein', 0, 60)
    opt.add_ingredient('Water', 0, 100)
    opt.set_formulation_total(100)
    opt.add_objective('Taste', 1, goal='max', min_val=0, max_val=10)
    opt.set_pending_batch([{'Protein': 40, 'Water': 60}], batch_no=1)
    with pytest.raises(ValueError) as refused:
        opt.edit_pending_formulations({1: {'Protein': 37}})
    assert str(refused.value) == (
        "Formulation 1 adds up to 97.00 g, not the 100 g batch size. "
        "Change its amounts to match.")


def test_an_edit_that_breaks_a_limit_names_the_limit(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer('limited', robust=False)
    opt.set_amount_unit('g')
    opt.add_ingredient('Protein', 0, 60)
    opt.add_ingredient('Oil', 0, 60)
    opt.add_ingredient('Water', 0, 100)
    opt.set_formulation_total(100)
    opt.set_formula('Water', '= rest')
    opt.add_quantity_constraint(['Protein', 'Oil'], max_val=50)
    opt.add_objective('Taste', 1, goal='max', min_val=0, max_val=10)
    opt.set_pending_batch([{'Protein': 20, 'Oil': 10, 'Water': 70}],
                          batch_no=1)
    with pytest.raises(ValueError) as refused:
        opt.edit_pending_formulations({1: {'Protein': 45}})
    said = str(refused.value)
    assert said.startswith("Formulation 1 breaks the limit "), said
    assert said.endswith(". Change its amounts."), said
    assert "Print at a smaller" not in said


def test_an_edit_that_leaves_the_rest_row_below_zero_names_it(editable):
    editable._var_by_name('Protein')['bounds'] = (5.0, 200.0)
    with pytest.raises(ValueError) as refused:
        editable.edit_pending_formulations({1: {'Protein': 150}})
    assert str(refused.value) == (
        "Formulation 1: Water would go below 0 g. Change the other amounts.")


# ------------------------------------------------------------------ #
#  Or upload results from a file
# ------------------------------------------------------------------ #

def test_a_file_that_is_not_a_workbook_is_named(open_batch, monkeypatch):
    _uploads(monkeypatch, "results_file", "notes.xlsx", b"not a workbook")
    at = _app(wording.TAB_BATCH)
    at.run()
    assert not at.exception
    assert ("notes.xlsx is not a workbook. Upload the one downloaded for "
            "this Round.") in _errors(at), _errors(at)


def test_text_in_a_result_cell_names_the_cell_and_what_was_typed(
        open_batch, monkeypatch):
    csv = b"Formulation,Juiciness,Firmness\n1,7,soft\n"
    _uploads(monkeypatch, "results_file", "results.csv", csv)
    at = _app(wording.TAB_BATCH)
    at.run()
    assert not at.exception
    assert ("Formulation 1: enter a number for Firmness, not “soft”."
            in _errors(at)), _errors(at)


def test_an_extra_row_names_the_rounds_own_numbers(open_batch, monkeypatch):
    csv = b"Formulation,Juiciness,Firmness\n9,7,6\n"
    _uploads(monkeypatch, "results_file", "results.csv", csv)
    at = _app(wording.TAB_BATCH)
    at.run()
    assert not at.exception
    assert ("Formulation 9 is not in Round 1, which has 1 and 2. Check its "
            "number.") in _errors(at), _errors(at)


def test_a_missing_column_says_which_to_add(open_batch):
    import pandas as pd
    frame = pd.DataFrame({"Formulation": [1], "Juiciness": [7]})
    with pytest.raises(ValueError) as refused:
        open_batch.parse_batch_results(frame, open_batch.pending_batch)
    assert str(refused.value) == "Add a Firmness column, then upload again."
    frame = pd.DataFrame({"Formulation": [1]})
    with pytest.raises(ValueError) as refused:
        open_batch.parse_batch_results(frame, open_batch.pending_batch)
    assert str(refused.value) == (
        "Add columns for Juiciness and Firmness, then upload again.")


def _filled_workbook(opt, round_no=None, **cells):
    data = opt.workbook_bytes(opt.pending_batch, opt.open_round_size(),
                              print_pack=False)
    book = openpyxl.load_workbook(io.BytesIO(data))
    sheet = book[wording.RESULTS_SHEET]
    return book, sheet


def _result_cell(sheet, label):
    row = next(c.row for r in sheet for c in r if c.column == 1
               and str(c.value or "").startswith(label))
    return sheet.cell(row, 2)


def _bytes(book):
    out = io.BytesIO()
    book.save(out)
    out.seek(0)
    return out


def test_a_result_typed_as_a_calculation_names_the_cell(open_batch):
    book, sheet = _filled_workbook(open_batch)
    cell = _result_cell(sheet, "Firmness")
    cell.value = "=3+4"
    with pytest.raises(ValueError) as refused:
        open_batch.results_from_workbook(_bytes(book))
    said = str(refused.value)
    assert (f"Cell {cell.coordinate} on Results has a calculation but no "
            "value." in said or
            f"Cell {cell.coordinate} on the Results sheet is a calculation."
            in said), said
    assert "formula" not in said.lower()


def test_a_workbook_for_another_round_names_both_rounds(open_batch):
    data = open_batch.workbook_bytes(open_batch.pending_batch, None,
                                     print_pack=False)
    open_batch.set_pending_batch([{"Pea protein": 5.0,
                                   "Methylcellulose": 1.0}], batch_no=4)
    with pytest.raises(ValueError) as refused:
        open_batch.results_from_workbook(io.BytesIO(data))
    assert str(refused.value) == (
        "This workbook is for Round 1; Round 4 is open. Upload Round 4's "
        "workbook.")


def test_a_workbook_downloaded_before_an_edit_names_the_formulation(
        editable):
    data = editable.workbook_bytes(editable.pending_batch, 100,
                                   print_pack=False)
    editable.edit_pending_formulations({2: {'Protein': 12}})
    with pytest.raises(ValueError) as refused:
        editable.results_from_workbook(io.BytesIO(data))
    assert str(refused.value) == (
        "Formulation 2 changed after this workbook was downloaded. Download "
        "it again and copy your results in.")


def test_a_workbook_without_its_results_sheet_names_it(open_batch):
    book, _ = _filled_workbook(open_batch)
    del book[wording.RESULTS_SHEET]
    with pytest.raises(ValueError) as refused:
        open_batch.results_from_workbook(_bytes(book))
    assert str(refused.value) == (
        "This workbook has no Results sheet. Download it again and copy "
        "your results in.")


def test_upload_row_refusals_say_the_fix():
    assert wording.formulation_twice_in_the_sheet(4) == (
        "Formulation 4 is in the file twice. Keep one.")
    assert wording.formulation_has_no_measurements(4) == (
        "Formulation 4 has no results. Fill them in, or tick Not scored.")
    assert wording.workbook_sheet_missing("Round 3") == (
        "This workbook has no Round 3 sheet. Upload the workbook downloaded "
        "for Round 3.")
    assert wording.workbook_actual_calculation(4, "Salt") == (
        "Formulation 4: the Actual cell for Salt is a calculation. Type what "
        "you weighed.")
    long = wording.formulation_value_not_a_number(1, "Firmness", "x" * 80)
    assert len(long) < 90 and long.endswith("”."), long


# ------------------------------------------------------------------ #
#  Additional records in an uploaded workbook
# ------------------------------------------------------------------ #

def test_an_additional_record_written_twice_names_field_and_row(open_batch):
    import custom_records
    import workbook_flow
    custom_records.add_field(open_batch, "Operator", "formulation")
    data = open_batch.workbook_bytes(open_batch.pending_batch, None)
    book = openpyxl.load_workbook(io.BytesIO(data))
    workbook_flow.restore_metadata(book)
    marked = [c for s in book for r in s for c in r if c.comment
              and c.comment.text.startswith(custom_records.MARKER)
              and '"subject": "1"' in c.comment.text]
    if len(marked) < 2:
        # One cell per record: write it as a calculation instead.
        marked[0].value = "=1+1"
        workbook_flow.hide_metadata(book)
        with pytest.raises(ValueError) as refused:
            custom_records.read_workbook(_bytes(book), open_batch, 1)
        assert str(refused.value) == (
            "The Operator cell for Formulation 1 is a calculation. Type the "
            "value, then upload again.")
        return
    marked[0].value, marked[1].value = "Ana", "Ben"
    workbook_flow.hide_metadata(book)
    with pytest.raises(ValueError) as refused:
        custom_records.read_workbook(_bytes(book), open_batch, 1)
    assert str(refused.value) == (
        "Formulation 1 has two different Operator entries. Make them match, "
        "then upload again.")


# ------------------------------------------------------------------ #
#  Results · Add past formulations from a file
# ------------------------------------------------------------------ #

def _check_past_file(monkeypatch, name, data):
    _uploads(monkeypatch, "import_file", name, data)
    at = _app(wording.TAB_RESULTS)
    at.run()
    at.radio(key="add_past_mode").set_value(wording.UPLOAD_A_FILE)
    at.run()
    next(b for b in at.button if b.key == "check_import").click()
    at.run()
    assert not at.exception
    return at


def test_a_past_file_missing_a_column_says_which(burger, monkeypatch):
    csv = b"Pea protein,Juiciness,Firmness\n10,7,6\n"
    at = _check_past_file(monkeypatch, "past.csv", csv)
    assert "Add a Methylcellulose column, then upload again." in _errors(at), \
        _errors(at)


def test_a_past_file_with_blank_amounts_says_which_column(burger,
                                                          monkeypatch):
    csv = (b"Pea protein,Methylcellulose,Juiciness,Firmness\n"
           b"10,,7,6\n")
    at = _check_past_file(monkeypatch, "past.csv", csv)
    assert ("Fill the blank cells in Methylcellulose; type 0 for none."
            in _errors(at)), _errors(at)


def test_an_unreadable_past_file_is_named(burger, monkeypatch):
    at = _check_past_file(monkeypatch, "past.xlsx", b"garbage")
    assert ("Could not read past.xlsx. Upload a spreadsheet with one row per "
            "formulation.") in _errors(at), _errors(at)
