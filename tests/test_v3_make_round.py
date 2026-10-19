"""v3 · Make a round (final-v3/2-make-round.png, wording.md §2): the Round
as one table — the ingredients down, Best so far and each formulation
across, each cell the amount alone, coloured by direction (orange more,
purple less) when it differs from Best so far by at least 5 % of the row's
allowed range (owner, v3 final cleanup; it replaced the bracketed change
and the grey shading), limit tags and a Total row — and no forecast
anywhere on the screen."""
import pytest
from streamlit.testing.v1 import AppTest

import sample_projects as sp
import storage as storage_backend
import ui_batch
import wording
from food_bo import FoodOptimizer, uploaded_parts
from tests.test_app_ui import APP_PATH
from tests.test_insights import _burger


@pytest.fixture(autouse=True)
def _samples(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    sp.ensure(storage_backend.LocalStorage())
    yield


def _round6():
    return FoodOptimizer(sp.ROUND6, storage=storage_backend.LocalStorage())


def _cells(table):
    """{column: {row: CompareCell}}."""
    return {c: dict(zip(table.rows, col)) for c, col in zip(table.columns, table.cells)}


def _record_round6(opt):
    """Upload the pilot's Round 6 results file, as the screen does."""
    with open(sp.ROUND6_RESULTS, "rb") as fh:
        upload = opt.results_from_workbook(fh)
    sheet, weighed, lots = uploaded_parts(upload)
    parsed, _ = opt.parse_batch_results(sheet, opt.pending_batch, with_skipped=True,
                                        weighed=weighed)
    batch_no = opt.pending_batch_no
    by_number = {r["formulation"]: r["recipe"] for r in opt.pending_batch}
    for number, results, note in parsed:
        opt.tell(opt.amounts_as_weighed(by_number[number], weighed.get(number)),
                 results, formulation_no=number, batch_no=batch_no, note=note)
    opt.set_pending_batch(None)
    return opt


# ------------------------------------------------------------------ #
#  The table, worked out
# ------------------------------------------------------------------ #

def test_round6_table_values_signs_tags_and_total():
    table = ui_batch.round_compare(_round6())
    assert table.columns == ["Best so far (F10 · 95.0)", "F16", "F17", "F18"]
    assert table.rows[:7] == ["Coconut cream", "Tap water", "Lemon juice", "Salt",
                              "Nutritional yeast", "Tapioca starch", "Xanthan gum"]
    assert table.rows[-1] == "Total (allowed 97.69–98.69 g)"
    assert table.row_header == "Ingredient (g)"
    assert table.best_no == 10 and table.tied == [] and table.numbers == []
    c = _cells(table)
    best = c["Best so far (F10 · 95.0)"]
    assert best["Coconut cream"].text == "51.84" and best["Tap water"].text == "41.22"
    # One number per cell: the amount alone, no bracketed change.
    assert c["F16"]["Coconut cream"].text == "51.27"
    assert c["F18"]["Coconut cream"].text == "52.27"
    assert c["F17"]["Coconut cream"].text == "50.24"
    assert not any("(" in cell.text for col in table.cells for cell in col)
    # At the upper limit: the tag on a line of its own.
    assert c["F16"]["Tap water"].text == "42.00\nat upper limit"
    assert c["F16"]["Tap water"].tag == wording.AT_UPPER_LIMIT
    assert all(c[f]["Tapioca starch"].tag == wording.AT_UPPER_LIMIT for f in ("F16", "F17", "F18"))
    assert c["F18"]["Xanthan gum"].text == "0.65\nat upper limit"
    assert c["F18"]["Tap water"].tag == ""
    # The share is of the row's allowed range: salt 0–1.8 g.
    assert c["F18"]["Salt"].share == pytest.approx(abs(0.551 - 0.21) / 1.8, abs=0.01)
    totals = [c[k][table.rows[-1]].text for k in table.columns]
    assert totals == ["98.69", "98.69", "97.69", "98.69"]


def test_a_set_total_reads_as_its_one_number():
    opt = _burger(rounds=1)
    ui_batch.ask_alone(opt, n_suggestions=2)
    assert ui_batch.round_compare(opt).rows[-1] == "Total (set to 100.00 g)"


def test_no_total_limit_reads_total_in_its_unit():
    opt = _burger(rounds=3, fill=False)
    opt.quantity_constraints = []
    opt.formulation_total = None
    ui_batch.ask_alone(opt, n_suggestions=2)
    table = ui_batch.round_compare(opt)
    assert table.rows[-1] == "Total (g)"


def test_tie_names_the_first_made_and_suggestions_are_numbered():
    """After the pilot's Round 6, F16 and F17 tie at 98.3: the column names
    F16 (first made) and the tie is said; an app-generated Round not yet
    made reads Suggestion 1-3, its numbers kept for the workbook."""
    opt = _record_round6(_round6())
    ui_batch.ask_alone(opt, n_suggestions=3)
    table = ui_batch.round_compare(opt)
    assert table.columns[0] == "Best so far (F16 · 98.3)"
    assert table.best_no == 16 and table.tied == [17]
    labels = [c.split(" · ")[0] for c in table.columns[1:]]
    assert labels == ["Suggestion 1", "Suggestion 2", "Suggestion 3"]
    numbers = [int(r["formulation"]) for r in opt.pending_batch]
    assert table.numbers == list(zip([1, 2, 3], numbers))
    assert numbers == [19, 20, 21]
    # The workbook and the uploader read the formulation numbers, never the
    # screen's labels.
    assert opt.batch_frame(opt.pending_batch)["Formulation"].tolist() == numbers
    assert "F16, F17 (tied)" in wording.improve_on_caption([16, 17])
    assert wording.best_tied_help(16, [17]).endswith("tied with F17.")


def test_own_rows_keep_their_numbers_beside_suggestions(monkeypatch):
    opt = _burger(rounds=3)
    ui_batch.ask_alone(opt, n_suggestions=2)
    rows = list(opt.pending_batch)
    rows[1] = dict(rows[1], note="My own")
    monkeypatch.setattr(opt, "pending_batch", rows)
    table = ui_batch.round_compare(opt)
    assert table.columns[1].startswith("Suggestion 1")
    assert table.columns[2].startswith(f"F{rows[1]['formulation']}")


def test_less_tried_label_only_for_explores(monkeypatch):
    opt = _burger(rounds=3)
    ui_batch.ask_alone(opt, n_suggestions=3)
    roles = iter([wording.SUGGESTION_CLOSE, wording.SUGGESTION_DIFFERENT,
                  wording.SUGGESTION_BALANCED])
    monkeypatch.setattr(opt, "suggestion_role", lambda recipe, before=None: next(roles))
    table = ui_batch.round_compare(opt)
    assert table.columns[1:] == ["Suggestion 1", "Suggestion 2 · Less-tried area",
                                 "Suggestion 3"]
    for c in table.columns:
        assert wording.SUGGESTION_CLOSE not in c and wording.SUGGESTION_BALANCED not in c


def test_cold_start_shows_amounts_only():
    opt = _burger(rounds=1)
    ui_batch.ask_alone(opt, n_suggestions=3)
    table = ui_batch.round_compare(opt)
    assert table.best_no is None
    assert not any(c.startswith("Best so far") for c in table.columns)
    assert all("(" not in cell.text for col in table.cells for cell in col)
    assert all(cell.direction == "" for col in table.cells for cell in col)
    styles = ui_batch.compare_styles(table, "light")
    assert not any("background" in css or "color:" in css for col in styles for css in col)


# ------------------------------------------------------------------ #
#  Direction: orange more, purple less, at 5 % of the allowed range
# ------------------------------------------------------------------ #

def _directions(table):
    return {col: {row: cell.direction for row, cell in rows.items()}
            for col, rows in _cells(table).items()}


def test_round6_directions_against_f10():
    """Changes from F10 as a share of each row's allowed range, at the
    Formulation total: coconut F17 −10.7 %; salt F16 +9.5 %, F18 +19.0 %;
    yeast −17.7, −14.9, −18.6 %; xanthan F16 +5.1 %, F18 +14.5 %. Plain:
    lemon (−1.6, +4.9, −4.9 %), water (+4.6, +4.6, −1.8 %), tapioca
    (+4.9 % each)."""
    d = _directions(ui_batch.round_compare(_round6()))
    assert d["F17"]["Coconut cream"] == "less"
    assert d["F16"]["Coconut cream"] == d["F18"]["Coconut cream"] == ""
    assert [d[f]["Nutritional yeast"] for f in ("F16", "F17", "F18")] == ["less"] * 3
    assert d["F16"]["Salt"] == d["F18"]["Salt"] == "more" and d["F17"]["Salt"] == ""
    assert d["F16"]["Xanthan gum"] == d["F18"]["Xanthan gum"] == "more"
    assert d["F17"]["Xanthan gum"] == ""
    for row in ("Lemon juice", "Tap water", "Tapioca starch"):
        assert [d[f][row] for f in ("F16", "F17", "F18")] == [""] * 3, row
    assert not any(d["Best so far (F10 · 95.0)"].values())
    assert not d["F16"]["Total (allowed 97.69–98.69 g)"]


def test_the_5_percent_edge_uses_a_tolerance():
    assert ui_batch.direction_of(0.05 - 1e-12, 1.0) == "more"
    assert ui_batch.direction_of(-(0.05 - 1e-12), 1.0) == "less"
    assert ui_batch.direction_of(0.0499, 1.0) == ""
    assert ui_batch.direction_of(0.2, 4.0) == "more"       # 5 % of a 4 g range


@pytest.mark.parametrize("which", ["light", "dark"])
def test_direction_colours_use_the_themes_tokens(which):
    table = ui_batch.round_compare(_round6())
    tokens = ui_batch._COMPARE_COLOURS[which]
    styles = ui_batch.compare_styles(table, which)
    c = {col: dict(zip(table.rows, css)) for col, css in zip(table.columns, styles)}
    import ui_insights
    more, less = (ui_insights.direction_colour(k, which) for k in ("above", "below"))
    assert f"background-color: {tokens['best']}" in c[table.columns[0]]["Salt"]
    assert f"color: {more}" in c["F18"]["Salt"]
    assert f"color: {less}" in c["F17"]["Coconut cream"]
    # Plain: no colour and no grey shading of its own.
    assert "color" not in c["F16"]["Lemon juice"] and "background" not in c["F16"]["Lemon juice"]
    other = ui_batch._COMPARE_COLOURS["dark" if which == "light" else "light"]
    assert other["best"] not in " ".join(c[table.columns[0]].values())
    assert {"light": ("#c2410c", "#7e3fbf"), "dark": ("#ff9f5a", "#c79bff")}[which] == (more, less)


def test_coloured_numbers_are_bold_and_the_limit_line_is_muted_small_text():
    table = ui_batch.round_compare(_round6())
    cells, _css = ui_batch.compare_frames(_round6(), table, "light")
    assert cells.loc["Salt", "F18"] == "**0.55**"
    assert cells.loc["Lemon juice", "F18"] == "0.45"
    # The limit line: never the number's colour, whatever it is.
    assert cells.loc["Xanthan gum", "F18"] == "**0.65**  \n*at upper limit*"
    assert cells.loc["Tap water", "F16"] == "42.00  \n*at upper limit*"
    import ui_insights
    rule = ui_batch.limit_tag_style("light")
    assert rule["selector"] == "td em"
    assert f"color: {ui_insights._MUTED_LIGHT}" in rule["props"]
    assert "font-size" in rule["props"] and "font-style: normal" in rule["props"]


# ------------------------------------------------------------------ #
#  The screen
# ------------------------------------------------------------------ #

def _app(name=sp.ROUND6):
    at = AppTest.from_file(APP_PATH, default_timeout=600)
    at.session_state["_loaded_project"] = name
    at.run()
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    assert not at.exception, [e.message for e in at.exception]
    return at


_FORECAST_WORDS = ("Predicted", "predicted", "forecast", "Forecast", "95 %", "Predicts",
                   "Partly", "Not yet", "worth trying")


def test_round6_screen_has_the_table_captions_and_no_forecast():
    at = _app()
    tab = at.tabs[1]
    captions = [c.value for c in tab.caption]
    assert "Best bets to improve on F10." in captions
    assert wording.CHOSEN_BY_EGBO in captions
    key = next(c for c in captions if "More than F10" in c)
    assert "Less than F10" in key and "plain: within 5 % of its range" in key
    assert "Sheets: Round overview · Results" in captions
    assert len(tab.table) == 1
    shown = tab.table[0].value
    assert list(shown.columns) == ["Best so far (F10 · 95.0)", "F16", "F17", "F18"]
    assert shown.loc["Tap water", "F16"] == "42.00  \n*at upper limit*"
    assert shown.loc["**Total (allowed 97.69–98.69 g)**", "F17"] == "**97.69**"
    texts = captions + [e.label for e in tab.expander] + [m.value for m in tab.markdown]
    for word in _FORECAST_WORDS:
        assert not any(word in t for t in texts), word
    assert not any(e.label == 'Why these suggestions' for e in tab.expander)
    assert not tab.dataframe        # the old formulation-per-row table is gone
    # The key's "?" says the 5 % rule.
    key = next(c for c in tab.caption if "More than F10" in c.value)
    assert "5 %" in key.help and "allowed range" in key.help


def test_app_generated_round_reads_suggestions_on_screen():
    opt = _record_round6(_round6())
    ui_batch.ask_alone(opt, n_suggestions=3)
    opt.save()
    at = _app()
    tab = at.tabs[1]
    shown = tab.table[0].value
    assert [c.split(" · ")[0] for c in shown.columns[1:]] == [
        "Suggestion 1", "Suggestion 2", "Suggestion 3"]
    assert shown.columns[0] == "Best so far (F16 · 98.3)"
    captions = [c.value for c in tab.caption]
    assert "Best bets to improve on F16, F17 (tied)." in captions
    assert wording.CHOSEN_BY_EGBO not in captions
    assert not any(c.startswith("F19") for c in shown.columns)


def test_the_record_grid_names_suggestions_as_the_table_does():
    """v3 integration: the grid under the table says 'Suggestion 1' for an
    app suggestion not yet made, as the table does; the workbook, the upload
    and the records keep Formulation 19-21. Once a row is recorded every
    label is the formulation number again, in the table and the grid."""
    opt = _record_round6(_round6())
    ui_batch.ask_alone(opt, n_suggestions=3)
    opt.save()
    assert ui_batch.suggestion_numbers(opt) == {19: 1, 20: 2, 21: 3}
    at = _app()
    heads = [m.value for m in at.tabs[1].markdown]
    assert [h for h in heads if h.startswith("**Suggestion")] == [
        wording.suggestion_heading(k) for k in (1, 2, 3)]
    assert not [h for h in heads if h.startswith("**Formulation 1") or
                h.startswith("**Formulation 2")]
    # The workbook still names the formulations by number, and its upload
    # lands on those rows.
    import io
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(opt.workbook_bytes(opt.pending_batch)))
    assert [s for s in book.sheetnames if s.startswith("Formulation")] == [
        "Formulation 19", "Formulation 20", "Formulation 21"]
    for row, value in zip(range(18, 22), (12, 5, 5, 13)):
        book["Formulation 20"].cell(row, 4, value)
    out = io.BytesIO()
    book.save(out)
    out.seek(0)
    sheet, weighed, _ = uploaded_parts(opt.results_from_workbook(out))
    parsed, _ = opt.parse_batch_results(sheet, opt.pending_batch, with_skipped=True,
                                        weighed=weighed)
    assert [number for number, _, _ in parsed] == [20]
    # One row recorded: the Round's labels are formulation numbers again.
    number, results, note = parsed[0]
    recipe = next(r["recipe"] for r in opt.pending_batch if r["formulation"] == 20)
    opt.tell(recipe, results, formulation_no=20, batch_no=opt.pending_batch_no, note=note)
    assert ui_batch.suggestion_numbers(opt) == {}
    assert [c.split(" · ")[0] for c in ui_batch.round_compare(opt).columns[1:]] == [
        "F19", "F20", "F21"]


def test_make_a_round_fits_nothing(monkeypatch):
    """No forecast on Make a round, so nothing on it fits a model (another
    tab's own fits are that tab's)."""
    import os
    import traceback

    import food_bo
    from_here = []
    real = food_bo.fit_gpytorch_mll

    def spy(*a, **k):
        if any(os.path.basename(f.filename) == "ui_batch.py" for f in traceback.extract_stack()):
            from_here.append(1)
        return real(*a, **k)
    monkeypatch.setattr(food_bo, "fit_gpytorch_mll", spy)
    _app()
    assert from_here == []


def test_a_process_only_round_reads_trial(tmp_path, monkeypatch):
    """Five scored trials and one open: the table in the project's own words."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("ferment_v3")
    opt.add_process_parameter("Incubation temperature", 30, 42, unit="°C")
    opt.add_process_parameter("Incubation time", 4, 16, unit="h")
    opt.add_objective("Acidity", 1.0, goal="target", target=4.5, min_val=3.0, max_val=7.0,
                      unit="pH")
    for no, (temp, hours) in enumerate(
            ((32.0, 5.0), (34.0, 7.0), (36.0, 9.0), (38.0, 11.0), (40.0, 13.0)), start=1):
        opt.tell({"Incubation temperature": temp, "Incubation time": hours},
                 {"Acidity": 3.0 + 0.1 * temp}, formulation_no=no, batch_no=1)
    opt.set_pending_batch([{"Incubation temperature": 37.0, "Incubation time": 16.0}])
    opt.save()
    at = _app(opt.project_name)
    tab = at.tabs[1]
    shown = tab.table[0].value
    assert shown.index.name == "Process setting"
    assert list(shown.index) == ["Incubation temperature (°C)", "Incubation time (h)"]
    assert shown.columns[0].startswith("Best so far (T") and shown.columns[1].startswith(
        "Suggestion 1")
    assert shown.iloc[1, 1].endswith("*at upper limit*")
    key = next(c for c in tab.caption if "More than T" in c.value)
    assert "setting" in key.help
    texts = [c.value for c in tab.caption] + [str(c.help or "") for c in tab.caption]
    assert not any("ormulation" in t for t in texts)


def test_the_v3_lines_pass_the_vocabulary_guard_and_the_length_rules():
    from tests.test_food_bo import _BANNED
    captions = [wording.improve_on_caption([10]), wording.improve_on_caption([16, 17]),
                wording.CHOSEN_BY_EGBO, wording.COVER_CAPTION, wording.workbook_sheets(),
                wording.workbook_sheets(True)]
    labels = [wording.best_column(10, "95.0"), wording.suggestion_column(1), wording.LESS_TRIED,
              wording.AT_UPPER_LIMIT, wording.AT_LOWER_LIMIT, wording.compare_row_header("g")]
    helps = [wording.improve_on_help([10]), wording.best_tied_help(16, [17]),
             wording.SUGGESTION_NUMBERS_HELP, wording.less_tried_help(30),
             wording.compare_key_more(10), wording.compare_key_less(10),
             wording.compare_key_plain(5), wording.compare_key_help(10, 5),
             wording.workbook_sheets_help(), wording.workbook_sheets_help(True),
             wording.total_row("97.69", "98.69", "g")]
    for text in captions + labels + helps:
        assert not any(p.search(text) for p in _BANNED), text
    words = lambda text: [w for w in text.split() if w != "·"]
    for text in captions:
        assert len(words(text)) <= 8, text
    for text in labels:
        assert len(words(text)) <= 5, text
    assert wording.improve_on_caption([10]) == "Best bets to improve on F10."
    assert wording.improve_on_help([10]) == (
        "Chosen for how much they could improve on F10, weighing how good they look and how "
        "unsure the model is.")
    assert wording.workbook_sheets() == "Sheets: Round overview · Results"
    assert "Preparation" in wording.workbook_sheets_help()


def test_the_suggestion_numbers_help_lists_no_numbers():
    """v3 final cleanup: a live Round 7 on the pilot would have mapped its
    suggestions to F19–F21, the pilot's own numbers; the "?" says when
    they get numbers instead of listing them."""
    text = wording.SUGGESTION_NUMBERS_HELP
    assert text == "Suggestions get formulation numbers when their results are recorded."
    assert not any(ch.isdigit() for ch in text)
