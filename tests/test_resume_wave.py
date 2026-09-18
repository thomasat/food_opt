"""Regression coverage for the interrupted optional-field and workbook changes."""
import io
import json
import os
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from food_bo import FoodOptimizer
import storage
import wording

APP = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'app.py'))

@pytest.fixture
def project(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer('recording', robust=False)
    opt.add_ingredient('Water', 0, 100)
    opt.variables[0].update(vendor='Supplier', sku='W1')
    opt.add_objective('Taste', 1, goal='max', min_val=0, max_val=10)
    opt.set_pending_batch([{'Water': 40}], batch_no=1)
    return opt


def test_optional_fields_default_off_and_hidden_values_survive_grid_save(project):
    opt = project
    assert not any(opt.records(f) for f in ('vendor', 'sku', 'lot', 'actual'))
    frame = opt.ingredient_grid_frame()
    assert 'Vendor' not in frame and 'SKU' not in frame
    frame.loc[1, 'Highest'] = '95'
    assert opt.apply_ingredient_grid(frame)[0] == []
    opt.set_records('vendor', True)
    assert opt.ingredient_grid_frame().iloc[0]['Vendor'] == 'Supplier'
    assert opt.variables[0]['sku'] == 'W1'
    opt.set_records('vendor', False)
    assert FoodOptimizer('recording').variables[0]['vendor'] == 'Supplier'


def test_optional_fields_migrate_only_when_old_state_needs_them(project):
    state = project.export_json()
    del state['recorded_fields']
    other = FoodOptimizer('older', robust=False)
    other.import_json(state)
    assert other.records('vendor') and other.records('sku')
    assert not other.records('lot')
    state['recorded_fields'] = {'lot': 'false'}
    with pytest.raises(ValueError):
        FoodOptimizer.validate_state(state)


def test_no_lot_or_actual_cells_and_measurements_still_print_and_import(project):
    book = load_workbook(io.BytesIO(project.workbook_bytes(project.pending_batch, None)))
    summary, page = book['Round 1'], book['Formulation 1']
    assert 'Lot' not in [c.value for row in summary for c in row]
    assert 'Actual (g)' not in [c.value for row in page for c in row]
    assert ':$D$' in page.print_area
    row = next(c.row for cells in page for c in cells if c.value == 'Taste')
    page.cell(row, 4, 8)
    stream = io.BytesIO(); book.save(stream); stream.seek(0)
    upload = project.results_from_workbook(stream)
    assert upload.frame['Taste'].tolist() == [8]
    assert upload.actual == {} and upload.lots == {}


def test_preparation_records_survive_upload_reopen_and_export(project):
    opt = project
    opt.add_premix('Seasoning', 'portioned')
    opt.set_premix_parts('Seasoning', [{'name': 'Salt', 'share': 100, 'unit': 'g'}])
    opt.add_ingredient('Seasoning', 1, 3)
    opt.set_records('actual', True); opt.set_records('lot', True)
    opt.set_pending_batch([{'Water': 40, 'Seasoning': 2}], batch_no=1)
    book = load_workbook(io.BytesIO(opt.workbook_bytes(opt.pending_batch, None)))
    prep = book.active
    prep['D4'] = 101.2
    for label, value in [(wording.PREMIX_LOT_LABEL, 'S-17'),
                         (wording.PREMIX_BLENDED_BY_LABEL, 'Alex')]:
        row = next(c.row for cells in prep for c in cells if c.value == label)
        prep.cell(row, 2, value)
    page = book[wording.formulation_sheet_name(opt.pending_batch[0]['formulation'])]
    row = next(c.row for cells in page for c in cells if c.value == 'Taste')
    page.cell(row, 4, 8)
    stream = io.BytesIO(); book.save(stream); stream.seek(0)
    upload = opt.results_from_workbook(stream)
    assert upload.lots['Seasoning'] == 'S-17'
    assert any(r['value'] == '101.2' for r in upload.bench_records)
    opt.store_bench_records(1, upload.bench_records)
    restored = FoodOptimizer('recording')
    assert restored.bench_records['1'] == upload.bench_records
    assert restored.premixes['Seasoning']['parts'][0]['share'] == 100
    assert FoodOptimizer.validate_state(json.loads(json.dumps(restored.export_json())))
    exported = load_workbook(io.BytesIO(restored.all_formulations_workbook()))
    assert wording.BENCH_RECORDS_SHEET in exported.sheetnames
    assert 'Alex' in [c.value for row in exported[wording.BENCH_RECORDS_SHEET] for c in row]


def test_optional_checkbox_saves_and_upload_opens_after_download(project):
    at = AppTest.from_file(APP, default_timeout=120).run()
    at.multiselect(key='record_fields').set_value(['vendor']).run()
    assert not at.exception
    assert FoodOptimizer('recording').records('vendor')
    at.session_state['main_tab'] = wording.TAB_BATCH
    at.session_state['_sheets_downloaded'] = True
    at.run()
    fold = next(e for e in at.expander if e.label == wording.UPLOAD_EXPANDER)
    assert fold.proto.expanded
    assert any(wording.UPLOAD_SHEETS_ARE_BACK_CAPTION == c.value for c in at.caption)


def test_saved_copy_cleanup_keeps_newest_three_and_recent_ones(project, tmp_path):
    backend = storage.LocalStorage()
    now = datetime.now().timestamp()
    for i in range(31):
        name = backend.archive('recording', 'pre_edit', copy=True)
        # Every copy is old; the three newest must nevertheless survive.
        when = now - (7 * 86400 + 12 * 3600 + i * 60)
        os.utime(tmp_path / f'{name}.pkl', (when, when))
    at = AppTest.from_file(APP, default_timeout=120).run()
    assert not at.exception
    assert 'Older copies (30)' in [e.label for e in at.expander]
    at.button(key='delete_old_copies__btn').click().run()
    assert len(backend.list_archives()) == 31
    next(b for b in at.button if b.label == wording.YES_DELETE).click().run()
    assert not at.exception
    assert len(backend.list_archives()) == 3
    assert (tmp_path / 'recording.pkl').exists()
    with pytest.raises(storage.StorageError):
        backend.delete_archive('recording')


def test_process_only_project_has_no_batch_size_or_ingredient_columns(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer('fermentation', robust=False)
    opt.add_process_parameter('Temperature', 20, 40, unit='°C')
    opt.add_objective('pH at 6 h', 1, goal='min', min_val=0, max_val=14)
    opt.set_pending_batch([{'Temperature': 30}], batch_no=1)
    at = AppTest.from_file(APP, default_timeout=120).run()
    assert not at.exception
    assert not any('batch size' in n.label.lower() for n in at.number_input)
    book = load_workbook(io.BytesIO(opt.workbook_bytes(opt.pending_batch, None)))
    values = [c.value for sheet in book for row in sheet for c in row]
    assert wording.PERCENT_COLUMN not in values and wording.LOT_COLUMN not in values
    assert 'Actual (g)' not in values
    assert 'Total (g)' not in opt.batch_frame(opt.pending_batch).columns


def test_cleanup_never_lists_another_projects_copies(project, tmp_path):
    other = FoodOptimizer('recording_other', robust=False)
    other.add_ingredient('Water', 0, 100)
    backend = storage.LocalStorage()
    for i in range(5):
        backend.archive('recording_other', 'pre_edit', copy=True)
    at = AppTest.from_file(APP, default_timeout=120)
    at.session_state['_loaded_project'] = 'recording'
    at.run()
    assert not at.exception
    assert not any(e.label.startswith('Older copies') for e in at.expander)



def test_scaling_ingredients_does_not_scale_process_cautions(project):
    project.add_process_parameter('Mixing time', 45, 150, unit='s')
    assert project.bounds_caution('Mixing time', 60, 2.5) == ''
    assert project.bounds_caution('Mixing time', 151, 2.5)
