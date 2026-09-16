"""UI-level regression tests for app.py, using Streamlit's AppTest.

These drive the real Streamlit script headlessly, so they catch the class of
bug where the backend raises a clean ValueError but the UI fails to catch it
and shows the user a raw traceback.
"""

import html
import io
import json
import os
import pathlib
import re
import time

import openpyxl
import pandas as pd
import pyarrow as pa
import pytest
from streamlit.testing.v1 import AppTest

import storage as storage_backend
import wording
from food_bo import FoodOptimizer, ingredients_template_workbook

APP_PATH = os.path.join(os.path.dirname(__file__), "..", "app.py")


def _submit_button(at, label):
    return next(b for b in at.button if b.label == label)


# ------------------------------------------------------------------ #
#  Tab 1's two editable grids
#
#  AppTest cannot click a cell. It does not have to: `st.data_editor`'s own
#  session-state value IS the record of what was typed — which cells
#  changed, which rows were added, which were taken out — so injecting that
#  record is exactly the input the browser sends.
#
#  It has to be injected before EVERY run that should see it. AppTest
#  replays the widgets it can drive itself, and a data editor is not one of
#  them, so a run it has not set the record for hands the editor back its
#  default. `_save_grid` does both halves.
# ------------------------------------------------------------------ #
ING_GRID = "ingredient_grid"
MEAS_GRID = "measurement_grid"
PROP_GRID = "property_grid"
_GRID_SAVE_KEYS = {ING_GRID: "save_ingredient_grid",
                   MEAS_GRID: "save_measurement_grid"}


def _grid_edits(at, base, edited=None, added=None, deleted=None):
    """Type into one of tab 1's grids. The key carries the grid's own
    counter, which Discard changes and every Save that lands turn over
    (ui_helpers.grid_key)."""
    nonce_key = f"_{base}_nonce"
    nonce = at.session_state[nonce_key] if nonce_key in at.session_state else 0
    at.session_state[f"{base}_{nonce}"] = {
        "edited_rows": {int(k): v for k, v in (edited or {}).items()},
        "added_rows": list(added or []),
        "deleted_rows": list(deleted or [])}


def _stale_grid_edits(at, base, nonce, edited=None, added=None, deleted=None):
    """The record the browser is still holding for a grid that has since
    been drawn under a NEW key. Injected at a fixed nonce, so it is the
    record of the editor that was on screen before the turnover — the thing
    a turnover has to make unreachable."""
    at.session_state[f"{base}_{nonce}"] = {
        "edited_rows": {int(k): v for k, v in (edited or {}).items()},
        "added_rows": list(added or []),
        "deleted_rows": list(deleted or [])}


class _FakeUpload(io.BytesIO):
    """What st.file_uploader hands back: a readable file with a name, a size
    and an id. AppTest cannot drive the real widget."""

    def __init__(self, name, data):
        super().__init__(data)
        self.name = name
        self.size = len(data)
        self.file_id = name


def _grid_save(at, base):
    """That grid's own `Save changes`, whether it is the plain button or the
    one confirm_action draws when a row has been taken out."""
    key = _GRID_SAVE_KEYS[base]
    return next(b for b in at.button
                if b.key in (f"{key}__save", f"{key}__btn"))


def _grid_discard(at, base):
    return next(b for b in at.button
                if b.key == f"{_GRID_SAVE_KEYS[base]}__discard")


def _save_grid(at, base, confirm=False, discard=False, **edits):
    """Type into a grid and save it. `confirm` answers the question a
    deleted row is asked about first; `discard` answers the one an edit
    that would take the open round away is asked."""
    _grid_edits(at, base, **edits)
    at.run()
    _grid_save(at, base).click()
    _grid_edits(at, base, **edits)
    at.run()
    if confirm or discard:
        _submit_button(at, wording.YES_DELETE if confirm
                       else wording.YES_SAVE_AND_DISCARD).click()
        _grid_edits(at, base, **edits)
        at.run()
    return at


def _ingredient_row(name, kind=None, low=0.0, high=100.0, unit="g", **extra):
    """A row typed on the empty line at the bottom of the ingredients grid."""
    row = {wording.NAME_LABEL: name,
           wording.TYPE_LABEL: kind or wording.KIND_INGREDIENT,
           wording.LOWEST_LABEL: low, wording.HIGHEST_LABEL: high,
           wording.UNIT_LABEL: unit}
    row.update(extra)
    return row


def _measurement_row(name, goal="max", low=0.0, high=10.0, unit="",
                     share=50.0, target=None):
    return {wording.MEASUREMENT_COLUMN: name,
            wording.GOAL_LABEL: wording.GOAL_LABELS[goal],
            wording.TARGET_LABEL: target,
            wording.LOWEST_MEASURABLE_LABEL: low,
            wording.HIGHEST_MEASURABLE_LABEL: high,
            wording.UNIT_LABEL: unit,
            wording.SHARE_COLUMN: share}


def _grid_frame(at, which=0):
    """One of tab 1's grids as it is drawn. Tab 1 renders first, so the
    ingredients grid is 0, the measurements grid is 1 and the properties
    grid inside More settings — drawn only while the project has a property
    and an ingredient — is 2."""
    return at.dataframe[which].value


def _save_properties(at, **edits):
    """Type into the properties grid inside More settings and save it. Its
    button is grey and applies on the click: there is no Discard beside it
    and the foot keeps the tab's one coloured button."""
    _grid_edits(at, PROP_GRID, **edits)
    at.run()
    next(b for b in at.button if b.key == "save_properties").click()
    _grid_edits(at, PROP_GRID, **edits)
    at.run()
    return at


@pytest.fixture
def project_with_history(tmp_path, monkeypatch):
    """A project on disk (in a temp cwd) that already has one experiment,
    so mid-run rules apply. Named my_project: the app's default."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("my_project")
    opt.add_ingredient("Water", 0, 100)
    opt.add_objective("Taste", 1.0, goal="max")
    opt.tell({"Water": 50.0}, {"Taste": 7.0})
    return opt


def test_hard_reset_targets_active_project_not_typed_name(project_with_history, tmp_path):
    """Typing another name in the sidebar without clicking Create must not
    redirect Hard Reset at that other project (audit: confirmed critical bug)."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.sidebar.text_input[0].set_value("other")   # the sidebar "Project Name" box (only text input in sidebar)
    _submit_button(at, "Start this project over").click()
    at.run()
    _submit_button(at, "Yes, start over").click()
    at.run()
    assert not at.exception
    assert (tmp_path / "my_project_archived.pkl").exists()
    assert not (tmp_path / "other.pkl").exists()
    assert not (tmp_path / "other_archived.pkl").exists()
    assert [h.value for h in at.sidebar.subheader] == ["Projects", "my_project"]
    assert FoodOptimizer("my_project").X_history == []


def test_restore_requires_confirmation_and_archives_current(project_with_history, tmp_path):
    """Restore must preview, confirm, and archive the current project first.
    We can't drive st.file_uploader in AppTest, so we prime the parsed state
    the way the uploader handler does."""
    other = FoodOptimizer("donor")
    other.add_ingredient("Flour", 0, 100)
    other.add_objective("Crunch", 1.0, goal="max")
    for i in range(3):
        other.tell({"Flour": 10.0 * i}, {"Crunch": 5.0})
    donor_state = other.export_json()

    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"  # pin: "donor" now exists on
    # disk with a later mtime, which would otherwise become the auto-loaded
    # "most recent" project (Task 8) and defeat this restore-flow test.
    at.session_state["_restore_candidate"] = donor_state
    at.run()
    assert any("3 formulations" in w.value for w in at.warning), [w.value for w in at.warning]
    assert len(FoodOptimizer("my_project").X_history) == 1   # nothing changed yet
    _submit_button(at, "Yes, replace").click()
    at.run()
    assert not at.exception
    assert (tmp_path / "my_project_pre_restore.pkl").exists()
    assert len(FoodOptimizer("my_project").X_history) == 3
    assert any("Restored 3 formulations" in s.value for s in at.success)


def test_restore_rejects_wrong_shaped_variables(project_with_history):
    """validate_state must catch a backup whose 'variables' list is the right
    type but contains elements of the wrong shape, before anything is
    applied to the live optimizer."""
    candidate = project_with_history.export_json()
    candidate["variables"] = [{"nope": 1}]
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.session_state["_restore_candidate"] = candidate
    at.run()
    assert not at.exception
    assert any("wrong shape" in e.value for e in at.error), [e.value for e in at.error]
    assert len(FoodOptimizer("my_project").X_history) == 1
    assert at.session_state["optimizer"].variables[0]["name"] == "Water"


def test_restore_rejects_empty_json(project_with_history):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_restore_candidate"] = {}
    at.run()
    assert any("not a Food Optimizer copy" in e.value for e in at.error), \
        [e.value for e in at.error]
    assert len(FoodOptimizer("my_project").X_history) == 1


def test_save_failure_shows_one_banner_with_a_copy_and_reload(project_with_history):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.session_state["optimizer"].save_error = "The server could not be reached — your last change was NOT saved."
    at.run()
    assert not at.exception
    assert sum(1 for w in at.warning if "not saved" in w.value) == 1, [w.value for w in at.warning]
    labels = [b.label for b in at.button]
    assert "Reload project" in labels, labels
    # The banner says "Save a copy now, then click Reload project." and the
    # two buttons under it are the two halves of that sentence, word for
    # word — a button reading "Save a copy" sent the reader looking for a
    # third control.
    assert any(w.value == wording.SAVE_ERROR_WARNING for w in at.warning), \
        [w.value for w in at.warning]
    assert wording.SAVE_COPY_NOW == "Save a copy now"
    assert _unknown(at.main, "download_button", wording.SAVE_COPY_NOW)


def test_first_run_creates_no_project_file(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert list(tmp_path.glob("*.pkl")) == []
    assert any("Create your first project" in m.value for m in at.markdown), \
        [m.value for m in at.markdown]
    # the welcome panel already offers the sample; the sidebar must not duplicate it on a first run
    assert not any(b.label == "Try the sample project" for b in at.sidebar.button), [b.label for b in at.sidebar.button]


def test_create_project_rejects_bad_name(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.text_input(key="new_project_name").set_value("bad/name")
    _submit_button(at, "Create project").click()
    at.run()
    assert any("Up to 64 characters" in e.value for e in at.error)
    assert list(tmp_path.glob("*.pkl")) == []


def test_create_project_writes_file_and_opens_it(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.text_input(key="new_project_name").set_value("Oat cookie v2")
    _submit_button(at, "Create project").click()
    at.run()
    assert (tmp_path / "Oat cookie v2.pkl").exists()
    assert any("Created Oat cookie v2" in s.value for s in at.success)


def test_returning_user_lands_in_most_recent_project(project_with_history):
    older = FoodOptimizer("older")
    older.add_ingredient("Sugar", 0, 50)
    time.sleep(0.01)
    project_with_history.save()   # touch my_project so it is newest on disk
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert "my_project" in [h.value for h in at.sidebar.subheader]
    assert "older" not in [h.value for h in at.sidebar.subheader]


def test_pending_confirm_is_cleared_on_project_switch(project_with_history, tmp_path):
    other = FoodOptimizer("second")
    other.add_ingredient("Flour", 0, 100)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"  # pin: "second" now exists on
    # disk with a later mtime, which would otherwise become the auto-loaded "most
    # recent" project (Task 8) before this test ever switches to it on purpose.
    at.run()
    _submit_button(at, "Start this project over").click()   # arm the confirmation
    at.run()
    assert any(b.label == "Yes, start over" for b in at.button)
    # Choosing a project reruns the script and enables Open; only then can a
    # user click it (Open is greyed out while the box shows the open project).
    at.selectbox(key="project_select").set_value("second")
    at.run()
    _submit_button(at, "Open").click()
    at.run()
    # Open is drawn at the foot of the sidebar, so the armed confirmation above
    # it is still in the aborted pass's element tree; the next plain run shows
    # what the browser would show.
    at.run()
    assert not at.exception
    assert not any(b.label == "Yes, start over" for b in at.button), [b.label for b in at.button]
    assert "second" in [h.value for h in at.sidebar.subheader]


def test_open_button_lights_up_only_when_it_would_switch(project_with_history, tmp_path):
    """Button rule: the button that moves you forward is coloured, and a
    button that would do nothing is greyed out. Open must be disabled while
    the box shows the project already on screen, and primary once a
    different project is chosen."""
    FoodOptimizer("second").add_ingredient("Flour", 0, 100)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    open_btn = _submit_button(at.sidebar, "Open")
    assert open_btn.disabled and open_btn.proto.type == "secondary"

    at.selectbox(key="project_select").set_value("second")
    at.run()
    open_btn = _submit_button(at.sidebar, "Open")
    assert not open_btn.disabled and open_btn.proto.type == "primary"

    open_btn.click()
    at.run()
    assert "second" in [h.value for h in at.sidebar.subheader]
    open_btn = _submit_button(at.sidebar, "Open")
    assert open_btn.disabled


def test_delete_project_confirms_archives_and_leaves_the_list(project_with_history, tmp_path):
    """Delete project must ask first, rename the file to an archive copy
    rather than erase it, drop the project from the list, and land on the
    welcome panel when nothing is left to open."""
    from storage import LocalStorage
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at.sidebar, "Delete this project").click()
    at.run()
    assert (tmp_path / "my_project.pkl").exists()               # not yet
    warn = next(w.value for w in at.warning if "my_project" in w.value)
    # Real grammar, never "1 formulation(s)".
    assert "its 1 formulation?" in warn and "(s)" not in warn, warn
    _submit_button(at.sidebar, "Yes, delete it").click()
    at.run()
    assert not at.exception
    assert not (tmp_path / "my_project.pkl").exists()
    assert (tmp_path / "my_project_deleted.pkl").exists()
    assert LocalStorage().list_projects() == []
    assert "my_project_deleted" in LocalStorage().list_archives()
    assert any("Deleted my_project" in i.value for i in at.info), [i.value for i in at.info]
    assert any("Create your first project" in m.value for m in at.markdown)


def test_delete_project_opens_the_most_recent_remaining_project(project_with_history, tmp_path):
    other = FoodOptimizer("second")
    other.add_ingredient("Flour", 0, 100)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    _submit_button(at.sidebar, "Delete this project").click()
    at.run()
    _submit_button(at.sidebar, "Yes, delete it").click()
    at.run()
    assert not at.exception
    assert "second" in [h.value for h in at.sidebar.subheader]
    assert at.selectbox(key="project_select").options == ["second"]


def test_lower_sidebar_heading_names_the_open_project_and_follows_switches(project_with_history):
    """Backup, restore, reset and delete act on the open project, so the
    heading above them carries its name and follows Open, not the box."""
    FoodOptimizer("second").add_ingredient("Flour", 0, 100)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    assert "my_project" in [h.value for h in at.sidebar.subheader]

    at.selectbox(key="project_select").set_value("second")   # chosen, not opened
    at.run()
    assert "my_project" in [h.value for h in at.sidebar.subheader]
    assert "second" not in [h.value for h in at.sidebar.subheader]

    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert "second" in [h.value for h in at.sidebar.subheader]
    assert "my_project" not in [h.value for h in at.sidebar.subheader]


def test_sample_project_button_creates_ready_project(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    # The sidebar also offers "Try the sample project", so target the
    # welcome-panel button explicitly (scoped to at.main).
    _submit_button(at.main, "Try the sample project").click()
    at.run()
    assert not at.exception
    opt = FoodOptimizer("Sample project")
    # Eight ingredients: enough for a real formulation, few enough to read at
    # a glance in the batch table and the printable sheets.
    assert len(opt.variables) == 8
    assert opt.amount_unit == "g"
    # The sample is the plant-based burger brief for a trained panel: two
    # intensity scores with targets (7 and 6 out of 10), firmness the more
    # important of the two.
    assert [o["name"] for o in opt.objectives] == ["Juiciness", "Firmness"]
    assert all(o["goal"] == "target" for o in opt.objectives)
    assert {o["name"]: o["target"] for o in opt.objectives} == {"Juiciness": 7,
                                                                "Firmness": 6}
    assert {o["name"]: o["unit"] for o in opt.objectives} == {"Juiciness": "/10",
                                                               "Firmness": "/10"}
    weights = {o["name"]: o["weight"] for o in opt.objectives}
    assert weights["Firmness"] > weights["Juiciness"]


def test_the_sample_lists_firmness_first(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at.main, "Try the sample project").click()
    at.run()
    # Set up is complete, but nothing has been made yet: the sample opens on
    # its set-up so the user sees what they are about to make.
    assert at.session_state["main_tab"] == wording.TAB_SETUP
    grid = _grid_frame(at, 1)
    # A panel score's "/10" is a column of its own on the grid: the name
    # cell is the name, because it is the cell a rename would be typed in.
    assert list(grid["Measurement"]) == ["Firmness", "Juiciness"]
    assert list(grid[wording.UNIT_LABEL]) == ["/10", "/10"]
    assert list(grid["Goal"]) == [wording.GOAL_LABELS["target"]] * 2
    assert list(grid["Target"]) == [6.0, 7.0]


def test_the_sample_names_its_targets_and_welcomes_the_first_visit(tmp_path, monkeypatch):
    """The sample sets its own targets_source and shows the caption it
    produces, plus the two-line welcome under the tab-1 title -- both gone
    once the sample is not what is open."""
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at.main, "Try the sample project").click()
    at.run()
    assert not at.exception
    opt = FoodOptimizer("Sample project")
    assert opt.targets_source == wording.SAMPLE_TARGETS_SOURCE
    captions = [c.value for c in at.tabs[0].caption]
    assert wording.SAMPLE_TAB1_DESCRIPTION in captions
    assert (wording.targets_from_caption(wording.SAMPLE_TARGETS_SOURCE)
            in captions)


def test_the_sample_welcome_is_gone_once_a_formulation_is_scored(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    pre = FoodOptimizer("Sample project")
    pre.add_ingredient("Water", 0, 100)
    pre.add_objective("Taste", 1.0, goal="max")
    pre.set_targets_source(wording.SAMPLE_TARGETS_SOURCE)
    pre.tell({"Water": 50.0}, {"Taste": 7.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "Sample project"
    at.session_state["_land_on_open"] = True
    at.run()
    assert not at.exception
    captions = [c.value for c in at.tabs[0].caption]
    assert wording.SAMPLE_TAB1_DESCRIPTION not in captions
    # The targets_source caption is unrelated to X_history and stays.
    assert (wording.targets_from_caption(wording.SAMPLE_TARGETS_SOURCE)
            in captions)


def test_the_sample_welcome_is_gone_once_a_formulation_is_not_scored(
        tmp_path, monkeypatch):
    """A batch whose only row was generated and then ticked Not scored has
    still been made -- X_history stays empty, but 'Next: make a batch' would
    be wrong about it."""
    monkeypatch.chdir(tmp_path)
    pre = FoodOptimizer("Sample project")
    pre.add_ingredient("Water", 0, 100)
    pre.add_objective("Taste", 1.0, goal="max")
    pre.record_skipped(1, 1, {"Water": 50.0})
    assert pre.X_history == [] and pre.skipped
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "Sample project"
    at.session_state["_land_on_open"] = True
    at.run()
    assert not at.exception
    captions = [c.value for c in at.tabs[0].caption]
    assert wording.SAMPLE_TAB1_DESCRIPTION not in captions


def test_the_welcome_never_shows_on_a_project_of_the_users_own(burger):
    """Only the sample gets the welcome line, matched by name -- a project
    the user named 'Sample project' of their own would be an odd coincidence,
    but the check is the one the app has always used to recognise its own
    sample."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    captions = [c.value for c in at.tabs[0].caption]
    assert wording.SAMPLE_TAB1_DESCRIPTION not in captions


def test_the_targets_source_button_opens_a_prefilled_box_and_saves(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert not any(c.value.startswith("Targets from:")
                  for c in at.tabs[0].caption)
    # With no note yet the button says what it will make, not "edit" of a
    # note nobody can see.
    _submit_button(at, wording.ADD_TARGETS_SOURCE_BUTTON).click()
    at.run()
    box = at.text_input(key="targets_source_box")
    assert box.value == ""   # nothing set yet: opens blank
    box.set_value("Benchmark burger, panel of 8.")
    _submit_button(at, "Save").click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").targets_source == "Benchmark burger, panel of 8."
    captions = [c.value for c in at.tabs[0].caption]
    assert "Targets from: Benchmark burger, panel of 8." in captions
    # The box reopens prefilled with what is stored, next time it is opened,
    # and the button is now an Edit.
    _submit_button(at, wording.TARGETS_SOURCE_BUTTON).click()
    at.run()
    assert (at.text_input(key="targets_source_box").value
            == "Benchmark burger, panel of 8.")


def test_the_targets_source_box_can_be_cancelled_without_saving(burger):
    """Cancel closes the box and throws away what was typed -- the same word
    and the same act as the measurement editor's own Cancel."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at, wording.ADD_TARGETS_SOURCE_BUTTON).click()
    at.run()
    at.text_input(key="targets_source_box").set_value("Half-typed note")
    _submit_button(at, "Cancel").click()
    at.run()
    assert not at.exception
    assert "_targets_source_open" not in at.session_state
    assert FoodOptimizer("burger").targets_source == ""
    # Reopening starts blank again, not from the discarded text.
    _submit_button(at, wording.ADD_TARGETS_SOURCE_BUTTON).click()
    at.run()
    assert at.text_input(key="targets_source_box").value == ""


def test_the_targets_source_save_is_never_the_lit_button(burger):
    """One lit button per tab: the foot's Continue stays it, whether or not
    this box is open."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at, wording.ADD_TARGETS_SOURCE_BUTTON).click()
    at.run()
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON], \
        _tab_primaries(at, 0)
    assert not _submit_button(at, "Save").disabled


def test_targets_source_box_is_cleared_on_a_project_switch(tmp_path, monkeypatch):
    """A half-typed note belongs to the project it was typed in -- the same
    rule every other set-up box on this tab already follows."""
    monkeypatch.chdir(tmp_path)
    first = FoodOptimizer("first")
    first.add_objective("Taste", 1.0, goal="max")
    second = FoodOptimizer("second")
    second.add_objective("Taste", 1.0, goal="max")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "first"
    at.run()
    _submit_button(at, wording.ADD_TARGETS_SOURCE_BUTTON).click()
    at.run()
    at.text_input(key="targets_source_box").set_value("Half-typed note")
    at.run()
    at.selectbox(key="project_select").set_value("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    # The box itself closes on a project switch, same as every other set-up
    # form -- and what it would open on next is the parked empty value, not
    # the half-typed note left behind in "first".
    assert "_targets_source_open" not in at.session_state
    _submit_button(at, wording.ADD_TARGETS_SOURCE_BUTTON).click()
    at.run()
    assert at.text_input(key="targets_source_box").value == ""


def test_sample_project_button_reopens_existing_without_recreating(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    # Pre-create "Sample project" with an experiment, simulating a shared host
    # where a previous visitor already ran the sample project.
    pre = FoodOptimizer("Sample project")
    pre.add_ingredient("Water", 0, 100)
    pre.add_objective("Taste", 1.0, goal="max")
    pre.tell({"Water": 50.0}, {"Taste": 7.0})

    # app.py never auto-opens the most-recently-used project on a shared host
    # (os.path.isdir("/mount/src")); force that guard so the welcome panel
    # renders here even though "Sample project" already exists on disk, which
    # is exactly the scenario where clicking the sample button again must
    # reopen the existing project instead of re-creating it.
    real_isdir = os.path.isdir
    monkeypatch.setattr(
        os.path, "isdir",
        lambda p: True if p == "/mount/src" else real_isdir(p),
    )

    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert any("Create your first project" in m.value for m in at.markdown)

    # Scope to the welcome-panel button explicitly: the sidebar now also has
    # a same-labelled button.
    _submit_button(at.main, "Try the sample project").click()
    at.run()
    assert not at.exception

    assert len(FoodOptimizer("Sample project").X_history) == 1
    assert "Sample project" in [h.value for h in at.sidebar.subheader]


def test_sidebar_sample_project_button_available_with_project_open(project_with_history):
    """The sample project must be reachable from the sidebar even once a
    project already exists (the welcome panel is gone at that point)."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception

    _submit_button(at.sidebar, "Try the sample project").click()
    at.run()
    assert not at.exception

    opt = FoodOptimizer("Sample project")
    assert len(opt.variables) >= 3
    assert "Sample project" in [h.value for h in at.sidebar.subheader]


def test_sidebar_sample_project_button_reopens_without_recreating(project_with_history):
    """Clicking the sidebar sample button when the sample already exists
    must reopen it, not recreate it (experiment count unchanged)."""
    pre = FoodOptimizer("Sample project")
    pre.add_ingredient("Water", 0, 100)
    pre.add_objective("Taste", 1.0, goal="max")
    pre.tell({"Water": 50.0}, {"Taste": 7.0})

    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception

    _submit_button(at.sidebar, "Try the sample project").click()
    at.run()
    assert not at.exception

    assert len(FoodOptimizer("Sample project").X_history) == 1


def test_restore_warns_when_another_window_edited_project(project_with_history, tmp_path):
    """Restore's 'Yes, replace' constructs a fresh FoodOptimizer against the
    current file, which would re-stamp the shared conflict tracker and hide a
    concurrent edit. The stale check must run against what THIS session
    loaded, before that reconstruction, and must leave the restore candidate
    in place so the user can Reload and retry."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()

    # Another window edits and saves the project after this session loaded it.
    other = FoodOptimizer("my_project")
    other.add_objective("Crunch", 1.0)
    time.sleep(0.01)  # coarse-mtime filesystems

    donor = FoodOptimizer("donor")
    donor.add_ingredient("Flour", 0, 100)
    donor.add_objective("Sweetness", 1.0)
    donor.tell({"Flour": 10.0}, {"Sweetness": 5.0})

    at.session_state["_restore_candidate"] = donor.export_json()
    at.run()
    _submit_button(at, "Yes, replace").click()
    at.run()

    assert not at.exception
    assert any("another window" in e.value for e in at.error), [e.value for e in at.error]
    reloaded = FoodOptimizer("my_project")
    assert len(reloaded.X_history) == 1   # original single experiment untouched
    assert [o["name"] for o in reloaded.objectives] == ["Taste", "Crunch"]


def _unknown(node, kind, label):
    """AppTest has no accessor for st.download_button or st.file_uploader: both
    arrive as UnknownElement carrying the raw proto. Walk the element tree and
    return the first one of `kind` with this label. For a download button
    proto.type is 'primary'/'secondary'; for an uploader it is the list of
    accepted extensions."""
    def walk(n):
        children = getattr(n, "children", None) or {}
        if hasattr(children, "values"):
            children = children.values()
        for child in children:
            if (type(child).__name__ == "UnknownElement"
                    and getattr(child, "type", None) == kind
                    and getattr(child, "label", None) == label):
                yield child
            yield from walk(child)
    return next(walk(node))


def _unknowns(node, kind):
    """Every UnknownElement of `kind` under `node`, in render order."""
    def walk(n):
        children = getattr(n, "children", None) or {}
        if hasattr(children, "values"):
            children = children.values()
        for child in children:
            if (type(child).__name__ == "UnknownElement"
                    and getattr(child, "type", None) == kind):
                yield child
            yield from walk(child)
    return list(walk(node))


def _tab_primaries(at, index):
    """The coloured ACTIONS inside one tab — st.button and st.download_button
    alike, because a lit download is as much the next action as a lit button.
    AppTest renders every tab's body on every run, so 'one primary per tab' can
    only be checked tab by tab: at.tabs[0] is Set up, [1] is Make a batch, [2]
    is Results."""
    tab = at.tabs[index]
    lit = [b.label for b in tab.button if b.proto.type == "primary"]
    lit += [d.label for d in _unknowns(tab, "download_button")
            if d.proto.type == "primary"]
    return lit


def test_tabs_are_the_loop_in_order(project_with_history):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert [t.label for t in at.tabs] == [wording.TAB_SETUP, wording.TAB_BATCH, wording.TAB_RESULTS]


def test_opening_a_project_lands_on_set_up_when_it_is_incomplete(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    FoodOptimizer("bare")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "bare"
    at.session_state["_land_on_open"] = True
    at.run()
    assert at.session_state["main_tab"] == wording.TAB_SETUP


def test_opening_a_project_lands_on_the_trial_when_one_is_open(project_with_history):
    project_with_history.set_pending_batch([{"Water": 10.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.session_state["_land_on_open"] = True
    at.run()
    assert at.session_state["main_tab"] == wording.TAB_BATCH


def test_opening_a_project_lands_on_results(project_with_history):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.session_state["_land_on_open"] = True
    at.run()
    assert at.session_state["main_tab"] == wording.TAB_RESULTS


def test_opening_a_ready_project_with_no_results_lands_on_set_up(tmp_path,
                                                                 monkeypatch):
    """Results would be an empty tab. The set-up is what there is to look at,
    and confirming it is the step before making anything."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("ready")
    opt.add_ingredient("Water", 0, 100)
    opt.add_objective("Taste", 1.0, goal="max")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "ready"
    at.session_state["_land_on_open"] = True
    at.run()
    assert at.session_state["main_tab"] == wording.TAB_SETUP


def test_a_project_whose_only_formulation_was_left_out_lands_on_results(
        tmp_path, monkeypatch):
    """A left-out formulation is a result the project holds: it has a number,
    amounts and a note, and tab 3 lists it."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("left_out")
    opt.add_ingredient("Water", 0, 100)
    opt.add_objective("Taste", 1.0, goal="max")
    opt.record_skipped(1, 1, {"Water": 10.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "left_out"
    at.session_state["_land_on_open"] = True
    at.run()
    assert at.session_state["main_tab"] == wording.TAB_RESULTS


def test_a_plain_rerun_never_moves_the_tab(project_with_history):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.session_state["main_tab"] = wording.TAB_SETUP
    at.run()
    at.run()
    assert at.session_state["main_tab"] == wording.TAB_SETUP


def test_the_sidebar_carries_no_coloured_button_except_open(project_with_history, tmp_path):
    """The lit thing on screen must be the tab's next action, never a sidebar
    control that would do nothing."""
    FoodOptimizer("second").add_ingredient("Flour", 0, 100)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    lit = [b.label for b in at.sidebar.button if b.proto.type == "primary"]
    assert lit == [], lit
    open_button = _submit_button(at.sidebar, "Open")
    assert open_button.disabled and open_button.proto.type == "secondary"
    at.selectbox(key="project_select").set_value("second")
    at.run()
    open_button = _submit_button(at.sidebar, "Open")
    assert not open_button.disabled and open_button.proto.type == "primary"
    # ...and it steps aside the moment a confirmation is armed: the Yes is the
    # one lit thing in the sidebar, even with a switch waiting to be made.
    _submit_button(at.sidebar, "Start this project over").click()
    at.run()
    lit = [b.label for b in at.sidebar.button if b.proto.type == "primary"]
    assert lit == ["Yes, start over"], lit
    open_button = _submit_button(at.sidebar, "Open")
    assert not open_button.disabled and open_button.proto.type == "secondary"


def test_arming_one_manage_project_confirmation_greys_the_other(project_with_history):
    """Two armed confirmations would put two coloured Yes buttons in the
    sidebar, each quietly keeping its own copy of the project."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    assert not at.button(key="hard_reset__btn").disabled
    assert not at.button(key="delete_project__btn").disabled
    _submit_button(at.sidebar, "Start this project over").click()
    at.run()
    assert at.button(key="delete_project__btn").disabled
    assert not at.button(key="hard_reset__btn").disabled
    _submit_button(at.sidebar, "Cancel").click()
    at.run()
    at.run()
    _submit_button(at.sidebar, "Delete this project").click()
    at.run()
    # Delete renders below Hard reset, so it can only grey the button above it
    # on the next run; a click in between is ignored by confirm_action anyway.
    at.run()
    assert at.button(key="hard_reset__btn").disabled
    assert not at.button(key="delete_project__btn").disabled


def test_an_armed_confirmation_is_the_one_lit_sidebar_button(project_with_history):
    """A confirmation button is transient and IS the next action while it is
    armed, so exactly one thing is lit — and only until it is answered."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    assert [b.label for b in at.sidebar.button if b.proto.type == "primary"] == []
    _submit_button(at.sidebar, "Start this project over").click()
    at.run()
    lit = [b.label for b in at.sidebar.button if b.proto.type == "primary"]
    assert lit == ["Yes, start over"], lit
    _submit_button(at.sidebar, "Cancel").click()
    at.run()
    assert at.session_state["hard_reset__pending"] is False
    # Cancel ends in st.rerun(), and AppTest leaves the aborted pass's buttons
    # in the element tree alongside the settled ones; the next plain run shows
    # what the browser would show.
    at.run()
    assert [b.label for b in at.sidebar.button if b.proto.type == "primary"] == []


def test_the_restore_confirmation_is_lit_like_every_other_confirmation(project_with_history):
    donor = FoodOptimizer("donor")
    donor.add_ingredient("Flour", 0, 100)
    donor.add_objective("Crunch", 1.0, goal="max")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.session_state["_restore_candidate"] = donor.export_json()
    at.run()
    lit = [b.label for b in at.sidebar.button if b.proto.type == "primary"]
    assert lit == ["Yes, replace"], lit


def test_removing_every_measurement_keeps_the_open_batch(project_with_history):
    """Measurements never discard a batch: a formulation is amounts, and it can
    still be made while the user reworks what they will score."""
    project_with_history.set_pending_batch([{"Water": 10.0}, {"Water": 20.0}])
    project_with_history.remove_objective("Taste")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    assert not at.exception
    assert FoodOptimizer("my_project").pending_batch is not None
    assert not any("ingredient list" in i.value for i in at.info), [i.value for i in at.info]


def test_restore_says_nothing_about_a_copy_when_none_was_kept(project_with_history, monkeypatch):
    """'kept as None' is not a sentence. With no file to archive, the message
    simply stops after naming what was restored."""
    from storage import LocalStorage
    monkeypatch.setattr(LocalStorage, "archive", lambda self, *a, **k: None)
    donor = FoodOptimizer("donor")
    donor.add_ingredient("Water", 0, 100)
    donor.add_objective("Taste", 1.0, goal="max")
    for value in (10.0, 20.0, 30.0):
        donor.tell({"Water": value}, {"Taste": 5.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.session_state["_restore_candidate"] = donor.export_json()
    at.run()
    _submit_button(at, "Yes, replace").click()
    at.run()
    assert not at.exception
    note = next(s.value for s in at.success if "Restored" in s.value)
    assert note == "Restored 3 formulations into my_project.", note
    assert "None" not in note


def test_the_saved_line_shows_for_a_project_opened_but_not_yet_saved(project_with_history):
    """last_saved_at is only set by a save in THIS session, so a returning user
    would otherwise see no saved line at all until their first edit."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    assert at.session_state["optimizer"].last_saved_at is None
    assert any("Saved automatically at " in c.value for c in at.sidebar.caption), \
        [c.value for c in at.sidebar.caption]


def test_manage_project_holds_empty_and_delete(project_with_history):
    """Two names for two outcomes: one empties the project, the other takes it
    out of the list. 'Hard reset' named the mechanism and read as the worse of
    the two."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(e.label == "Manage project" for e in at.sidebar.expander), \
        [e.label for e in at.sidebar.expander]
    labels = [b.label for b in at.sidebar.button]
    assert "Start this project over" in labels and "Delete this project" in labels, \
        labels


def test_restore_accepts_any_file_name(project_with_history):
    """Contents decide, not the extension: the uploader must not filter."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    restore = _unknown(at.sidebar, "file_uploader", "Open a saved copy")
    assert list(restore.proto.type) == []


def test_restoring_a_real_0_2_x_backup_numbers_it_on_the_way_in(project_with_history, tmp_path):
    """The whole compatibility promise, exercised through the sidebar the user
    actually clicks, not through import_json."""
    donor = FoodOptimizer("donor")
    donor.add_ingredient("Water", 0, 100)
    donor.add_objective("Taste", 1.0, goal="max")
    for value in (10.0, 20.0, 30.0):
        donor.tell({"Water": value}, {"Taste": 5.0})
    legacy = donor.export_json()
    for key in ("formulation_ids", "batch_history", "notes_history", "skipped",
                "next_formulation_no", "pending_batch_no",
                "pending_batch_created", "pending_batch_discarded"):
        legacy.pop(key, None)
    legacy["CLASS_VERSION"] = 6

    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.session_state["_restore_candidate"] = legacy
    at.run()
    _submit_button(at, "Yes, replace").click()
    at.run()
    assert not at.exception
    restored = FoodOptimizer("my_project")
    assert restored.formulation_ids == [1, 2, 3]
    assert restored.batch_history == [None, None, None]
    assert restored.next_formulation_no == 4


def test_a_batch_that_no_longer_fits_the_ingredients_is_discarded_with_a_notice(project_with_history):
    """Replaces the 0.2.x stale-batch test: a batch saved before the ingredient
    list changed cannot be made, and must not sit there looking makeable."""
    project_with_history.set_pending_batch([{"Water": 10.0, "Ghost": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    assert not at.exception
    assert FoodOptimizer("my_project").pending_batch is None
    assert any("your set-up changed" in i.value for i in at.info), \
        [i.value for i in at.info]


def test_the_trial_line_counts_what_is_left_to_make(project_with_history):
    project_with_history.set_pending_batch([{"Water": 10.0}, {"Water": 20.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    assert any(c.value == wording.batch_line_open(1, 2) for c in at.caption), \
        [c.value for c in at.caption]


def test_an_uploaded_import_sheet_does_not_survive_a_hard_reset(project_with_history):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.session_state["_import_rows"] = pd.DataFrame({"Water": [1.0], "Taste": [5.0]})
    at.run()
    assert "_import_rows" in at.session_state
    _submit_button(at, "Start this project over").click()
    at.run()
    _submit_button(at, "Yes, start over").click()
    at.run()
    assert not at.exception
    assert "_import_rows" not in at.session_state


def test_an_uploaded_import_sheet_does_not_follow_a_project_switch(project_with_history, tmp_path):
    FoodOptimizer("second").add_ingredient("Flour", 0, 100)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.session_state["_import_rows"] = pd.DataFrame({"Water": [1.0], "Taste": [5.0]})
    at.run()
    at.selectbox(key="project_select").set_value("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert "_import_rows" not in at.session_state


@pytest.fixture
def burger(tmp_path, monkeypatch):
    """A project shaped like the sample: units everywhere, two targets,
    firmness the more important of the two."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("burger")
    opt.set_amount_unit("g")
    opt.add_ingredient("Pea protein", 0, 25)
    opt.add_ingredient("Methylcellulose", 0, 3)
    opt.add_objective("Juiciness", 1.0, goal="target", target=7,
                      min_val=0, max_val=10, unit="/10")
    opt.add_objective("Firmness", 1.5, goal="target", target=6,
                      min_val=0, max_val=10, unit="N")
    return opt


def _labels(at):
    return [b.label for b in at.button]
def test_a_unit_typed_with_a_trailing_space_is_saved_once(burger):
    """The unit is stripped before it is stored, and a plain rerun after it
    writes nothing: a re-save on every rerun bounced the script into a loop."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.UNIT_LABEL: "kg "}})
    assert FoodOptimizer("burger").unit_of("Methylcellulose") == "kg"
    saved_at = at.session_state["optimizer"].last_saved_at
    at.run()
    assert at.session_state["optimizer"].last_saved_at == saved_at   # no re-save
    assert not at.exception


def test_measurements_table_is_sorted_by_importance(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    grid = _grid_frame(at, 1)
    # No Priority column: it was the row's position in a grid already sorted
    # by what each measurement is worth, which is the same fact twice. And
    # no Importance column: the share IS the number that is typed now.
    assert [c for c in grid.columns if c != "_id"] == [
        "Measurement", "Goal", "Target", "Lowest measurable",
        "Highest measurable", "Unit", wording.SHARE_COLUMN]
    assert list(grid["Measurement"]) == ["Firmness", "Juiciness"]
    assert list(grid["Goal"]) == [wording.GOAL_LABELS["target"]] * 2
    assert list(grid["Target"]) == [6.0, 7.0]
    assert list(grid[wording.SHARE_COLUMN]) == [60.0, 40.0]


def test_the_score_function_is_written_out_under_the_table(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(c.value == (
        "Overall score = 60 % × Firmness closeness + 40 % × "
        "Juiciness closeness. A formulation that hits every goal scores "
        "100."
    ) for c in at.caption), [c.value for c in at.caption]
    # How closeness works is said once, in the expander, per goal.
    assert not any("falls evenly with distance" in c.value for c in at.caption)
def test_the_share_column_is_whole_percents_that_add_up_to_a_hundred(burger):
    """Share of score is what the reader types (spec 1.3). There is no
    Importance box anywhere, and no slider: the column adds up to 100 and
    the app derives the importances from it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    shares = list(_grid_frame(at, 1)[wording.SHARE_COLUMN])
    assert shares == [60.0, 40.0] and sum(shares) == 100
    assert not at.slider, [s.label for s in at.slider]
    assert not any("Importance" in str(c) for c in _grid_frame(at, 1).columns)


def test_no_share_line_follows_what_is_typed(burger):
    """w / Σw is not a measurement's influence on the overall score: a target
    goal's closeness never spans the full 0 to 1, so two equally weighted
    measurements are not equally influential. The line is gone."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not any("% of the overall score" in m.value for m in at.markdown), \
        [m.value for m in at.markdown]


def test_continue_is_grey_and_names_what_is_missing(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("bare2")
    opt.add_ingredient("Water", 0, 100)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    button = _submit_button(at, wording.NEXT_MAKE_BATCH_BUTTON)
    assert button.disabled
    assert button.proto.type == "secondary"
    assert any(c.value == "Add at least one measurement." for c in at.caption), \
        [c.value for c in at.caption]


def test_continue_is_lit_and_moves_to_the_trial_tab(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    button = _submit_button(at, wording.NEXT_MAKE_BATCH_BUTTON)
    assert not button.disabled and button.proto.type == "primary"
    button.click()
    at.run()
    assert not at.exception
    assert at.session_state["main_tab"] == wording.TAB_BATCH


def test_set_up_has_exactly_one_coloured_button(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON], _tab_primaries(at, 0)


def test_an_armed_confirmation_takes_the_colour_off_continue(burger):
    """Two lit buttons ask two questions at once. While a confirmation is on
    screen, answering it is the only coloured thing to do."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    edits = dict(deleted=[0])
    _grid_edits(at, MEAS_GRID, **edits)
    at.run()
    _grid_save(at, MEAS_GRID).click()
    _grid_edits(at, MEAS_GRID, **edits)
    at.run()
    assert _tab_primaries(at, 0) == ["Yes, delete"], _tab_primaries(at, 0)
    assert _submit_button(at, wording.NEXT_MAKE_BATCH_BUTTON).disabled
    _submit_button(at, "Cancel").click()
    _grid_edits(at, MEAS_GRID, **edits)
    at.run()
    at.run()   # Cancel reruns from inside the handler; AppTest keeps the
               # aborted pass's elements until the next run
    assert {o["name"] for o in FoodOptimizer("burger").objectives} == {
        "Firmness", "Juiciness"}


def test_a_measurements_name_typed_over_is_a_rename(burger):
    """Every result already recorded is filed under the name, so the rename
    moves the two together — nothing is rescored and no copy is kept."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    before = list(FoodOptimizer("burger").Y_history)
    _save_grid(at, MEAS_GRID,
               edited={0: {wording.MEASUREMENT_COLUMN: "Bite"}})
    assert not at.exception
    assert [e.value for e in at.error] == []
    saved = FoodOptimizer("burger")
    assert {o["name"] for o in saved.objectives} == {"Bite", "Juiciness"}
    assert saved.results_history[0] == {"Bite": 6.0, "Juiciness": 7.0}
    assert saved.Y_history == before
    assert any(s.value == "Bite saved." for s in at.success), \
        [s.value for s in at.success]


def test_a_measurement_renamed_onto_a_taken_name_is_refused_by_row(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, MEAS_GRID,
               edited={0: {wording.MEASUREMENT_COLUMN: "Pea protein"}})
    assert [e.value for e in at.error] == [wording.row_error(
        1, "Pea protein is already the name of an ingredient. Choose "
           "another name.")]
    assert {o["name"] for o in FoodOptimizer("burger").objectives} == {
        "Firmness", "Juiciness"}


def test_changing_a_share_keeps_a_copy_and_reports_the_move(burger, tmp_path):
    # Chosen so Formulation 3 wins while Firmness is worth 60 % and
    # Formulation 7 wins once it is worth 67 % — the move the sentence
    # reports.
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 1.0}, formulation_no=3, batch_no=1)
    burger.tell({"Pea protein": 20.0, "Methylcellulose": 2.0},
                {"Juiciness": 0.0, "Firmness": 7.0}, formulation_no=7, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_SETUP
    at.run()
    _save_grid(at, MEAS_GRID, edited={0: {wording.SHARE_COLUMN: 67.0}})
    assert not at.exception
    assert (tmp_path / "burger_pre_edit.pkl").exists()
    # The copy is named in the sentence: this save has no confirmation
    # before it, so nothing else tells the user there is a way back.
    assert any(wording.best_moved(3, 7) in v.value and wording.COPY_KEPT
               in v.value and wording.RECALCULATED_SUFFIX.strip() in v.value
               for v in at.success), [v.value for v in at.success]
    assert at.session_state["main_tab"] == wording.TAB_SETUP   # a set-up edit never moves


def test_editing_a_measurement_leaves_the_open_batch_alone(burger):
    """A formulation is a set of amounts; changing what you measure cannot
    invalidate it, and silently retiring its numbers would break the promise
    that a number is never reissued."""
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, MEAS_GRID, edited={0: {wording.SHARE_COLUMN: 70.0}})
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert [r["formulation"] for r in reloaded.pending_batch] == [1]
    assert reloaded.pending_batch_no == 1


def test_adding_and_deleting_a_measurement_leaves_the_open_trial_alone(burger):
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, MEAS_GRID, added=[_measurement_row("Chewiness", share=20.0)])
    assert FoodOptimizer("burger").pending_batch is not None
    _save_grid(at, MEAS_GRID, confirm=True, deleted=[2])
    assert not at.exception
    assert [o["name"] for o in FoodOptimizer("burger").objectives
            if o["name"] == "Chewiness"] == []
    assert FoodOptimizer("burger").pending_batch is not None


def test_adding_an_ingredient_does_discard_the_open_trial(burger):
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID, discard=True,
               added=[_ingredient_row("Beet juice powder", high=2.0)])
    assert not at.exception
    assert FoodOptimizer("burger").pending_batch is None
    # food_bo drops the batch inside add_ingredient, so app.py's makeability
    # check never sees a mismatch: the handler has to say so itself.
    assert any(wording.batch_discarded_notice(1) == i.value
               for i in at.info), [i.value for i in at.info]


def test_deleting_a_measurement_keeps_a_copy_and_recalculates(burger, tmp_path):
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 1.0}, formulation_no=1, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, MEAS_GRID, confirm=True, deleted=[1])
    assert not at.exception
    assert (tmp_path / "burger_pre_edit.pkl").exists()
    assert [o["name"] for o in FoodOptimizer("burger").objectives] == ["Firmness"]
    assert any("Every overall score was recalculated" in s.value for s in at.success), \
        [s.value for s in at.success]


def test_a_measurement_is_added_on_the_empty_row_at_the_bottom(burger):
    """Its share is what it is worth: 20 % of the score, and the two already
    there give way proportionally to make room for it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, MEAS_GRID, added=[
        _measurement_row("Chewiness", unit="N", high=20.0, share=20.0)])
    assert not at.exception
    opt = FoodOptimizer("burger")
    added = next(o for o in opt.objectives if o["name"] == "Chewiness")
    assert added["unit"] == "N" and added["max_val"] == 20.0
    assert opt.share_percents() == {"Firmness": 48, "Juiciness": 32,
                                    "Chewiness": 20}


def test_the_ingredient_uploader_help_names_the_real_columns(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    uploader = _unknown(at.main, "file_uploader", "Upload ingredients (Excel or CSV)")
    # The caption lists the columns; the tooltip says the one thing it does
    # not — what a blank Unit cell means — and says it in this project's unit.
    assert uploader.proto.help == ("Leave the Unit column blank and the app "
                                   "reads it as g."), uploader.proto.help
    assert any(c.value.startswith("A file with the columns Name, Lowest, Highest")
               for c in at.caption), [c.value for c in at.caption]
def test_an_ingredient_added_mid_run_may_keep_a_lowest_of_its_own(burger):
    """0.5.0: the grid lets a new row start above 0 — "5 g of salt in every
    formulation from here on" is a thing a formulator asks for — and says
    once that the formulations already made contain none of it."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, added=[
        _ingredient_row("Beet juice powder", low=2.0, high=2.0)])
    assert not at.exception
    opt = FoodOptimizer("burger")
    assert opt._var_by_name("Beet juice powder")["bounds"] == (2.0, 2.0)
    assert any(i.value == wording.formulations_contain_none_of(
        "Beet juice powder") for i in at.info), [i.value for i in at.info]


def test_process_setting_baseline_error_is_shown_not_raised(burger):
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, added=[
        _ingredient_row("Cook temperature", kind=wording.KIND_SETTING,
                        low=150.0, high=220.0, unit="°C",
                        **{wording.BASELINE_LABEL: 100.0})])  # outside the range
    assert not at.exception
    assert any("must be between" in str(e.value) for e in at.error), \
        [str(e.value) for e in at.error]


def test_deleting_a_process_setting_confirms_and_archives(burger, tmp_path):
    """0.2.x safety behaviour that must survive the rebuild: a removal that
    rewrites recorded formulations keeps a copy first."""
    burger.add_process_parameter("Cook temperature", 150, 220)
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0,
                 "Cook temperature": 180.0},
                {"Juiciness": 7.0, "Firmness": 6.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, deleted=[2])
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[2])
    at.run()
    assert any("Cook temperature" in w.value for w in at.warning), \
        [w.value for w in at.warning]
    assert any(v["name"] == "Cook temperature" for v in FoodOptimizer("burger").variables)
    _submit_button(at, "Yes, delete").click()
    _grid_edits(at, ING_GRID, deleted=[2])
    at.run()
    assert not at.exception
    assert (tmp_path / "burger_pre_delete.pkl").exists()
    assert not any(v["name"] == "Cook temperature"
                   for v in FoodOptimizer("burger").variables)


def test_limit_fields_start_blank_and_there_is_no_maximum_tick_box(burger):
    """A pre-filled 100 g maximum on a 400 g burger made every future batch
    infeasible in one click; a tick box the user forgot then threw away the
    number they did type. Blank means no limit, everywhere."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    for key in ("qc_max", "qc_min"):
        assert at.number_input(key=key).value is None, key
        assert at.number_input(key=key).placeholder == "no limit", key
    keys = [c.key for c in at.checkbox]
    assert "qc_use_max" not in keys and "qc_use_min" not in keys, keys


def test_a_new_limit_says_past_formulations_are_kept(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.multiselect(key="qty_pick").select("Pea protein")
    at.run()
    at.number_input(key="qc_max").set_value(400.0)
    _submit_button(at, "Add ingredient limit").click()
    at.run()
    assert not at.exception
    assert any(wording.LIMIT_KEPT in s.value for s in at.success), [s.value for s in at.success]


def test_a_limit_that_excludes_everything_made_so_far_warns(burger):
    burger.tell({"Pea protein": 20.0, "Methylcellulose": 2.0},
                {"Juiciness": 7.0, "Firmness": 6.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.multiselect(key="qty_pick").select("Pea protein")
    at.multiselect(key="qty_pick").select("Methylcellulose")
    at.run()
    at.number_input(key="qc_max").set_value(5.0)     # the one formulation is 22 g
    _submit_button(at, "Add ingredient limit").click()
    at.run()
    assert not at.exception
    assert any(w.value == "No formulation you have made fits this limit."
               for w in at.warning), [w.value for w in at.warning]
    assert FoodOptimizer("burger").quantity_constraints    # still added


def test_conflicting_save_shows_the_banner_immediately(burger, tmp_path):
    """Another window changed the file first: the conflict must appear on THIS
    run, not wait for the next click."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    other = FoodOptimizer("burger")
    other.add_objective("Crunch", 1.0)
    time.sleep(0.01)   # coarse-mtime filesystems
    _save_grid(at, MEAS_GRID, added=[_measurement_row("Zing", share=20.0)])
    assert not at.exception
    assert any("another window" in e.value for e in at.error), [e.value for e in at.error]


def test_adding_a_measurement_that_already_exists_is_refused(burger):
    """add_objective REPLACES by name: typing Firmness into the add form
    silently dropped its unit, target and scale and rescored every result
    with no copy kept."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    # case-insensitive
    _save_grid(at, MEAS_GRID, added=[_measurement_row("firmness", share=20.0)])
    assert not at.exception
    assert any(wording.MEASUREMENT_EXISTS_ERROR in e.value
               for e in at.error), [e.value for e in at.error]
    kept = next(o for o in FoodOptimizer("burger").objectives
                if o["name"] == "Firmness")
    # 60 %, not 1.5: a saved project reads its importances as shares of 100.
    assert kept["unit"] == "N" and kept["weight"] == 60.0
    assert kept["goal"] == "target" and kept["target"] == 6.0
    assert not any("Added" in m.value for m in at.success), [m.value for m in at.success]


def test_fixing_a_row_discards_the_round_and_says_so(burger):
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID, discard=True,
               edited={1: {wording.LOWEST_LABEL: 1.0,
                           wording.HIGHEST_LABEL: 1.0}})
    assert not at.exception
    assert FoodOptimizer("burger").pending_batch is None
    assert any(f"{wording.ROUND_CAP} 1 was discarded" in i.value
               for i in at.info), [i.value for i in at.info]


def test_changing_the_range_also_keeps_a_copy_and_reports_it(burger, tmp_path):
    """A share is not the only field that rescores history: the goal, the
    target and either end of the range all feed closeness."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, MEAS_GRID,
               edited={0: {wording.HIGHEST_MEASURABLE_LABEL: 20.0}})
    assert not at.exception
    assert (tmp_path / "burger_pre_edit.pkl").exists()
    assert any(m.value == "Firmness saved. Every overall score was "
               "recalculated. " + wording.COPY_KEPT
               for m in at.success), [m.value for m in at.success]
    assert FoodOptimizer("burger").objectives[1]["max_val"] == 20.0


def test_only_the_unit_changed_means_no_copy_and_no_recalculation_claim(burger):
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, MEAS_GRID, edited={0: {wording.UNIT_LABEL: "kPa"}})
    assert not at.exception
    # A unit is how a number is written, not the number, so nothing is
    # rescored and nothing claims to have been.
    assert any(m.value == "Firmness saved." for m in at.success), \
        [m.value for m in at.success]


def test_deleting_a_measurement_keeps_a_copy_even_with_no_results(burger, tmp_path):
    """A copy is kept before every removal on this tab, history or not."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, MEAS_GRID, confirm=True, deleted=[1])
    assert not at.exception
    assert (tmp_path / "burger_pre_edit.pkl").exists()
    assert [o["name"] for o in FoodOptimizer("burger").objectives] == ["Firmness"]


def test_deleting_a_process_setting_confirms_even_with_no_results(burger, tmp_path):
    burger.add_process_parameter("Cook temperature", 150, 220)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, deleted=[2])
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[2])
    at.run()
    assert any("Cook temperature" in w.value for w in at.warning), \
        [w.value for w in at.warning]
    assert any(v["name"] == "Cook temperature" for v in FoodOptimizer("burger").variables)
    _submit_button(at, "Yes, delete").click()
    _grid_edits(at, ING_GRID, deleted=[2])
    at.run()
    assert not at.exception
    assert (tmp_path / "burger_pre_delete.pkl").exists()
    assert not any(v["name"] == "Cook temperature"
                   for v in FoodOptimizer("burger").variables)


def test_a_failed_save_shows_red_and_never_a_green_added(burger):
    """A green 'Added Beet juice powder.' over a change that never reached the
    file is the worst message this tab could show."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    other = FoodOptimizer("burger")          # another window writes first
    other.add_objective("Crunch", 1.0)
    time.sleep(0.01)                          # coarse-mtime filesystems
    _save_grid(at, ING_GRID,
               added=[_ingredient_row("Beet juice powder", high=2.0)])
    assert not at.exception
    assert any("another window" in e.value for e in at.error), [e.value for e in at.error]
    assert not any("Beet juice powder" in m.value for m in at.success), \
        [m.value for m in at.success]


def test_the_range_labels_are_sentence_case(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    columns = list(_grid_frame(at, 1).columns)
    assert "Lowest measurable" in columns and "Highest measurable" in columns


def test_the_ingredient_grid_has_no_status_column_at_all(burger):
    """0.5.0: a row pinned at one amount says so in a Lowest that is its
    Highest, which is where the reader already looks."""
    burger.add_ingredient("Methylcellulose", 1.0, 1.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    grid = _grid_frame(at, 0)
    assert [c for c in grid.columns if c != "_id"] == [
        "Name", "Type", "Lowest", "Highest", "Unit", "Vendor", "SKU"], \
        list(grid.columns)
    row = grid[grid["Name"] == "Methylcellulose"].iloc[0]
    assert (row["Lowest"], row["Highest"]) == (1.0, 1.0)


@pytest.fixture
def open_batch(burger):
    """Round 1 of two formulations, numbered 1 and 2, nothing recorded yet."""
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0},
                              {"Pea protein": 20.0, "Methylcellulose": 2.0}])
    return burger


@pytest.fixture
def room_to_scale(burger):
    """A round that 200 g is a REACHABLE size for — the three ingredients
    add up to 228 g between them — and at which two of the three are past
    their own Highest. A size the ingredients cannot add up to at all is
    refused outright now, so the caution is about the sizes they can."""
    burger.add_ingredient("Water", 0, 200)
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0,
                               "Water": 5.0}])
    return burger


def test_batch_tab_is_grey_and_names_what_is_missing(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("bare3")
    opt.add_ingredient("Water", 0, 100)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    generate = _submit_button(at, "Generate formulations")
    back = _submit_button(at, "Back to set up")
    assert generate.disabled and generate.proto.type == "secondary"
    assert back.proto.type == "secondary"
    assert any(c.value == "Add at least one measurement." for c in at.caption)


def test_generate_opens_a_numbered_batch_and_stays_on_the_tab(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    # Pin the project so the landing rule does not fire, and seed the tab:
    # st.tabs(key=) does not write session_state on render.
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    _submit_button(at, "Generate 3 formulations").click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert [r["formulation"] for r in reloaded.pending_batch] == [1, 2, 3]
    assert reloaded.pending_batch_no == 1
    assert at.session_state["main_tab"] == wording.TAB_BATCH   # never auto-moves


def test_the_getting_started_sentence_is_the_how_it_works_bullet(burger):
    """Word for word How it works bullet 3, and the number in it is the one
    ask() uses (n_init_random). The caption's two regimes are one sentence
    now; test_the_getting_started_caption_is_the_bullet_word_for_word shows
    it does not change after the fifth result."""
    from ui_setup import HOW_IT_WORKS
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert wording.HOW_CHOSEN == (
        "Until five formulations have results, new ones are spread out to "
        "learn the space. After that, each round aims closer to your "
        "targets."), wording.HOW_CHOSEN
    assert wording.HOW_CHOSEN in HOW_IT_WORKS, HOW_IT_WORKS
    assert any(c.value == wording.HOW_CHOSEN for c in at.caption), \
        [c.value for c in at.caption]


def test_only_one_batch_at_a_time_is_said_where_it_helps(open_batch):
    """The line belongs to the no-batch state, not under a batch that is
    plainly open."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not any(c.value == "Only one batch is open at a time."
                   for c in at.caption), [c.value for c in at.caption]


def test_own_formulation_starts_a_batch_when_none_is_open(burger):
    """A formulation the scientist chose is a formulation like any other: it
    opens the batch when nothing is open, and draws the next number."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    assert any(c.value == wording.ADD_OWN_NO_BATCH_CAPTION
               for c in at.caption), [c.value for c in at.caption]
    # Every typed amount carries its unit, the same header the batch table
    # prints above that column.
    assert at.number_input(key="own_Pea protein").label == "Pea protein (g)"
    at.number_input(key="own_Pea protein").set_value(12.0)
    at.number_input(key="own_Methylcellulose").set_value(1.5)
    at.run()
    _submit_button(at, wording.ADD_TO_THIS_BATCH).click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert len(reloaded.pending_batch) == 1
    assert reloaded.pending_batch[0]["formulation"] == 1
    assert reloaded.pending_batch[0]["recipe"] == {"Pea protein": 12.0,
                                                   "Methylcellulose": 1.5}
    assert reloaded.pending_batch[0]["note"] == wording.OWN_FORMULATION_NOTE
    assert reloaded.pending_batch_no == 1
    assert len(reloaded.pending_batch_created) == 10        # an ISO date
    assert any(s.value == "Formulation 1 added to Round 1."
               for s in at.success), [s.value for s in at.success]
    # The boxes empty again, so the next one is typed into a clean form.
    assert at.session_state["own_Pea protein"] is None
    assert at.session_state["own_note"] == ""


def test_own_formulation_joins_the_open_batch(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="how_many").set_value(2)
    at.run()
    _submit_button(at, "Generate 2 formulations").click()
    at.run()
    at.number_input(key="own_Pea protein").set_value(12.0)
    at.number_input(key="own_Methylcellulose").set_value(1.5)
    at.run()
    _submit_button(at, wording.ADD_TO_THIS_BATCH).click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert [r["formulation"] for r in reloaded.pending_batch] == [1, 2, 3]
    assert reloaded.pending_batch[-1]["recipe"] == {"Pea protein": 12.0,
                                                    "Methylcellulose": 1.5}
    assert reloaded.pending_batch_no == 1
    # The table above says three, and the row says what it is.
    assert any(m.value == wording.make_these(1, 3)
               for m in at.markdown), [m.value for m in at.markdown]
    table = next(d.value for d in at.dataframe if "Formulation" in d.value.columns)
    assert list(table["Note"]) == ["", "", wording.OWN_FORMULATION_NOTE]
    assert any(s.value == "Formulation 3 added to Round 1."
               for s in at.success), [s.value for s in at.success]


def test_the_own_formulation_note_placeholder_gives_an_example(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="how_many").set_value(2)
    at.run()
    _submit_button(at, "Generate 2 formulations").click()
    at.run()
    assert (at.text_input(key="own_note").proto.placeholder
           == "e.g. Repeat of 4 with more salt")


def test_start_from_the_best_prefills_and_notes_the_repeat(burger):
    """The old Repeat checkbox is now two clicks: fill the boxes from the best
    formulation, then add it. The note says which formulation it repeats."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    _submit_button(at, wording.START_FROM_BEST).click()
    at.run()
    assert not at.exception
    assert at.session_state["own_Pea protein"] == 10.0
    assert at.session_state["own_Methylcellulose"] == 1.0
    assert at.session_state["own_note"] == wording.repeat_of_formulation(1)
    assert at.number_input(key="own_Pea protein").value == 10.0
    # ...and adding it carries that note onto the batch.
    _submit_button(at, wording.ADD_TO_THIS_BATCH).click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert reloaded.pending_batch[-1]["recipe"] == {"Pea protein": 10.0,
                                                    "Methylcellulose": 1.0}
    assert reloaded.pending_batch[-1]["note"] == wording.repeat_of_formulation(1)


def test_blank_amount_is_refused(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="own_Pea protein").set_value(12.0)   # the other is blank
    at.run()
    _submit_button(at, wording.ADD_TO_THIS_BATCH).click()
    at.run()
    assert not at.exception
    assert any(e.value == wording.ENTER_EVERY_AMOUNT
               for e in at.error), [e.value for e in at.error]
    # Nothing opened, and what was typed is still there to finish.
    assert not FoodOptimizer("burger").pending_batch
    assert at.session_state["own_Pea protein"] == 12.0


def test_an_own_amount_outside_the_allowed_ones_is_a_caution(burger):
    """A formulation the user means to make is never refused for being
    outside the allowed amounts; the caution says so and the row is kept."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="own_Pea protein").set_value(40.0)      # allowed: 0-25
    at.number_input(key="own_Methylcellulose").set_value(1.0)
    at.run()
    _submit_button(at, wording.ADD_TO_THIS_BATCH).click()
    at.run()
    assert not at.exception
    assert any("Pea protein 40 g is outside its allowed amounts of 0 to 25 g."
               == w.value for w in at.warning), [w.value for w in at.warning]
    assert FoodOptimizer("burger").pending_batch[0]["recipe"]["Pea protein"] == 40.0


def test_no_repeat_checkbox_remains(burger):
    """The checkbox is retired: a formulation of your own does that job, and
    says so in its note."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    assert not any(c.label.startswith("Repeat") for c in at.checkbox), \
        [c.label for c in at.checkbox]
    # The box still asks for NEW formulations, and the button counts them.
    assert any(n.label == wording.FORMULATIONS_TO_GENERATE
               for n in at.number_input), [n.label for n in at.number_input]
    assert any(b.label == "Generate 3 formulations" for b in at.button), \
        _labels(at)


def test_the_trial_table_carries_units_and_a_total(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(m.value == wording.make_these(1, 2)
               for m in at.markdown), [m.value for m in at.markdown]
    table = next(d.value for d in at.dataframe if "Formulation" in d.value.columns)
    # No what-is-it-trying column during the cold start: every cell under it
    # said the header back, word for word.
    assert list(table.columns) == ["Formulation", "Pea protein (g)",
                                   "Methylcellulose (g)", "Total (g)"]
    box = at.number_input(key="scale_total")
    # Blank: no stored size and no default, so the box has no number to open
    # at. It opened at the mean of the rows' own sums for a while, which put
    # a size on screen that no formulation weighed.
    assert box.value is None
    assert box.proto.placeholder == "e.g. 100"


def test_a_batch_size_rescales_the_screen_the_sheet_and_the_record(open_batch):
    """What is downloaded must equal what is on screen: the whole point of
    the box is that the lab weighs those amounts out. Since 0.5.0 the stored
    row moves with them, so what is recorded is what was made."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(22.0)
    at.run()
    table = next(d.value for d in at.dataframe if "Formulation" in d.value.columns)
    assert table["Pea protein (g)"].iloc[0] == pytest.approx(20.0)
    # One caption about the size on this tab, not two: it names the number
    # the files were written for, under the files.
    assert any(c.value == "Sheets show each formulation made to 22 g."
               for c in at.caption), [c.value for c in at.caption]
    assert not any(c.value == "Amounts shown for this total."
                   for c in at.caption), [c.value for c in at.caption]
    stored = FoodOptimizer("burger").pending_batch
    assert stored[0]["recipe"]["Pea protein"] == pytest.approx(20.0)
    assert stored[1]["recipe"]["Pea protein"] == pytest.approx(20.0)


def test_one_download_is_offered_and_it_is_lit_until_a_result_is_typed(open_batch):
    """One file leaves this tab: the workbook. It was three — a sheet to fill
    in, a set of sheets to print and a preview of those — and a bench had to
    choose between them before it could weigh anything out."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert [d.label for d in _unknowns(at.tabs[1], "download_button")] == \
        [wording.DOWNLOAD_BATCH_SHEETS], \
        [d.label for d in _unknowns(at.tabs[1], "download_button")]
    assert _unknown(at.main, "download_button",
                    "Download the round sheets (Excel)").proto.type == "primary"
    # The download IS the coloured thing here, and it is the only one.
    assert _tab_primaries(at, 1) == ["Download the round sheets (Excel)"], _tab_primaries(at, 1)
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.run()
    assert _unknown(at.main, "download_button",
                    "Download the round sheets (Excel)").proto.type == "secondary"


def test_leaving_a_row_out_does_not_grey_the_batch_sheets(open_batch):
    """A disabled number_input still returns its stored value, so a naive
    'has anything been typed' check flipped the download grey on a tick."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.run()
    at.checkbox(key="f1_leave_out").check()
    at.run()
    assert _unknown(at.main, "download_button",
                    "Download the round sheets (Excel)").proto.type == "primary"


def test_regenerating_names_the_numbers_it_discards(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at, wording.GENERATE_DIFFERENT_BATCH).click()
    at.run()
    assert any(w.value == wording.regenerate_warning(1, "1 and 2", 3)
               for w in at.warning), [w.value for w in at.warning]
    # The number is read off the project, not guessed from the rows, and the
    # sentence says NEW FORMULATIONS: the batch keeps its own number.
    assert wording.regenerate_warning(1, "1 and 2", 3) == (
        "Discard Round 1 and Formulations 1 and 2? New formulations start at "
        "Formulation 3.")
    assert at.session_state["optimizer"].next_formulation_no == 3
    _submit_button(at, "Yes, discard").click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert [r["formulation"] for r in reloaded.pending_batch] == [3, 4]
    assert reloaded.pending_batch_no == 1     # the batch keeps its number
    assert any(c.value == wording.batch_discarded_caption("1 and 2")
               for c in at.caption), [c.value for c in at.caption]


def test_regenerating_asks_for_the_generated_rows_only(burger):
    """A formulation of the user's own is theirs, not a slot the model filled:
    discarding a batch of three generated rows plus one own asks for three."""
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0},
                              {"Pea protein": 15.0, "Methylcellulose": 1.5},
                              {"Pea protein": 20.0, "Methylcellulose": 2.0}])
    burger.add_to_pending_batch({"Pea protein": 5.0, "Methylcellulose": 0.5},
                                note=wording.OWN_FORMULATION_NOTE)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    _submit_button(at, wording.GENERATE_DIFFERENT_BATCH).click()
    at.run()
    _submit_button(at, "Yes, discard").click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert len(reloaded.pending_batch) == 3
    assert [r["formulation"] for r in reloaded.pending_batch] == [5, 6, 7]
    assert reloaded.pending_batch_no == 1      # the batch keeps its number
    # All four numbers retired, the own one included: the whole batch went.
    assert any(c.value == wording.batch_discarded_caption("1, 2, 3 and 4")
               for c in at.caption), [c.value for c in at.caption]


def test_each_sheet_names_its_formulation_and_its_batch(open_batch):
    """The page a technician carries says which formulation it is, which
    batch it belongs to and whose project it came out of — the screen is two
    rooms away by then."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    book = _workbook(at)
    assert book.sheetnames == ["Round 1", "Formulation 1", "Formulation 2"], \
        book.sheetnames
    texts = _sheet_text(at)
    assert f"{wording.FORMULATION_CAP} 1 · {wording.ROUND_CAP} 1 · burger" \
        in texts, texts
    assert wording.NOT_SCORED_CHECKBOX_SHEET in texts, texts
    # The measurement is named, what a good number looks like is beside it,
    # and the cell to write the reading in is empty.
    assert "Firmness" in texts and "target 6 N" in texts, texts
    assert wording.MEASURED_COLUMN in texts, texts


def test_result_inputs_carry_the_goal_and_unit_and_are_ordered_by_importance(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert at.number_input(key="f1_Firmness").label == "Firmness · target 6 N"
    # The panel unit is on the label, once, and never after the number.
    assert (at.number_input(key="f1_Juiciness").label
            == "Juiciness (/10) · target 7")
    labels = [n.label for n in at.number_input]
    assert (labels.index("Firmness · target 6 N")
            < labels.index("Juiciness (/10) · target 7"))
    assert any(c.value == ("One number per measurement; the panel mean where "
                           "a panel scored it. Leave blank if it was not "
                           "measured.") for c in at.caption), \
        [c.value for c in at.caption]


def test_result_inputs_do_not_clamp_and_refuse_out_of_range_on_save(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    field = at.number_input(key="f1_Firmness")
    # An unbounded st.number_input reports ±DBL_MAX rather than None, so what
    # is checked is that neither end is the 0-to-10 scale: nothing is clamped.
    assert field.min < 0.0 and field.max > 10.0
    field.set_value(12.0)
    at.number_input(key="f2_Firmness").set_value(6.0)
    at.run()
    assert at.number_input(key="f1_Firmness").value == 12.0   # not clamped to 10
    _submit_button(at, "Save results").click()
    at.run()
    assert not at.exception
    assert any(e.value == ("Firmness 12 N is outside your range of 0 to 10 N. "
                           "Widen the range in Set up, or check the value.")
               for e in at.error), [e.value for e in at.error]
    assert FoodOptimizer("burger").X_history == []


def test_leave_out_comes_after_the_note_field(open_batch):
    """The first control in a row must be a box the user came to fill, not a
    checkbox for not making it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()

    def positions(node, out):
        children = getattr(node, "children", None) or {}
        if hasattr(children, "values"):
            children = children.values()
        for child in children:
            key = getattr(child, "key", None)
            if key in ("f1_Firmness", "f1_note", "f1_leave_out"):
                out.append(key)
            positions(child, out)
        return out

    order = positions(at.main, [])
    assert order.index("f1_Firmness") < order.index("f1_note") < order.index("f1_leave_out"), order


def test_save_lights_only_when_every_kept_row_has_a_value(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _submit_button(at, "Save results").disabled
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.run()
    assert _submit_button(at, "Save results").disabled
    # "complete", not "to record": every other screen uses "to record" for
    # the rows that still have no number ("Back to Round 1 · 2 to record"),
    # and this line counts the opposite. A row is complete only when EVERY
    # measurement has a number: one of two typed is not half a result, so it
    # is counted as partly filled instead. The Save results button is the
    # next thing on the screen, so the line no longer names it.
    assert any(c.value == "0 of 2 complete · 1 partly filled"
               for c in at.caption), [c.value for c in at.caption]
    at.number_input(key="f1_Juiciness").set_value(4.0)
    at.run()
    assert any(c.value == "1 of 2 complete"
               for c in at.caption), [c.value for c in at.caption]
    at.number_input(key="f2_Juiciness").set_value(4.0)
    at.run()
    # Save stays allowed with a partly filled row: a partial result is a
    # result, and the model is told so.
    save = _submit_button(at, "Save results")
    assert not save.disabled and save.proto.type == "primary"
    assert any(c.value == "1 of 2 complete · 1 partly filled"
               for c in at.caption), [c.value for c in at.caption]
    assert _tab_primaries(at, 1) == ["Save results"], _tab_primaries(at, 1)


def test_the_counter_and_the_foot_do_not_contradict_each_other(open_batch):
    """`to record` means a row with no value yet, everywhere it is written:
    the foot of tab 3 and the line under the title both use it that way. The
    results line counts the opposite — the rows that are finished — so it
    says `complete` instead. Saying `1 of 2 to record` beside `2 to record`
    put two different meanings on one screen."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.run()
    assert any(c.value.startswith("0 of 2 complete · 1 partly filled")
               for c in at.caption), [c.value for c in at.caption]
    # Nothing is saved yet, so both open rows are still to record, and both
    # other screens say so in those words.
    assert any(c.value == wording.batch_line_open(1, 2) for c in at.caption), \
        [c.value for c in at.caption]
    assert _submit_button(at, wording.back_to_batch_label(1, 2)).label == \
        wording.back_to_batch_label(1, 2)
    assert not any("to record" in c.value and "filled in" not in c.value
                   and c.value.startswith(("1 of", "2 of"))
                   for c in at.caption), [c.value for c in at.caption]


def test_the_counter_counts_every_row_of_the_batch(open_batch):
    """The denominator is the sheets in the technician's hand. Counting only
    the rows still being scored turned a batch of two into "1 of 1"."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.number_input(key="f1_Juiciness").set_value(7.0)
    at.checkbox(key="f2_leave_out").check()
    at.run()
    assert any(c.value == "1 of 2 complete · 1 not scored"
               for c in at.caption), [c.value for c in at.caption]
    assert not _submit_button(at, "Save results").disabled


def test_saving_records_partials_notes_and_moves_to_results(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.number_input(key="f1_Juiciness").set_value(7.0)
    at.text_input(key="f1_note").set_value("held together")
    at.number_input(key="f2_Firmness").set_value(2.0)     # Juiciness left blank
    at.run()
    _submit_button(at, "Save results").click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert reloaded.formulation_ids == [1, 2]
    assert reloaded.batch_history == [1, 1]
    assert reloaded.notes_history[0] == "held together"
    assert reloaded.results_history[1] == {"Firmness": 2.0}
    assert reloaded.pending_batch is None
    assert at.session_state["main_tab"] == wording.TAB_RESULTS
    # A real save also puts the saved line in the sidebar and no save nag.
    assert any("Saved automatically at " in c.value for c in at.sidebar.caption), \
        [c.value for c in at.sidebar.caption]
    assert not any("Save a copy now" in w.value for w in at.warning)


def test_a_failed_save_shows_the_banner_and_does_not_move_tabs(open_batch, monkeypatch):
    """go_to_tab reruns, which would outrun app.py's end-of-script save-error
    check and hide the fact that nothing was written."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.number_input(key="f2_Firmness").set_value(4.0)
    at.run()
    at.session_state["optimizer"].storage.save = lambda name, state: (_ for _ in ()).throw(
        __import__("storage").StorageError("The disk is full — your last change was NOT saved.")
    )
    _submit_button(at, "Save results").click()
    at.run()
    assert not at.exception
    assert at.session_state["main_tab"] == wording.TAB_BATCH
    assert any("NOT saved" in e.value for e in at.error), [e.value for e in at.error]


def test_an_uploaded_sheet_whose_final_write_fails_stays_on_the_batch(open_batch):
    """Twin of the typed path: closing the batch is a write like any other, and
    a green "Round 1 recorded." over a batch still open on disk is a lie."""
    import storage as storage_backend

    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.session_state["_results_upload"] = pd.DataFrame(
        {"Formulation": [1, 2], "Firmness": [6.0, 4.0],
         "Juiciness": [7.0, 5.0]})
    at.run()
    real_save = at.session_state["optimizer"].storage.save
    calls = []

    def flaky(name, state):
        calls.append(name)
        if len(calls) == 3:      # the write that closes the batch
            raise storage_backend.StorageError(
                "The disk is full — your last change was NOT saved.")
        return real_save(name, state)

    at.session_state["optimizer"].storage.save = flaky
    _submit_button(at, "Save uploaded results").click()
    at.run()
    assert not at.exception
    assert any("NOT saved" in e.value for e in at.error), [e.value for e in at.error]
    assert not any("recorded" in s.value for s in at.success), \
        [s.value for s in at.success]
    assert at.session_state["main_tab"] == wording.TAB_BATCH
    assert FoodOptimizer("burger").pending_batch is not None


def test_leave_out_disables_the_row_but_keeps_its_values(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="f2_Firmness").set_value(3.0)
    at.checkbox(key="f2_leave_out").check()
    at.run()
    assert at.number_input(key="f2_Firmness").disabled is True
    assert at.number_input(key="f2_Firmness").value == 3.0
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.run()
    _submit_button(at, "Save results").click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert reloaded.formulation_ids == [1]
    assert reloaded.skipped == [{"formulation": 2, "batch": 1,
                                 "recipe": {"Pea protein": 20.0,
                                            "Methylcellulose": 2.0},
                                 "note": "Not scored"}]


def test_every_row_left_out_says_there_is_nothing_to_save(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.checkbox(key="f1_leave_out").check()
    at.checkbox(key="f2_leave_out").check()
    at.run()
    assert any(i.value == "Nothing to save — at least one formulation needs results."
               for i in at.info), [i.value for i in at.info]
    assert _submit_button(at, "Save results").disabled
    assert FoodOptimizer("burger").pending_batch is not None


def test_a_partly_recorded_batch_reopens_with_recorded_rows_read_only(open_batch):
    open_batch.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                    {"Juiciness": 7.0, "Firmness": 6.0},
                    formulation_no=1, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    keys = [n.key for n in at.number_input]
    assert "f1_Firmness" not in keys
    assert "f2_Firmness" in keys
    assert any(c.value == "Firmness 6 N · Juiciness (/10) 7" for c in at.caption), \
        [c.value for c in at.caption]


def test_uploading_a_sheet_matches_global_formulation_numbers(open_batch):
    """AppTest cannot drive st.file_uploader, so prime the parsed sheet the way
    the uploader handler does."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.session_state["_results_upload"] = pd.DataFrame(
        {"Formulation": [2], "Firmness": [6.0], "Juiciness": [7.0],
         "Note": ["from the sheet"]})
    at.run()
    assert any(f"{wording.FORMULATION_CAP} 2" in i.value for i in at.info), \
        [i.value for i in at.info]
    _submit_button(at, "Save uploaded results").click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert reloaded.formulation_ids == [2]
    assert reloaded.notes_history == ["from the sheet"]
    assert [r["formulation"] for r in reloaded.pending_batch] == [1, 2]
    assert at.session_state["main_tab"] == wording.TAB_BATCH   # one row still open


def test_a_stale_upload_is_cleared_by_hard_reset(open_batch, tmp_path):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_results_upload"] = pd.DataFrame(
        {"Formulation": [1], "Firmness": [6.0], "Juiciness": [7.0]})
    at.run()
    assert "_results_upload" in at.session_state
    _submit_button(at, "Start this project over").click()
    at.run()
    _submit_button(at, "Yes, start over").click()
    at.run()
    assert not at.exception
    assert "_results_upload" not in at.session_state


def test_generate_is_the_one_lit_button_before_a_batch_exists(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _tab_primaries(at, 1) == ["Generate 3 formulations"], _tab_primaries(at, 1)


def test_a_confirmation_takes_the_colour_from_the_batch_sheets(open_batch):
    """While a confirmation is armed, answering it is the one thing to do: the
    lit download steps aside rather than competing with the Yes."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _unknown(at.main, "download_button",
                    "Download the round sheets (Excel)").proto.type == "primary"
    _submit_button(at, wording.GENERATE_DIFFERENT_BATCH).click()
    at.run()
    assert _unknown(at.main, "download_button",
                    "Download the round sheets (Excel)").proto.type == "secondary"
    assert _tab_primaries(at, 1) == ["Yes, discard"], _tab_primaries(at, 1)


def test_a_confirmation_greys_save_results(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.number_input(key="f2_Firmness").set_value(4.0)
    at.run()
    assert _tab_primaries(at, 1) == ["Save results"], _tab_primaries(at, 1)
    _submit_button(at, wording.GENERATE_DIFFERENT_BATCH).click()
    at.run()
    save = _submit_button(at, "Save results")
    assert save.disabled and save.proto.type == "secondary"
    assert _tab_primaries(at, 1) == ["Yes, discard"], _tab_primaries(at, 1)



def test_a_confirmation_greys_both_add_your_own_buttons(burger):
    """The expander sits under the batch table but is drawn AFTER the
    confirmations, so a Yes armed on this very run greys its two buttons —
    arming does not rerun, and a live button under a warning is a second
    thing to do."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1, batch_no=1)
    burger.set_pending_batch([{"Pea protein": 12.0, "Methylcellulose": 1.2},
                              {"Pea protein": 18.0, "Methylcellulose": 2.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    assert not _submit_button(at, wording.ADD_TO_THIS_BATCH).disabled
    assert not _submit_button(at, wording.START_FROM_BEST).disabled
    _submit_button(at, wording.GENERATE_DIFFERENT_BATCH).click()
    at.run()
    assert _submit_button(at, wording.ADD_TO_THIS_BATCH).disabled
    assert _submit_button(at, wording.START_FROM_BEST).disabled
    assert _tab_primaries(at, 1) == ["Yes, discard"], _tab_primaries(at, 1)


def _displayed(element):
    """The strings a st.dataframe actually puts on screen. A Styler's number
    formatting rides along in the proto as display values, so this is what the
    user reads, not the raw floats behind it."""
    return pa.ipc.open_stream(
        io.BytesIO(element.proto.arrow_data.styler.display_values)
    ).read_pandas()


def test_amounts_stay_as_generated_until_a_total_is_typed(open_batch):
    """The two formulations weigh 11 g and 22 g. Nothing may quietly rewrite
    the second one to the first one's size just because the tab was opened."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    table = next(d.value for d in at.dataframe if "Formulation" in d.value.columns)
    assert list(table["Total (g)"]) == [11.0, 22.0], table.to_dict()
    assert (at.number_input(key="scale_total").proto.placeholder
           == "e.g. 100")
    amounts = _sheet_amounts(at)
    # The sheets carry the amounts as generated, to the balance's own two
    # decimals. The box is blank until the bench types a size, and a size
    # nobody typed rewrites nothing.
    assert ("Pea protein", 20.0) in amounts, amounts
    assert ("Methylcellulose", 1.0) in amounts, amounts


def test_typing_the_first_rows_own_total_still_scales_the_others(open_batch):
    """11 g is what the first formulation already weighs, so a box that opened
    pre-filled with it made typing 11 a silent no-op. It must scale."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(11.0)
    at.run()
    table = next(d.value for d in at.dataframe if "Formulation" in d.value.columns)
    assert list(table["Total (g)"]) == pytest.approx([11.0, 11.0])
    assert table["Pea protein (g)"].iloc[1] == pytest.approx(10.0)
    assert any(c.value == "Sheets show each formulation made to 11 g."
               for c in at.caption), [c.value for c in at.caption]
    # Both sheets now carry the same amounts, so the download followed.
    assert _sheet_amounts(at).count(("Pea protein", 10.0)) == 2
    assert FoodOptimizer("burger").pending_batch[1]["recipe"]["Pea protein"] \
        == pytest.approx(10.0)


def test_the_batch_table_never_dresses_a_setting_as_an_amount(burger):
    burger.add_process_parameter("Cook temperature", 100, 200)
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0,
                               "Cook temperature": 180.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    shown = _displayed(next(d for d in at.dataframe
                            if "Formulation" in d.value.columns))
    assert shown["Pea protein (g)"].iloc[0] == "10.00"
    assert shown["Cook temperature"].iloc[0] == "180", shown.to_dict()


def test_a_one_formulation_trial_reads_as_one(burger):
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(m.value == wording.make_these(1, 1)
               for m in at.markdown), [m.value for m in at.markdown]


def test_adding_a_formulation_of_your_own_keeps_the_typed_results(open_batch):
    """Both buttons in `Add a formulation of your own` end in st.rerun(), and
    Streamlit discards the session-state entry of every widget the run did
    not create. Drawn above the result grid, they took every measurement,
    note and Not-made tick already typed into the open batch with them."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.number_input(key="f1_Juiciness").set_value(7.0)
    at.text_input(key="f1_note").set_value("held together")
    at.checkbox(key="f2_leave_out").check()
    at.run()
    # `Add to this batch`, with every amount typed so the row really is added.
    at.number_input(key="own_Pea protein").set_value(15.0)
    at.number_input(key="own_Methylcellulose").set_value(2.0)
    at.run()
    _submit_button(at, wording.ADD_TO_THIS_BATCH).click()
    at.run()
    assert not at.exception
    assert at.session_state["f1_Firmness"] == 6.0
    assert at.session_state["f1_Juiciness"] == 7.0
    assert at.session_state["f1_note"] == "held together"
    assert at.session_state["f2_leave_out"] is True


def test_start_from_the_best_keeps_the_typed_results(scored_open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="f4_Firmness").set_value(5.0)
    at.text_input(key="f4_note").set_value("second try")
    at.run()
    _submit_button(at, wording.START_FROM_BEST).click()
    at.run()
    assert not at.exception
    assert at.session_state["f4_Firmness"] == 5.0
    assert at.session_state["f4_note"] == "second try"


@pytest.fixture
def scored_open_batch(scored):
    """One recorded batch and a second batch open, so `Start from the best so
    far` has a best to start from and the grid has a row to type into."""
    scored.set_pending_batch([{"Pea protein": 15.0, "Methylcellulose": 2.0}])
    return scored


def test_a_note_typed_on_a_row_that_is_left_out_is_kept(open_batch):
    """Why it was not scored is often typed before the box is ticked, and it is
    the only record of what went wrong."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.text_input(key="f2_note").set_value("burner failed")
    at.checkbox(key="f2_leave_out").check()
    at.run()
    _submit_button(at, "Save results").click()
    at.run()
    assert not at.exception
    # "Not scored" first: the row is otherwise blank, and nothing else on it
    # says the formulation was left out.
    assert FoodOptimizer("burger").skipped[0]["note"] == "Not scored · burner failed"


def test_a_recorded_row_shows_its_note(open_batch):
    open_batch.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                    {"Juiciness": 7.0, "Firmness": 6.0},
                    formulation_no=1, batch_no=1, note="held together")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(c.value == "Firmness 6 N · Juiciness (/10) 7 · Note: held together"
               for c in at.caption), [c.value for c in at.caption]


@pytest.fixture
def scored(burger):
    """Two recorded formulations in batch 1 and one left out, so the results
    screens have a best, a progress line and a left-out row to show."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 1.0},
                formulation_no=1, batch_no=1, note="crumbly")
    burger.tell({"Pea protein": 20.0, "Methylcellulose": 2.0},
                {"Juiciness": 7.0, "Firmness": 8.0},
                formulation_no=2, batch_no=1)
    burger.record_skipped(3, 1, {"Pea protein": 25.0, "Methylcellulose": 3.0})
    return burger


def test_results_empty_state_offers_the_first_trial_and_the_import(burger):
    """A fresh project is exactly when someone imports past work, so the
    import must not hide behind the empty state."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert any(m.value == "No results yet." for m in at.markdown), \
        [m.value for m in at.markdown]
    assert any(e.label == wording.EDIT_PAST_FORMULATIONS_EXPANDER
               for e in at.expander), [e.label for e in at.expander]
    button = _submit_button(at, wording.MAKE_YOUR_FIRST_BATCH_BUTTON)
    assert button.proto.type == "primary"
    button.click()
    at.run()
    assert at.session_state["main_tab"] == wording.TAB_BATCH


def test_edit_past_formulations_holds_the_three_parts(scored):
    """Three controls became one section: the correction picker, the delete
    list and the only way into work done before the project existed, which
    used to hide in an expander of its own."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    labels = [e.label for e in at.expander]
    assert wording.EDIT_PAST_FORMULATIONS_EXPANDER in labels, labels
    marks = [m.value for m in at.markdown]
    for heading in (wording.CORRECT_A_FORMULATION_HEADING,
                    wording.DELETE_FORMULATIONS_HEADING,
                    wording.ADD_PAST_FORMULATION_HEADING):
        assert heading in marks, marks
    # None of the three retired controls is still on screen.
    assert not any(s.label == "Correct a result" for s in at.selectbox), \
        [s.label for s in at.selectbox]
    assert not any(b.label.startswith("Delete the last") for b in at.button), \
        _labels(at)
    assert not any(l.startswith("Import past") for l in labels), labels
    assert at.radio(key="add_past_mode").options == ["Type it in",
                                                     "Upload a file"]


def test_correcting_amounts_and_a_result_in_one_save(scored, tmp_path):
    """What went into the bowl is as correctable as what came off the panel,
    and the encoded row goes with it or the model keeps steering towards a
    formulation nobody made."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    # The boxes open on the amounts the row was recorded with, each labelled
    # with its own unit.
    assert at.number_input(key="correct_amount_1_Pea protein").value == 10.0
    assert at.number_input(key="correct_amount_1_Pea protein").label \
        == "Pea protein (g)"
    at.number_input(key="correct_amount_1_Pea protein").set_value(15.0)
    at.number_input(key="correct_1_Firmness").set_value(6.0)
    at.run()
    _submit_button(at, "Save correction").click()
    at.run()
    assert not at.exception
    assert (tmp_path / "burger_pre_edit.pkl").exists(), \
        [f.name for f in tmp_path.glob("*.pkl")]
    reloaded = FoodOptimizer("burger")
    assert reloaded.recipe_history[0]["Pea protein"] == 15.0
    assert reloaded.results_history[0]["Firmness"] == 6.0
    assert reloaded.X_history[0] == reloaded._encode(reloaded.recipe_history[0])
    assert any(s.value.startswith(wording.formulation_corrected(1))
               for s in at.success), [s.value for s in at.success]


def test_a_corrected_amount_outside_the_allowed_ones_is_a_caution(scored):
    """An amount outside what the project allows is a fact about work already
    done, exactly as an imported one is: it is stored, and cautioned."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    at.number_input(key="correct_amount_1_Pea protein").set_value(99.0)
    at.run()
    _submit_button(at, "Save correction").click()
    at.run()
    assert not at.exception
    assert any(w.value == ("Pea protein 99 g is outside its allowed amounts "
                           "of 0 to 25 g.") for w in at.warning), \
        [w.value for w in at.warning]
    assert FoodOptimizer("burger").recipe_history[0]["Pea protein"] == 99.0


def test_type_in_a_past_formulation(burger):
    """The typed-in half of `Add a formulation you already made`: a row from
    before this project existed, belonging to no batch, partial where a
    measurement was never taken."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert at.text_input(key="past_note").value == wording.IMPORTED_NOTE
    assert at.number_input(key="past_Pea protein").label == "Pea protein (g)"
    at.number_input(key="past_Pea protein").set_value(12.0)
    at.number_input(key="past_Methylcellulose").set_value(1.2)
    at.number_input(key="past_m_Firmness").set_value(5.0)
    at.run()
    _submit_button(at, wording.ADD_THIS_FORMULATION).click()
    at.run()
    assert not at.exception
    assert any(s.value == wording.formulation_added(1) for s in at.success), \
        [s.value for s in at.success]
    reloaded = FoodOptimizer("burger")
    assert reloaded.formulation_ids == [1]
    assert reloaded.batch_history == [None]          # it belongs to no batch
    assert reloaded.notes_history == [wording.IMPORTED_NOTE]
    assert reloaded.results_history[0] == {"Firmness": 5.0}
    table = next(d.value for d in at.dataframe if "Best" in d.value.columns)
    assert table["Note"].iloc[0] == wording.IMPORTED_NOTE
    assert str(table["Round"].iloc[0]) in ("", "None", "nan", "<NA>"), \
        table["Round"].iloc[0]


def test_a_typed_past_formulation_needs_every_amount(burger):
    """A blank amount is a refusal, and what was typed stays in the boxes."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.number_input(key="past_Pea protein").set_value(12.0)
    at.run()
    _submit_button(at, wording.ADD_THIS_FORMULATION).click()
    at.run()
    assert not at.exception
    assert any(e.value == wording.ENTER_EVERY_AMOUNT for e in at.error), \
        [e.value for e in at.error]
    assert at.number_input(key="past_Pea protein").value == 12.0
    assert FoodOptimizer("burger").X_history == []


def test_a_typed_past_measurement_outside_its_range_is_refused(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.number_input(key="past_Pea protein").set_value(12.0)
    at.number_input(key="past_Methylcellulose").set_value(1.2)
    at.number_input(key="past_m_Firmness").set_value(99.0)
    at.run()
    _submit_button(at, wording.ADD_THIS_FORMULATION).click()
    at.run()
    assert not at.exception
    assert any(e.value == ("Firmness 99 N is outside your range of 0 to 10 N. "
                           "Widen the range in Set up, or check the value.")
               for e in at.error), [e.value for e in at.error]
    assert FoodOptimizer("burger").X_history == []


def test_best_heading_off_by_table_and_score_caption(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(h.value == wording.best_so_far_heading(2, 1)
               for h in at.subheader), [h.value for h in at.subheader]
    off_by = next(d.value for d in at.dataframe if "Off by" in d.value.columns)
    assert list(off_by["Measurement"]) == ["Firmness", "Juiciness (/10)"]
    assert off_by["Off by"].iloc[0] == "2 N too high"
    assert off_by["Off by"].iloc[1] == "On target"
    assert any(m.value == "**Amounts to make it**" for m in at.markdown)
    # The best here is 2 N off target, so a caption saying "every
    # measurement on target" would be a false claim about the formulation
    # the table above it is describing. What the ceiling means is said once,
    # on Set up. The whole caption, word for word:
    assert any(c.value == ("Overall score 88.00 of 100.00. Scores only "
                           "compare within this project. Change a share, a "
                           "goal or a range and every score is worked out "
                           "again.")
               for c in at.caption), [c.value for c in at.caption]


def test_not_used_lists_ingredients_only(scored):
    """A cook temperature legitimately set to 0 is not an unused ingredient."""
    # Added mid-run, a process setting needs the value all prior batches ran
    # at; zero is the point of this test.
    scored.add_process_parameter("Cook temperature", 0, 220, baseline=0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not any("Cook temperature" in c.value for c in at.caption
                   if c.value.startswith("Not used:")), [c.value for c in at.caption]


def test_the_first_batch_is_not_announced_twice_on_one_screen(scored):
    """Saving flashes "Round 1 recorded." above the tabs. A caption under the
    heading saying the same thing is the same sentence twice on one screen,
    and there is nothing to compare a first batch with anyway."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not any(c.value == wording.batch_recorded_flash(1) for c in at.caption), \
        [c.value for c in at.caption]


def test_the_progress_line_reports_an_improvement_and_a_flat_trial(scored):
    scored.tell({"Pea protein": 12.0, "Methylcellulose": 1.2},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=4, batch_no=2)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(c.value == wording.batch_recorded_progress(2, 88.00, 100.00)
               for c in at.caption), [c.value for c in at.caption]
    scored.tell({"Pea protein": 30.0, "Methylcellulose": 0.5},
                {"Juiciness": 1.0, "Firmness": 1.0}, formulation_no=5, batch_no=3)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(c.value == wording.batch_recorded_no_improvement(3)
               for c in at.caption), [c.value for c in at.caption]


def test_all_formulations_table_stars_the_best_and_marks_the_left_out(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    table = next(d.value for d in at.dataframe if "Best" in d.value.columns)
    assert list(table.columns) == ["Best", "Round", "Formulation", "Firmness (N)",
                                   "Juiciness (/10)", "Overall score", "Recorded",
                                   "Note"]
    assert list(table["Formulation"]) == [2, 1, 3]
    # Best is a star or nothing. That a row has no result is said in the
    # Note column, which is where a fact about the row belongs.
    assert list(table["Best"]) == ["★", "", ""]
    assert table["Note"].iloc[2] == "Not scored"
    assert at.selectbox(key="results_order").options == ["Best first",
                                                         "Newest first",
                                                         "Round order"]


def test_show_amounts_adds_the_amount_columns(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.toggle(key="show_amounts").set_value(True)
    at.run()
    table = next(d.value for d in at.dataframe if "Best" in d.value.columns)
    assert "Pea protein (g)" in table.columns


def test_download_all_formulations_is_grey(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    element = _unknown(at.main, "download_button",
                       "Download all formulations (Excel)")
    assert element.proto.type == "secondary"


def test_correcting_a_result_reports_the_change_and_the_move(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    at.number_input(key="correct_1_Firmness").set_value(6.0)
    at.run()
    _submit_button(at, "Save correction").click()
    at.run()
    assert not at.exception
    note = next(s.value for s in at.success if "corrected" in s.value)
    # One sentence for the whole correction — a row can now change its
    # amounts and its measurements in one save — and the move it caused.
    assert note == (wording.formulation_corrected(1) + " "
                    + wording.best_moved(2, 1) + " " + wording.COPY_KEPT)
    assert at.session_state["main_tab"] == wording.TAB_RESULTS    # no auto-move


def test_a_correction_can_be_closed_again(scored):
    """A selectbox opened with index=None cannot be cleared by the user, so
    the edit row would sit there for the rest of the session."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    assert "correct_1_Firmness" in [n.key for n in at.number_input]
    _submit_button(at, "Close").click()
    at.run()
    assert not at.exception
    assert "correct_1_Firmness" not in [n.key for n in at.number_input]


def test_a_correction_outside_the_range_is_refused(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    field = at.number_input(key="correct_1_Firmness")
    field.set_value(12.0)
    at.run()
    # No clamping: the box keeps the reading it was given (an unbounded
    # st.number_input reports its limits as +/-DBL_MAX, never None), and the
    # refusal comes on save.
    assert at.number_input(key="correct_1_Firmness").value == 12.0
    _submit_button(at, "Save correction").click()
    at.run()
    assert not at.exception
    assert any(e.value == ("Firmness 12 N is outside your range of 0 to 10 N. "
                           "Widen the range in Set up, or check the value.")
               for e in at.error), [e.value for e in at.error]
    assert FoodOptimizer("burger").results_history[0]["Firmness"] == 1.0


def test_a_correction_with_every_measurement_blank_changes_nothing(burger):
    """A row whose only measurement was removed has nothing left to retype,
    and its amounts are already what was recorded: saving writes nothing and
    must not claim it did."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Firmness": 5.0}, formulation_no=1, batch_no=1)
    burger.remove_objective("Firmness")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    assert at.number_input(key="correct_1_Juiciness").value is None
    _submit_button(at, "Save correction").click()
    at.run()
    assert not at.exception
    assert any(s.value == wording.formulation_unchanged(1) for s in at.success), \
        [s.value for s in at.success]


def test_a_correction_with_a_blank_amount_is_refused(scored):
    """An amount box typed empty cannot be encoded, and a formulation with a
    missing ingredient is not a correction of anything. AppTest cannot type a
    box empty, so the row opens on the state a blanked box leaves behind."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["correct_formulation"] = 1
    at.session_state["correct_amount_1_Pea protein"] = None
    at.run()
    assert at.number_input(key="correct_amount_1_Pea protein").value is None
    _submit_button(at, "Save correction").click()
    at.run()
    assert not at.exception
    assert any(e.value == wording.ENTER_EVERY_AMOUNT for e in at.error), \
        [e.value for e in at.error]
    assert FoodOptimizer("burger").recipe_history[0]["Pea protein"] == 10.0


def test_results_renders_when_the_last_measurement_is_gone(scored):
    """History outlives measurements: st.columns(0) and an empty closeness
    table must not take the tab down."""
    scored.remove_objective("Firmness")
    scored.remove_objective("Juiciness")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert any("Add a measurement in Set up" in i.value for i in at.info), \
        [i.value for i in at.info]


def test_results_foot_starts_the_next_trial(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    button = _submit_button(at, wording.START_NEXT_BATCH)
    assert button.proto.type == "primary"
    button.click()
    at.run()
    assert at.session_state["main_tab"] == wording.TAB_BATCH


def test_results_foot_points_back_at_an_unrecorded_trial(scored):
    scored.set_pending_batch([{"Pea protein": 5.0, "Methylcellulose": 0.5},
                              {"Pea protein": 6.0, "Methylcellulose": 0.6}],
                             batch_no=2)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(b.label == wording.back_to_batch_label(2, 2) for b in at.button), \
        _labels(at)


def test_results_has_exactly_one_coloured_button(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _tab_primaries(at, 2) == [wording.START_NEXT_BATCH], _tab_primaries(at, 2)


def test_the_results_foot_steps_aside_for_a_confirmation(scored):
    """A tab never shows two coloured buttons: while a confirmation is armed
    its Yes is the one lit button and the foot is grey and unclickable."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.multiselect(key="delete_formulations").set_value([1])
    at.run()
    _submit_button(at, wording.delete_formulation_button(1)).click()
    at.run()
    assert _tab_primaries(at, 2) == ["Yes, delete"], _tab_primaries(at, 2)
    foot = _submit_button(at, wording.START_NEXT_BATCH)
    assert foot.proto.type == "secondary" and foot.disabled


def test_change_setup_from_results_jumps_to_tab_1(scored):
    """Directly under the best-so-far block: the one way back to Set up from
    a formulation on screen, for a range too narrow or a limit too tight."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    button = _submit_button(at, wording.CHANGE_SETUP_FROM_RESULTS_BUTTON)
    assert button.proto.type == "secondary"   # never the foot's colour
    button.click()
    at.run()
    assert not at.exception
    assert at.session_state["main_tab"] == wording.TAB_SETUP


def test_change_setup_from_results_steps_aside_for_a_confirmation(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.multiselect(key="delete_formulations").set_value([1])
    at.run()
    _submit_button(at, wording.delete_formulation_button(1)).click()
    at.run()
    assert _tab_primaries(at, 2) == ["Yes, delete"], _tab_primaries(at, 2)
    button = _submit_button(at, wording.CHANGE_SETUP_FROM_RESULTS_BUTTON)
    assert button.proto.type == "secondary" and button.disabled


def test_the_empty_state_button_steps_aside_for_a_confirmation(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    _submit_button(at, "Start this project over").click()
    at.run()
    button = _submit_button(at, wording.MAKE_YOUR_FIRST_BATCH_BUTTON)
    assert button.proto.type == "secondary" and button.disabled


def test_results_collapsed_sections_are_the_spec_list(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    labels = [e.label for e in at.expander]
    # One section holds every way the record is fixed after the fact.
    for label in (wording.PROGRESS_CHART_EXPANDER,
                  wording.EDIT_PAST_FORMULATIONS_EXPANDER):
        assert label in labels, labels


def test_whole_batch_pick_fills_the_selection(scored):
    """The pick is a shortcut into the list beside it, never a delete of its
    own: every number the batch issued — the one nobody made included — lands
    where it can be read and changed before anything is confirmed."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert at.selectbox(key="delete_whole_batch").options == ["1"]
    at.selectbox(key="delete_whole_batch").set_value(1)
    at.run()
    assert not at.exception
    assert at.session_state["delete_formulations"] == [1, 2, 3]
    # The box empties again, so the shortcut can be used a second time.
    assert at.selectbox(key="delete_whole_batch").value is None
    assert FoodOptimizer("burger").formulation_ids == [1, 2]   # nothing gone


def test_deleting_a_whole_batch_says_it_once_and_keeps_a_copy(scored, tmp_path):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="delete_whole_batch").set_value(1)
    at.run()
    # Formulations, not results: one of the three was left out and has none.
    sentence = wording.delete_formulations_warning("1, 2 and 3")
    # The confirmation carries the sentence; a caption above it would say the
    # same thing twice.
    assert not any(c.value == sentence for c in at.caption), \
        [c.value for c in at.caption]
    _submit_button(at, "Delete Formulations 1, 2 and 3").click()
    at.run()
    assert any(w.value == sentence for w in at.warning), [w.value for w in at.warning]
    _submit_button(at, "Yes, delete").click()
    at.run()
    assert not at.exception
    # pre_delete, like every other Delete: the act is not called undo any
    # more. pre_undo stays readable by is_archive_name for older copies.
    assert (tmp_path / "burger_pre_delete.pkl").exists()
    assert not (tmp_path / "burger_pre_undo.pkl").exists()
    reloaded = FoodOptimizer("burger")
    assert reloaded.X_history == [] and reloaded.skipped == []
    assert reloaded.next_formulation_no == 4      # 1, 2 and 3 retire
    assert any(s.value == wording.formulations_deleted("1, 2 and 3")
               for s in at.success), [s.value for s in at.success]


def test_delete_several_formulations_with_one_confirmation(scored, tmp_path):
    """Three rows used to mean three confirmations and three copies of the
    project; one question now names every number that is about to go."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.multiselect(key="delete_formulations").set_value([1, 3])
    at.run()
    _submit_button(at, "Delete Formulations 1 and 3").click()
    at.run()
    assert any(w.value == ("Delete Formulations 1 and 3? Later formulations "
                           "keep their numbers. " + wording.COPY_KEPT)
               for w in at.warning), [w.value for w in at.warning]
    # One Cancel on this tab, not the confirmation's and a picker's side by
    # side: a multiselect empties itself.
    tab_labels = [b.label for b in at.tabs[2].button]
    assert tab_labels.count("Cancel") == 1, tab_labels
    _submit_button(at, "Yes, delete").click()
    at.run()
    assert not at.exception
    assert [p.name for p in tmp_path.glob("burger_pre_*.pkl")] == \
        ["burger_pre_delete.pkl"]
    reloaded = FoodOptimizer("burger")
    assert reloaded.formulation_ids == [2] and reloaded.skipped == []
    assert any(s.value == wording.formulations_deleted("1 and 3")
               for s in at.success), [s.value for s in at.success]


def test_emptying_the_delete_list_takes_the_question_down(scored):
    """A confirmation is armed by one click and answered on a later run, so
    emptying the list it was asked about left the armed flag set with nothing
    on screen to answer it — and every coloured button in the app grey behind
    a question the user could no longer reach."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.multiselect(key="delete_formulations").set_value([1, 2])
    at.run()
    _submit_button(at, "Delete Formulations 1 and 2").click()
    at.run()
    assert _tab_primaries(at, 2) == ["Yes, delete"], _tab_primaries(at, 2)
    at.multiselect(key="delete_formulations").set_value([])
    at.run()
    assert not at.exception
    # On this run, not the next one: the button that would have cleared it
    # was one of the grey ones.
    assert _tab_primaries(at, 2) == [wording.START_NEXT_BATCH], _tab_primaries(at, 2)
    assert _tab_primaries(at, 0) and _tab_primaries(at, 1), \
        (_tab_primaries(at, 0), _tab_primaries(at, 1))
    assert not _submit_button(at.sidebar, wording.START_OVER_LABEL).disabled
    assert "delete_formulations__pending" not in at.session_state
    assert FoodOptimizer("burger").formulation_ids == [1, 2]   # nothing gone


def test_deleting_a_formulation_leaves_an_open_batch_alone(scored):
    """The list only ever offers recorded and not-made numbers, so a delete
    here cannot reach the batch on the bench — and must not take it down as a
    side effect the way `Delete the last batch` could."""
    scored.set_pending_batch([{"Pea protein": 5.0, "Methylcellulose": 0.5}],
                             batch_no=2)
    open_no = int(scored.pending_batch[0]["formulation"])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert str(open_no) not in at.multiselect(key="delete_formulations").options
    at.multiselect(key="delete_formulations").set_value([1])
    at.run()
    _submit_button(at, wording.delete_formulation_button(1)).click()
    at.run()
    _submit_button(at, "Yes, delete").click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert reloaded.formulation_ids == [2]
    assert [r["formulation"] for r in reloaded.pending_batch] == [open_no]
    assert reloaded.pending_batch_no == 2


def test_a_delete_keeps_the_open_batchs_uploaded_sheet(scored):
    """A parsed batch sheet belongs to the OPEN batch, whose rows are not in
    the delete list at all; throwing it away made the user upload it again."""
    scored.set_pending_batch([{"Pea protein": 5.0, "Methylcellulose": 0.5}],
                             batch_no=2)
    sheet = pd.DataFrame({"Formulation": [int(scored.pending_batch[0]["formulation"])],
                          "Firmness": [6.0], "Juiciness": [7.0]})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_results_upload"] = sheet
    at.run()
    at.multiselect(key="delete_formulations").set_value([1])
    at.run()
    _submit_button(at, wording.delete_formulation_button(1)).click()
    at.run()
    _submit_button(at, "Yes, delete").click()
    at.run()
    assert not at.exception
    assert "_results_upload" in at.session_state
    assert list(at.session_state["_results_upload"]["Firmness"]) == [6.0]


def test_correcting_a_row_older_than_an_ingredient_says_unchanged(scored, tmp_path):
    """An ingredient added after a formulation was recorded has no entry in
    that row's amounts. The box opens at the value the model already reads it
    at, so saving an untouched row writes nothing — and must not claim a
    correction, or keep a copy of the project for one."""
    scored.add_ingredient("Oat fibre", 0, 10)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    assert at.number_input(key="correct_amount_1_Oat fibre").value == 0.0
    _submit_button(at, "Save correction").click()
    at.run()
    assert not at.exception
    assert any(s.value == wording.formulation_unchanged(1) for s in at.success), \
        [s.value for s in at.success]
    assert not (tmp_path / "burger_pre_edit.pkl").exists(), \
        [f.name for f in tmp_path.glob("*.pkl")]


def test_deleting_one_formulation_keeps_later_numbers(scored, tmp_path):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.multiselect(key="delete_formulations").set_value([1])
    at.run()
    _submit_button(at, wording.delete_formulation_button(1)).click()
    at.run()
    assert any("Later formulations keep their numbers. " + wording.COPY_KEPT
               in w.value for w in at.warning), [w.value for w in at.warning]
    _submit_button(at, "Yes, delete").click()
    at.run()
    assert not at.exception
    assert (tmp_path / "burger_pre_delete.pkl").exists()
    assert FoodOptimizer("burger").formulation_ids == [2]


def test_the_import_caption_names_this_projects_columns(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.run()
    caption = next(c.value for c in at.caption
                   if c.value.startswith("One row per formulation"))
    assert "Pea protein" in caption and "Firmness" in caption, caption


def test_importing_past_formulations_marks_them_and_accepts_partials(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.session_state["_import_rows"] = pd.DataFrame({
        "Pea protein": [12.0, 13.0], "Methylcellulose": [1.2, 1.3],
        "Juiciness": [6.0, None], "Firmness": [5.0, 5.5],
    })
    at.run()
    _submit_button(at, "Import all rows").click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert reloaded.formulation_ids == [1, 2]
    assert reloaded.batch_history == [None, None]
    assert reloaded.notes_history == [wording.IMPORTED_NOTE] * 2
    assert reloaded.results_history[1] == {"Firmness": 5.5}


def test_the_importer_reads_the_downloaded_files_own_headers(burger):
    """`Download all formulations (CSV)` heads each column with its unit, as
    the screen does. The importer reads both that spelling and the bare
    name the caption lists, so a file downloaded from this project goes
    straight back into it."""
    burger.tell({"Pea protein": 12.0, "Methylcellulose": 1.2},
                {"Juiciness": 6.0, "Firmness": 5.0}, formulation_no=1,
                batch_no=1)
    downloaded = pd.read_csv(io.StringIO(FoodOptimizer("burger").history_csv()))
    assert "Pea protein (g)" in downloaded.columns
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.session_state["_import_rows"] = downloaded
    at.run()
    assert not at.error, [e.value for e in at.error]
    _submit_button(at, "Import all rows").click()
    at.run()
    assert not at.exception
    assert not at.error, [e.value for e in at.error]
    reloaded = FoodOptimizer("burger")
    assert reloaded.formulation_ids == [1, 2]
    assert reloaded.recipe_history[1]["Pea protein"] == 12.0
    assert reloaded.results_history[1] == {"Juiciness": 6.0, "Firmness": 5.0}


def test_the_best_amounts_table_carries_each_rows_own_unit(burger):
    """One Amount column formatted per row: a cook temperature is not grams,
    and settings are listed after the ingredients, however large they are."""
    burger.add_process_parameter("Cook temperature", 0, 220)
    burger.tell({"Pea protein": 20.0, "Methylcellulose": 2.0,
                 "Cook temperature": 200.0},
                {"Juiciness": 7.0, "Firmness": 6.0},
                formulation_no=1, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    table = next(t.value for t in at.table
                 if "Ingredient or process setting" in t.value.columns)
    assert list(table.columns) == ["Ingredient or process setting", "Amount"]
    names = list(table["Ingredient or process setting"])
    assert names == ["Pea protein", "Methylcellulose", "Cook temperature"], names
    amounts = dict(zip(names, table["Amount"]))
    # Two decimals on an amount, and a setting is not an amount.
    assert amounts["Pea protein"] == "20.00 g"
    assert amounts["Cook temperature"] == "200"


def test_only_one_confirmation_can_be_armed_at_a_time(scored):
    """Two armed confirmations would put two coloured Yes buttons on the tab
    and leave the user guessing which copy is kept."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.multiselect(key="delete_formulations").set_value([1])
    at.run()
    _submit_button(at.sidebar, wording.START_OVER_LABEL).click()
    at.run()
    assert _tab_primaries(at, 2) == [], _tab_primaries(at, 2)
    assert _submit_button(at, wording.delete_formulation_button(1)).disabled


def test_a_correction_box_carries_no_tooltip_of_its_own(scored):
    """One caption above the form says it for every box on it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    assert not at.number_input(key="correct_1_Firmness").help


def test_an_import_outside_the_range_is_refused_naming_the_row(burger):
    """A 99 typed into a 0-10 column would sail in as the best formulation."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.session_state["_import_rows"] = pd.DataFrame({
        "Pea protein": [12.0, 13.0], "Methylcellulose": [1.2, 1.3],
        "Juiciness": [6.0, 6.5], "Firmness": [5.0, 99.0],
    })
    at.run()
    _submit_button(at, "Import all rows").click()
    at.run()
    assert not at.exception
    assert any(e.value == ("Row 2: Firmness 99 N is outside your range of 0 to "
                           "10 N. Widen the range in Set up, or check the "
                           "value.") for e in at.error), [e.value for e in at.error]
    assert FoodOptimizer("burger").X_history == []   # the whole file is refused


def test_an_import_outside_an_ingredient_range_warns_and_still_imports(burger):
    """An amount outside the range is a fact about work already done, not a
    mistake to refuse: the model needs it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.session_state["_import_rows"] = pd.DataFrame({
        "Pea protein": [12.0, 999.0], "Methylcellulose": [1.2, 1.3],
        "Juiciness": [6.0, 6.5], "Firmness": [5.0, 5.5],
    })
    at.run()
    _submit_button(at, "Import all rows").click()
    at.run()
    assert not at.exception
    assert any(w.value == ("Row 2: Pea protein 999 g is outside its allowed amounts of "
                           "0 to 25 g.") for w in at.warning), \
        [w.value for w in at.warning]
    assert len(FoodOptimizer("burger").X_history) == 2


def test_a_finished_import_does_not_offer_to_import_again(burger):
    """The parse used to run on every rerun, so a second click after a
    successful import recorded every row twice."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.session_state["_import_rows"] = pd.DataFrame({
        "Pea protein": [12.0, 13.0], "Methylcellulose": [1.2, 1.3],
        "Juiciness": [6.0, 6.5], "Firmness": [5.0, 5.5],
    })
    at.run()
    _submit_button(at, "Import all rows").click()
    at.run()
    assert "_import_rows" not in at.session_state
    at.run()
    assert not any(b.label == "Import all rows" for b in at.button), _labels(at)
    assert len(FoodOptimizer("burger").X_history) == 2


def test_deleting_takes_a_last_trial_that_was_entirely_left_out(burger):
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0},
                formulation_no=1, batch_no=1)
    burger.record_skipped(2, 2, {"Pea protein": 20.0, "Methylcellulose": 2.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="delete_whole_batch").set_value(2)
    at.run()
    _submit_button(at, wording.delete_formulation_button(2)).click()
    at.run()
    assert any(w.value == wording.delete_formulation_warning(2)
               for w in at.warning), \
        [w.value for w in at.warning]
    _submit_button(at, "Yes, delete").click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert reloaded.skipped == [] and reloaded.formulation_ids == [1]


def test_the_desktop_bundle_ships_every_module():
    root = pathlib.Path(APP_PATH).resolve().parent
    build = (root / "desktop" / "build_dmg.sh").read_text()
    e2e = (root / "desktop" / "test_e2e.sh").read_text()
    for module in ("ui_setup.py", "ui_batch.py", "ui_results.py"):
        assert module in build, module
        assert module in e2e, module


_FIRST_RUN_SENTENCE = ("This usually takes under a minute; "
                       "on a slow network, a few minutes.")


def _flowed(text):
    """A file's prose as one line. A sentence is wrapped differently in each
    file — Swift splits it across concatenated literals, Markdown across
    blockquote lines, Start Here across indented ones — and the sentence, not
    the wrapping, is what has to be identical everywhere."""
    import re
    joined = re.sub(r'"\s*\+\s*"', "", text)          # Swift: "a " + "b"
    joined = re.sub(r"^[>\s]+", " ", joined, flags=re.M)  # quote / indent
    return re.sub(r"\s+", " ", joined)


def test_first_run_copy_gives_the_honest_timing():
    """One sentence, verbatim, in all four places a first-run user reads it."""
    root = pathlib.Path(APP_PATH).resolve().parent
    for name in ("desktop/FoodOptimizerApp.swift", "desktop/start_here.txt",
                 "desktop/README.md", "README.md"):
        text = (root / name).read_text()
        assert _FIRST_RUN_SENTENCE in _flowed(text), name
        assert "a few minutes, up to 15" not in text, name


def test_the_last_trial_line_counts_a_trial_that_was_entirely_left_out(burger):
    """Nobody managed to make batch 2, but it is still the last batch: the line
    under the title must not fall back to batch 1."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0},
                formulation_no=1, batch_no=1)
    burger.record_skipped(2, 2, {"Pea protein": 20.0, "Methylcellulose": 2.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert any(c.value == wording.batch_line_recorded(2) for c in at.caption), \
        [c.value for c in at.caption]


def test_a_stale_trial_is_reported_above_the_tabs_not_in_the_sidebar(open_batch,
                                                                     tmp_path):
    """The notice is about the batch the user is looking at, so it belongs in
    the fixed message container above the tabs."""
    import json
    notice = wording.batch_discarded_notice(1)
    state = json.loads((tmp_path / "burger.pkl").read_text())
    for row in state["pending_batch"]:
        row["recipe"].pop("Methylcellulose")
    (tmp_path / "burger.pkl").write_text(json.dumps(state))
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    assert not at.exception
    assert any(i.value == notice for i in at.info), [i.value for i in at.info]
    assert not any(i.value == notice for i in at.sidebar.info), \
        [i.value for i in at.sidebar.info]
    assert FoodOptimizer("burger").pending_batch is None


def test_advanced_and_limit_forms_do_not_follow_you_to_another_project(burger):
    """prop_ and bo_ keys belong to the project they were typed in, exactly as
    the measurement and ingredient forms do."""
    FoodOptimizer("second").save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    at.session_state["bo_json"] = '{"kernel": "matern52"}'
    at.session_state["prop_min"] = 3.0
    at.run()
    at.sidebar.selectbox(key="project_select").select("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert "bo_json" not in at.session_state
    # A box that is still on screen in the new project is emptied rather than
    # popped: popping does not reach the browser, which posts the old value
    # straight back on the next click.
    assert at.session_state["prop_min"] is None


def test_edit_past_formulations_does_not_follow_you_to_another_project(scored):
    """A half-typed past formulation, a selection waiting for a confirmation
    and a correction in progress all belong to the project they were typed
    in. Popping alone does not reach the browser, so the boxes that are still
    on screen are emptied rather than dropped."""
    second = FoodOptimizer("second")
    second.set_amount_unit("g")
    second.add_ingredient("Pea protein", 0, 25)     # the same boxes, empty
    second.add_objective("Firmness", 1.0, goal="max", min_val=0, max_val=10)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.multiselect(key="delete_formulations").set_value([2])
    at.number_input(key="past_Pea protein").set_value(12.0)
    at.text_input(key="past_note").set_value("from the 2024 bench book")
    at.radio(key="add_past_mode").set_value(wording.UPLOAD_A_FILE)
    at.run()
    assert at.session_state["correct_amount_1_Pea protein"] == 10.0
    at.sidebar.selectbox(key="project_select").select("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert at.session_state["correct_formulation"] is None
    assert at.number_input(key="past_Pea protein").value is None
    assert at.text_input(key="past_note").value == wording.IMPORTED_NOTE
    assert at.radio(key="add_past_mode").value == "Type it in"
    # The correction amounts and the delete list have no widget left in this
    # project, so the parked empty value is all there is to check.
    for key in ("correct_amount_1_Pea protein", "delete_formulations"):
        assert key not in at.session_state or not at.session_state[key], key


def test_a_correction_keeps_a_copy_first_and_says_so(scored, tmp_path):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    at.number_input(key="correct_1_Firmness").set_value(6.0)
    at.run()
    _submit_button(at, "Save correction").click()
    at.run()
    assert not at.exception
    assert (tmp_path / "burger_pre_edit.pkl").exists(), \
        [p.name for p in tmp_path.glob("*.pkl")]
    assert any(s.value == (wording.formulation_corrected(1) + " "
                           + wording.best_moved(2, 1) + " " + wording.COPY_KEPT)
               for s in at.success), [s.value for s in at.success]
    assert FoodOptimizer("burger").results_history[0]["Firmness"] == 6.0


def test_a_no_op_correction_says_so_and_keeps_no_copy(scored, tmp_path):
    """Saving a row nobody changed writes nothing, so it must not claim a copy
    was kept — or keep one."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    _submit_button(at, "Save correction").click()      # the boxes hold what
    at.run()                                           # was already recorded
    assert not at.exception
    assert any(s.value == wording.formulation_unchanged(1) for s in at.success), \
        [s.value for s in at.success]
    assert not (tmp_path / "burger_pre_edit.pkl").exists(), \
        [p.name for p in tmp_path.glob("*.pkl")]


def _type_a_past_formulation(at):
    """Half of `Add a formulation you already made`, typed in."""
    at.number_input(key="past_Pea protein").set_value(12.0)
    at.number_input(key="past_m_Firmness").set_value(4.0)
    at.text_input(key="past_note").set_value("from the old notebook")
    at.run()


def test_the_whole_batch_pick_keeps_the_typed_past_formulation(scored):
    """The pick only fills the list beside it, but it reruns to do so — and
    the form it is drawn above went blank, because Streamlit discards every
    widget the run did not reach."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    _type_a_past_formulation(at)
    at.selectbox(key="delete_whole_batch").set_value(1)
    at.run()
    assert not at.exception
    assert at.session_state["delete_formulations"] == [1, 2, 3]
    assert at.session_state["past_Pea protein"] == 12.0
    assert at.session_state["past_m_Firmness"] == 4.0
    assert at.session_state["past_note"] == "from the old notebook"


def test_emptying_the_delete_list_keeps_the_typed_past_formulation(scored):
    """Taking the last number back out disarms the question and redraws, and
    that redraw emptied the form below it too."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.multiselect(key="delete_formulations").set_value([2])
    at.run()
    _submit_button(at, wording.delete_formulation_button(2)).click()
    at.run()
    _type_a_past_formulation(at)
    at.multiselect(key="delete_formulations").set_value([])
    at.run()
    assert not at.exception
    assert "_armed_confirmation" not in at.session_state
    assert at.session_state["past_Pea protein"] == 12.0
    assert at.session_state["past_note"] == "from the old notebook"


@pytest.mark.parametrize("label", [wording.START_OVER_LABEL,
                                   wording.DELETE_PROJECT_LABEL])
def test_a_sidebar_cancel_leaves_every_tab_form_as_it_was(scored_open_batch,
                                                          label):
    """The sidebar runs before all three tabs, so a Cancel up there reruns
    without ever reaching them — and took the result grid, the formulation
    of your own and the open correction with it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="f4_Firmness").set_value(5.0)
    at.text_input(key="f4_note").set_value("second try")
    at.number_input(key="own_Pea protein").set_value(15.0)
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    at.number_input(key="correct_1_Firmness").set_value(3.0)
    at.run()
    _submit_button(at, label).click()
    at.run()
    _submit_button(at, wording.CANCEL).click()
    at.run()
    assert not at.exception
    assert at.session_state["f4_Firmness"] == 5.0
    assert at.session_state["f4_note"] == "second try"
    assert at.session_state["own_Pea protein"] == 15.0
    assert at.session_state["correct_formulation"] == 1
    assert at.session_state["correct_1_Firmness"] == 3.0


def test_a_sidebar_cancel_keeps_the_formulation_total_and_the_batch_size(
        open_batch):
    """Losing the formulation total is not a blank box: the batch table and
    the printed sheets silently go back to as-generated, and the bench weighs
    out different numbers from the ones that were on screen a click ago."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="scale_total").set_value(25.0)
    at.run()
    assert any(c.value == "Sheets show each formulation made to 25 g."
               for c in at.caption), [c.value for c in at.caption]
    _submit_button(at, wording.START_OVER_LABEL).click()
    at.run()
    _submit_button(at, wording.CANCEL).click()
    at.run()
    assert not at.exception
    assert at.session_state["scale_total"] == 25.0
    assert any(c.value == "Sheets show each formulation made to 25 g."
               for c in at.caption), [c.value for c in at.caption]


def test_a_sidebar_cancel_keeps_how_many_formulations_to_generate(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="how_many").set_value(7)
    at.run()
    _submit_button(at, wording.DELETE_PROJECT_LABEL).click()
    at.run()
    _submit_button(at, wording.CANCEL).click()
    at.run()
    assert not at.exception
    assert at.session_state["how_many"] == 7
    assert _submit_button(at, wording.generate_button_label(7)).label == \
        wording.generate_button_label(7)


def test_the_downloaded_file_imports_whole_minus_the_rows_nobody_made(burger):
    """`Download all formulations (CSV)` carries the formulations nobody made
    — they have a number, their amounts and their note. They have no result
    to teach the model, so the import leaves them out and says so, rather
    than stopping at the first one and saving half the file."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    burger.tell({"Pea protein": 12.0, "Methylcellulose": 1.5},
                {"Firmness": 5.0}, formulation_no=2, batch_no=1)   # partial
    burger.record_skipped(3, 1, {"Pea protein": 14.0, "Methylcellulose": 2.0})
    burger.record_skipped(4, 1, {"Pea protein": 16.0, "Methylcellulose": 2.5})
    downloaded = pd.read_csv(io.StringIO(FoodOptimizer("burger").history_csv()))
    assert len(downloaded) == 4

    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.session_state["_import_rows"] = downloaded
    at.run()
    _submit_button(at, "Import all rows").click()
    at.run()
    assert not at.exception
    assert not at.error, [e.value for e in at.error]
    assert any(m.value == ("Imported 2 formulations. 2 rows had no "
                           "measurements, so they were not imported.")
               for m in at.success), [m.value for m in at.success]
    reloaded = FoodOptimizer("burger")
    assert reloaded.formulation_ids == [1, 2, 5, 6]
    assert reloaded.recipe_history[2]["Pea protein"] == 10.0
    assert reloaded.results_history[3] == {"Firmness": 5.0}     # partial kept


def test_one_row_with_nothing_measured_is_said_in_the_singular(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.session_state["_import_rows"] = pd.DataFrame({
        "Pea protein": [12.0, 13.0], "Methylcellulose": [1.2, 1.3],
        "Juiciness": [6.0, None], "Firmness": [5.0, None],
    })
    at.run()
    _submit_button(at, "Import all rows").click()
    at.run()
    assert not at.exception
    assert any(m.value == ("Imported 1 formulation. 1 row had no "
                           "measurements, so it was not imported.")
               for m in at.success), [m.value for m in at.success]


def test_the_open_batchs_recorded_rows_are_not_offered_for_deletion(scored):
    """A batch recorded one sheet at a time keeps its recorded rows while the
    batch stays open. Deleting one took the result and put the row straight
    back on tab 2's grid — a delete that puts work back on the bench."""
    scored.set_pending_batch([{"Pea protein": 15.0, "Methylcellulose": 2.0},
                              {"Pea protein": 18.0, "Methylcellulose": 2.5}])
    assert scored.pending_batch_no == 2
    scored.tell({"Pea protein": 15.0, "Methylcellulose": 2.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=4,
                batch_no=2)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception
    # Batch 1 is finished and every one of its rows is offered; batch 2 is
    # on the bench, so neither its recorded row nor its number is.
    assert list(at.multiselect(key="delete_formulations").proto.options) == \
        ["1", "2", "3"]
    assert list(at.selectbox(key="delete_whole_batch").proto.options) == ["1"]


def test_a_batch_of_your_own_formulations_says_how_to_add_generated_ones(burger):
    """Adding one with nothing open opens the batch, and from then on there
    is no Generate control on the screen at all."""
    burger.add_to_pending_batch({"Pea protein": 10.0, "Methylcellulose": 1.0},
                                note=wording.OWN_FORMULATION_NOTE)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    assert not at.exception
    assert any(c.value == wording.ONLY_OWN_FORMULATIONS_CAPTION
               for c in at.caption), [c.value for c in at.caption]
    # A generated row in the batch takes the caption away again.
    burger.set_pending_batch(
        list(burger.pending_batch) + [{"Pea protein": 20.0,
                                       "Methylcellulose": 2.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    assert not any(c.value == wording.ONLY_OWN_FORMULATIONS_CAPTION
                   for c in at.caption)


def test_a_confirmation_takes_the_colour_from_save_correction(scored):
    """Two coloured buttons on one tab, even for one frame: arming a
    confirmation below the correction row must take the colour from it on the
    same run, which is why the button row is drawn last."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    assert _tab_primaries(at, 2) == ["Save correction"], _tab_primaries(at, 2)
    at.multiselect(key="delete_formulations").set_value([2])
    at.run()
    _submit_button(at, wording.delete_formulation_button(2)).click()
    at.run()
    assert _tab_primaries(at, 2) == ["Yes, delete"], _tab_primaries(at, 2)
    save = _submit_button(at, "Save correction")
    assert save.proto.type == "secondary" and save.disabled
    _submit_button(at, "Cancel").click()
    at.run()
    assert _tab_primaries(at, 2) == ["Save correction"], _tab_primaries(at, 2)


def test_save_correction_is_the_one_lit_action_while_the_row_is_open(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert _tab_primaries(at, 2) == [wording.START_NEXT_BATCH], _tab_primaries(at, 2)
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    assert _tab_primaries(at, 2) == ["Save correction"], _tab_primaries(at, 2)
    foot = _submit_button(at, wording.START_NEXT_BATCH)
    assert foot.proto.type == "secondary" and foot.disabled


def test_a_half_typed_measurement_does_not_follow_you_to_another_project(burger):
    """Streamlit keeps a widget's value under its key for the whole session, so
    without a clean-up on switch, project B's form opens holding A's typing."""
    FoodOptimizer("second").save()              # somewhere to switch to
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _grid_edits(at, MEAS_GRID, added=[_measurement_row("Chewiness")])
    at.run()
    assert at.session_state["measurement_grid_0"]["added_rows"]
    at.sidebar.selectbox(key="project_select").select("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert at.session_state["_loaded_project"] == "second"
    # The grid's key turned over, so the half-typed row is on no widget the
    # new project draws: a data editor's value cannot be assigned at all, so
    # a fresh editor is the only way to empty one.
    assert at.session_state["_measurement_grid_nonce"] == 1
    assert wording.SAVE_CHANGES_BUTTON not in _labels(at), _labels(at)


def test_only_one_confirmation_can_be_armed_on_the_set_up_tab(burger):
    """Two armed confirmations would put two coloured Yes buttons on one tab,
    each quietly keeping its own copy of the project."""
    burger.add_process_parameter("Cook temp", 150, 200)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    # The ingredients grid renders first, so arming ITS deletion is the case
    # that has to hold within one render: confirmation_open() is read by
    # every later site, the measurements grid's Save among them.
    edits = {ING_GRID: dict(deleted=[0]),
             MEAS_GRID: dict(deleted=[0])}
    for base, e in edits.items():
        _grid_edits(at, base, **e)
    at.run()
    assert not at.exception
    assert not _grid_save(at, ING_GRID).disabled
    _grid_save(at, ING_GRID).click()
    for base, e in edits.items():
        _grid_edits(at, base, **e)
    at.run()
    assert _grid_save(at, MEAS_GRID).disabled, "two armed questions at once"
    assert _tab_primaries(at, 0) == ["Yes, delete"], _tab_primaries(at, 0)


def test_an_uploaded_sheet_stops_at_the_first_row_that_did_not_save(open_batch,
                                                                    monkeypatch):
    """Half a sheet on disk under a green 'Round 1 recorded.' is the worst
    outcome: the save loop must stop and say so at the first failure."""
    import storage as storage_backend

    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.session_state["_results_upload"] = pd.DataFrame(
        {"Formulation": [1, 2], "Firmness": [6.0, 4.0],
         "Juiciness": [7.0, 5.0]})
    at.run()
    real_save = at.session_state["optimizer"].storage.save
    calls = []

    def flaky(name, state):
        calls.append(name)
        if len(calls) == 2:
            raise storage_backend.StorageError(
                "The disk is full — your last change was NOT saved.")
        return real_save(name, state)

    at.session_state["optimizer"].storage.save = flaky
    _submit_button(at, "Save uploaded results").click()
    at.run()
    assert not at.exception
    assert any("NOT saved" in e.value for e in at.error), [e.value for e in at.error]
    assert not any("recorded" in s.value for s in at.success), \
        [s.value for s in at.success]
    # One row reached the disk; the second never did, and the batch stays open.
    reloaded = FoodOptimizer("burger")
    assert reloaded.formulation_ids == [1]
    assert at.session_state["main_tab"] == wording.TAB_BATCH


# ------------------------------------------------------------------ #
#  Verification round: the fixes the four-agent pass asked for
# ------------------------------------------------------------------ #

def test_the_line_under_the_title_is_on_tab_one_only(open_batch):
    """Tab 2 carries the batch's own heading. The caption sat directly above
    `Batch 1 · make these 2 formulations`, saying the same thing twice."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert any(c.value == wording.batch_line_open(1, 2) for c in at.tabs[0].caption), \
        [c.value for c in at.tabs[0].caption]
    assert not any(c.value == wording.batch_line_open(1, 2) for c in at.tabs[1].caption), \
        [c.value for c in at.tabs[1].caption]
    assert any(m.value == wording.make_these(1, 2)
               for m in at.tabs[1].markdown)


def test_a_hard_reset_leaves_the_project_open_and_in_the_list(project_with_history,
                                                              tmp_path):
    """The archive is a copy of what was there; the project itself stays,
    empty, under its own name. It used to be renamed away, so it vanished
    from Open project while the sidebar still named it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at, "Start this project over").click()
    at.run()
    _submit_button(at, "Yes, start over").click()
    at.run()
    assert not at.exception
    assert (tmp_path / "my_project_archived.pkl").exists()
    assert (tmp_path / "my_project.pkl").exists(), list(tmp_path.iterdir())
    assert FoodOptimizer("my_project").X_history == []
    assert at.session_state["_loaded_project"] == "my_project"
    assert at.sidebar.selectbox(key="project_select").value == "my_project"
    # One more run: AppTest keeps the answered confirmation's buttons in its
    # local element tree until the next full render.
    at.run()
    # Nothing in the sidebar is lit: the project on screen is the one open.
    assert [b.label for b in at.sidebar.button if b.proto.type == "primary"] == []
    assert _submit_button(at.sidebar, "Open").disabled


def test_every_confirmation_counts_the_left_out_formulations(scored):
    """Two scored and one left out is three formulations to archive."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at, "Start this project over").click()
    at.run()
    assert any("Its 3 formulations, ingredients and measurements all go"
               in w.value
               for w in at.warning), [w.value for w in at.warning]
    _submit_button(at, "Cancel").click()
    at.run()
    _submit_button(at, "Delete this project").click()
    at.run()
    assert any("and its 3 formulations?" in w.value for w in at.warning), \
        [w.value for w in at.warning]


def test_save_changes_is_the_lit_action_while_a_grid_has_an_edit_in_it(burger):
    """Next: make a round would leave the tab and throw the edit away,
    exactly as tab 3's correction row already refuses to let happen."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, MEAS_GRID, edited={0: {wording.SHARE_COLUMN: 70.0}})
    at.run()
    assert _tab_primaries(at, 0) == ["Save changes"], _tab_primaries(at, 0)
    assert _submit_button(at, wording.NEXT_MAKE_BATCH_BUTTON).disabled


def test_only_the_topmost_unsaved_grid_lights_its_save(burger):
    """Two grids can be mid-edit at once, and a tab shows one coloured
    button. The one higher up the page is the one the reader is looking
    at."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, edited={0: {wording.HIGHEST_LABEL: 40.0}})
    _grid_edits(at, MEAS_GRID, edited={0: {wording.SHARE_COLUMN: 70.0}})
    at.run()
    assert _tab_primaries(at, 0) == ["Save changes"], _tab_primaries(at, 0)
    assert _grid_save(at, ING_GRID).proto.type == "primary"
    assert _grid_save(at, MEAS_GRID).proto.type == "secondary"


def test_saving_one_grid_leaves_the_other_s_edit_where_it_was(burger):
    """The cold reader edited both tables, clicked the Ingredients Save,
    and watched BOTH banners and both Save buttons go while the
    Measurements table still showed the 70 they had typed and the project
    still held 60. Each grid's banner, pending edit and Save are its own."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    ing = {0: {wording.HIGHEST_LABEL: 40.0}}
    meas = {0: {wording.SHARE_COLUMN: 70.0}}
    _grid_edits(at, ING_GRID, edited=ing)
    _grid_edits(at, MEAS_GRID, edited=meas)
    at.run()
    captions = [c.value for c in at.caption]
    assert "Ingredients and process settings — not saved yet." in captions
    assert "Measurements and targets — not saved yet." in captions
    # Save the ingredients only. Its rerun happens above the measurements
    # grid, so the browser's record for that grid is dropped — the frame it
    # parked is what has to come back.
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, edited=ing)
    _grid_edits(at, MEAS_GRID, edited=meas)
    at.run()
    assert not at.exception
    saved = FoodOptimizer("burger")
    assert saved._var_by_name("Pea protein")["bounds"] == (0.0, 40.0)
    # The measurements grid still has its edit, its banner and its Save —
    # and the project has not seen it.
    captions = [c.value for c in at.caption]
    assert "Ingredients and process settings — not saved yet." not in captions
    assert "Measurements and targets — not saved yet." in captions, captions
    assert list(_grid_frame(at, 1)[wording.SHARE_COLUMN])[0] == 70.0
    assert saved.share_percents() == {"Firmness": 60, "Juiciness": 40}
    # One coloured button on the tab: the measurements Save, now that the
    # grid above it has nothing left to save.
    assert _tab_primaries(at, 0) == [wording.SAVE_CHANGES_BUTTON], \
        _tab_primaries(at, 0)
    assert _grid_save(at, MEAS_GRID).proto.type == "primary"
    # And it lands when it is clicked.
    _grid_save(at, MEAS_GRID).click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").share_percents() == {"Firmness": 70,
                                                        "Juiciness": 30}


def test_discard_changes_puts_the_grid_back(burger):
    """Discard turns the grid's key over, which is the only way to drop what
    the browser is holding: a data editor's value cannot be assigned from
    session state at all. The stale record is handed back afterwards, as the
    browser would hand it back — the new editor must not read it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, edited={0: {wording.HIGHEST_LABEL: 40.0}})
    at.run()
    _grid_discard(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, edited={0: {wording.HIGHEST_LABEL: 40.0}})
    at.run()
    assert not at.exception
    assert at.session_state["_ingredient_grid_nonce"] == 1
    _stale_grid_edits(at, ING_GRID, 0, edited={0: {wording.HIGHEST_LABEL: 40.0}})
    at.run()
    assert at.session_state["_ingredient_grid_nonce"] == 1
    assert wording.SAVE_CHANGES_BUTTON not in _labels(at), _labels(at)
    assert list(_grid_frame(at, 0)[wording.HIGHEST_LABEL]) == [25.0, 3.0]
    assert FoodOptimizer("burger")._var_by_name("Pea protein")["bounds"] == (
        0.0, 25.0)


def test_saving_the_ingredients_grid_turns_its_key_over(burger):
    """A save that left the record standing would add the row a second time
    on the next run, or write the same cell again."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={0: {wording.HIGHEST_LABEL: 40.0}})
    assert not at.exception
    assert at.session_state["_ingredient_grid_nonce"] == 1
    _stale_grid_edits(at, ING_GRID, 0, edited={0: {wording.HIGHEST_LABEL: 60.0}})
    at.run()
    assert wording.SAVE_CHANGES_BUTTON not in _labels(at), _labels(at)
    assert FoodOptimizer("burger")._var_by_name("Pea protein")["bounds"] == (
        0.0, 40.0)


def test_saving_the_measurements_grid_turns_its_key_over(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, MEAS_GRID, edited={0: {wording.HIGHEST_MEASURABLE_LABEL: 12.0}})
    assert not at.exception
    assert at.session_state["_measurement_grid_nonce"] == 1
    _stale_grid_edits(at, MEAS_GRID, 0,
                      edited={0: {wording.HIGHEST_MEASURABLE_LABEL: 99.0}})
    at.run()
    assert wording.SAVE_CHANGES_BUTTON not in _labels(at), _labels(at)
    # Row 0 of the measurements grid is the largest share: Firmness.
    saved = {o["name"]: o for o in FoodOptimizer("burger").objectives}
    assert saved["Firmness"]["max_val"] == 12.0
    assert saved["Juiciness"]["max_val"] == 10.0


def test_opening_a_saved_copy_turns_the_grids_over(project_with_history):
    """The reviewer's reproduction: a Lowest typed against Water in one
    project, and a saved copy of another project opened over it. The data
    editor's record is positional — "row 0's Lowest changed" — so left
    standing it writes that number onto whatever the copy puts in row 0."""
    donor = FoodOptimizer("donor")
    donor.add_ingredient("Cocoa", 0, 40)
    donor.add_ingredient("Butter", 0, 60)
    donor.add_objective("Snap", 1.0, goal="max", min_val=0, max_val=10)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    typed = {"edited": {0: {wording.LOWEST_LABEL: 7.0}}}
    _grid_edits(at, ING_GRID, **typed)
    at.run()
    assert wording.SAVE_CHANGES_BUTTON in _labels(at), _labels(at)
    at.session_state["_restore_candidate"] = donor.export_json()
    _grid_edits(at, ING_GRID, **typed)
    at.run()
    _submit_button(at, wording.YES_REPLACE).click()
    _grid_edits(at, ING_GRID, **typed)
    at.run()
    assert not at.exception
    assert at.session_state["_ingredient_grid_nonce"] == 1
    _stale_grid_edits(at, ING_GRID, 0, **typed)
    at.run()
    assert wording.SAVE_CHANGES_BUTTON not in _labels(at), _labels(at)
    assert FoodOptimizer("my_project")._var_by_name("Cocoa")["bounds"] == (
        0.0, 40.0)


def test_replacing_the_ingredients_from_a_file_turns_the_grids_over(
        burger, monkeypatch):
    """The file replaces the list outright, so every row on the tab belongs
    to a different ingredient afterwards. AppTest cannot drive a file
    uploader, so the ingredients one is stood in for."""
    import streamlit as st_module
    csv = (b"Name,Lowest,Highest,Unit\n"
           b"Cocoa,0,40,g\nButter,0,60,g\n")
    real_uploader = st_module.file_uploader

    def _uploader(label, *a, **k):
        if str(k.get("key") or "").startswith("ingredients_file"):
            return _FakeUpload("ingredients.csv", csv)
        return real_uploader(label, *a, **k)

    monkeypatch.setattr(st_module, "file_uploader", _uploader)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    typed = {"edited": {0: {wording.LOWEST_LABEL: 7.0}}}
    _grid_edits(at, ING_GRID, **typed)
    at.run()
    assert wording.SAVE_CHANGES_BUTTON in _labels(at), _labels(at)
    next(b for b in at.button if b.key == "load_ingredients").click()
    _grid_edits(at, ING_GRID, **typed)
    at.run()
    assert not at.exception
    assert at.session_state["_ingredient_grid_nonce"] == 1
    _stale_grid_edits(at, ING_GRID, 0, **typed)
    at.run()
    assert wording.SAVE_CHANGES_BUTTON not in _labels(at), _labels(at)
    assert FoodOptimizer("burger")._var_by_name("Cocoa")["bounds"] == (0.0, 40.0)


def test_the_empty_results_button_sends_an_unready_project_to_set_up(tmp_path,
                                                                     monkeypatch):
    """A project with no ingredients cannot make a batch: the lit button used
    to land on a tab holding a greyed Generate."""
    monkeypatch.chdir(tmp_path)
    FoodOptimizer("empty").save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "empty"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not any(b.label == wording.MAKE_YOUR_FIRST_BATCH_BUTTON for b in at.button), \
        _labels(at)
    _submit_button(at, "Set up this project").click()
    at.run()
    assert not at.exception
    assert at.session_state["main_tab"] == wording.TAB_SETUP


def test_a_project_saved_before_units_existed_says_the_g_is_a_guess(tmp_path,
                                                                    monkeypatch):
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("legacy")
    opt.add_ingredient("Water", 0, 100)
    opt.add_objective("Taste", 1.0, goal="max")
    state = opt.export_json()
    state.pop("amount_unit")                 # a 0.2.x file has no unit at all
    opt.storage.save("legacy", state)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "legacy"
    at.run()
    assert not at.exception
    assert any(c.value == ("These amounts were recorded before units "
                           "existed, so they are read as g. Set the right "
                           "unit below.")
               for c in at.caption), [c.value for c in at.caption]
    # ...and it is a caption like every other on the tab: one line, said once.
    _assert_captions_read_once(at)


def test_an_imported_formulation_can_still_be_deleted(burger):
    """A row that belongs to no batch used to face a section whose first
    control was `Delete the last batch`; it is in the same list as every
    other number now."""
    burger.import_formulation({"Pea protein": 10.0, "Methylcellulose": 1.0},
                              {"Juiciness": 7.0, "Firmness": 6.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    assert at.multiselect(key="delete_formulations").options == ["1"]
    # No batch to pick from: the row was made before this project existed.
    assert at.selectbox(key="delete_whole_batch").options == []


def test_a_process_setting_carries_its_own_unit_everywhere(burger):
    """Sheets printed `Cook temperature: 180` and the batch column was bare:
    a setting is not an amount, so it needs a unit of its own."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, added=[
        _ingredient_row("Cook temperature", kind=wording.KIND_SETTING,
                        low=160.0, high=200.0, unit="°C")])
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    setting = next(v for v in reloaded.variables if v.get('category') == 'process')
    assert setting["unit"] == "°C"
    grid = _grid_frame(at, 0)
    row = grid[grid["Name"] == "Cook temperature"].iloc[0]
    assert (row["Type"], row["Lowest"], row["Highest"], row["Unit"]) == \
        ("Process setting", 160.0, 200.0, "°C")
    # ... on the batch table, the printable sheet and the amounts table.
    reloaded.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0,
                                 "Cook temperature": 180.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    table = next(d.value for d in at.dataframe if "Formulation" in d.value.columns)
    assert "Cook temperature (°C)" in table.columns, list(table.columns)
    assert ("Cook temperature (°C)", 180.0) in _sheet_amounts(at), \
        _sheet_amounts(at)


def test_the_baseline_of_a_process_setting_is_shown_with_its_unit(burger):
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1, batch_no=1)
    burger.add_process_parameter("Cook temperature", 160, 200, baseline=180,
                                 unit="°C")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    # Baseline is a NUMBER column on the grid — it is typed into, and a
    # setting added mid-run must carry one — so the unit rides in the row's
    # own Unit cell rather than inside the figure.
    grid = _grid_frame(at, 0)
    baselines = dict(zip(grid["Name"], grid[wording.BASELINE_LABEL]))
    assert baselines["Cook temperature"] == 180.0
    assert all(pd.isna(baselines[n])
               for n in ("Pea protein", "Methylcellulose"))
    row = grid[grid["Name"] == "Cook temperature"].iloc[0]
    assert row[wording.UNIT_LABEL] == "°C"


def test_the_limits_caption_covers_a_property_named_in_the_app(with_properties):
    """A property is named in the app as often as it arrives in a file, and
    the grid that names one is inside the same fold as this caption.

    It no longer carries the blank-figure rule: the properties grid a few
    lines below says it, and each limit's own line names the ingredients it
    is reading as zeroes."""
    with_properties.add_property("Sodium mg per 100 g")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(c.value == ("Limits are hard rules for every formulation the "
                           "app suggests. A formulation of your own is "
                           "recorded as you typed it. A limit is on an amount "
                           "you weigh out or a property of your ingredients; "
                           "measurements have goals and targets instead.")
               for c in at.caption), [c.value for c in at.caption]
    said = [c.value for c in at.caption if "counts as 0 in any limit" in c.value]
    assert said == [wording.properties_grid_caption(True)], said


def test_a_project_with_no_properties_still_offers_to_name_one(burger):
    """It used to send the user off to build a CSV; a property can now be
    named here, so the section offers the box instead of an errand."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not any("Upload an ingredient CSV with extra columns" in c.value
                   for c in at.caption), [c.value for c in at.caption]
    assert at.text_input(key="prop_new").label == \
        "Add a property"
    # Nothing to limit yet, so no property picker and no limit button.
    assert "prop_metric" not in [b.key for b in at.selectbox]
    assert "Add property limit" not in _labels(at)


def test_the_sample_project_is_created_not_opened(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at.main, "Try the sample project").click()
    at.run()
    assert not at.exception
    assert any(s.value == "Created Sample project." for s in at.success), \
        [s.value for s in at.success]
    # And opening it again later is an opening.
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at.sidebar, "Try the sample project").click()
    at.run()
    assert any(s.value == "Opened Sample project." for s in at.success), \
        [s.value for s in at.success]


def test_a_partly_uploaded_sheet_says_what_it_recorded(open_batch):
    """The whole-batch sentence over a batch that is still open claimed work
    nobody had done."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.session_state["_results_upload"] = pd.DataFrame(
        {"Formulation": [2], "Firmness": [6.0], "Juiciness": [7.0]})
    at.run()
    _submit_button(at, "Save uploaded results").click()
    at.run()
    assert not at.exception
    assert any(s.value == wording.upload_partial_flash(1, 2, 1, 1)
               for s in at.success), [s.value for s in at.success]
    assert at.session_state["main_tab"] == wording.TAB_BATCH
    assert FoodOptimizer("burger").pending_batch is not None


def test_an_edit_with_nothing_recorded_claims_no_recalculation(burger):
    """There is no overall score to recalculate, and saying otherwise
    invents a history. No copy either: nothing could have been lost."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, MEAS_GRID, edited={0: {wording.SHARE_COLUMN: 70.0}})
    assert not at.exception
    assert any(s.value == "Firmness saved." for s in at.success), \
        [s.value for s in at.success]


def test_a_half_typed_ingredient_does_not_follow_you_to_another_project(burger):
    """A grid's session-state value is the record of what was typed into it,
    and Streamlit refuses to let a script assign one, so the only way to
    empty it is a fresh editor: opening another project turns the key over.
    Left standing, the record would have applied burger's half-typed row to
    the project next opened."""
    FoodOptimizer("second").save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _grid_edits(at, ING_GRID,
                added=[_ingredient_row("Half typed ingredient", high=7.0)])
    at.run()
    assert wording.SAVE_CHANGES_BUTTON in _labels(at)
    at.sidebar.selectbox(key="project_select").select("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert at.session_state["_ingredient_grid_nonce"] == 1
    assert wording.SAVE_CHANGES_BUTTON not in _labels(at), _labels(at)
    assert FoodOptimizer("second").variables == []


def test_the_result_grid_does_not_follow_you_to_another_project(open_batch):
    """f1_Firmness is a live key in every project that has a formulation 1."""
    FoodOptimizer("second").save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="f1_Firmness").set_value(9.0)
    at.text_input(key="f1_note").set_value("half typed")
    at.run()
    at.sidebar.selectbox(key="project_select").select("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert at.session_state["f1_Firmness"] is None
    assert at.session_state["f1_note"] == ""


def test_get_help_gives_an_address_to_write_to():
    """A user who threw away the disk image has no other route."""
    root = pathlib.Path(APP_PATH).resolve().parent
    swift = _flowed((root / "desktop" / "FoodOptimizerApp.swift").read_text())
    assert "https://github.com/thomasat/food_opt/issues" in swift
    assert ("Describe the problem in words, and do not attach project files, "
            "saved copies or formulations, because that page is public.") in swift


def test_emptying_the_forms_raises_no_widget_warning_on_screen(open_batch):
    """Streamlit puts a yellow box on the page when a widget carries both a
    `value=` and a session-state entry. Parking empty values across a project
    switch touches a dozen widgets, so none of them may declare a default."""
    FoodOptimizer("second").save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="f1_Firmness").set_value(9.0)
    at.number_input(key="scale_total").set_value(50.0)
    at.run()
    at.sidebar.selectbox(key="project_select").select("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    noisy = [w.value for w in at.warning if "Session State API" in w.value]
    assert noisy == [], noisy
    assert at.session_state["scale_total"] is None


def test_a_deleted_project_leaves_the_sidebar_box(project_with_history):
    """The box kept the deleted project's name — popping a widget's key does
    not reach the browser — so Open lit up over a project that was gone."""
    FoodOptimizer("keeper").save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.run()
    _submit_button(at, "Delete this project").click()
    at.run()
    _submit_button(at, "Yes, delete it").click()
    at.run()
    assert not at.exception
    assert at.session_state["_loaded_project"] == "keeper"
    assert at.sidebar.selectbox(key="project_select").value == "keeper"
    at.run()      # AppTest keeps the answered confirmation's buttons a run
    assert [b.label for b in at.sidebar.button if b.proto.type == "primary"] == []


def test_a_process_setting_is_rounded_wherever_it_is_shown(burger):
    """188.494 °C is a precision no oven dial has. The batch table, the
    printable sheet and the amounts table must round it the same way."""
    burger.add_process_parameter("Cook temperature", 100, 220, unit="°C")
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0,
                               "Cook temperature": 188.4936}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    shown = _displayed(next(d for d in at.dataframe
                            if "Formulation" in d.value.columns))
    assert shown["Cook temperature (°C)"].iloc[0] == "188.49", shown.to_dict()
    assert ("Cook temperature (°C)", 188.49) in _sheet_amounts(at), \
        _sheet_amounts(at)
    # ... and on tab 3, once it is recorded.
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0,
                 "Cook temperature": 188.4936},
                {"Juiciness": 7.0, "Firmness": 6.0},
                formulation_no=1, batch_no=1)
    burger.set_pending_batch(None)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    table = next(t.value for t in at.table
                 if "Ingredient or process setting" in t.value.columns)
    amounts = dict(zip(table["Ingredient or process setting"], table["Amount"]))
    assert amounts["Cook temperature"] == "188.49 °C", amounts


# ------------------------------------------------------------------ #
#  Units per ingredient (owner amendment, 2026-09-10)
# ------------------------------------------------------------------ #

@pytest.fixture
def mixed_units(tmp_path, monkeypatch):
    """A project whose ingredients are not all in one unit: the powder is
    weighed in grams, the water is measured in millilitres."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("mixed")
    opt.add_ingredient("Pea protein", 0, 25)
    opt.add_ingredient("Water", 0, 60, unit="ml")
    opt.add_objective("Firmness", 1.0, goal="target", target=6,
                      min_val=0, max_val=10, unit="N")
    return opt
def test_the_ingredient_grid_has_plain_headers_and_a_unit_column(mixed_units):
    """Min (g) over a row measured in ml was a lie. The headers are plain and
    each row says what it is measured in."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    grid = _grid_frame(at, 0)
    assert [c for c in grid.columns if c != "_id"] == [
        "Name", "Type", "Lowest", "Highest", "Unit", "Vendor", "SKU"]
    assert dict(zip(grid["Name"], grid["Unit"])) == {"Pea protein": "g",
                                                      "Water": "ml"}


def test_a_new_ingredient_is_added_with_its_own_unit(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID,
               added=[_ingredient_row("Water", high=60.0, unit="ml")])
    assert not at.exception
    assert FoodOptimizer("burger").unit_of("Water") == "ml"


def test_a_unit_typed_into_the_grid_changes_one_row_and_says_so(burger):
    """Nothing is converted and nothing is rescored, and this sentence is
    the only place the app says so — so a unit moved on the grid keeps it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.UNIT_LABEL: "mg"}})
    assert not at.exception
    assert any("Methylcellulose is now in mg. The amounts you already "
               "recorded were not converted — check them." in s.value
               for s in at.success), [s.value for s in at.success]
    reloaded = FoodOptimizer("burger")
    assert reloaded.unit_of("Methylcellulose") == "mg"
    assert reloaded.unit_of("Pea protein") == "g"


def test_the_uploader_says_the_unit_column_is_optional(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    uploader = _unknown(at.main, "file_uploader", "Upload ingredients (Excel or CSV)")
    assert "Unit" in uploader.proto.help, uploader.proto.help


def test_the_batch_table_carries_each_unit_and_a_per_unit_total(mixed_units):
    mixed_units.set_pending_batch([{"Pea protein": 10.0, "Water": 40.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    frame = next(d for d in at.dataframe if "Formulation" in d.value.columns)
    assert list(frame.value.columns) == ["Formulation", "Pea protein (g)",
                                         "Water (ml)", "Total"]
    assert _displayed(frame)["Total"].iloc[0] == "10.00 g · 40.00 ml"


def test_scaling_is_offered_only_when_the_ingredients_share_a_unit(mixed_units):
    """Scaling 10 g of powder and 40 ml of water to a total of 400 means
    nothing, so the box is not offered and one line says why."""
    mixed_units.set_pending_batch([{"Pea protein": 10.0, "Water": 40.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert [n.key for n in at.number_input if n.key == "scale_total"] == []
    assert any(c.value == ("To make each formulation to a set amount, "
                           "every ingredient needs the same unit.")
               for c in at.caption), [c.value for c in at.caption]
    # ... and it comes back the moment they do share one.
    mixed_units.set_ingredient_unit("Water", "g")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert at.number_input(key="scale_total").label == \
        wording.batch_size_label("g")


def test_a_scale_left_behind_does_not_rewrite_a_mixed_unit_batch(mixed_units):
    """The box is gone, but Streamlit keeps a hidden widget's last value. It
    must not go on scaling amounts nobody can see it acting on."""
    mixed_units.set_pending_batch([{"Pea protein": 10.0, "Water": 40.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["scale_total"] = 400.0
    at.run()
    frame = next(d for d in at.dataframe if "Formulation" in d.value.columns)
    assert _displayed(frame)["Total"].iloc[0] == "10.00 g · 40.00 ml"


def test_the_sheet_writes_every_amount_in_its_own_unit(mixed_units):
    """A sheet mixing grams and millilitres carries each unit on its own
    heading, and a total that cannot be one number is written out as two."""
    mixed_units.set_pending_batch([{"Pea protein": 10.0, "Water": 40.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    # Every row carries its own unit, because the column header cannot.
    assert _sheet_amounts(at) == [("Pea protein (g)", 10.0),
                                  ("Water (ml)", 40.0)], _sheet_amounts(at)
    texts = _sheet_text(at)
    assert "10.00 g · 40.00 ml" in texts, texts
    assert wording.AMOUNT_COLUMN in texts, texts
    # And there is no `%` column at all: 10 g of powder is a share of no
    # total that also holds 40 ml of water.
    assert wording.PERCENT_COLUMN not in texts, texts
    rows = _summary_rows(at)
    assert list(rows)[2:] == ["Ingredient", "Pea protein (g)",
                              "Water (ml)", "Total",
                              wording.MEASURED_COLUMN,
                              wording.SHEET_WRITE_IN_NOTE,
                              "Firmness · target 6 N",
                              wording.NOT_SCORED_CHECKBOX_SHEET, "Note",
                              wording.SUMMARY_TICK_NOTE], list(rows)
    # Under the title, the line that says which of this sheet's cells take
    # a number.
    assert list(rows)[1] == wording.SUMMARY_SHADED_NOTE
    # One column per formulation, not two — and the Lot cell after it.
    assert rows["Pea protein (g)"] == [10.0, None], rows["Pea protein (g)"]


def test_the_amounts_to_make_it_table_uses_each_ingredients_unit(mixed_units):
    mixed_units.tell({"Pea protein": 10.0, "Water": 40.0}, {"Firmness": 6.0},
                     formulation_no=1, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    table = next(t.value for t in at.table
                 if "Ingredient or process setting" in t.value.columns)
    amounts = dict(zip(table["Ingredient or process setting"], table["Amount"]))
    assert amounts == {"Water": "40.00 ml", "Pea protein": "10.00 g"}


def test_an_amount_limit_across_units_is_refused_on_screen(mixed_units):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.multiselect(key="qty_pick").select("Pea protein")
    at.multiselect(key="qty_pick").select("Water")
    at.number_input(key="qc_max").set_value(50.0)
    _submit_button(at, "Add ingredient limit").click()
    at.run()
    assert not at.exception
    assert [e.value for e in at.error] == ["A limit adds amounts, so these ingredients need one unit; enter Water in g instead of ml."]
    assert FoodOptimizer("mixed").quantity_constraints == []


def test_a_fixed_ingredient_reads_as_one_amount_in_its_own_unit(mixed_units):
    """The table says it in Lowest and Highest; the workbook's Set-up sheet
    says it in words, in the row's own unit."""
    mixed_units.add_process_parameter("Cook temperature", 160, 200, unit="°C")
    mixed_units.add_ingredient("Water", 30.0, 30.0, unit="ml")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    grid = _grid_frame(at, 0)
    assert "Status" not in grid.columns
    row = grid[grid["Name"] == "Water"].iloc[0]
    assert (row["Lowest"], row["Highest"], row["Unit"]) == (30.0, 30.0, "ml")
    saved = FoodOptimizer("mixed")
    assert saved.fixed_at_text(saved._var_by_name("Water")) == "30.00 ml"


def test_the_grid_opens_on_the_project_that_is_open(mixed_units, tmp_path):
    """The rows a grid shows are the rows of the project the sidebar names,
    and nothing half-typed in the last one is still on it."""
    other = FoodOptimizer("other")
    other.set_amount_unit("ml")
    other.add_ingredient("Flour", 0, 100)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "mixed"
    at.run()
    _grid_edits(at, ING_GRID, edited={1: {wording.UNIT_LABEL: "kg"}})
    at.run()
    at.sidebar.selectbox(key="project_select").select("other")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert list(_grid_frame(at, 0)["Name"]) == ["Flour"]
    assert wording.SAVE_CHANGES_BUTTON not in _labels(at), _labels(at)
    assert FoodOptimizer("mixed").unit_of("Water") == "ml"


def test_an_amount_limit_is_listed_in_the_unit_it_limits(mixed_units):
    """A limit is a sum of ingredients that share a unit, so the number it
    holds them to is written in that unit."""
    mixed_units.add_quantity_constraint(["Water"], max_val=45)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(t.value == "Water: at most 45 ml" for t in at.text), \
        [t.value for t in at.text]


# ------------------------------------------------------------------ #
#  Projects made of process settings alone (owner ruling, 2026-09-10)
# ------------------------------------------------------------------ #

@pytest.fixture
def ferment(tmp_path, monkeypatch):
    """A fermentation project: nothing is weighed out, two settings are
    dialled in, and one measurement is scored."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("ferment")
    opt.add_process_parameter("Incubation temperature", 30, 42, unit="°C")
    opt.add_process_parameter("Incubation time", 4, 16, unit="h")
    opt.add_objective("Acidity", 1.0, goal="target", target=4.5,
                      min_val=3.0, max_val=7.0, unit="pH")
    return opt


def test_a_project_of_settings_alone_is_complete_and_lights_continue(ferment):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON]
    assert not any(c.value.startswith("Add at least one") for c in at.caption), \
        [c.value for c in at.caption]


def test_the_foot_names_a_setting_as_well_as_an_ingredient(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    FoodOptimizer("empty_one")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "empty_one"
    at.run()
    assert any(c.value == "Add at least one ingredient or process setting."
               for c in at.caption), [c.value for c in at.caption]


def test_nothing_that_belongs_to_ingredients_shows_without_any(ferment):
    """No amounts are weighed out, so there is no total to show, nothing to
    scale to a total, and no amount limits to set."""
    ferment.set_pending_batch([{"Incubation temperature": 37.0,
                                "Incubation time": 8.0},
                               {"Incubation temperature": 40.0,
                                "Incubation time": 12.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    frame = next(d for d in at.dataframe if "Formulation" in d.value.columns)
    assert list(frame.value.columns) == ["Formulation",
                                         "Incubation temperature (°C)",
                                         "Incubation time (h)"]
    assert [n.key for n in at.number_input if n.key == "scale_total"] == []
    assert not any("Scaling needs" in c.value for c in at.caption), \
        [c.value for c in at.caption]
    # The sheet lists the settings and claims no total.
    assert ("Incubation temperature (°C)", 37.0) in _sheet_amounts(at), \
        _sheet_amounts(at)
    assert wording.TOTAL_LABEL not in _sheet_text(at), _sheet_text(at)
    # Both kinds of limit are about what you weigh out, so the whole section
    # is gone: a setting's own Lowest and Highest are its bounds.
    assert [e.label for e in at.expander if e.label == "Limits (optional)"] == []
    assert [b.label for b in at.button
            if b.label == "Add ingredient limit"] == []
    assert [b.label for b in at.button if b.label == "Add property limit"] == []


def test_a_settings_only_project_goes_round_the_whole_loop(ferment):
    """Generate, record two formulations, save, and read the best back."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at, wording.NEXT_MAKE_BATCH_BUTTON).click()
    at.run()
    at.number_input(key="how_many").set_value(2)
    at.run()
    _submit_button(at, "Generate 2 formulations").click()
    at.run()
    assert not at.exception
    made = FoodOptimizer("ferment")
    assert len(made.pending_batch) == 2
    at.number_input(key="f1_Acidity").set_value(4.5)
    at.number_input(key="f2_Acidity").set_value(5.5)
    at.run()
    assert _tab_primaries(at, 1) == ["Save results"], _tab_primaries(at, 1)
    _submit_button(at, "Save results").click()
    at.run()
    assert not at.exception
    assert at.session_state["main_tab"] == wording.TAB_RESULTS
    assert any(h.value == wording.best_so_far_heading(1, 1)
               for h in at.subheader), [h.value for h in at.subheader]
    table = next(t.value for t in at.table
                 if "Ingredient or process setting" in t.value.columns)
    shown = dict(zip(table["Ingredient or process setting"], table["Amount"]))
    assert set(shown) == {"Incubation temperature", "Incubation time"}
    assert shown["Incubation temperature"].endswith(" °C"), shown
    assert not any(c.value.startswith("Not used:") for c in at.caption), \
        [c.value for c in at.caption]


def test_a_unit_change_that_breaks_an_amount_limit_removes_it_and_says_so(
        burger):
    """A limit is a sum. Once one of its ingredients is measured in another
    unit the sum means nothing, so the limit goes and the line says which."""
    burger.add_ingredient("Salt", 0, 3)
    burger.add_quantity_constraint(["Pea protein", "Methylcellulose"],
                                   max_val=20)
    burger.add_total_mass_constraint(max_val=100)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.UNIT_LABEL: "ml"}})
    assert not at.exception
    assert [w.value for w in at.warning] == [
        "The limit on Pea protein + Methylcellulose was deleted because "
        "those ingredients no longer share a unit.",
        "The limit on All ingredients was deleted because "
        "those ingredients no longer share a unit.",
    ], [w.value for w in at.warning]
    assert FoodOptimizer("burger").quantity_constraints == []
    # A limit whose ingredients still share a unit is left alone. (Read the
    # project back first: the fixture's own copy is now behind the screen's.)
    reloaded = FoodOptimizer("burger")
    reloaded.add_quantity_constraint(["Pea protein", "Salt"], max_val=20)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.UNIT_LABEL: "mg"}})  # not in it
    assert [w.value for w in at.warning] == [], [w.value for w in at.warning]
    assert len(FoodOptimizer("burger").quantity_constraints) == 1


def test_a_blank_unit_cell_is_refused_by_row(burger):
    """A blank would rewrite the ingredient to no unit at all — and could
    take an amount limit with it — on a save that looks like a no-op."""
    burger.add_quantity_constraint(["Pea protein", "Methylcellulose"],
                                   max_val=20)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.UNIT_LABEL: ""}})
    assert [e.value for e in at.error] == [
        wording.row_error(2, wording.UNIT_REQUIRED_ERROR)]
    reloaded = FoodOptimizer("burger")
    assert reloaded.unit_of("Methylcellulose") == "g"
    assert len(reloaded.quantity_constraints) == 1
def test_the_formulations_download_says_what_is_in_it(burger):
    """The file carries the screen's own units and every formulation the
    project holds, so the help says that rather than the two things that
    used to be true of it."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    download = _unknown(at.main, "download_button",
                        "Download all formulations (Excel)")
    assert download.proto.help == ("One row per formulation, with the same "
                                   "units the screen shows, and a second "
                                   "sheet holding the set-up they were made "
                                   "under. Formulations marked not scored "
                                   "are included, with their measurements "
                                   "blank.")


# A reloaded ingredient file is the third edit that can empty a limit of
# meaning, and the only one AppTest cannot drive (st.file_uploader takes no
# file from a test). The backend path is covered in test_food_bo; this pins
# the lines the handler writes, including the one no other path can produce.
REMOVED_LIMIT_LINES = """
import streamlit as st
from food_bo import FoodOptimizer
from ui_helpers import render_flash
from ui_setup import _flash_removed_limits


class _Opt:
    variables = [{"name": "Water", "category": "ingredient"},
                 {"name": "Oil", "category": "ingredient"},
                 {"name": "Salt", "category": "ingredient"}]
    # How a limit is named, and the sentence that says one has gone, both
    # live on the optimizer — four doors owe the same line — so the stub
    # borrows them rather than spelling the rule twice.
    _ingredients = FoodOptimizer._ingredients
    limit_label = FoodOptimizer.limit_label
    limit_removed_messages = FoodOptimizer.limit_removed_messages


_flash_removed_limits(_Opt(), [
    {"ingredients": ["Water", "Oil"], "reason": "unit"},
    {"ingredients": ["Water", "Coconut oil"], "reason": "missing",
     "missing": ["Coconut oil"]},
    {"ingredients": ["Water", "Coconut oil", "Beet juice powder"],
     "reason": "missing", "missing": ["Coconut oil", "Beet juice powder"]},
    {"ingredients": ["Water", "Oil", "Salt"], "reason": "unit"},
])
render_flash()
"""


def test_every_removed_limit_is_named_in_one_line():
    at = AppTest.from_string(REMOVED_LIMIT_LINES)
    at.run()
    assert not at.exception
    assert [w.value for w in at.warning] == [
        "The limit on Water + Oil was deleted because those ingredients no "
        "longer share a unit.",
        "The limit on Water + Coconut oil was deleted because Coconut oil is "
        "no longer an ingredient.",
        "The limit on Water + Coconut oil + Beet juice powder was deleted "
        "because Coconut oil and Beet juice powder are no longer ingredients.",
        # Every ingredient in it: named as the Limits list names it.
        "The limit on All ingredients was deleted because "
        "those ingredients no longer share a unit.",
    ], [w.value for w in at.warning]


# ------------------------------------------------------------------ #
#  Verification round 2: fresh forms per project, an armed restore,
#  a unit on every number, and Total reserved
# ------------------------------------------------------------------ #
def test_a_unit_change_that_splits_the_units_says_the_trial_is_unscaled(burger):
    """Scaling needs one unit. A unit change that takes it away used to leave
    the table silently back at the generated amounts."""
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}],
                             batch_no=2)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(400.0)
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.UNIT_LABEL: "ml"}})
    assert not at.exception
    assert any(s.value == (wording.unit_changed("Methylcellulose", "ml", True)
                           + " " + wording.unscaled_tail(2))
               for s in at.success), [s.value for s in at.success]
    # ...and the box that held the total is empty, not quietly meaning nothing.
    assert ("scale_total" not in at.session_state
            or at.session_state["scale_total"] is None)


def test_a_limit_on_all_ingredients_is_refused_in_words_the_screen_can_obey(
        mixed_units):
    """A limit picked across every ingredient is still a sum, and the refusal
    names the ingredient to re-enter and the unit to enter it in — not just
    that the units differ."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.multiselect(key="qty_pick").select("Pea protein")
    at.multiselect(key="qty_pick").select("Water")
    at.run()
    at.number_input(key="qc_max").set_value(300.0)
    _submit_button(at, "Add ingredient limit").click()
    at.run()
    assert not at.exception
    assert [e.value for e in at.error] == [
        "A limit adds amounts, so these ingredients need one unit; "
        "enter Water in g instead of ml."]
    assert FoodOptimizer("mixed").quantity_constraints == []


def test_leaving_a_formulation_out_keeps_the_note_box_open(open_batch):
    """The note is the only record of what went wrong, and it is typed after
    the tick as often as before it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.checkbox(key="f1_leave_out").check()
    at.run()
    assert not at.exception
    note = next(t for t in at.text_input if t.key == "f1_note")
    assert not note.disabled
    tick = next(c for c in at.checkbox if c.key == "f1_leave_out")
    assert tick.label == "Not scored"
    assert tick.help == ("Ticked wins over any number typed in this row. Say "
                         "why in Note; it stays with the formulation.")
    # The measurement boxes still grey out: a left-out formulation has no
    # results, and only the note it leaves behind.
    assert next(n for n in at.number_input if n.key == "f1_Firmness").disabled


def test_a_left_out_formulation_keeps_the_note_typed_after_the_tick(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.checkbox(key="f1_leave_out").check()
    at.run()
    at.text_input(key="f1_note").set_value("Mixer jammed")
    at.run()
    at.number_input(key="f2_Firmness").set_value(6.0)
    at.number_input(key="f2_Juiciness").set_value(7.0)
    at.run()
    _submit_button(at, "Save results").click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").skipped[0]["note"] == "Not scored · Mixer jammed"


def test_a_measurement_recorded_for_the_first_time_is_stored(burger):
    """A measurement added after the fact has no value on the row, so the box
    opens blank and the number typed into it is a correction like any other."""
    burger.add_objective("Serving temperature", 1.0, goal="target", target=65,
                         min_val=0, max_val=100, unit="°C")
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    assert at.number_input(key="correct_1_Serving temperature").value is None
    at.number_input(key="correct_1_Serving temperature").set_value(66.0)
    at.run()
    _submit_button(at, "Save correction").click()
    at.run()
    assert not at.exception
    assert any(s.value.startswith(wording.formulation_corrected(1))
               for s in at.success), [s.value for s in at.success]
    assert FoodOptimizer("burger").results_history[0]["Serving temperature"] == 66.0


def test_the_best_score_says_partial_when_a_measurement_was_not_scored(burger):
    """The All formulations row already says (partial); the caption above it
    said the same score as a full one."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Firmness": 6.0}, formulation_no=1, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception
    # The missing measurement is NAMED after the ceiling, and nothing after
    # it claims the measurements were on target.
    assert any(c.value == ("Overall score 60.00 of 100.00 · Juiciness not "
                           "measured. Scores only compare within this "
                           "project. Change a share, a goal or a range and "
                           "every score is worked out again.")
               for c in at.caption), [c.value for c in at.caption]
    # ...and one line under it says what the missing measurement costs, in
    # the same words the All formulations table uses. A dropped measurement
    # is arithmetically a zero closeness, and the row is fitted at that
    # depressed score, so the model is taught that region is bad.
    assert any(c.value == wording.PARTIAL_SCORES_CAPTION
               for c in at.caption), [c.value for c in at.caption]


def test_there_is_only_one_unit_column_per_grid(burger):
    """There were two boxes labelled Unit on this tab — the add form's and
    the control row's — and the reader had to work out which was which. A
    grid has one Unit cell per row, in the row it belongs to."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert list(_grid_frame(at, 0).columns).count(wording.UNIT_LABEL) == 1
    assert not [b for b in at.text_input if b.label == "New unit"]


def test_loading_over_a_list_says_it_replaces_it(burger, tmp_path,
                                                 monkeypatch):
    """The file does not add to the list, it replaces it."""
    from ui_setup import _load_label
    assert _load_label(burger) == "Replace ingredients"
    monkeypatch.chdir(tmp_path)
    assert _load_label(FoodOptimizer("empty_list")) == "Load ingredients"


def test_the_grid_speaks_for_settings_as_well_as_ingredients(ferment):
    """In a fermentation project every row of the grid is a process
    setting, and the grid is still the one place they are edited."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    grid = _grid_frame(at, 0)
    assert list(grid["Name"]) == ["Incubation temperature", "Incubation time"]
    assert list(grid["Type"]) == [wording.KIND_SETTING] * 2
    # And a setting is fixed the same way an ingredient is: one number in
    # both cells, written in the setting's own unit.
    _save_grid(at, ING_GRID, edited={0: {wording.LOWEST_LABEL: 30.0,
                                         wording.HIGHEST_LABEL: 30.0}})
    saved = FoodOptimizer(ferment.project_name)
    assert saved.fixed_at_text(
        saved._var_by_name("Incubation temperature")) == "30 °C"


def test_the_csv_template_has_one_name_on_both_screens(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()                                   # the welcome panel: no project
    assert _unknown(at.main, "download_button", "Download ingredients template (Excel)")
    FoodOptimizer("named").save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _unknown(at.main, "download_button", "Download ingredients template (Excel)")


def test_the_tab_three_foot_never_offers_a_first_trial_over_an_open_one(burger):
    """A batch on the bench is not a first batch to make, and the foot of
    every other screen already knows the words."""
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0},
                              {"Pea protein": 20.0, "Methylcellulose": 2.0},
                              {"Pea protein": 5.0, "Methylcellulose": 0.5}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception
    assert _tab_primaries(at, 2) == [wording.back_to_batch_label(1, 3)], \
        _tab_primaries(at, 2)


def test_the_compared_with_column_reads_the_amounts_on_the_table(burger):
    """While the batch is scaled, a change read off the generated amounts is
    a number nothing on screen shows."""
    for i in range(5):
        burger.tell({"Pea protein": 10.0, "Methylcellulose": 10.0 + i},
                    {"Juiciness": 7.0, "Firmness": 6.0 if i == 0 else 1.0},
                    formulation_no=i + 1, batch_no=1)
    burger.set_pending_batch([{"Pea protein": 30.0, "Methylcellulose": 10.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    column = "Compared with Formulation 1"
    # As generated, the only change is +20.00 g of protein.
    table = next(d.value for d in at.dataframe if column in d.value.columns)
    assert table[column].iloc[0] == \
        "Trying something different · Pea protein +20.00 g"
    at.number_input(key="scale_total").set_value(20.0)
    at.run()
    # Scaled to 20 g the table reads 15.00 / 5.00 against 10.00 / 10.00, and
    # the cell reads the table: the best is scaled to the same total.
    table = next(d.value for d in at.dataframe if column in d.value.columns)
    assert table[column].iloc[0] == (
        "Trying something different · Pea protein +5.00 g, "
        "Methylcellulose −5.00 g"), table[column].iloc[0]


def _uploader_keys(at):
    """The user key of every file uploader on the tabs, read off its widget
    id (AppTest gives no accessor for st.file_uploader)."""
    return sorted(u.proto.id.split("-")[-1]
                  for u in _unknowns(at.main, "file_uploader"))


def test_each_project_gets_its_own_file_uploaders(burger):
    """A file uploader cannot be emptied from session state — assigning None
    is refused and popping the key leaves the mounted widget holding the file
    — so each is keyed to its project and a new project renders an empty one."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    burger.set_pending_batch([{"Pea protein": 20.0, "Methylcellulose": 2.0}],
                             batch_no=2)
    FoodOptimizer("second").save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    # The formulations uploader is the CSV half of `Add a formulation you
    # already made`; the radio opens on the typed-in half.
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.run()
    assert _uploader_keys(at) == ["import_file_burger", "ingredients_file_burger",
                                  "results_file_burger"], _uploader_keys(at)
    at.sidebar.selectbox(key="project_select").select("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert at.session_state["_loaded_project"] == "second"
    # Not one widget of the last project's is still on screen.
    assert all(k.endswith("_second") for k in _uploader_keys(at)), \
        _uploader_keys(at)


def test_opening_a_project_renames_the_sidebar_box_and_greys_open(burger):
    """The box kept the previous project's name and Open lit up over the
    project the user was typing into."""
    FoodOptimizer("second").save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    at.sidebar.selectbox(key="project_select").select("second")
    at.run()
    _submit_button(at.sidebar, "Open").click()
    at.run()
    assert not at.exception
    assert at.sidebar.selectbox(key="project_select").value == "second"
    assert [b.label for b in at.sidebar.button if b.proto.type == "primary"] == []
    assert _submit_button(at.sidebar, "Open").disabled


def test_creating_a_project_renames_the_sidebar_box_and_greys_open(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.sidebar.text_input[0].set_value("Probe proj")
    at.sidebar.button[0].click()          # Create project (a form submit)
    at.run()
    assert not at.exception
    assert at.session_state["_loaded_project"] == "Probe proj"
    assert at.sidebar.selectbox(key="project_select").value == "Probe proj"
    assert [b.label for b in at.sidebar.button if b.proto.type == "primary"] == []


def test_a_checked_backup_is_an_armed_confirmation(scored):
    """Yes, replace keeps a copy and replaces the project, so it is the one
    coloured button on screen until it is answered."""
    donor = FoodOptimizer("donor")
    donor.add_ingredient("Flour", 0, 100)
    donor.add_objective("Crunch", 1.0, goal="max")
    donor.tell({"Flour": 10.0}, {"Crunch": 5.0})
    FoodOptimizer("spare").save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["_restore_candidate"] = donor.export_json()
    at.run()
    assert not at.exception
    # One lit button in the sidebar, and none on any tab.
    assert [b.label for b in at.sidebar.button
            if b.proto.type == "primary"] == ["Yes, replace"]
    assert _tab_primaries(at, 0) == [], _tab_primaries(at, 0)
    assert _tab_primaries(at, 2) == [], _tab_primaries(at, 2)
    assert _submit_button(at, wording.NEXT_MAKE_BATCH_BUTTON).disabled
    # Cancel puts everything back.
    _submit_button(at.sidebar, "Cancel").click()
    at.run()
    at.run()
    assert "_restore_candidate" not in at.session_state
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON]


def test_the_restore_warning_counts_every_formulation_and_names_settings(scored):
    """Both halves of the sentence count the same way, or replacing a project
    with its own backup reads as losing one."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["_restore_candidate"] = scored.export_json()
    at.run()
    assert not at.exception
    assert any(w.value == ("This copy holds **burger**: 3 formulations, "
                           "2 ingredients. Replace **burger**, which has "
                           "3 formulations? " + wording.COPY_KEPT)
               for w in at.warning), [w.value for w in at.warning]


def test_the_restore_warning_names_a_settings_only_project(ferment):
    ferment.tell({"Incubation temperature": 38.0, "Incubation time": 6.0},
                 {"Acidity": 4.5}, formulation_no=1, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_restore_candidate"] = ferment.export_json()
    at.run()
    assert not at.exception
    assert any(w.value.startswith("This copy holds **ferment**: "
                                  "1 formulation, 0 ingredients, "
                                  "2 process settings.")
               for w in at.warning), [w.value for w in at.warning]


def test_the_printed_sheet_says_a_repeat_is_a_repeat(burger):
    """Two sheets with identical amounts and nothing printed to say why is
    how a batch gets made twice."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    burger.set_pending_batch([{"Pea protein": 20.0, "Methylcellulose": 2.0}],
                             batch_no=2)
    burger.add_to_pending_batch({"Pea protein": 10.0, "Methylcellulose": 1.0},
                                note=wording.repeat_of_formulation(1))
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    texts = _sheet_text(at)
    assert wording.repeat_of_formulation(1) in texts, texts
    # Every sheet carries the label and a box to write a note in, whether or
    # not the app already put one there.
    assert texts.count(wording.NOTE) == 2, texts
    # ... and so does the summary sheet, on the row it keeps for notes.
    # Two formulations, each with an amount column and a share column, then
    # the Lot column past the end of them.
    assert _summary_rows(at)[wording.NOTE] == [
        None, None, wording.repeat_of_formulation(1), None, None], \
        _summary_rows(at)[wording.NOTE]


def test_a_settings_only_project_reads_as_one_section(ferment):
    """A fermentation scientist weighs nothing out. One section covers both
    kinds, so there is nothing to reorder and no empty Ingredients fold."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    labels = [e.label for e in at.expander]
    assert "Ingredients (optional)" not in labels, labels
    assert "Process settings" not in labels, labels
    table = next(d.value for d in at.dataframe if "Type" in d.value.columns)
    assert list(table["Type"]) == ["Process setting", "Process setting"]
    assert list(table["Unit"]) == ["°C", "h"]


def test_an_ingredient_list_puts_the_ingredients_first(burger):
    """One table, ingredients above settings, each in the order they were
    added."""
    burger.add_process_parameter("Cook temperature", 150, 200, unit="°C")
    burger.add_ingredient("Salt", 0, 3)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert list(_grid_frame(at, 0)["Name"]) == [
        "Pea protein", "Methylcellulose", "Salt", "Cook temperature"]


def test_a_row_edited_into_another_unit_names_the_limit_it_broke(burger):
    """Every other unit edit prunes the limits it breaks and says which. A
    row whose range moves as well goes the long way round, through the add
    path, and still owes the same line."""
    burger.add_quantity_constraint(["Pea protein", "Methylcellulose"],
                                   min_val=5, max_val=150)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.UNIT_LABEL: "ml",
                                         wording.HIGHEST_LABEL: 3.0}})
    assert not at.exception
    assert any(w.value == ("The limit on All ingredients was "
                           "deleted because those ingredients no longer share "
                           "a unit.") for w in at.warning), \
        [w.value for w in at.warning]
    assert FoodOptimizer("burger").quantity_constraints == []


def test_an_ingredient_named_total_is_refused_by_row(burger):
    """Total is the round table's own column: two of them break the table and
    put the round total on the sheet where the ingredient's amount belongs."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, added=[_ingredient_row("Total")])
    assert not at.exception
    assert any("Total is a column name Food Optimizer uses" in e.value
               for e in at.error), [e.value for e in at.error]
    assert [v["name"] for v in FoodOptimizer("burger").variables] == [
        "Pea protein", "Methylcellulose"]


# ------------------------------------------------------------------ #
#  A property limit is a limit on the finished formulation, per 100 g
# ------------------------------------------------------------------ #


@pytest.fixture
def with_properties(tmp_path, monkeypatch):
    """A project loaded from a file with a property column, so the limits
    section has something to limit."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("props")
    opt.load_ingredients_from_csv(pd.DataFrame({
        "Name": ["Coconut oil", "Water"],
        "Min": [0, 20],
        "Max": [15, 60],
        "Unit": ["g", "g"],
        "Fat per 100 g": [99.0, 0.0],
    }))
    opt.add_objective("Firmness", 1.0, goal="target", target=6,
                      min_val=0, max_val=10, unit="N")
    return opt


def test_a_property_limit_says_it_is_per_100_g_of_the_formulation(with_properties):
    """It used to be a total that grew with the batch, and nothing on screen
    said which it was."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert any(m.value == "**Finished-product limit**"
               for m in at.markdown), [m.value for m in at.markdown]
    assert at.selectbox(key="prop_metric").label == "Ingredient property"
    assert any(c.value == ("Per 100 g of formulation, from the properties "
                           "of your ingredients.")
               for c in at.caption), [c.value for c in at.caption]


def test_the_caption_names_the_unit_the_ingredients_are_in(with_properties):
    with_properties.set_ingredient_unit("Coconut oil", "ml")
    with_properties.set_ingredient_unit("Water", "ml")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(c.value.startswith("Per 100 ml of formulation")
               for c in at.caption), [c.value for c in at.caption]


def test_adding_a_property_limit_lists_it_and_keeps_what_was_made(with_properties):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="prop_max").set_value(20.0)
    at.run()
    _submit_button(at, "Add property limit").click()
    at.run()
    assert not at.exception
    assert any(s.value == (wording.limit_added_on("Fat per 100 g") + " " + wording.LIMIT_KEPT)
               for s in at.success), \
        [s.value for s in at.success]
    assert any(t.value == "Fat per 100 g: at most 20" for t in at.text), \
        [t.value for t in at.text]
    assert FoodOptimizer("props").constraints[0]['max'] == 20.0


def test_a_property_limit_is_refused_while_the_units_differ(with_properties):
    with_properties.set_ingredient_unit("Water", "ml")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="prop_max").set_value(20.0)
    at.run()
    _submit_button(at, "Add property limit").click()
    at.run()
    assert not at.exception
    # The refusal names the fix, not just the fault.
    assert any(e.value == ("Property limits are per 100 g, so every ingredient "
                           "needs a mass unit; enter Water in g instead of "
                           "ml.") for e in at.error), \
        [e.value for e in at.error]
    assert FoodOptimizer("props").constraints == []


def test_a_limit_from_an_older_version_says_how_it_is_read_now(with_properties):
    """A 0.2.x file's limit was a total. It is now read per 100 g, which is a
    different number, so the screen says so once."""
    with_properties.constraints = [{'metric': "Fat per 100 g", 'min': None,
                                    'max': 25.0}]
    with_properties.save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert sum(1 for c in at.caption
               if c.value == ("An older limit is now read per 100 g of "
                              "formulation.")) == 1, \
        [c.value for c in at.caption]


def test_a_limit_set_now_says_nothing_extra(with_properties):
    with_properties.add_constraint("Fat per 100 g", max_val=25.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not any("before this version" in c.value for c in at.caption), \
        [c.value for c in at.caption]


# ------------------------------------------------------------------ #
#  One section for everything you can vary
# ------------------------------------------------------------------ #


def _tab1(at):
    return at.tabs[0]


def _more_settings(at):
    """Tab 1's middle tier, by name."""
    return next(e for e in _tab1(at).expander
                if e.label == wording.MORE_SETTINGS_EXPANDER)


def _advanced(at):
    """Tab 1's bottom tier, by name."""
    return next(e for e in _tab1(at).expander
                if e.label == wording.ADVANCED_EXPANDER)


def _advanced_text(at):
    """Everything the Advanced fold says, as one block: the model settings
    and the two explanations are plain markdown inside it now, because
    Streamlit cannot nest an expander in an expander."""
    return "\n".join(m.value for m in _advanced(at).markdown)


def _advanced_block(at, heading):
    """One of the three blocks inside Advanced: the markdown that follows a
    heading, up to the next heading."""
    values = [m.value for m in _advanced(at).markdown]
    out = []
    for value in values[values.index(heading) + 1:]:
        if value.startswith("**"):
            break
        out.append(value)
    return "\n".join(out)


def _tab_outline(at, index):
    """The tab's headings and folds, in the order they are drawn. This is
    what "the tab reads in one order" means, and it is read off the element
    tree rather than off a list of labels that happen to be on screen
    somewhere.

    Plain containers are walked through rather than stopped at: tab 2
    reserves each step's position in one (a confirmation has to be asked
    before the grid below it is drawn), and a container is not something a
    reader can see. An expander IS, so it is named and not entered."""
    out = []

    def walk(node):
        children = getattr(node, "children", None) or {}
        if hasattr(children, "values"):
            children = children.values()
        for element in children:
            kind = element.__class__.__name__
            if kind == "Subheader":
                out.append(element.value)
            elif kind == "Expander":
                out.append(element.label)
            elif kind == "Markdown" and element.value[:1] in ("#", "*"):
                out.append(element.value)
            elif kind not in ("Expander", "Tab", "Tabs"):
                walk(element)

    walk(at.tabs[index])
    return out


def _tab1_captions(at):
    """The captions a reader sees on tab 1 without opening anything: the
    contents of a collapsed expander are explanations, and they are allowed
    to be longer."""
    tab = at.tabs[0]
    folded = {id(c) for e in tab.expander for c in e.caption}
    return [c.value for c in tab.caption if id(c) not in folded]


def test_ingredients_and_settings_are_one_section_for_both_types(burger):
    """Four places used to hold the ingredients and the settings: a subheader,
    an expander, another expander and a fold. They are one section now."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert [h.value for h in _tab1(at).subheader] == [
        "Ingredients and process settings", "Measurements and targets"], \
        [h.value for h in _tab1(at).subheader]
    # ...and the section has no caption of its own: the one it carried only
    # restated the heading it now wears.
    # No "Add a measurement" fold: a measurement is typed on the empty row
    # at the bottom of its own grid, like an ingredient.
    # Three folds on the whole tab (spec 1.5): the ingredients grid's own
    # alternative, then the two tiers. Limits and the two explanations are
    # inside them, as headings — Streamlit cannot nest one expander in
    # another, and a tab that folded the optional half away twice made the
    # reader open two things to reach one.
    labels = [e.label for e in _tab1(at).expander]
    assert labels == ["Or upload an ingredients file",
                      "More settings", "Advanced"], labels


def test_the_ingredients_grid_is_the_first_thing_on_the_tab(burger):
    """No add form, no control row, no per-row editor: one grid, typed where
    it is read, and it is not folded away inside anything."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    grid = _grid_frame(at, 0)
    assert [c for c in grid.columns if c != "_id"] == [
        "Name", "Type", "Lowest", "Highest", "Unit", "Vendor", "SKU"]
    assert list(grid["Name"]) == ["Pea protein", "Methylcellulose"]
    assert wording.SAVE_CHANGES_BUTTON not in _labels(at), _labels(at)
    folded = {id(d) for e in _tab1(at).expander for d in e.dataframe}
    assert id(_tab1(at).dataframe[0]) not in folded


def test_a_row_typed_at_the_bottom_adds_an_ingredient(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID,
               added=[_ingredient_row("Beet juice powder", high=2.0)])
    assert not at.exception
    assert any(s.value == "Beet juice powder added." for s in at.success), \
        [s.value for s in at.success]
    saved = FoodOptimizer("burger")
    assert [v["name"] for v in saved.variables][-1] == "Beet juice powder"
    assert saved.unit_of("Beet juice powder") == "g"


def test_a_row_typed_at_the_bottom_adds_a_process_setting(burger):
    """Type is a cell, not a second form: the same seven answers, with
    Process setting chosen in the row's own Type cell."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, added=[
        _ingredient_row("Cook temperature", kind=wording.KIND_SETTING,
                        low=150.0, high=200.0, unit="°C")])
    assert not at.exception
    assert any(s.value == "Cook temperature added." for s in at.success), \
        [s.value for s in at.success]
    saved = FoodOptimizer("burger")
    setting = saved._var_by_name("Cook temperature")
    assert setting["category"] == "process" and setting["unit"] == "°C"


def test_vendor_and_sku_ride_on_the_row_and_reach_the_sheet(burger):
    """Optional text the app never reads: it is printed so the bench knows
    what to reach for."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={0: {wording.VENDOR_LABEL: "Roquette",
                                         wording.SKU_LABEL: "NUTRALYS F85M"}})
    assert not at.exception
    saved = FoodOptimizer("burger")
    var = saved._var_by_name("Pea protein")
    assert (var["vendor"], var["sku"]) == ("Roquette", "NUTRALYS F85M")
    sheet = openpyxl.load_workbook(
        io.BytesIO(saved.all_formulations_workbook()))["Set-up"]
    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    head = next(r for r in rows if r[0] == wording.NAME_LABEL)
    assert head[6] == wording.VENDOR_LABEL and head[7] == wording.SKU_LABEL
    line = next(r for r in rows if r[0] == "Pea protein")
    assert (line[6], line[7]) == ("Roquette", "NUTRALYS F85M")


def test_a_vendor_typed_on_a_setting_is_refused_by_row(burger):
    """A cook temperature has no supplier. Refused rather than quietly
    dropped: a cell the reader filled in is an answer."""
    burger.add_process_parameter("Cook temperature", 150, 200, unit="°C")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={2: {wording.VENDOR_LABEL: "Acme"}})
    assert [e.value for e in at.error] == [
        wording.row_error(3, wording.only_an_ingredient_has(
            wording.VENDOR_LABEL))]


def test_a_baseline_typed_on_an_ingredient_is_refused_by_row(burger):
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={0: {wording.BASELINE_LABEL: 5.0}})
    assert [e.value for e in at.error] == [
        wording.row_error(1, wording.only_a_setting_has(
            wording.BASELINE_LABEL))]


def test_the_baseline_column_belongs_to_a_project_with_results(burger):
    """It is the value every formulation already made is read at, so a
    project with nothing recorded has nothing to read — and no column."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert wording.BASELINE_LABEL not in _grid_frame(at, 0).columns
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert wording.BASELINE_LABEL in _grid_frame(at, 0).columns
    # ...and a setting added now must carry one: those formulations ran at
    # some setting, not at zero.
    _save_grid(at, ING_GRID, added=[
        _ingredient_row("Cook temperature", kind=wording.KIND_SETTING,
                        low=150.0, high=200.0, unit="°C")])
    assert [e.value for e in at.error] == [
        wording.row_error(3, wording.ADD_BASELINE_ERROR)]


def test_the_grid_names_the_type_of_every_row(burger):
    burger.add_process_parameter("Cook temperature", 150, 200, unit="°C")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    grid = _grid_frame(at, 0)
    assert list(grid["Name"]) == ["Pea protein", "Methylcellulose",
                                  "Cook temperature"]
    assert list(grid["Type"]) == ["Ingredient", "Ingredient",
                                  "Process setting"]
    assert list(grid["Unit"]) == ["g", "g", "°C"]


def test_a_fixed_row_adds_no_column_to_the_grid(burger):
    burger.add_ingredient("Methylcellulose", 1.0, 1.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    grid = _grid_frame(at, 0)
    assert "Status" not in grid.columns
    assert list(grid["Lowest"]) == [0.0, 1.0]
    assert list(grid["Highest"]) == [25.0, 1.0]


def test_saving_the_grid_changes_the_allowed_amounts_and_says_so(burger):
    burger.set_formulation_total(25.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.HIGHEST_LABEL: 6.0}})
    assert not at.exception
    assert any(s.value == "Methylcellulose saved." for s in at.success), \
        [s.value for s in at.success]
    assert FoodOptimizer("burger")._var_by_name("Methylcellulose")["bounds"] \
        == (0.0, 6.0)
    # Saved, so there is nothing left to save: the buttons go and the foot
    # takes the colour back.
    assert wording.SAVE_CHANGES_BUTTON not in _labels(at), _labels(at)
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON]


def test_saving_the_grid_retires_the_open_round_with_the_notice(burger):
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID, discard=True,
               edited={1: {wording.HIGHEST_LABEL: 6.0}})
    assert FoodOptimizer("burger").pending_batch is None
    assert any(f"{wording.ROUND_CAP} 1 was discarded" in i.value
               for i in at.info), [i.value for i in at.info]


def test_a_save_that_would_discard_a_round_asks_first_and_counts_your_own(
        burger):
    """The cold reader lost a formulation they had typed in by hand, with
    their own note on it, to a toast that arrived after it was gone. The
    question comes first, it names the round, it counts what the round
    holds and it says how much of that is theirs — and nothing is written
    until they answer it."""
    burger.set_pending_batch([
        {"Pea protein": 10.0, "Methylcellulose": 1.0},
        {"Pea protein": 20.0, "Methylcellulose": 2.0},
    ])
    burger.add_to_pending_batch({"Pea protein": 15.0, "Methylcellulose": 1.5},
                                note="my control, salt to 3 g")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _grid_edits(at, ING_GRID, edited={1: {wording.HIGHEST_LABEL: 6.0}})
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, edited={1: {wording.HIGHEST_LABEL: 6.0}})
    at.run()
    assert not at.exception
    assert any(w.value == ("Saving will discard Round 1: 3 formulations, "
                           "1 of them added by you. A formulation you added "
                           "goes with the round — a set-up change can make "
                           "it invalid.")
               for w in at.warning), [w.value for w in at.warning]
    # Nothing is written while the question is up, and the tab still shows
    # exactly one coloured button.
    assert FoodOptimizer("burger").pending_batch is not None
    assert FoodOptimizer("burger")._var_by_name(
        "Methylcellulose")["bounds"] == (0.0, 3.0)
    assert _tab_primaries(at, 0) == [wording.YES_SAVE_AND_DISCARD], \
        _tab_primaries(at, 0)
    _submit_button(at, wording.YES_SAVE_AND_DISCARD).click()
    _grid_edits(at, ING_GRID, edited={1: {wording.HIGHEST_LABEL: 6.0}})
    at.run()
    assert not at.exception
    saved = FoodOptimizer("burger")
    assert saved.pending_batch is None
    assert saved._var_by_name("Methylcellulose")["bounds"] == (0.0, 6.0)


def test_the_round_that_was_discarded_is_said_under_the_grid_not_only_in_a_toast(
        burger):
    """A toast fades and the round does not come back. The same sentence
    stands under the grid that took it away until there is a round again."""
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID, discard=True,
               edited={1: {wording.HIGHEST_LABEL: 6.0}})
    assert not at.exception
    line = wording.batch_discarded_notice(1)
    assert any(i.value == line for i in at.info), [i.value for i in at.info]
    # The flash is drained on the next run; the standing line is not.
    at.run()
    assert any(i.value == line for i in at.info), [i.value for i in at.info]
    at.button(key="generate").click()
    at.run()
    assert not at.exception
    assert not any(i.value == line for i in at.info), [i.value for i in at.info]


def test_a_name_typed_over_another_is_a_rename_and_keeps_what_was_recorded(
        burger):
    """The hidden row identity is what makes this a rename rather than a new
    row beside a deleted one: the results already recorded are filed under
    the name being changed."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    burger.add_quantity_constraint(["Methylcellulose"], max_val=2.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID,
               edited={1: {wording.NAME_LABEL: "Methyl cellulose"}})
    assert not at.exception
    assert any(s.value == "Methyl cellulose saved." for s in at.success), \
        [s.value for s in at.success]
    saved = FoodOptimizer("burger")
    assert [v["name"] for v in saved.variables] == ["Pea protein",
                                                    "Methyl cellulose"]
    assert saved.recipe_history[0] == {"Pea protein": 10.0,
                                       "Methyl cellulose": 1.0}
    assert saved.quantity_constraints[0]["ingredients"] == ["Methyl cellulose"]
    assert "Methyl cellulose" in list(_grid_frame(at, 0)["Name"])


def test_a_rename_onto_a_name_that_is_taken_refuses_and_changes_nothing(burger):
    """The refusal comes before anything is written: a half-saved grid would
    leave the allowed amounts changed under the old name."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.NAME_LABEL: "Pea protein",
                                         wording.HIGHEST_LABEL: 6.0}})
    assert [e.value for e in at.error] == [wording.row_error(
        2, "Pea protein is already the name of an ingredient. Choose "
           "another name.")]
    saved = FoodOptimizer("burger")
    assert [v["name"] for v in saved.variables] == ["Pea protein",
                                                    "Methylcellulose"]
    assert saved._var_by_name("Methylcellulose")["bounds"] == (0.0, 3.0)


def test_a_setting_is_edited_in_its_own_words(ferment):
    """Type, Lowest, Highest and the setting's own unit, in the row itself."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    grid = _grid_frame(at, 0)
    row = grid[grid["Name"] == "Incubation time"].iloc[0]
    assert (row["Type"], row["Unit"], row["Lowest"]) == (
        wording.KIND_SETTING, "h", 4.0)
    _save_grid(at, ING_GRID, edited={1: {wording.HIGHEST_LABEL: 24.0}})
    assert not at.exception
    assert any(s.value == "Incubation time saved." for s in at.success), \
        [s.value for s in at.success]
    assert FoodOptimizer("ferment")._var_by_name("Incubation time")["bounds"] \
        == (4.0, 24.0)


@pytest.fixture
def mid_run(tmp_path, monkeypatch):
    """A project with one formulation recorded and a setting added after it,
    so the setting carries a baseline the recorded formulation is read at."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("mid_run")
    opt.set_amount_unit("g")
    opt.add_ingredient("Water", 0, 100)
    opt.add_objective("Taste", 1.0, goal="max", min_val=0, max_val=10)
    opt.tell({"Water": 50.0}, {"Taste": 7.0})
    opt.add_process_parameter("Cook temperature", 150, 200, baseline=175,
                              unit="°C")
    return opt


def test_the_grid_shows_the_baseline_a_setting_already_has(mid_run):
    """Changing it moves the formulations already made with it, because the
    baseline is what they were read at."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "mid_run"
    at.run()
    grid = _grid_frame(at, 0)
    row = grid[grid["Name"] == "Cook temperature"].iloc[0]
    assert row[wording.BASELINE_LABEL] == 175.0
    _save_grid(at, ING_GRID, edited={1: {wording.BASELINE_LABEL: 180.0}})
    assert not at.exception, at.exception
    saved = FoodOptimizer("mid_run")
    assert saved._var_by_name("Cook temperature")["_absent_value"] == 180.0
    assert saved.X_history == [saved._encode(r) for r in saved.recipe_history]


def test_an_ingredient_that_already_exists_may_set_its_own_lowest(mid_run):
    """It was recorded at its own amounts, so its Lowest is its own."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "mid_run"
    at.run()
    _save_grid(at, ING_GRID, edited={0: {wording.LOWEST_LABEL: 5.0}})
    assert not at.exception
    assert FoodOptimizer("mid_run")._var_by_name("Water")["bounds"] == (5.0, 100.0)


def test_a_row_is_fixed_by_typing_one_number_in_both_cells(burger):
    """There is no Hold button: a row is pinned at one amount by saying so
    in Lowest and Highest, and the search obeys the two cells."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID, edited={0: {wording.LOWEST_LABEL: 5.0,
                                         wording.HIGHEST_LABEL: 5.0}})
    assert not at.exception
    assert any(s.value == "Pea protein saved." for s in at.success), \
        [s.value for s in at.success]
    saved = FoodOptimizer("burger")
    var = saved._var_by_name("Pea protein")
    assert var["bounds"] == (5.0, 5.0) and saved._fixed_value(var) == 5.0
    grid = _grid_frame(at, 0)
    row = grid[grid["Name"] == "Pea protein"].iloc[0]
    assert (row["Lowest"], row["Highest"]) == (5.0, 5.0)


def test_a_fixed_row_is_given_a_range_again_in_the_same_cells(burger):
    burger.add_ingredient("Pea protein", 20.0, 20.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID, edited={0: {wording.HIGHEST_LABEL: 40.0}})
    assert not at.exception
    assert any(s.value == "Pea protein saved." for s in at.success), \
        [s.value for s in at.success]
    assert FoodOptimizer("burger").fixed_variables() == []


def test_fixing_a_row_that_breaks_the_batch_size_is_refused_over_the_grid(
        burger):
    """Asked ONCE, over the finished grid, and answered above the rows
    rather than against any one of them — the refusal is about the grid as a
    whole, and it names the rows that are fixed."""
    burger.set_formulation_total(26.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID, edited={0: {wording.LOWEST_LABEL: 5.0,
                                         wording.HIGHEST_LABEL: 5.0},
                                     1: {wording.LOWEST_LABEL: 1.0,
                                         wording.HIGHEST_LABEL: 1.0}})
    assert not at.exception
    trouble = [e.value for e in at.error]
    assert len(trouble) == 1, trouble
    assert not trouble[0].startswith("Row "), trouble
    assert wording.fixed_rows_tail("Pea protein and Methylcellulose") \
        in trouble[0], trouble
    saved = FoodOptimizer("burger")
    assert saved._var_by_name("Pea protein")["bounds"] == (0.0, 25.0)


def test_a_limit_already_impossible_is_not_blamed_on_this_save(burger):
    """A range narrowed last week can strand a limit, and always could. A
    save that fixes an unrelated row must not be refused for it — there
    would be no way to edit back out."""
    burger.add_quantity_constraint(["Pea protein"], min_val=40)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.LOWEST_LABEL: 2.0,
                                         wording.HIGHEST_LABEL: 2.0}})
    assert not at.exception
    assert [e.value for e in at.error] == []
    assert FoodOptimizer("burger")._var_by_name(
        "Methylcellulose")["bounds"] == (2.0, 2.0)


def test_a_rename_on_its_own_keeps_the_open_round(burger):
    """Nothing about the question the model was asked has moved, so the round
    on the bench is still the answer to it — and its rows are rewritten to
    the new name rather than thrown away."""
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID,
               edited={1: {wording.NAME_LABEL: "Methyl cellulose"}})
    assert not at.exception
    assert [s.value for s in at.success] == ["Methyl cellulose saved."]
    assert not at.info, [i.value for i in at.info]
    saved = FoodOptimizer("burger")
    assert saved.pending_batch[0]["recipe"] == {"Pea protein": 10.0,
                                                "Methyl cellulose": 1.0}


def test_a_vendor_typed_on_its_own_keeps_the_open_round_too(burger):
    """A vendor is printed on the sheets and read by nothing, so it is not a
    change to the question the model was asked."""
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.VENDOR_LABEL: "Acme"}})
    assert not at.exception
    assert FoodOptimizer("burger").pending_batch is not None


def test_a_grid_saved_unchanged_has_nothing_to_save(burger):
    """The detector reads the two frames, not the editor's record of what
    was typed: a cell typed back to what it held is not a change, and a
    Save over one would retire the open round for nothing."""
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    _grid_edits(at, ING_GRID, edited={1: {wording.HIGHEST_LABEL: 3.0}})
    at.run()
    assert not at.exception
    assert wording.SAVE_CHANGES_BUTTON not in _labels(at), _labels(at)
    assert FoodOptimizer("burger").pending_batch is not None


def test_the_grid_has_no_hold_or_vary_button(burger):
    """0.5.0: pinning a row at one amount is Lowest and Highest, typed in
    the two cells. There is no control row left at all."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    labels = _labels(at)
    assert not [b for b in labels
                if b.startswith(("Hold ", "Vary ", "Edit ", "Delete Pea",
                                 "Set unit", "Save Pea"))
                or " again" in b], labels
    assert not [t for t in at.text_input if t.label in ("New unit", "Unit")]


def test_a_project_with_one_row_can_still_fix_it(tmp_path, monkeypatch):
    """The old Hold button was greyed on the last active row. A range is not
    a button and nothing greys it: the refusal, when it comes, is Generate
    saying every row is fixed."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("one_only")
    opt.set_amount_unit("g")
    opt.add_ingredient("Water", 0, 100)
    opt.add_objective("Taste", 1.0, goal="max", min_val=0, max_val=10)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={0: {wording.LOWEST_LABEL: 50.0,
                                         wording.HIGHEST_LABEL: 50.0}})
    assert not at.exception
    assert FoodOptimizer("one_only")._var_by_name("Water")["bounds"] == \
        (50.0, 50.0)


def test_a_unit_set_on_a_process_setting_breaks_no_limit(burger):
    """A setting carries its own unit, so the same cell sets it — and it is
    part of no sum and no average, so no limit is touched."""
    burger.add_process_parameter("Cook temperature", 150, 200, unit="°C")
    burger.add_quantity_constraint(["Pea protein", "Methylcellulose"],
                                   max_val=20)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={2: {wording.UNIT_LABEL: "°F"}})
    assert not at.exception
    # A setting carries numbers, not amounts: the sentence says which.
    assert any(s.value == ("Cook temperature is now in °F. The numbers you "
                           "already recorded were not converted — check "
                           "them.")
               for s in at.success), [s.value for s in at.success]
    reloaded = FoodOptimizer("burger")
    assert reloaded.unit_of("Cook temperature") == "°F"
    assert len(reloaded.quantity_constraints) == 1
    assert not at.warning, [w.value for w in at.warning]


def test_a_blank_unit_is_allowed_for_a_setting(burger):
    """A mixer speed of 3 is a 3. An ingredient's amounts are added up,
    averaged and printed, so a blank there is a sum of nothing — and that
    one is refused (test_a_blank_unit_cell_is_refused_by_row)."""
    burger.add_process_parameter("Mixer speed", 1, 5, unit="rpm")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={2: {wording.UNIT_LABEL: ""}})
    assert not at.exception
    assert [e.value for e in at.error] == []
    assert FoodOptimizer("burger").unit_of("Mixer speed") == ""


def test_a_unit_change_that_breaks_a_property_limit_says_so(with_properties):
    """A property limit is an average over the amounts, so it needs one unit
    just as a sum does; when a unit change takes that away it is removed, in
    the same one line an amount limit gets."""
    with_properties.add_constraint("Fat per 100 g", max_val=25.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    names = list(_grid_frame(at, 0)["Name"])
    _save_grid(at, ING_GRID,
               edited={names.index("Water"): {wording.UNIT_LABEL: "ml"}})
    assert not at.exception
    assert [w.value for w in at.warning] == [
        "The limit on Fat per 100 g was deleted because the ingredients no "
        "longer share a unit."], [w.value for w in at.warning]
    assert FoodOptimizer("props").constraints == []


def test_deleting_a_row_confirms_names_it_and_keeps_a_copy(burger, tmp_path):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    assert any(w.value == wording.delete_rows_warning("Methylcellulose")
               for w in at.warning), [w.value for w in at.warning]
    assert len(FoodOptimizer("burger").variables) == 2   # not yet
    _submit_button(at, "Yes, delete").click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    assert not at.exception
    assert (tmp_path / "burger_pre_delete.pkl").exists()
    assert [v["name"] for v in FoodOptimizer("burger").variables] == \
        ["Pea protein"]
    assert any(s.value == "Methylcellulose deleted." for s in at.success), \
        [s.value for s in at.success]


def test_a_deletion_leaves_the_tab_exactly_one_lit_button(burger):
    """Through the whole question, not only at the ends of it. The run that
    puts the question up used to draw the lit Save changes AND the lit Yes
    below it, which is the one thing this tab promises never to do."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON]
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    # An edit in hand: Save changes is the lit one and the foot demotes.
    assert _tab_primaries(at, 0) == [wording.SAVE_CHANGES_BUTTON]
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    # The question is up: its Yes is the lit one, and Save changes demotes.
    assert _tab_primaries(at, 0) == [wording.YES_DELETE], _tab_primaries(at, 0)
    assert _grid_save(at, ING_GRID).proto.type == "secondary"
    _submit_button(at, "Cancel").click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    assert _tab_primaries(at, 0) == [wording.SAVE_CHANGES_BUTTON]


def test_a_measurement_deletion_takes_the_colour_the_same_way(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, MEAS_GRID, deleted=[1])
    at.run()
    assert _tab_primaries(at, 0) == [wording.SAVE_CHANGES_BUTTON]
    _grid_save(at, MEAS_GRID).click()
    _grid_edits(at, MEAS_GRID, deleted=[1])
    at.run()
    assert _tab_primaries(at, 0) == [wording.YES_DELETE], _tab_primaries(at, 0)


def test_putting_a_measurement_row_back_disarms_its_deletion(burger):
    """The grid above has always done this; the one below it now does too."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, MEAS_GRID, deleted=[1])
    at.run()
    _grid_save(at, MEAS_GRID).click()
    _grid_edits(at, MEAS_GRID, deleted=[1])
    at.run()
    assert "Yes, delete" in _labels(at)
    _grid_edits(at, MEAS_GRID, deleted=[0])
    at.run()
    assert "Yes, delete" not in _labels(at), _labels(at)
    assert len(FoodOptimizer("burger").objectives) == 2


def test_a_row_rubbed_out_is_refused_not_quietly_deleted(burger):
    """It carries the grid's hidden identity, so it is a row of the project
    with its answers rubbed out — not the empty line at the bottom. Skipped,
    it left the grid, and the save applied a deletion with no question asked
    and no copy kept."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={1: {
        wording.NAME_LABEL: "", wording.TYPE_LABEL: "",
        wording.LOWEST_LABEL: None, wording.HIGHEST_LABEL: None,
        wording.UNIT_LABEL: ""}})
    assert not at.exception
    assert [e.value for e in at.error] == [
        wording.row_error(2, wording.NAME_REQUIRED_ERROR)]
    assert "Yes, delete" not in _labels(at), _labels(at)
    assert [v["name"] for v in FoodOptimizer("burger").variables] == [
        "Pea protein", "Methylcellulose"]


def test_a_name_given_up_in_the_same_save_can_be_taken(burger):
    """Pea protein becomes Pea protein isolate and Methylcellulose takes the
    name it gave up. One edit, one Save, and the write order frees each name
    before it is taken."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID,
               edited={0: {wording.NAME_LABEL: "Pea protein isolate"},
                       1: {wording.NAME_LABEL: "Pea protein"}})
    assert not at.exception
    assert [e.value for e in at.error] == []
    assert [v["name"] for v in FoodOptimizer("burger").variables] == [
        "Pea protein isolate", "Pea protein"]


def test_a_clash_is_blamed_on_the_row_the_reader_typed_into(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID,
               edited={1: {wording.NAME_LABEL: "Pea protein"}})
    assert [e.value for e in at.error] == [wording.row_error(
        2, "Pea protein is already the name of an ingredient. Choose "
           "another name.")]


def _any_armed(at):
    """True while a confirmation is on the books anywhere in the app. What
    it costs is every coloured button on every tab, so a question nothing on
    screen can answer is the worst state this app has."""
    return at.session_state["_armed_confirmation"] \
        if "_armed_confirmation" in at.session_state else None


def test_discard_while_the_question_is_up_takes_the_question_down(burger):
    """Discarding is the reader answering "none of it". Left armed, the
    question was off screen with no Yes to reach and every coloured button
    in the app — on every tab — stayed grey behind it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    assert _tab_primaries(at, 0) == [wording.YES_DELETE]
    _grid_discard(at, ING_GRID).click()
    at.run()
    assert not at.exception
    assert _any_armed(at) is None, _any_armed(at)
    assert "Yes, delete" not in _labels(at), _labels(at)
    # The foot has its colour back, and nothing anywhere is greyed behind a
    # question that is not on screen.
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON]
    assert not _submit_button(at, wording.NEXT_MAKE_BATCH_BUTTON).disabled
    assert len(FoodOptimizer("burger").variables) == 2


def test_putting_the_row_back_while_the_question_is_up_takes_it_down(burger):
    """The commonest way for the question to stop being the question that is
    up: the grid asks for no deletion at all. The early return for that case
    used to skip the disarming altogether."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    assert "Yes, delete" in _labels(at)
    # The row is back, and an ordinary edit is left in its place.
    _grid_edits(at, ING_GRID, edited={0: {wording.HIGHEST_LABEL: 40.0}})
    at.run()
    assert _any_armed(at) is None, _any_armed(at)
    assert "Yes, delete" not in _labels(at), _labels(at)
    assert _tab_primaries(at, 0) == [wording.SAVE_CHANGES_BUTTON]


def test_the_edit_taken_away_entirely_takes_the_question_down(burger):
    """And the other way out of it: nothing left to save, so nothing left to
    ask about."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    assert "Yes, delete" in _labels(at)
    at.run()                      # the grid comes back empty of edits
    assert _any_armed(at) is None, _any_armed(at)
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON]


def test_each_grid_remembers_its_own_armed_deletion(burger):
    """One shared record, and the second grid's bookkeeping popped it out
    from under the first on every run: the first grid's question then
    re-worded itself to a set of rows nobody had confirmed."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    edits = {ING_GRID: dict(deleted=[1]), MEAS_GRID: dict(deleted=[1])}

    def retype(**over):
        for base, e in {**edits, **over}.items():
            _grid_edits(at, base, **e)

    retype()
    at.run()
    _grid_save(at, ING_GRID).click()
    retype()
    at.run()
    assert any("Delete Methylcellulose?" in w.value for w in at.warning), \
        [w.value for w in at.warning]
    # Both grids still have a deletion pending; the record of what THIS
    # question is about must survive the other grid's run.
    retype()
    at.run()
    assert any("Delete Methylcellulose?" in w.value for w in at.warning), \
        [w.value for w in at.warning]
    # Now swap which row the armed grid deletes: the question comes down.
    retype(**{ING_GRID: dict(deleted=[0])})
    at.run()
    assert _any_armed(at) is None, _any_armed(at)
    assert "Yes, delete" not in _labels(at), _labels(at)
    assert len(FoodOptimizer("burger").variables) == 2


def test_a_measurement_name_given_up_in_the_same_save_can_be_taken(burger):
    """The grid above has been able to do this since round 1; the one below
    it now can too."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, MEAS_GRID,
               edited={0: {wording.MEASUREMENT_COLUMN: "Bite"},
                       1: {wording.MEASUREMENT_COLUMN: "Firmness"}})
    assert not at.exception
    assert [e.value for e in at.error] == []
    assert sorted(o["name"] for o in FoodOptimizer("burger").objectives) == [
        "Bite", "Firmness"]


def test_delete_still_offers_the_used_ingredient_path(burger):
    """The permanent-deletion tick box lives inside the confirmation, and
    the refusal comes BEFORE anything is written."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not [c for c in at.checkbox if c.key == "delete_ing_force"]
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    assert at.checkbox(key="delete_ing_force").label == (
        "Delete even though formulations used it — those amounts go too")
    _submit_button(at, "Yes, delete").click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    assert any("was used" in e.value for e in at.error), \
        [e.value for e in at.error]
    assert len(FoodOptimizer("burger").variables) == 2
    # Asked again, this time with the box ticked.
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    at.checkbox(key="delete_ing_force").set_value(True)
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    _submit_button(at, "Yes, delete").click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    assert not at.exception
    assert [v["name"] for v in FoodOptimizer("burger").variables] == \
        ["Pea protein"]


def test_delete_takes_a_process_setting_too(burger, tmp_path):
    burger.add_process_parameter("Cook temperature", 150, 200, unit="°C")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, confirm=True, deleted=[2])
    assert not at.exception
    assert any(s.value == "Cook temperature deleted." for s in at.success), \
        [s.value for s in at.success]
    assert (tmp_path / "burger_pre_delete.pkl").exists()
    assert [v["name"] for v in FoodOptimizer("burger").variables] == \
        ["Pea protein", "Methylcellulose"]


def test_putting_the_row_back_disarms_the_deletion(burger):
    """An armed Save belongs to the rows it was armed over. Put one back and
    the question on screen is about a different set of rows."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[1])
    at.run()
    assert "Yes, delete" in _labels(at)
    _grid_edits(at, ING_GRID, deleted=[0])
    at.run()
    assert "Yes, delete" not in _labels(at), _labels(at)
    assert len(FoodOptimizer("burger").variables) == 2


def test_the_upload_is_folded_away_beneath(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    fold = next(e for e in _tab1(at).expander if e.label == "Or upload an ingredients file")
    assert any(c.value == ("A file with the columns Name, Lowest, Highest "
                           "and, optionally, Unit. Extra columns become "
                           "properties you can set limits on.")
               for c in fold.caption), \
        [c.value for c in fold.caption]
    assert [b.label for b in _unknowns(fold, "download_button")] == \
        ["Download ingredients template (Excel)"]


def test_the_tab_has_no_project_wide_unit_box(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not [t for t in at.text_input if t.key == "amount_unit"], \
        [t.key for t in at.text_input]


def test_the_target_bullet_does_not_claim_a_floor_of_zero(burger):
    """u = max(0, 1 - |norm - target|) with norm clamped to [0, 1], so the
    largest attainable distance is max(t, 1 - t): a mid-range target bottoms
    out at 0.5 and only a target on a range end ever reaches 0."""
    from food_bo import FoodOptimizer as _FO
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    text = _advanced_block(at, wording.HOW_CLOSENESS_HEADING)
    assert "a full range away scores 0" not in text, text
    assert ("by one point per full range; the lowest score depends on how "
            "far the target sits from the ends of your range") in text, text
    # ...and the claim the bullet now makes is the one the code makes: a
    # target of 6 on a 0 to 10 range scores 0.4 closeness at 0, never 0.
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Firmness": 0.0}, formulation_no=1, batch_no=1)
    reloaded = _FO(burger.project_name)
    # Firmness is worth 60 % of the score, and a measured 0 against a target
    # of 6 on a 0 to 10 range is 0.4 closeness — never 0.
    assert float(reloaded.Y_history[0]) == pytest.approx(60.0 * 0.4)


def test_how_it_works_says_what_the_model_does_in_five_lines(burger):
    """It is collapsed, it sits under the measurements table, and it says
    what the model varies, what it aims for, how the next batch is chosen,
    what it will never break and what each suggestion is trying — five lines,
    no arithmetic. The arithmetic is the fold directly beneath, which is the
    one place a banned word is said."""
    from ui_setup import HOW_CLOSENESS, HOW_IT_WORKS
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not _advanced(at).proto.expanded
    text = _advanced_block(at, wording.HOW_IT_WORKS_HEADING)
    for line in HOW_IT_WORKS:
        assert line in text, line
    assert len(HOW_IT_WORKS) == 5
    assert "the model varies" in text and "aims for" in text
    assert "hard rules for every formulation the app suggests" in text
    assert "close to the best or tries something different" in text
    # Nothing here is arithmetic. Share of score IS said — it is the column
    # the reader types into now, so the fold that explains the screen has to
    # name it — but a bare "Share" is still nobody's word.
    assert "closeness =" not in text, text
    assert "Share of score says how much each measurement counts" in text
    assert any("its share suggests" in line for line in HOW_CLOSENESS), \
        HOW_CLOSENESS


def _assert_captions_read_once(at):
    """Sparse: one line each, and never the same line twice. The score
    function is generated from the measurements, not written here."""
    captions = [c for c in _tab1_captions(at)
                if not c.startswith("Overall score = ")]
    assert len(captions) == len(set(captions)), captions
    # The two grid captions are each one sentence longer than the rest: a
    # grid has to say where a new row is typed and what makes a row fixed,
    # and neither has a control of its own to say it any more.
    long_ones = {wording.INGREDIENT_GRID_CAPTION, wording.LIMITS_CAPTION}
    assert all(len(c) <= 100 for c in captions if c not in long_ones), \
        [c for c in captions if len(c) > 100 and c not in long_ones]


def test_no_caption_is_said_twice_on_the_tab(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _assert_captions_read_once(at)


def test_the_tab_reads_in_one_order(burger):
    """Three tiers and nothing else (spec 1.5): the ingredients grid with
    its collapsed file alternative, the measurements grid, one More settings
    and one Advanced. Read off the element tree, top to bottom."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _tab_outline(at, 0) == [
        wording.VARIABLES_HEADER,
        wording.UPLOAD_INGREDIENTS_EXPANDER,
        wording.MEASUREMENTS_HEADER,
        wording.MORE_SETTINGS_EXPANDER,
        wording.ADVANCED_EXPANDER,
    ], _tab_outline(at, 0)
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON]


def test_more_settings_holds_the_four_optional_things_in_order(burger):
    """Default batch size, where the targets came from, the limits, then the
    properties those limits read — everything tab 1 asks at most once, in
    one fold, in that order."""
    burger.add_property("Cost")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    fold = _more_settings(at)
    assert not fold.proto.expanded
    # The box the default batch size is typed into is the first thing in it.
    assert fold.number_input[0].label == "Default batch size (g)"
    headings = [m.value for m in fold.markdown if m.value.startswith("**")]
    assert headings == [wording.LIMITS_HEADING,
                        wording.FINISHED_PRODUCT_LIMIT_HEADING,
                        wording.LIMIT_ON_CHOSEN_INGREDIENTS_HEADING,
                        wording.PROPERTIES_HEADING], headings
    # The note about where the targets came from sits between the box and
    # the limits, and it is the only button above them.
    assert wording.ADD_TARGETS_SOURCE_BUTTON in [b.label for b in fold.button]
    # ...and nothing that belongs to the two grids came with it.
    assert not fold.dataframe or len(fold.dataframe) == 1     # the properties grid
    assert wording.SAVE_CHANGES_BUTTON not in [b.label for b in fold.button]


def test_nothing_else_is_on_the_tab(burger):
    """Spec 1.5 says "nothing else": every element tab 1 draws at its top
    level is one of the five the order test names."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    kinds = [e.__class__.__name__ for e in _tab1(at).children.values()]
    # A subheader and a grid apiece, the error slot under each, the file
    # fold, the two tiers, the dividers and the foot's own button — and no
    # second section, no add form, no control row.
    assert kinds.count("Subheader") == 2
    assert kinds.count("Expander") == 3
    assert "Markdown" not in kinds, kinds


# ------------------------------------------------------------------ #
#  Properties named in the app (no ingredient CSV needed)
# ------------------------------------------------------------------ #

def _prop_name_box(at):
    return at.text_input(key="prop_new")


def test_a_property_is_named_on_the_limits_section_and_the_box_empties(burger):
    """A property used to need an ingredient CSV with an extra column, so a
    project typed in by hand could not limit sodium at all."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _prop_name_box(at).label == "Add a property"
    _prop_name_box(at).set_value("Sodium per 100 g")
    at.run()
    next(b for b in at.button if b.key == "add_property").click()
    at.run()
    assert not at.exception
    assert any("Sodium per 100 g added." in s.value for s in at.success), \
        [s.value for s in at.success]
    assert FoodOptimizer("burger").properties() == ["Sodium per 100 g"]
    # The box empties, or the next click adds the same name again.
    assert _prop_name_box(at).value == ""
    assert at.selectbox(key="prop_metric").options == ["Sodium per 100 g"]


def test_a_property_name_that_clashes_is_refused_on_screen(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _prop_name_box(at).set_value("Firmness")
    at.run()
    next(b for b in at.button if b.key == "add_property").click()
    at.run()
    assert not at.exception
    assert any("already the name of a measurement" in e.value for e in at.error), \
        [e.value for e in at.error]
    assert FoodOptimizer("burger").properties() == []


def test_the_ingredients_grid_carries_no_property_columns(burger):
    """Properties are not a column of the ingredients grid: a project with
    six of them would put six columns of mostly-blank numbers between Name
    and Unit. They have a grid of their own inside More settings."""
    burger.add_property("Sodium per 100 g")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert "Sodium per 100 g" not in _grid_frame(at, 0).columns
    assert "Sodium per 100 g" in _grid_frame(at, 2).columns
    # And the interim fold it was set in is gone.
    assert not hasattr(wording, "SET_PROPERTIES_BUTTON")
    assert "prop_pick" not in at.session_state


def test_the_properties_grid_has_a_row_per_ingredient_only(burger):
    """A process setting is weighed into nothing, so it has no properties
    and no row on this grid."""
    burger.add_property("Cost")
    burger.add_process_parameter("Cook temperature", 150, 200, unit="°C")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    grid = _grid_frame(at, 2)
    assert list(grid.columns) == ["Ingredient", "Cost"]
    assert list(grid["Ingredient"]) == ["Pea protein", "Methylcellulose"]


def test_the_properties_grid_writes_the_figures_typed_into_it(burger):
    burger.add_property("Cost")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_properties(at, edited={1: {"Cost": 42.0}})
    assert not at.exception
    assert any(s.value == wording.PROPERTIES_SAVED for s in at.success), \
        [s.value for s in at.success]
    assert FoodOptimizer("burger").property_value("Methylcellulose", "Cost") == 42.0
    # The row above it is untouched: a figure is written for the cell that
    # moved and for no other.
    assert not FoodOptimizer("burger").has_property_value("Pea protein", "Cost")


def test_the_properties_grid_opens_on_the_figures_already_stored(burger):
    """Every ingredient's figure is readable down one column, which is the
    whole reason the picker and its one box per property went."""
    burger.add_property("Cost")
    burger.set_property_value("Pea protein", "Cost", 3.5)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    figures = list(_grid_frame(at, 2)["Cost"])
    assert figures[0] == 3.5
    # ...and an ingredient with no figure opens on an empty cell, never on
    # the row above it and never on a 0.
    assert pd.isna(figures[1]), figures


def test_a_grid_cell_emptied_clears_the_figure(burger):
    """No figure is not a 0: an ingredient with none counts as 0 in a limit
    and the limit's own line says so, which is a different thing from one
    the reader has measured at 0."""
    burger.add_property("Cost")
    burger.set_property_value("Pea protein", "Cost", 3.5)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_properties(at, edited={0: {"Cost": None}})
    assert not at.exception
    assert not FoodOptimizer("burger").has_property_value("Pea protein", "Cost")


def test_a_properties_grid_saved_unchanged_says_nothing(burger):
    """A green line for a save that wrote nothing is a line about nothing."""
    burger.add_property("Cost")
    burger.set_property_value("Pea protein", "Cost", 3.5)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_properties(at)
    assert not at.exception
    assert not [s.value for s in at.success], [s.value for s in at.success]


def test_the_properties_grid_keeps_the_tab_one_lit_button(burger):
    """`Save properties` is grey, in every state: it lives inside a
    collapsed fold, and the tab's one coloured button is the foot's."""
    burger.add_property("Cost")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON]
    _grid_edits(at, PROP_GRID, edited={0: {"Cost": 3.5}})
    at.run()
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON], \
        _tab_primaries(at, 0)
    # ...and the line that says nothing is written while typing is still
    # there, so an edit in hand is not a silent one.
    assert wording.unsaved_grid_caption(wording.PROPERTIES_NAME) in [
        c.value for c in _more_settings(at).caption]


def test_a_property_is_deleted_from_a_picker_and_one_question(burger):
    burger.add_property("Cost")
    burger.add_property("Fat per 100 g")
    burger.set_property_value("Pea protein", "Cost", 3.5)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert at.selectbox(key="prop_delete").options == ["Cost", "Fat per 100 g"]
    _submit_button(at, wording.delete_button("Cost")).click()
    at.run()
    assert any("Delete Cost?" in w.value for w in at.warning), \
        [w.value for w in at.warning]
    # While the question is up it is the tab's one coloured button.
    assert _tab_primaries(at, 0) == [wording.YES_DELETE], _tab_primaries(at, 0)
    _submit_button(at, wording.YES_DELETE).click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").properties() == ["Fat per 100 g"]
    assert list(_grid_frame(at, 2).columns) == ["Ingredient", "Fat per 100 g"]


def test_a_legacy_ingredient_named_property_is_hidden_but_deletable(burger):
    """A property called "Ingredient" can only ever have arrived as a
    column of an old ingredient file: `add_property` refuses that name
    outright, since the properties grid's own row column already carries
    it. It has to stay off that grid and off the limit picker built from
    the same list — the grid can never show or write a figure for it — but
    the delete picker still has to offer it, or it could never be taken
    out."""
    burger.add_property("Cost")
    burger.ingredient_properties["Pea protein"] = {"Ingredient": 1.0}
    burger.save()
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").properties() == ["Cost", "Ingredient"]
    assert at.selectbox(key="prop_metric").options == ["Cost"]
    assert at.selectbox(key="prop_delete").options == ["Cost", "Ingredient"]
    at.selectbox(key="prop_delete").set_value("Ingredient")
    at.run()
    next(b for b in at.button if b.key == "rm_prop_Ingredient__btn").click()
    at.run()
    assert any("Delete Ingredient?" in w.value for w in at.warning), \
        [w.value for w in at.warning]
    next(b for b in at.button if b.key == "rm_prop_Ingredient__yes").click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").properties() == ["Cost"]


def test_a_limit_names_the_ingredients_that_have_no_value(burger):
    """An ingredient with no value counts as 0 in the average, and a limit
    that looks satisfied for that reason is the one way this lies."""
    burger.add_property("Sodium per 100 g")
    burger.set_property_value("Pea protein", "Sodium per 100 g", 20)
    burger.add_constraint("Sodium per 100 g", max_val=450)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    line = next(t.value for t in at.text if t.value.startswith("Sodium per 100 g:"))
    assert line == ("Sodium per 100 g: at most 450 · Methylcellulose has no "
                    "figure for it and counts as 0.")
    at.session_state["optimizer"].set_property_value(
        "Methylcellulose", "Sodium per 100 g", 0)
    at.run()
    line = next(t.value for t in at.text if t.value.startswith("Sodium per 100 g:"))
    assert line == "Sodium per 100 g: at most 450"


def test_deleting_a_property_asks_first_keeps_a_copy_and_takes_its_limit(
        burger, tmp_path):
    burger.add_property("Sodium per 100 g")
    burger.set_property_value("Pea protein", "Sodium per 100 g", 20)
    burger.add_constraint("Sodium per 100 g", max_val=450)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    next(b for b in at.button
         if b.key == "rm_prop_Sodium per 100 g__btn").click()
    at.run()
    assert any("Delete Sodium per 100 g and its 1 limit?" in w.value
               for w in at.warning), [w.value for w in at.warning]
    assert any(wording.COPY_KEPT in w.value
               for w in at.warning), [w.value for w in at.warning]
    assert FoodOptimizer("burger").properties() == ["Sodium per 100 g"]
    next(b for b in at.button
         if b.key == "rm_prop_Sodium per 100 g__yes").click()
    at.run()
    assert not at.exception
    saved = FoodOptimizer("burger")
    assert saved.properties() == []
    assert saved.constraints == []
    assert saved.ingredient_properties["Pea protein"] == {}
    assert (tmp_path / "burger_pre_delete.pkl").exists()
    assert any("Sodium per 100 g deleted. Its 1 limit went with it." in s.value
               for s in at.success), [s.value for s in at.success]


def test_deleting_an_ingredient_turns_over_the_properties_grid(burger):
    """The properties grid is drawn ingredient by ingredient (row 0 is
    whichever ingredient is first), and a pending, unsaved edit in it is
    positional — Streamlit's data editor knows only that "row 0's Cost"
    changed, not which ingredient that was. Deleting an ingredient above it
    changes who sits in that row, so the properties grid must be reset
    alongside the ingredients grid: a save that lands here must not have a
    stale edit land on whoever inherited the row."""
    burger.add_property("Cost")
    burger.set_property_value("Methylcellulose", "Cost", 4.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    # Delete Pea protein (row 0 of the ingredients grid) and save.
    _save_grid(at, ING_GRID, confirm=True, deleted=[0])
    assert not at.exception
    # A pending edit sitting in the properties grid's ORIGINAL widget, as if
    # typed before the deletion: row 0's Cost set to 99. If the properties
    # grid turned over with the ingredients grid, this key belongs to a
    # widget nothing on screen reads any more.
    stale = {"edited_rows": {0: {"Cost": 99.0}}, "added_rows": [],
             "deleted_rows": []}
    at.session_state["property_grid_0"] = stale
    at.run()
    next(b for b in at.button if b.key == "save_properties").click()
    # AppTest drops widget state it is not handed again on the next run,
    # so the stale edit is put back for the run that processes the click —
    # without this line the test passes with or without the fix.
    at.session_state["property_grid_0"] = stale
    at.run()
    assert not at.exception
    # Methylcellulose is now row 0 of the properties grid. Its own figure
    # must be unchanged, not overwritten by the stale edit meant for Pea
    # protein's old row.
    assert FoodOptimizer("burger").property_value(
        "Methylcellulose", "Cost") == 4.0


def test_a_hand_made_property_limit_reaches_the_batch(burger):
    """The whole point: name a property, give it values, limit it, and the
    formulations the app chooses respect the limit."""
    burger.add_property("Sodium per 100 g")
    burger.set_property_value("Pea protein", "Sodium per 100 g", 1000)
    burger.set_property_value("Methylcellulose", "Sodium per 100 g", 0)
    at = AppTest.from_file(APP_PATH, default_timeout=300)
    at.run()
    at.selectbox(key="prop_metric").select("Sodium per 100 g")
    at.number_input(key="prop_max").set_value(200.0)
    at.run()
    _submit_button(at, "Add property limit").click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").constraints[0]['max'] == 200.0
    at.button(key="generate").click()
    at.run()
    opt = at.session_state["optimizer"]
    assert opt.pending_batch, [e.value for e in at.error]
    for row in opt.pending_batch:
        assert opt.property_per_100(row['recipe'], "Sodium per 100 g") <= 200.0 + 1e-6


# ------------------------------------------------------------------ #
#  The coherence wave: one word per concept, one control per idea
# ------------------------------------------------------------------ #

def test_the_amount_limit_picker_needs_an_ingredient_chosen(burger):
    """The picker's empty state used to mean every ingredient, which wrote a
    limit on all of them while showing none. The total over every ingredient
    now has its own box under the ingredients table, so this picker asks for
    a choice and its button stays dark until one is made."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    picker = at.multiselect(key="qty_pick")
    assert picker.value == []
    assert picker.proto.placeholder != "All ingredients"
    assert at.number_input(key="qc_min").label == "At least (g)"
    assert at.number_input(key="qc_max").label == "At most (g)"
    # One button, not three: there is no second ingredient-limit control.
    assert [b.label for b in at.button
            if "ingredient limit" in b.label] == ["Add ingredient limit"]
    assert _submit_button(at, "Add ingredient limit").disabled
    # Nothing is written by a click on it, and nothing is refused in words
    # either: a disabled button has nothing to say.
    at.number_input(key="qc_max").set_value(100.0)
    at.run()
    assert _submit_button(at, "Add ingredient limit").disabled
    assert FoodOptimizer("burger").quantity_constraints == []
    # A group is picked, and reads in the one list under its own names.
    at.multiselect(key="qty_pick").select("Pea protein")
    at.run()
    assert not _submit_button(at, "Add ingredient limit").disabled
    at.number_input(key="qc_max").set_value(50.0)
    at.run()
    _submit_button(at, "Add ingredient limit").click()
    at.run()
    assert not at.exception
    assert any(s.value.startswith("Limit added on Pea protein.")
               for s in at.success), [s.value for s in at.success]
    assert "Pea protein: at most 50 g" in [t.value for t in at.text]


def test_an_amount_limit_needs_a_number(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.multiselect(key="qty_pick").select("Pea protein")
    at.run()
    _submit_button(at, "Add ingredient limit").click()
    at.run()
    assert [e.value for e in at.error] == ["Enter at least, at most, or both."]
    assert FoodOptimizer("burger").quantity_constraints == []


def test_the_type_column_says_what_the_two_types_are(burger):
    """It decides whether the row is weighed into the batch size, scaled,
    and eligible for an amount limit, and nothing on screen said so. The
    tooltip rides on the column now, not on a radio."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert wording.VARIABLE_TYPE_HELP == (
        "Ingredients are weighed into the formulation and count towards its "
        "total. Process settings, such as temperature or time, are set on "
        "the equipment.")


def test_the_model_settings_in_use_line_shows_only_under_standard(burger):
    """With the four boxes on screen the line repeated them; it is what the
    project is using while the boxes are away."""
    burger.set_bo_config({"kernel": "matern52", "lengthscale_prior": "default",
                          "noise": "default", "acquisition": "qlognei"})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(e.label == wording.ADVANCED_EXPANDER
               for e in at.expander), [e.label for e in at.expander]
    assert at.radio(key="bo_cfg_mode").value == "Expert-selected"
    assert not any(c.value.startswith("In use:") for c in at.caption), \
        [c.value for c in at.caption]
    at.radio(key="bo_cfg_mode").set_value("Standard (default)")
    at.run()
    assert any(c.value.startswith("In use: kernel: matern52")
               for c in at.caption), [c.value for c in at.caption]


def test_off_by_appears_only_when_a_measurement_has_a_target(tmp_path,
                                                             monkeypatch):
    """Off by is a distance from a target; over Higher is better rows it was
    a column of dashes."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("no_target")
    opt.add_ingredient("Water", 0, 100)
    opt.add_objective("Crunch", 1.0, goal="max", min_val=0, max_val=10, unit="N")
    opt.tell({"Water": 50.0}, {"Crunch": 7.0}, formulation_no=1, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    table = next(d.value for d in at.dataframe if "Measured" in d.value.columns)
    assert list(table.columns) == ["Measurement", "Goal", "Measured"]
    # ...and it comes back the moment a measurement has a target.
    opt.add_objective("Firmness", 1.0, goal="target", target=6, min_val=0,
                      max_val=10, unit="N")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    table = next(d.value for d in at.dataframe if "Measured" in d.value.columns)
    assert list(table.columns) == ["Measurement", "Goal", "Measured", "Off by"]


def test_the_partial_sentence_is_said_once_on_the_tab(burger):
    """It belongs under the score or under the table, never twice on one
    screen."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Firmness": 6.0}, formulation_no=1, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    sentence = wording.PARTIAL_SCORES_CAPTION
    assert sentence == ("A formulation missing a measurement scores it as "
                        "zero, so its overall score is low. Record the "
                        "missing number to fix it."), sentence
    assert sum(1 for c in at.caption if c.value == sentence) == 1, \
        [c.value for c in at.caption]
    # With a complete best above it, the table is the one that says it.
    burger.tell({"Pea protein": 12.0, "Methylcellulose": 1.5},
                {"Firmness": 6.0, "Juiciness": 7.0}, formulation_no=2,
                batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert sum(1 for c in at.caption if c.value == sentence) == 1, \
        [c.value for c in at.caption]


def test_the_correct_picker_offers_not_scored_formulations(scored):
    """Formulation 3 was never scored. It used to be missing from the picker
    with a caption saying so; it is now offered, in number order with the
    rest and saying what it is, and the caption says what picking it does."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert at.selectbox(key="correct_formulation").options == [
        "1", "2", "3 · not scored"]
    assert any(c.value == wording.NOT_SCORED_CAN_BE_SCORED_CAPTION
               for c in at.caption), \
        [c.value for c in at.caption]


def test_scoring_a_not_scored_formulation_from_results(scored, tmp_path):
    """The amounts it was generated with are shown, not offered for editing —
    nobody weighed them out twice — and the measurements open empty. Saving
    moves the row into the scored history under its own number and batch."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(3)
    at.run()
    keys = [n.key for n in at.number_input]
    assert "correct_amount_3_Pea protein" not in keys, keys
    assert "correct_3_Firmness" in keys, keys
    assert at.number_input(key="correct_3_Firmness").value is None
    at.number_input(key="correct_3_Firmness").set_value(6.0)
    at.number_input(key="correct_3_Juiciness").set_value(7.0)
    at.run()
    _submit_button(at, wording.SAVE_RESULT_BUTTON).click()
    at.run()
    assert not at.exception
    said = next(s.value for s in at.success if "scored" in s.value)
    # A copy is kept first, exactly as it is for a correction, and the flash
    # ends with the same sentence.
    assert said == (wording.formulation_scored(3) + " "
                    + wording.best_moved(2, 3) + " " + wording.COPY_KEPT)
    opt = FoodOptimizer("burger")
    assert opt.skipped == []
    assert opt.formulation_ids == [1, 2, 3]
    assert opt.batch_history == [1, 1, 1]
    assert opt.recipe_history[2] == {"Pea protein": 25.0,
                                     "Methylcellulose": 3.0}
    # The row was ticked with no reason typed against it, so the note it
    # carried was the marker alone and the scored row carries nothing.
    assert opt.notes_history[2] == ""
    assert opt.next_formulation_no == 4
    # The row is done, so it closes itself exactly as a correction does.
    assert at.session_state["correct_formulation"] is None
    assert (tmp_path / "burger_pre_edit.pkl").exists(), \
        [f.name for f in tmp_path.glob("*.pkl")]


@pytest.fixture
def not_scored_at_a_total(burger):
    """Round 1 printed to 100 g: one formulation scored, one left not scored
    with a reason typed against it."""
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0},
                              {"Pea protein": 15.0, "Methylcellulose": 3.0}],
                             batch_no=1)
    burger.set_pending_batch_total(100.0)
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 5.0}, formulation_no=1,
                batch_no=1)
    burger.record_skipped(2, 1, {"Pea protein": 15.0, "Methylcellulose": 3.0},
                          note=wording.not_scored_with_note("burner failed"))
    burger.set_pending_batch(None)
    return burger


def _amount_tables(at):
    return [t.value for t in at.table
            if wording.AMOUNT_COLUMN in t.value.columns]


def test_the_scoring_row_shows_the_amounts_its_batch_was_made_to(
        not_scored_at_a_total):
    """The stored amounts are as generated; this batch's sheets were printed
    to 100 g, and those are the numbers the bench had in front of it. The
    best block eight lines up already says it that way."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    caution = next(c.value for c in at.caption
                   if c.value.startswith("At 100 g,"))
    assert sum(1 for c in at.caption if c.value == caution) == 1
    at.selectbox(key="correct_formulation").set_value(2)
    at.run()
    assert not at.exception
    assert sum(1 for m in at.markdown
               if m.value == "**Amounts to make it (100 g)**") == 2, \
        [m.value for m in at.markdown]
    table = _amount_tables(at)[-1]
    amounts = dict(zip(table[wording.INGREDIENT_OR_SETTING_LABEL],
                       table[wording.AMOUNT_COLUMN]))
    # 15 g and 3 g in the same proportion, made up to 100 g.
    assert amounts == {"Pea protein": "83.33 g", "Methylcellulose": "16.67 g"}
    # A total the amounts were never chosen for pushes them past what the
    # project allows, and the row says so under the table, in the one
    # sentence tab 2 and the best block both use.
    assert sum(1 for c in at.caption if c.value == caution) == 2, \
        [c.value for c in at.caption]


def test_the_scoring_row_opens_on_the_reason_and_keeps_it_as_the_note(
        not_scored_at_a_total):
    """Why it was not scored is the only thing anyone typed about the row.
    The box holds the reason, without the marker the screen put in front of
    it, and what is left there is the scored row's note."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(2)
    at.run()
    assert at.text_input(key="correct_note_2").value == "burner failed"
    at.number_input(key="correct_2_Firmness").set_value(6.0)
    at.number_input(key="correct_2_Juiciness").set_value(7.0)
    at.run()
    _submit_button(at, wording.SAVE_RESULT_BUTTON).click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").notes_history[-1] == "burner failed"


def test_the_reason_can_be_rewritten_as_the_row_is_scored(
        not_scored_at_a_total):
    """It was typed to explain a missing result. The result is here now, so
    the line that explains the row may be a different one."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(2)
    at.run()
    at.text_input(key="correct_note_2").set_value(
        "made and panelled a day late")
    at.number_input(key="correct_2_Firmness").set_value(6.0)
    at.number_input(key="correct_2_Juiciness").set_value(7.0)
    at.run()
    _submit_button(at, wording.SAVE_RESULT_BUTTON).click()
    at.run()
    assert not at.exception
    opt = FoodOptimizer("burger")
    assert opt.notes_history[-1] == "made and panelled a day late"
    assert opt.skipped == []


def test_scoring_a_not_scored_formulation_with_nothing_typed_is_refused(scored):
    """The same refusal tab 2 gives, and the row stays not scored."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(3)
    at.run()
    _submit_button(at, wording.SAVE_RESULT_BUTTON).click()
    at.run()
    assert not at.exception
    assert any("at least one measurement" in e.value for e in at.error), \
        [e.value for e in at.error]
    assert [s["formulation"] for s in FoodOptimizer("burger").skipped] == [3]


def test_the_one_download_on_the_batch_names_its_format(open_batch):
    """One button, and it says what comes out of it. The preview expander is
    gone with it: a sheet is looked at in the spreadsheet it opens in, not in
    a fold halfway down a tab."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    names = [b.label for b in _unknowns(at.tabs[1], "download_button")]
    assert names == ["Download the round sheets (Excel)"], names
    assert "Preview the printed sheets" not in [e.label for e in at.expander], \
        [e.label for e in at.expander]


# ------------------------------------------------------------------ #
#  Coherence follow-up: every Remove names its object, one shape for
#  the uploaders, one set of property boxes at a time
# ------------------------------------------------------------------ #

def test_every_delete_on_set_up_names_what_it_takes_out(burger):
    """A row of bare `Delete` buttons asks the reader which one is theirs."""
    burger.add_property("Fat per 100 g")
    burger.add_constraint("Fat per 100 g", max_val=20)
    burger.add_quantity_constraint(["Pea protein"], max_val=10)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    labels = _labels(at)
    # A row is deleted by taking it off its grid, and the question Save then
    # asks names it; what is left of the old row of bare Delete buttons is
    # the property list and the limits, and both name what they take out.
    assert "Delete Fat per 100 g" in labels, labels      # the property list
    # A limit's Delete names the limit too: two of them, drawn identically,
    # left the reader counting rows to work out which was which.
    assert "Delete limit on Fat per 100 g" in labels, labels
    assert "Delete limit on Pea protein" in labels, labels
    assert labels.count("Delete limit") == 0, labels
    assert "Delete" not in labels, labels
    _grid_edits(at, ING_GRID, deleted=[0])
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[0])
    at.run()
    assert any("Delete Pea protein?" in w.value for w in at.warning), \
        [w.value for w in at.warning]


def test_deleting_a_limit_is_confirmed_and_keeps_a_copy(burger, tmp_path):
    """It was the one destructive button in the app with no question and no
    copy behind it — and two of them, drawn identically, sat side by side."""
    burger.add_property("Fat per 100 g")
    burger.add_constraint("Fat per 100 g", max_val=20)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at, "Delete limit on Fat per 100 g").click()
    at.run()
    assert any(w.value == ("Delete the limit on Fat per 100 g? The next "
                           "round no longer has to obey it. " + wording.COPY_KEPT)
               for w in at.warning), [w.value for w in at.warning]
    assert FoodOptimizer("burger").constraints != []   # not yet
    _submit_button(at, wording.YES_DELETE).click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").constraints == []
    assert (tmp_path / "burger_pre_delete.pkl").exists(), \
        [f.name for f in tmp_path.glob("*.pkl")]
    assert any(s.value == "Limit on Fat per 100 g deleted. The next round no "
                          "longer has to obey it." for s in at.success), \
        [s.value for s in at.success]


def test_an_ingredient_limits_delete_names_the_ingredients(burger):
    burger.add_ingredient("Water", 0, 60)      # so the pair is not all of them
    burger.add_quantity_constraint(["Pea protein", "Methylcellulose"],
                                   max_val=20)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    labels = _labels(at)
    assert "Delete limit on Pea protein + Methylcellulose" in labels, labels


def test_a_limit_over_every_ingredient_is_named_all_ingredients(burger):
    burger.add_total_mass_constraint(max_val=30)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert "Delete limit on all ingredients" in _labels(at), _labels(at)


def test_the_foot_of_results_steps_aside_when_no_batch_can_be_made(burger):
    """Every measurement deleted: the project holds results but cannot
    generate, so a lit button would lead to a tab holding a greyed
    Generate."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    burger.remove_objective("Juiciness")
    burger.remove_objective("Firmness")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception
    foot = _submit_button(at, wording.START_NEXT_BATCH)
    assert foot.proto.type == "secondary" and foot.disabled
    # And the score caption is suppressed: "0.00 of 0.00" is a sentence
    # about nothing, and the banner above already says what to do.
    assert not any(c.value.startswith("Overall score ") for c in at.caption), \
        [c.value for c in at.caption]


def test_the_correction_picker_says_what_picking_one_does(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    picker = at.selectbox(key="correct_formulation")
    assert picker.label == "Formulation to correct"
    assert picker.proto.placeholder == wording.CHOOSE_A_FORMULATION_PLACEHOLDER


def test_the_third_part_of_edit_past_says_record_not_add(scored):
    """Add puts a formulation into the open batch, to be made; this one was
    made already, and the two doors wore the same verb."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert any(m.value == "##### Record a formulation you already made"
               for m in at.markdown), [m.value for m in at.markdown]
    assert "Record this formulation" in _labels(at), _labels(at)
    assert "Add this formulation" not in _labels(at), _labels(at)


def test_the_whole_batch_pick_says_what_it_does(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    picker = at.selectbox(key="delete_whole_batch")
    assert picker.label == "Add a whole round to the list"
    assert picker.proto.placeholder == "Choose a round"


def test_the_delete_multiselect_placeholder_says_choose_one_or_more(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert (at.multiselect(key="delete_formulations").proto.placeholder
           == "Choose one or more")


def test_the_progress_chart_caption_reads_a_flat_stretch(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert any(c.value == ("The top line only rises. A few flat rounds are "
                           "normal; a long flat stretch suggests this "
                           "ingredient list is close to the best it can do.")
               for c in at.caption), [c.value for c in at.caption]


def test_the_batch_size_help_says_what_it_scales(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert at.number_input(key="scale_total").help == (
        "Every formulation in this round adds up to this. Change it and the "
        "sheets scale with it.")


def test_a_scaled_amount_outside_the_allowed_amounts_is_flagged(room_to_scale):
    """The sheets are printed from the amounts on the table, so a total the
    formulations were never chosen for can send the bench out to weigh an
    amount the project says it does not allow."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(200.0)
    at.run()
    assert not at.exception
    cautions = [c.value for c in at.caption
                if wording.AMOUNTS_YOU_ALLOWED in c.value]
    # ONE line, however many ingredients on however many of the batch's rows
    # are outside: the total did it, and the fix is the same one every time.
    assert cautions == ["At 200 g, Pea protein and Methylcellulose go past "
                        "the amounts you allowed. Print at a smaller total, "
                        "or widen them in Set up."], cautions
    # And none at all once the round is made to a size the project allows.
    # Emptying the box does NOT undo it: the amounts have moved, and the way
    # back is another size, not a blank.
    at.number_input(key="scale_total").set_value(20.0)
    at.run()
    assert not any(wording.AMOUNTS_YOU_ALLOWED in c.value
                   for c in at.caption), [c.value for c in at.caption]


def test_the_properties_grid_says_what_an_empty_cell_holds(burger):
    burger.add_property("Fat per 100 g")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    # The rule is in the caption above the grid, said once. The grid's own
    # head names the property, so nothing inside it has to.
    assert wording.PROPERTY_BLANK_RULE == ("An empty cell counts as 0 in "
                                           "any limit.")
    said = [c.value for c in at.caption
            if wording.PROPERTY_BLANK_RULE in c.value]
    assert said == [wording.properties_grid_caption(True)], said


def test_a_target_is_a_cell_of_its_own_and_is_refused_outside_the_range(
        burger):
    """A grid cannot hide a cell for some rows and not others, so Target is
    always a column — and a target outside the range is refused by row,
    which is the answer the form used to give from a box that came and
    went."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert list(_grid_frame(at, 1)["Target"]) == [6.0, 7.0]
    _save_grid(at, MEAS_GRID, edited={0: {wording.TARGET_LABEL: 50.0}})
    assert [e.value for e in at.error] == [wording.row_error(
        1, "Target 50 must be between the range's lowest and highest "
           "(0 to 10).")]
    assert FoodOptimizer("burger").objectives[1]["target"] == 6.0


def test_a_goal_with_no_target_ignores_the_target_cell(burger):
    """Higher is better has no target, so whatever is in the cell is not
    read: the column is there for the rows that do have one."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, MEAS_GRID,
               edited={0: {wording.GOAL_LABEL: wording.GOAL_LABELS["max"]}})
    assert not at.exception
    assert [e.value for e in at.error] == []
    saved = next(o for o in FoodOptimizer("burger").objectives
                 if o["name"] == "Firmness")
    assert saved["goal"] == "max" and saved["target"] is None


def test_the_correction_amount_boxes_carry_no_tooltip(scored):
    """One caption above the form covers every box on it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    assert not at.number_input(key="correct_amount_1_Pea protein").help
    # The typed-in past formulation has nothing recorded to keep either.
    assert not at.number_input(key="past_Pea protein").help


def test_the_import_caption_says_extra_columns_are_ignored(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.run()
    assert any(c.value.endswith(" Extra columns are ignored.")
               for c in at.caption), [c.value for c in at.caption]


def test_an_extra_column_really_is_ignored(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.session_state["_import_rows"] = pd.DataFrame({
        "Pea protein": [12.0], "Methylcellulose": [1.2],
        "Juiciness": [6.0], "Firmness": [5.0], "Batch": ["whatever"],
    })
    at.run()
    _submit_button(at, "Import all rows").click()
    at.run()
    assert not at.exception
    assert not at.error, [e.value for e in at.error]
    assert FoodOptimizer("burger").formulation_ids == [1]


def test_show_amounts_rounds_like_every_other_table(burger):
    """The batch sheet and the downloaded file both say 11.88; this table
    printed 11.875, which reads as a third number."""
    burger.add_process_parameter("Cook temperature", 100, 200, unit="\u00b0C")
    burger.tell({"Pea protein": 11.875, "Methylcellulose": 1.0,
                 "Cook temperature": 180.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    import ui_results
    frame = burger.history_frame(include_amounts=True)
    rendered = frame.style.format(ui_results._amount_format(burger, frame))
    cells = rendered._translate(False, False)["body"]
    headers = list(frame.columns)

    def cell(column):
        return cells[0][headers.index(column) + 1]["display_value"]

    # The column header carries the unit, exactly as the batch table's does,
    # so only the number is formatted.
    assert cell("Pea protein (g)") == "11.88"
    # A setting is dialled in, not weighed: 180, never 180.00.
    assert cell("Cook temperature (\u00b0C)") == "180"
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.session_state["show_amounts"] = True
    at.run()
    assert not at.exception


def test_a_formulation_recorded_later_is_noted_as_made_earlier(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert at.text_input(key="past_note").value == "Made earlier"
    at.number_input(key="past_Pea protein").set_value(12.0)
    at.number_input(key="past_Methylcellulose").set_value(1.0)
    at.number_input(key="past_m_Firmness").set_value(5.0)
    at.run()
    _submit_button(at, wording.ADD_THIS_FORMULATION).click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").notes_history == ["Made earlier"]


def test_the_restore_uploader_says_what_to_give_it(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(c.value == ("A copy you saved yourself, or one the app saved "
                           "before a change. The current project is copied "
                           "first.")
               for c in at.sidebar.caption), \
        [c.value for c in at.sidebar.caption]


def test_restore_accepts_the_copy_every_confirmation_promises(burger, tmp_path):
    """Every destructive confirmation says a copy is kept and that Open a
    saved copy can load it, so it has to be able to."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    archived = storage_backend.LocalStorage().archive("burger", "pre_delete",
                                                      copy=True)
    assert archived == "burger_pre_delete"
    kept = json.loads((tmp_path / "burger_pre_delete.pkl").read_text())
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["_restore_candidate"] = kept
    at.run()
    assert not at.exception
    assert not at.sidebar.error, [e.value for e in at.sidebar.error]
    assert any("This copy holds **burger**" in w.value
               for w in at.warning), [w.value for w in at.warning]


def test_the_new_project_box_empties_after_creating_one(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.sidebar.text_input[0].set_value("Mine")
    _submit_button(at, wording.CREATE_PROJECT).click()
    at.run()
    assert not at.exception
    assert at.session_state["new_project_name"] == ""
    assert at.sidebar.text_input[0].value == ""


def test_opening_the_sample_after_creating_a_project_warns_about_nothing(
        tmp_path, monkeypatch):
    """A widget given BOTH a default and a session-state entry makes
    Streamlit print a warning on the page, and every handler that opens a
    project assigns project_select."""
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=300)
    at.run()
    at.sidebar.text_input[0].set_value("Mine")
    _submit_button(at, wording.CREATE_PROJECT).click()
    at.run()
    _submit_button(at, wording.TRY_SAMPLE_LABEL).click()
    at.run()
    assert not at.exception
    assert [w.value for w in at.warning] == []


def test_deleting_a_project_asks_and_promises_the_copy(burger):
    """What Delete does to the Open project list is what Delete means."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at, wording.DELETE_PROJECT_LABEL).click()
    at.run()
    assert any(w.value == ("Delete **burger**? It has no formulations yet. "
                           + wording.COPY_KEPT)
               for w in at.warning), [w.value for w in at.warning]
    assert not any("Open project list" in w.value for w in at.warning), \
        [w.value for w in at.warning]


def test_deleting_an_ingredient_says_what_the_records_keep(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, deleted=[0])
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[0])
    at.run()
    assert any("Later formulations keep their numbers." in w.value
               for w in at.warning), [w.value for w in at.warning]


def test_enter_every_amount_says_what_to_do_about_an_ingredient_left_out(
        open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="own_Pea protein").set_value(15.0)   # the other blank
    at.run()
    _submit_button(at, wording.ADD_TO_THIS_BATCH).click()
    at.run()
    assert [e.value for e in at.error] == [
        "Enter every amount. Type 0 for an ingredient you are leaving out."]


def test_the_question_names_the_rows_taken_off_the_grid(burger):
    """Two rows taken out in one edit is one question about both of them,
    not two questions in a row."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, deleted=[0, 1])
    at.run()
    _grid_save(at, ING_GRID).click()
    _grid_edits(at, ING_GRID, deleted=[0, 1])
    at.run()
    assert [w.value for w in at.warning] == [wording.delete_rows_warning(
        "Pea protein and Methylcellulose")], [w.value for w in at.warning]


def test_the_property_button_says_what_it_adds(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert "Add property" in _labels(at), _labels(at)
    # The name carries the unit, and the placeholder shows one.
    assert at.text_input(key="prop_new").proto.placeholder == \
        "e.g. Sodium mg per 100 g"


def test_the_per_100_caption_waits_for_one_unit(mixed_units):
    """`Per 100 g` over a project of grams and millilitres named a hundred of
    nothing. NEEDS_ONE_UNIT already says the ingredients must share a unit
    (the Default batch size box says it first, higher up in More settings),
    so the limit's own caption says nothing at all rather than repeating it
    in different words — the two sentences used to show together."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    captions = [c.value for c in _more_settings(at).caption]
    assert captions.count(wording.NEEDS_ONE_UNIT) == 1, captions
    assert not any(c.startswith("Per 100") for c in captions), captions
    mixed_units.set_ingredient_unit("Water", "g")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(c.value == ("Per 100 g of formulation, from the properties "
                           "of your ingredients.") for c in at.caption), \
        [c.value for c in at.caption]


def test_one_place_holds_every_property_figure(burger):
    """There used to be two — one on the add form, one in Set properties —
    and two boxes for one property on one screen is two answers to one
    question. There is one grid now, and no box anywhere."""
    burger.add_property("Cost")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    keys = [n.key for n in at.number_input]
    assert not [k for k in keys if k.startswith("setprop_")], keys
    assert not [k for k in keys if k.startswith("var_prop_")], keys
    assert [g for g in at.dataframe if "Cost" in g.value.columns], \
        [list(g.value.columns) for g in at.dataframe]
    assert wording.SAVE_PROPERTIES_BUTTON in _labels(at), _labels(at)


def test_both_upload_doors_read_the_same_way(open_batch):
    """One shape for the three uploaders, and one verb for the button that
    reads whatever was put in them."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.run()
    labels = [u.label for u in _unknowns(at.main, "file_uploader")]
    assert "Upload ingredients (Excel or CSV)" in labels, labels
    assert "Upload results (Excel or CSV)" in labels, labels
    assert "Upload formulations (Excel or CSV)" in labels, labels
    assert not any(l == "Results sheet" for l in labels), labels


# ------------------------------------------------------------------ #
#  Tab 2 flow wave: the total each formulation is made to, and the
#  workbook the bench carries away.
# ------------------------------------------------------------------ #
def _workbook(at):
    """The workbook behind `Download the round sheets`, opened.

    AppTest cannot reach a download button's bytes, so the file is built from
    the same optimizer, the same round and the same size the screen drew
    with — which is exactly what the button hands over. open_round_size, not
    the box's own value: a prefill the bench has not changed is not an
    answer, and the sheets carry the amounts the round actually holds."""
    opt = at.session_state["optimizer"]
    return openpyxl.load_workbook(
        io.BytesIO(opt.workbook_bytes(opt.pending_batch,
                                      opt.open_round_size())))


def _formulation_sheets(at):
    """Every per-formulation sheet, in order. Sheet one is the summary."""
    book = _workbook(at)
    return [book[name] for name in book.sheetnames[1:]]


def _sheet_amounts(at):
    """Every (name, number) pair the formulation sheets print: the
    ingredients in set-up order, then the settings under them. The name
    carries its unit wherever the variable has one."""
    out = []
    for sheet in _formulation_sheets(at):
        for row in sheet.iter_rows(values_only=True):
            if (isinstance(row[1], str)
                    and isinstance(row[2], (int, float))
                    and not isinstance(row[2], bool)):
                out.append((row[1], float(row[2])))
    return out


def _sheet_text(at):
    """Every word the formulation sheets carry, in order — what the reader
    would read off the printed pages."""
    out = []
    for sheet in _formulation_sheets(at):
        for row in sheet.iter_rows(values_only=True):
            out += [value for value in row if isinstance(value, str)]
    return out


def _summary_rows(at):
    """The summary sheet as {row label: [cell, ...]}, the labels being the
    ingredients, the total, the measurements, Not scored and Note."""
    sheet = _workbook(at).worksheets[0]
    return {row[0]: list(row[1:]) for row in sheet.iter_rows(values_only=True)
            if isinstance(row[0], str)}


def test_the_batch_size_box_asks_what_one_formulation_weighs(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    box = at.number_input(key="scale_total")
    # Not "Default batch size": that is Set up's box, for every round still
    # to come. This one is the round on screen.
    assert box.label == "Batch size (g)"
    assert box.proto.placeholder == "e.g. 100"
    assert box.help == ("Every formulation in this round adds up to this. "
                        "Change it and the sheets scale with it.")
    box.set_value(25.0)
    at.run()
    # One line about what the sheets hold, and it is the only line about the
    # size: the round is made to 25 g, the two ingredients reach it, and
    # nothing is outside what the project allows.
    assert [c.value for c in at.caption if "25 g" in c.value] == [
        "Sheets show each formulation made to 25 g."], \
        [c.value for c in at.caption]


def test_a_batch_size_the_ingredients_cannot_make_is_refused(open_batch):
    """The app printed round sheets for 250 g directly under a line saying
    250 g was impossible, and the bench would have been sent out to weigh
    amounts the project says it does not allow. The size is refused: the
    round keeps the one it has, the table and the download stay at that
    size, and the box comes back holding it on the next run."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(25.0)
    at.run()
    assert FoodOptimizer("burger").pending_batch_total == 25.0
    at.number_input(key="scale_total").set_value(250.0)
    at.run()
    assert not at.exception
    assert any(c.value == ("A batch size of 250 g is not reachable: the most "
                           "these ingredients can make is 28 g.")
               for c in at.caption), [c.value for c in at.caption]
    # Not made to it: the round, the table and the sheets are all still 25 g.
    saved = FoodOptimizer("burger")
    assert saved.pending_batch_total == 25.0
    table = next(d.value for d in at.dataframe
                 if "Formulation" in d.value.columns)
    assert list(table["Total (g)"]) == pytest.approx([25.0, 25.0])
    assert any(c.value == "Sheets show each formulation made to 25 g."
               for c in at.caption), [c.value for c in at.caption]
    # No extra button: the box itself comes back to the round's own size.
    assert wording.YES_CONTINUE not in _labels(at), _labels(at)
    at.run()
    assert at.number_input(key="scale_total").value == 25.0
    assert FoodOptimizer("burger").pending_batch_total == 25.0


def test_the_box_is_blank_and_refuses_nothing_when_no_size_is_in_force(
        tmp_path, monkeypatch):
    """A prefilled guess was worse than a blank box in two ways: it named a
    size no formulation weighed, and it fired the reach refusal at a number
    nobody had typed. Here the two rows weigh 30 g and 40 g while the allowed
    amounts reach 20 g, so the old mean of 35 g would have been refused out
    loud on the very first render."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("unreachable_mean")
    opt.set_amount_unit("g")
    opt.add_ingredient("Pea protein", 0, 10)
    opt.add_ingredient("Water", 0, 10)
    opt.add_objective("Firmness", 1.0, goal="target", target=6,
                      min_val=0, max_val=10, unit="N")
    opt.set_pending_batch([{"Pea protein": 20.0, "Water": 10.0},
                           {"Pea protein": 20.0, "Water": 20.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert at.number_input(key="scale_total").value is None
    said = [c.value for c in at.caption] + [e.value for e in at.error]
    assert not any("is not reachable" in s for s in said), said
    # ...and nothing was rewritten behind the blank box.
    assert [r['recipe'] for r in FoodOptimizer("unreachable_mean").pending_batch] \
        == [{"Pea protein": 20.0, "Water": 10.0},
            {"Pea protein": 20.0, "Water": 20.0}]
    assert FoodOptimizer("unreachable_mean").pending_batch_total is None


def test_there_is_no_change_the_total_button_any_more(open_batch):
    """The Batch size box is on this screen, above the table. A button that
    left the tab to answer a question the tab already asks was the long way
    round, and it named a number that was not even this round's."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(20.0)
    at.run()
    assert not at.exception
    labels = [b.label for b in at.button]
    assert "Change the total" not in labels, labels


def test_the_size_is_kept_with_the_round_and_then_with_its_number(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(25.0)
    at.run()
    assert FoodOptimizer("burger").pending_batch_total == 25.0
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.number_input(key="f2_Firmness").set_value(7.0)
    at.run()
    _submit_button(at, wording.SAVE_RESULTS).click()
    at.run()
    assert not at.exception
    reloaded = FoodOptimizer("burger")
    assert reloaded.pending_batch_total is None      # the round closed
    assert reloaded.batch_total(1) == 25.0
    # What is recorded is what was MADE. Until 0.5.0 the box scaled the
    # sheets and the model was told the amounts the bench had not weighed.
    assert reloaded.ingredient_total(reloaded.recipe_history[0]) == \
        pytest.approx(25.0)
    assert reloaded.recipe_history[0]["Pea protein"] == pytest.approx(
        25.0 * 10.0 / 11.0)


@pytest.fixture
def resized_and_recorded(tmp_path, monkeypatch):
    """One round, made to 200 g by the Batch size box and recorded there.
    Water's 120 g is past the 100 g the project allows, so tab 2 said so
    before the bench weighed it."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("resized")
    opt.set_amount_unit("g")
    opt.add_ingredient("Water", 0, 100)
    opt.add_ingredient("Flour", 0, 100)
    opt.add_objective("Taste", 1.0, goal="max", min_val=0, max_val=10)
    opt.set_formulation_total(100)
    opt.set_pending_batch([{"Water": 60.0, "Flour": 40.0}])
    opt.scale_round(200.0)
    return opt


def test_a_round_the_box_resized_still_says_so_on_results(resized_and_recorded):
    """The line is on the screen the bench weighs from and must be on the
    screen it is read back on: under a project default the row was never
    rewritten, so without the flag tab 3 checked it against nothing."""
    opt = resized_and_recorded
    assert opt.pending_batch[0]['recipe']["Water"] == pytest.approx(120.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    # Tab 2 says it while the round is on the bench.
    assert [c.value for c in at.caption
            if "Water goes past the amounts you allowed" in c.value], \
        [c.value for c in at.caption]
    at.number_input(key="f1_Taste").set_value(7.0)
    at.run()
    _submit_button(at, wording.SAVE_RESULTS).click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("resized").recorded_total(1) == 200.0
    # The round is closed, so tab 2 has nothing to say any more: whatever
    # says it now is tab 3's own line, under the amounts it is about.
    said = [c.value for c in at.caption
            if "Water goes past the amounts you allowed" in c.value]
    assert said, [c.value for c in at.caption]


def test_discarding_the_batch_forgets_the_total(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(150.0)
    at.run()
    _submit_button(at, wording.GENERATE_DIFFERENT_BATCH).click()
    at.run()
    _submit_button(at, wording.YES_DISCARD).click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").pending_batch_total is None


@pytest.fixture
def made_to_a_total(burger):
    """Round 1 recorded after its sheets were printed to 150 g."""
    burger.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0}],
                             batch_no=1)
    burger.set_pending_batch_total(150.0)
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    burger.set_pending_batch(None)
    return burger


def test_the_best_is_shown_at_the_total_its_batch_was_made_to(made_to_a_total):
    """The stored amounts are as generated; the bench weighed out 150 g of
    them. The heading names the total so the two numbers cannot be confused."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception
    assert any(m.value == "**Amounts to make it (150 g)**"
               for m in at.markdown), [m.value for m in at.markdown]
    table = next(t.value for t in at.table
                 if wording.AMOUNT_COLUMN in t.value.columns)
    amounts = dict(zip(table[wording.INGREDIENT_OR_SETTING_LABEL],
                       table[wording.AMOUNT_COLUMN]))
    # 10 g and 1 g scale to 150 g in the same proportion.
    assert amounts["Pea protein"] == "136.36 g", amounts
    assert amounts["Methylcellulose"] == "13.64 g", amounts


def test_a_batch_made_as_generated_still_reads_as_before(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert any(m.value == "**Amounts to make it**" for m in at.markdown), \
        [m.value for m in at.markdown]


def test_start_from_the_best_still_offers_the_recorded_amounts(made_to_a_total):
    """The boxes ask for amounts the project allows, and the model was told
    the recorded ones: starting from 136.36 g of protein would be a repeat of
    a formulation that was never recorded."""
    made_to_a_total.set_pending_batch([{"Pea protein": 20.0,
                                        "Methylcellulose": 2.0}], batch_no=2)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    _submit_button(at, wording.START_FROM_BEST).click()
    at.run()
    assert not at.exception
    assert at.session_state["own_Pea protein"] == 10.0
    assert at.session_state["own_Methylcellulose"] == 1.0


def test_every_sheet_has_boxes_to_write_in_and_a_line_to_sign(open_batch):
    """A cell with a border is something a pen can fill in. One per
    measurement per sheet — two formulations, two measurements each — plus
    the tick, the note and the two blanks at the foot."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    boxed = 0
    for sheet in _formulation_sheets(at):
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None and cell.border.left.style:
                    boxed += 1
    # Two measurements, the Not scored box, a note box and its overflow, and
    # one Actual cell per ingredient (two) and per setting (none here), on
    # each of the two sheets. The tick cells carry a printed box, so they
    # are bordered AND written in — they are not counted here.
    assert boxed == 2 * (2 + 1 + 2 + 2), boxed
    texts = _sheet_text(at)
    assert texts.count(wording.MADE_BY_FOOTER) == 2, texts
    assert texts.count(wording.NOT_SCORED_CHECKBOX_SHEET) == 2, texts


def test_the_workbook_is_named_for_the_project_and_the_batch(open_batch):
    """A downloaded file is found in a Downloads folder a month later,
    beside eleven others, so it carries both names."""
    assert wording.workbook_file_name("burger", 1) == "burger · Round 1.xlsx"
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    book = _workbook(at)
    # The summary sheet is named for the batch, which is how an uploaded
    # workbook is matched back to the one on the bench.
    assert book.sheetnames[0] == wording.batch_sheet_name(1)
    assert book.sheetnames[1:] == ["Formulation 1", "Formulation 2"]


# ------------------------------------------------------------------ #
#  Tab 2 in work order: make, print, record — with the folded extras
#  after Save and `Generate a different batch` last of all.
# ------------------------------------------------------------------ #
def _tab_flow(at, index=1):
    """Everything with a name on one tab, top to bottom: headings, captions,
    tables, widgets, buttons and expanders, in the order a reader meets them.
    Nested children follow their parent, so an expander's contents stay where
    the expander is."""
    out = []

    def named(node):
        proto = getattr(node, "proto", None)
        for attr in ("label", "value", "body"):
            # A widget whose value was never set raises rather than returning
            # None (tab 3's uploader is one), and it has no words on it
            # either way.
            try:
                text = getattr(node, attr, None)
            except Exception:
                text = None
            if not isinstance(text, str):
                text = getattr(proto, attr, None)
            if isinstance(text, str) and text:
                return text
        return ""

    def walk(node):
        children = getattr(node, "children", None) or {}
        if hasattr(children, "values"):
            children = children.values()
        for child in children:
            text = named(child)
            if text:
                out.append(text)
            walk(child)

    walk(at.tabs[index])
    return out


def _first(order, text):
    assert text in order, (text, order)
    return order.index(text)


def test_tab_two_reads_as_the_three_steps_of_the_work(open_batch):
    """Make them, print the sheets, record what you measured. The screen is
    in that order, and each step says which one it is."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    order = _tab_flow(at)
    assert (_first(order, wording.make_these(1, 2))
            < _first(order, wording.STEP_MAKE_HEADING)
            < _first(order, wording.STEP_PRINT_HEADING)
            < _first(order, wording.STEP_RECORD_HEADING)), order
    # Step 1 is the Batch size box and then the table; step 2 the download;
    # step 3 the grid, the counter and Save results.
    assert (_first(order, wording.STEP_MAKE_HEADING)
            < _first(order, wording.batch_size_label("g"))
            < _first(order, wording.STEP_PRINT_HEADING)
            < _first(order, wording.DOWNLOAD_BATCH_SHEETS)
            < _first(order, wording.STEP_RECORD_HEADING)
            < _first(order, wording.formulation_heading(1))
            < _first(order, "0 of 2 complete")
            < _first(order, wording.SAVE_RESULTS)), order
    # The two folded extras come after Save, in one order, and the one
    # button that throws the batch away is last on the screen. (The preview
    # of the printed sheets was the third: the workbook is looked at in the
    # spreadsheet it opens in.)
    assert (_first(order, wording.SAVE_RESULTS)
            < _first(order, wording.ADD_OWN_EXPANDER)
            < _first(order, wording.UPLOAD_EXPANDER)
            < _first(order, wording.GENERATE_DIFFERENT_BATCH)), order
    assert order[-1] == wording.GENERATE_DIFFERENT_BATCH, order[-4:]
    # Step 1 carries no count of its own: the title above it already has one.
    assert wording.STEP_MAKE_HEADING == "##### 1 · Make the formulations"
    assert wording.STEP_PRINT_HEADING == "##### 2 · Print the sheets"
    assert wording.STEP_RECORD_HEADING == "##### 3 · Record the results"


def test_the_ready_flash_says_what_to_do_next(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    _submit_button(at, wording.generate_button_label(3)).click()
    at.run()
    assert not at.exception
    assert any(s.value == ("Round 1 is ready to make. Print the sheets, then "
                           "record the results below when you have them.")
               for s in at.success), [s.value for s in at.success]


def test_arming_the_discard_still_greys_everything_above_it(open_batch):
    """`Generate a different batch` renders last now, so the question has to
    be asked before the downloads and the grid are drawn into their reserved
    slots — arming a confirmation does not rerun."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.number_input(key="f2_Firmness").set_value(4.0)
    at.run()
    assert _tab_primaries(at, 1) == [wording.SAVE_RESULTS], _tab_primaries(at, 1)
    _submit_button(at, wording.GENERATE_DIFFERENT_BATCH).click()
    at.run()
    assert _tab_primaries(at, 1) == [wording.YES_DISCARD], _tab_primaries(at, 1)
    assert _submit_button(at, wording.SAVE_RESULTS).disabled
    assert _submit_button(at, wording.ADD_TO_THIS_BATCH).disabled
    assert _unknown(at.main, "download_button",
                    wording.DOWNLOAD_BATCH_SHEETS).proto.type == "secondary"
    # The warning is still on screen, above the Yes.
    assert any(w.value == wording.regenerate_warning(1, "1 and 2", 3)
               for w in at.warning), [w.value for w in at.warning]


def test_adding_your_own_formulation_keeps_what_was_typed_into_the_grid(
        scored_open_batch):
    """The expander reruns, and Streamlit discards the session-state entry of
    every widget the run did not create. It renders after the grid now, so
    the grid exists by the time Add reruns."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="f4_Firmness").set_value(5.0)
    at.text_input(key="f4_note").set_value("second try")
    at.run()
    at.number_input(key="own_Pea protein").set_value(15.0)
    at.number_input(key="own_Methylcellulose").set_value(1.5)
    at.run()
    _submit_button(at, wording.ADD_TO_THIS_BATCH).click()
    at.run()
    assert not at.exception
    assert at.session_state["f4_Firmness"] == 5.0
    assert at.session_state["f4_note"] == "second try"
    # The row that was just added has no result in it yet, so nothing is lit:
    # the batch sheet stepped aside when the first value was typed, and Save
    # waits for every kept row. Filling the new row lights it again.
    assert _tab_primaries(at, 1) == [], _tab_primaries(at, 1)
    at.number_input(key="f5_Firmness").set_value(7.0)
    at.run()
    assert _tab_primaries(at, 1) == [wording.SAVE_RESULTS], _tab_primaries(at, 1)


def test_how_it_works_is_five_lines(burger):
    """One bullet per thing the model does. The formulas moved to a fold of
    their own: they answered a question the first four raise. The fifth is
    what each suggestion is trying, which is the column tab 2 now carries."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert wording.HOW_IT_WORKS == [
        "Ingredients and settings are what the model varies. Measurements "
        "and goals are what it aims for.",
        "Share of score says how much each measurement counts, out of 100. "
        "Closeness is a 0 to 1 score for how near a result is to its goal.",
        "Until five formulations have results, new ones are spread out to "
        "learn the space. After that, each round aims closer to your targets.",
        "Limits are hard rules for every formulation the app suggests. "
        "A formulation of your own is recorded as you typed it.",
        "Each suggestion says whether it stays close to the best or tries "
        "something different, and what it changes.",
    ], wording.HOW_IT_WORKS
    assert not _advanced(at).proto.expanded
    text = _advanced_block(at, wording.HOW_IT_WORKS_HEADING)
    for line in wording.HOW_IT_WORKS:
        assert f"- {line}" in text, line     # five flat bullets, none nested
    assert not hasattr(wording, "HOW_IT_WORKS_NESTED")
    # It is a heading inside Advanced now, not a fold of its own: three
    # things to open before the first sentence is what the tiers end.
    assert not hasattr(wording, "HOW_IT_WORKS_EXPANDER")


def test_the_closeness_formulas_are_the_block_below(burger):
    """The arithmetic answers a question the four bullets above it raise, so
    it reads directly beneath them inside the same fold."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    headings = [m.value for m in _advanced(at).markdown
                if m.value.startswith("**")]
    assert headings == [wording.HOW_FORMULATIONS_CHOSEN_HEADING,
                        wording.HOW_IT_WORKS_HEADING,
                        wording.HOW_CLOSENESS_HEADING], headings
    assert wording.HOW_CLOSENESS_HEADING == "**How closeness is calculated**"
    text = _advanced_block(at, wording.HOW_CLOSENESS_HEADING)
    for line in wording.HOW_CLOSENESS:
        assert line in text, line
    # The three goals, the floor and the share it costs, the re-scoring,
    # and what a repeat teaches the model. The property rule left this fold
    # for the Limits caption: a property never touches closeness.
    joined = " ".join(wording.HOW_CLOSENESS)
    assert "Higher is better" in joined and "Lower is better" in joined
    assert "Hit a target" in joined
    assert ("by one point per full range; the lowest score depends on how "
            "far the target sits from the ends of your range") in joined
    assert "a little less than its share suggests" in joined
    assert ("The model learns the one overall score, so changing a share, "
            "a goal or a range re-scores every past formulation.") in joined
    assert "property" not in joined, joined
    assert "how noisy your measurements are" in joined


def test_the_getting_started_caption_is_the_bullet_word_for_word(burger):
    """The same fact cannot be two sentences, and it cannot be two captions
    either: one line covers both sides of the fifth formulation."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    line = wording.HOW_IT_WORKS[2]
    assert any(c.value == line for c in at.tabs[1].caption), \
        [c.value for c in at.tabs[1].caption]
    for i in range(5):
        burger.tell({"Pea protein": 5.0 + i, "Methylcellulose": 1.0},
                    {"Juiciness": 6.0, "Firmness": 6.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    assert any(c.value == line for c in at.tabs[1].caption), \
        [c.value for c in at.tabs[1].caption]


# ------------------------------------------------------------------ #
#  Fix round 1: the stored total opens the box, and a Cancel keeps
#  what was typed into the grid.
# ------------------------------------------------------------------ #
def test_the_stored_total_opens_the_box_in_a_new_session(open_batch):
    """A reopened window knows nothing of what was typed. The box has to open
    at the total the batch is stored with — an empty box on that first render
    wrote its own blank straight over the saved number."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(25.0)
    at.run()
    assert FoodOptimizer("burger").pending_batch_total == 25.0

    fresh = AppTest.from_file(APP_PATH, default_timeout=180)
    fresh.session_state["_loaded_project"] = "burger"
    fresh.session_state["main_tab"] = wording.TAB_BATCH
    fresh.run()
    assert not fresh.exception
    assert fresh.number_input(key="scale_total").value == 25.0
    table = next(d.value for d in fresh.dataframe
                 if "Formulation" in d.value.columns)
    assert list(table["Total (g)"]) == pytest.approx([25.0, 25.0])
    assert any(c.value == "Sheets show each formulation made to 25 g."
               for c in fresh.caption), [c.value for c in fresh.caption]
    # ...and nothing about that render touched the file.
    assert FoodOptimizer("burger").pending_batch_total == 25.0


def test_emptying_the_box_leaves_the_round_at_the_size_it_was_made_to(
        open_batch):
    """An emptied box is not an undo. Since 0.5.0 the size moves the amounts
    themselves, so there is no as-generated to go back to — the round stays
    at 25 g and the way to another size is to type one. And an emptied box
    is not an answer either: it comes back holding the size the round is
    made to, rather than sitting blank over amounts it says nothing about."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(25.0)
    at.run()
    # An emptied box comes back as 0, the box's own floor.
    at.number_input(key="scale_total").set_value(0.0)
    at.run()
    assert not at.exception
    assert at.number_input(key="scale_total").value == 25.0
    # Nothing is rewritten: the round is still the 25 g one.
    table = next(d.value for d in at.dataframe
                 if "Formulation" in d.value.columns)
    assert list(table["Total (g)"]) == pytest.approx([25.0, 25.0])
    assert FoodOptimizer("burger").pending_batch_total == 25.0


def test_a_project_switch_back_opens_the_other_batch_at_its_own_total(
        open_batch):
    """Two projects, two open batches, two totals. A switch parks the box
    empty, so the arriving project has to fill it from its own record."""
    other = FoodOptimizer("second")
    other.set_amount_unit("g")
    other.add_ingredient("Flour", 0, 100)
    other.add_objective("Crunch", 1.0, goal="max", min_val=0, max_val=10)
    other.set_pending_batch([{"Flour": 10.0}])
    other.set_pending_batch_total(80.0)

    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="scale_total").set_value(25.0)
    at.run()
    at.sidebar.selectbox(key="project_select").select("second")
    at.run()
    _submit_button(at.sidebar, wording.OPEN_BUTTON).click()
    at.run()
    assert not at.exception
    assert at.number_input(key="scale_total").value == 80.0
    assert FoodOptimizer("burger").pending_batch_total == 25.0
    assert FoodOptimizer("second").pending_batch_total == 80.0


def test_cancelling_the_discard_keeps_what_was_typed_into_the_grid(open_batch):
    """The question is asked above the grid, so its Cancel reruns before the
    grid exists — and Streamlit discards the session-state entry of every
    widget a run did not create. A Cancel that empties the sheet you have
    already half-recorded is worse than no Cancel at all."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.number_input(key="f1_Juiciness").set_value(7.0)
    at.text_input(key="f1_note").set_value("looked good")
    at.number_input(key="f2_Firmness").set_value(4.0)
    at.number_input(key="f2_Juiciness").set_value(5.0)
    at.number_input(key="scale_total").set_value(150.0)
    at.run()
    assert _tab_primaries(at, 1) == [wording.SAVE_RESULTS], _tab_primaries(at, 1)
    _submit_button(at, wording.GENERATE_DIFFERENT_BATCH).click()
    at.run()
    _submit_button(at, wording.CANCEL).click()
    at.run()
    assert not at.exception
    assert at.session_state["f1_Firmness"] == 6.0
    assert at.session_state["f1_Juiciness"] == 7.0
    assert at.session_state["f1_note"] == "looked good"
    assert at.session_state["f2_Firmness"] == 4.0
    assert at.session_state["scale_total"] == 150.0
    # A second run to read the settled screen: the tree AppTest hands back
    # after a click that ends in st.rerun() is the interrupted run's, which
    # still carries the Yes the Cancel was answering.
    at.run()
    assert _tab_primaries(at, 1) == [wording.SAVE_RESULTS], _tab_primaries(at, 1)
    assert wording.YES_DISCARD not in _labels(at), _labels(at)


def test_the_scaled_cautions_read_under_the_box_that_caused_them(
        room_to_scale):
    """The total is what pushed the amount out of range, so the line belongs
    with the box, not under a table three steps above it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(200.0)
    at.run()
    order = _tab_flow(at)
    caution = next(t for t in order if wording.AMOUNTS_YOU_ALLOWED in t)
    assert (_first(order, wording.batch_size_label("g"))
            < order.index(caution)
            < _first(order, wording.STEP_RECORD_HEADING)), order


# ------------------------------------------------------------------ #
#  Fix round 2: a total that changed under the session — a restore, a
#  reload — still opens the box.
# ------------------------------------------------------------------ #
def _burger_donor(name, total=None):
    """A backup of a project shaped like burger, with its own open batch."""
    donor = FoodOptimizer(name)
    donor.set_amount_unit("g")
    donor.add_ingredient("Pea protein", 0, 25)
    donor.add_ingredient("Methylcellulose", 0, 3)
    donor.add_objective("Juiciness", 1.0, goal="target", target=7,
                        min_val=0, max_val=10, unit="/10")
    donor.add_objective("Firmness", 1.5, goal="target", target=6,
                        min_val=0, max_val=10, unit="N")
    donor.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0},
                             {"Pea protein": 20.0, "Methylcellulose": 2.0}],
                            batch_no=1)
    if total is not None:
        donor.set_pending_batch_total(total)
    return donor.export_json()


def test_a_restored_total_opens_a_box_the_session_had_seen_empty(open_batch):
    """The session had already rendered batch 1 with no total, so a seed
    keyed on the project and the batch number alone was skipped — and the
    empty box then wrote itself over the 150 the restore had just put there."""
    state = _burger_donor("donor_total", total=150.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    assert at.number_input(key="scale_total").value is None
    at.session_state["_restore_candidate"] = state
    at.run()
    _submit_button(at, wording.YES_REPLACE).click()
    at.run()
    assert not at.exception
    assert at.number_input(key="scale_total").value == 150.0
    assert FoodOptimizer("burger").pending_batch_total == 150.0
    table = next(d.value for d in at.dataframe
                 if "Formulation" in d.value.columns)
    assert list(table["Total (g)"]) == pytest.approx([150.0, 150.0])


def test_a_reloaded_total_opens_a_box_the_session_had_seen_empty(open_batch):
    """Reload project re-reads the file under a session that had already
    rendered the same batch number. The box has to follow what came back."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    assert at.number_input(key="scale_total").value is None
    # The file says 150 now; this session never saw it typed.
    FoodOptimizer("burger").set_pending_batch_total(150.0)
    at.session_state["optimizer"].save_error = (
        "The server could not be reached — your last change was NOT saved.")
    at.run()
    _submit_button(at, wording.RELOAD_PROJECT).click()
    at.run()
    assert not at.exception
    assert at.number_input(key="scale_total").value == 150.0
    assert FoodOptimizer("burger").pending_batch_total == 150.0


def test_the_sessions_own_write_is_not_read_as_a_change_from_elsewhere(
        open_batch):
    """The mark carries the size the box was last put on screen holding, so
    it has to be re-stamped after every write of the session's own —
    otherwise the session's own 25 reads as a size that changed somewhere
    else, and scale_round runs again on every rerun."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="scale_total").set_value(25.0)
    at.run()
    saved_at = at.session_state["optimizer"].last_saved_at
    at.run()
    at.run()
    assert not at.exception
    assert at.number_input(key="scale_total").value == 25.0
    # No write per rerun: the size the session itself typed is the size the
    # mark already carries.
    assert at.session_state["optimizer"].last_saved_at == saved_at
    assert FoodOptimizer("burger").pending_batch_total == 25.0


def test_a_delete_with_no_batch_open_empties_the_total_box_next_run(burger):
    """The clear must come AFTER preserve_tab_forms(): that call parks every
    tab-form key at its live value, scale_total included, so a blank parked
    first would be overwritten by the old total and the box would keep it."""
    opt = burger
    for i in range(4):
        opt.import_formulation({"Pea protein": 5.0 + i, "Methylcellulose": 0.5},
                               {"Juiciness": 6.0, "Firmness": 5.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.run()
    at.session_state["scale_total"] = 150.0     # a value left in the session
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    numbers = at.multiselect(key="delete_formulations").options
    at.multiselect(key="delete_formulations").select(numbers[0])
    at.run()
    _submit_button(at, wording.delete_formulation_button(int(numbers[0]))).click()
    at.run()
    _submit_button(at, wording.YES_DELETE).click()
    at.run()
    at.run()                       # the delete handler reruns
    assert not at.exception
    assert at.session_state["scale_total"] is None, at.session_state["scale_total"]


# ------------------------------------------------------------------ #
#  Fixes from the browser run: one caution for the scaled amounts on
#  every surface, an honest restore count, an out-of-range result said
#  as it is typed, a picker that clears itself, and notes that survive
#  a CSV round trip.
# ------------------------------------------------------------------ #
def test_one_caution_names_every_ingredient_the_total_pushes_out(
        room_to_scale):
    """One line under the box, not one line per ingredient per row: the fix
    is the same one every time, and eight captions of raw floats buried the
    Record step under them."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(200.0)
    at.run()
    assert not at.exception
    said = [c.value for c in at.caption
            if wording.AMOUNTS_YOU_ALLOWED in c.value]
    assert said == ["At 200 g, Pea protein and Methylcellulose go past "
                    "the amounts you allowed. Print at a smaller total, or "
                    "widen them in Set up."], said


def test_the_scaled_caution_is_singular_for_one_ingredient(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="scale_total").set_value(28.0)
    at.run()
    assert not at.exception
    said = [c.value for c in at.caption
            if wording.AMOUNTS_YOU_ALLOWED in c.value]
    assert said == ["At 28 g, Pea protein goes past the amounts you allowed. "
                    "Print at a smaller total, or widen them in "
                    "Set up."], said


def test_the_best_says_when_its_own_total_pushes_an_amount_out(made_to_a_total):
    """Tab 3 shows the best at the total its batch was made to, so it is the
    same warning about the same numbers — in the same words, under the table
    that shows them."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception
    caution = ("At 150 g, Pea protein and Methylcellulose go past the "
               "amounts you allowed. Print at a smaller total, or widen them "
               "in Set up.")
    order = _tab_flow(at, 2)
    assert order.count(caution) == 1, order
    assert (_first(order, wording.amounts_to_make_it_heading("150 g"))
            < order.index(caution)), order


def test_every_sheet_ends_with_the_caution(open_batch):
    """The sheet is what the bench weighs out from, and it leaves the app: a
    caution only on screen is not on the page in the technician's hand."""
    caution = ("At 200 g, Pea protein and Methylcellulose go past the "
               "amounts you allowed. Print at a smaller total, or widen them "
               "in Set up.")
    book = openpyxl.load_workbook(io.BytesIO(
        open_batch.workbook_bytes(open_batch.pending_batch, 200.0)))
    sheets = [book[name] for name in book.sheetnames[1:]]
    assert len(sheets) == 2, book.sheetnames
    for sheet in sheets:
        text = [value for row in sheet.iter_rows(values_only=True)
                for value in row if isinstance(value, str)]
        # Directly under the amounts it is about — under the Total line, not
        # at the foot of the page under the signature.
        assert text[text.index(wording.TOTAL_LABEL) + 1] == caution, text
        assert text.count(caution) == 1, text
    # ... and once on the summary sheet, under the amounts block there.
    summary = book[book.sheetnames[0]]
    summary_text = [value for row in summary.iter_rows(values_only=True)
                    for value in row if isinstance(value, str)]
    assert summary_text.count(caution) == 1, summary_text
    # A batch made as generated has nothing to caution about.
    plain = openpyxl.load_workbook(io.BytesIO(
        open_batch.workbook_bytes(open_batch.pending_batch, None)))
    assert not any(isinstance(value, str)
                   and wording.AMOUNTS_YOU_ALLOWED in value
                   for name in plain.sheetnames
                   for row in plain[name].iter_rows(values_only=True)
                   for value in row)


def test_the_restore_flash_counts_the_not_scored_rows_too(project_with_history,
                                                           tmp_path):
    """The preview counts scored and not-scored alike; a flash that counted
    only the scored ones reported losing formulations the restore had just
    put back."""
    donor = FoodOptimizer("donor_counts")
    donor.add_ingredient("Flour", 0, 100)
    donor.add_objective("Crunch", 1.0, goal="max", min_val=0, max_val=10)
    for i in range(5):
        donor.tell({"Flour": 10.0 * i}, {"Crunch": 5.0})
    donor.record_skipped(6, 1, {"Flour": 20.0})
    donor.record_skipped(7, 1, {"Flour": 30.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "my_project"
    at.session_state["_restore_candidate"] = donor.export_json()
    at.run()
    assert any("7 formulations" in w.value for w in at.warning), \
        [w.value for w in at.warning]
    _submit_button(at, wording.YES_REPLACE).click()
    at.run()
    assert not at.exception
    assert any("Restored 7 formulations" in s.value for s in at.success), \
        [s.value for s in at.success]


def test_an_out_of_range_result_is_said_where_it_was_typed(open_batch):
    """Save greys until every row has a value, so an out-of-range reading
    left the button grey with nothing to explain it. The refusal on Save
    stays; this is the same sentence, said as it is typed."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.number_input(key="f1_Firmness").set_value(12.0)
    at.run()
    assert not at.exception
    said = [c.value for c in at.caption if "outside your range" in c.value]
    assert said == ["Firmness 12 N is outside your range of 0 to 10 N."
                    " Widen the range in Set up, or check the value."], said
    order = _tab_flow(at)
    assert (order.index(said[0])
            < _first(order, wording.formulation_heading(2))), order
    save = _submit_button(at, wording.SAVE_RESULTS)
    assert save.disabled and save.proto.type == "secondary"


def test_the_picker_clears_after_a_saved_correction(scored):
    """A correction that stayed open kept the tab's lit button on Save
    correction, so the foot of the tab had no next action to offer."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    at.number_input(key="correct_1_Firmness").set_value(6.0)
    at.run()
    _submit_button(at, wording.SAVE_CORRECTION_BUTTON).click()
    at.run()
    at.run()        # the settled screen, after the rerun the save asked for
    assert not at.exception
    assert at.session_state["correct_formulation"] is None
    assert _tab_primaries(at, 2) == [wording.START_NEXT_BATCH], \
        _tab_primaries(at, 2)


def test_a_note_survives_a_csv_round_trip(burger):
    """`Download all formulations (CSV)` writes the Note column; an importer
    that ignored it turned every note in the file into `Made earlier`."""
    burger.tell({"Pea protein": 12.0, "Methylcellulose": 1.2},
                {"Juiciness": 6.0, "Firmness": 5.0}, formulation_no=1,
                batch_no=1, note="from the 2024 bench book")
    burger.tell({"Pea protein": 14.0, "Methylcellulose": 1.4},
                {"Juiciness": 6.5, "Firmness": 5.5}, formulation_no=2,
                batch_no=1)                                   # no note at all
    downloaded = pd.read_csv(io.StringIO(FoodOptimizer("burger").history_csv()))
    assert "Note" in downloaded.columns, list(downloaded.columns)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["add_past_mode"] = wording.UPLOAD_A_FILE
    at.session_state["_import_rows"] = downloaded
    at.run()
    _submit_button(at, wording.IMPORT_ALL_ROWS_BUTTON).click()
    at.run()
    assert not at.exception
    assert not at.error, [e.value for e in at.error]
    reloaded = FoodOptimizer("burger")
    assert reloaded.notes_history == ["from the 2024 bench book", "",
                                      "from the 2024 bench book",
                                      wording.IMPORTED_NOTE], \
        reloaded.notes_history


# ------------------------------------------------------------------ #
#  Copy wave (2026-09-13): the behaviour behind the words. A button
#  that names what it deletes, a counter that says "complete", the
#  property controls named after the properties, the damaged-file
#  banner, and the caution that counts instead of listing.
# ------------------------------------------------------------------ #
def test_the_delete_button_names_the_formulations_it_will_delete(scored):
    """"Delete 2 formulations" and "Delete Formulation 2" are one keystroke
    apart in meaning: a user who picked rows 2 and 3 read the count as the
    row number. Up to three rows the button prints the numbers."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.multiselect(key="delete_formulations").set_value([2, 3])
    at.run()
    assert not at.exception
    assert "Delete Formulations 2 and 3" in _labels(at), _labels(at)
    at.multiselect(key="delete_formulations").set_value([2])
    at.run()
    assert "Delete Formulation 2" in _labels(at), _labels(at)


def test_the_delete_button_counts_above_three_formulations(burger):
    """Four names in a button is a sentence, not a label."""
    for n in range(1, 6):
        burger.tell({"Pea protein": float(n), "Methylcellulose": 1.0},
                    {"Juiciness": 7.0, "Firmness": 6.0},
                    formulation_no=n, batch_no=1)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.multiselect(key="delete_formulations").set_value([1, 2, 3])
    at.run()
    assert "Delete Formulations 1, 2 and 3" in _labels(at), _labels(at)
    at.multiselect(key="delete_formulations").set_value([1, 2, 3, 4, 5])
    at.run()
    assert not at.exception
    assert "Delete 5 formulations" in _labels(at), _labels(at)


def test_the_results_counter_counts_the_rows_that_are_complete(open_batch):
    """A row with one of two measurements typed is not half a result, and
    the line beneath the boxes says so in one word."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="f1_Juiciness").set_value(7.0)
    at.number_input(key="f1_Firmness").set_value(6.0)
    at.number_input(key="f2_Juiciness").set_value(7.0)
    at.run()
    assert not at.exception
    said = [c.value for c in at.caption if "complete" in c.value]
    assert said == ["1 of 2 complete · 1 partly filled"], \
        [c.value for c in at.caption]
    # ...and the button directly beneath it is no longer named in the line.
    assert not any("saved when you press" in c.value for c in at.caption), \
        [c.value for c in at.caption]


def test_the_properties_grid_names_the_properties_across_its_head(burger):
    """A property name is the project's own and can be long, so the grid's
    columns do the naming and its caption says only what the head cannot:
    what one figure is per, and what an empty cell means."""
    burger.add_property("Cost")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert list(_grid_frame(at, 2).columns) == ["Ingredient", "Cost"]
    assert any(c.value == ("Each ingredient's figure, per 100 g. An empty "
                           "cell counts as 0 in any limit.")
               for c in at.caption), [c.value for c in at.caption]
    _save_properties(at, edited={0: {"Cost": 42.0}})
    assert any(s.value == wording.PROPERTIES_SAVED for s in at.success), \
        [s.value for s in at.success]
    assert FoodOptimizer("burger").property_value("Pea protein", "Cost") == 42.0


def test_the_caption_does_not_say_per_100_g_twice_over(burger):
    """Property names carry the basis as often as not, and a caption that
    then added ", per 100 g" said it twice in one line."""
    burger.add_property("Fat per 100 g")
    burger.add_property("Sodium per 100 g")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert any(c.value == ("Each ingredient's figure. An empty cell counts "
                           "as 0 in any limit.")
               for c in at.caption), [c.value for c in at.caption]
    # ...and one column whose name does NOT carry it puts the basis back.
    burger.add_property("Cost")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(c.value == ("Each ingredient's figure, per 100 g. An empty "
                           "cell counts as 0 in any limit.")
               for c in at.caption), [c.value for c in at.caption]


def test_the_damaged_file_banner_says_the_two_ways_out(tmp_path, monkeypatch):
    """The one screen where the user is already alarmed: two short sentences,
    each naming a control that is on screen beside them."""
    monkeypatch.chdir(tmp_path)
    (tmp_path / "broken.pkl").write_bytes(b"\x80\x04 not json not pickle \xff\xfe")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "broken"
    at.run()
    assert not at.exception
    assert wording.project_load_error_info() == (
        "This project file is damaged, so editing is off. Two ways out, both "
        "in the sidebar: Open a saved copy, if you saved one. Or "
        "Manage project › Start this project over — the damaged file is "
        "copied first.")
    assert any(i.value == wording.project_load_error_info() for i in at.info), \
        [i.value for i in at.info]


@pytest.fixture
def eight_ingredients(tmp_path, monkeypatch):
    """A project wide enough for the scaled caution to have to choose between
    naming the ingredients and counting them."""
    monkeypatch.chdir(tmp_path)
    opt = FoodOptimizer("wide")
    opt.set_amount_unit("g")
    # Four narrow rows and four wide ones: 400 g is a size these eight can
    # add up to (448 g between them), and exactly four of them are past
    # their own Highest at it.
    for i in range(8):
        opt.add_ingredient(f"Ingredient {i + 1}", 0,
                           12 if i < 4 else 100, unit="g")
    opt.add_objective("Firmness", 1.0, goal="target", target=6,
                      min_val=0, max_val=10, unit="N")
    opt.set_pending_batch([{f"Ingredient {i + 1}": 10.0 for i in range(8)}],
                          batch_no=1)
    return opt


def test_the_scaled_caution_counts_instead_of_listing_eight_names(eight_ingredients):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.number_input(key="scale_total").set_value(400.0)
    at.run()
    assert not at.exception
    said = [c.value for c in at.caption
            if wording.AMOUNTS_YOU_ALLOWED in c.value]
    assert said == ["At 400 g, 4 of 8 ingredients go past the amounts you "
                    "allowed. Print at a smaller total, or widen them in "
                    "Set up."], said


def test_the_correction_form_says_it_once_above_the_boxes(scored):
    """The boxes open pre-filled, so "leave blank" described a state the user
    was not in — and two tooltips said one thing twice."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    at.selectbox(key="correct_formulation").select(1)
    at.run()
    assert not at.exception
    assert any(c.value == wording.CORRECTION_CAPTION for c in at.caption), \
        [c.value for c in at.caption]
    boxes = [n for n in at.number_input
             if n.key in ("correct_1_Firmness", "correct_amount_1_Pea protein")]
    assert len(boxes) == 2, [n.key for n in at.number_input]
    assert not any(n.help for n in boxes), [n.help for n in boxes]


def test_the_goal_column_offers_its_three_options_and_no_more(burger):
    """Higher is better, Lower is better, Hit a target — said once, as the
    cell's own choices. The tooltip that used to repeat them is gone with
    the box it hung on."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert list(wording.GOAL_LABELS.values()) == [
        "Higher is better", "Lower is better", "Hit a target"]
    assert set(_grid_frame(at, 1)["Goal"]) <= set(wording.GOAL_LABELS.values())


def test_all_formulations_shows_every_number_to_two_decimals(burger):
    """A panel score typed as 5 read 5.000000 beside an amount written 11.88:
    pandas' default for a column nobody formatted. Every number in the table
    reads to two decimals; a blank measurement stays blank."""
    burger.tell({"Pea protein": 11.875, "Methylcellulose": 1.0},
                {"Juiciness": 5.0, "Firmness": 6.25}, formulation_no=1,
                batch_no=1)
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0}, formulation_no=2, batch_no=1)
    import ui_results
    frame = burger.history_frame(include_amounts=True)
    rendered = frame.style.format(ui_results._amount_format(burger, frame))
    cells = rendered._translate(False, False)["body"]
    headers = list(frame.columns)

    def cell(row, column):
        return cells[row][headers.index(column) + 1]["display_value"]

    by_no = {int(cells[i][headers.index(wording.FORMULATION_CAP) + 1]
                 ["display_value"]): i for i in range(len(cells))}
    assert cell(by_no[1], "Juiciness (/10)") == "5.00"
    assert cell(by_no[1], "Firmness (N)") == "6.25"
    assert cell(by_no[1], "Pea protein (g)") == "11.88"
    # Formulation 2 has no Firmness reading: blank, never "nan".
    assert cell(by_no[2], "Firmness (N)") == ""
    # The numbers that name things are not decimals.
    assert cell(by_no[1], wording.FORMULATION_CAP) == "1"
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_RESULTS
    at.run()
    assert not at.exception


# ------------------------------------------------------------------ #
#  Default batch size (0.4.0 §C)
# ------------------------------------------------------------------ #

def _total_box(at):
    return next((n for n in at.number_input if n.key == "formulation_total"),
                None)


def test_the_total_box_sits_under_the_ingredients_and_says_what_it_does(burger):
    """One number says how big a formulation is. It lives with the
    ingredients, not in Limits: the Limits list only shows what it wrote."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    box = _total_box(at)
    assert box is not None
    assert box.label == "Default batch size (g)"
    assert box.help == ("Every suggested formulation adds up to this. Set it "
                        "to what your mixer or your panel needs.")
    assert box.placeholder == "e.g. 100"
    assert box.value is None


def test_typing_a_total_holds_every_suggestion_to_it(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _total_box(at).set_value(20.0)       # 0-25 g plus 0-3 g reaches 20 g
    at.run()
    assert not at.exception
    saved = FoodOptimizer("burger")
    assert saved.formulation_total == 20.0
    assert saved.quantity_constraints[0]['source'] == 'formulation_total'
    # It reads back into the box on the next open, rather than blanking the
    # saved total on the first render.
    again = AppTest.from_file(APP_PATH, default_timeout=180)
    again.run()
    assert _total_box(again).value == 20.0


def test_a_total_the_amounts_cannot_reach_is_refused_on_screen(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _total_box(at).set_value(90.0)       # the two ingredients reach 28 g
    at.run()
    assert not at.exception
    assert [e.value for e in at.error] == [
        "A default batch size of 90 g is not reachable: the most these "
        "ingredients can make is 28 g."]
    assert FoodOptimizer("burger").formulation_total is None


def test_the_total_reads_once_in_the_limits_list_and_is_taken_off_above(burger):
    """It is over every ingredient by definition, so it is named for the box
    that wrote it rather than listed as a limit on a chosen few — and it is
    a READING of that box, not a second control for the same rule."""
    burger.set_formulation_total(20)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    lines = [t.value for t in at.text]
    assert "Default batch size · 20 g" in lines, lines
    assert not any(line.startswith("All ingredients") for line in lines), lines
    assert not any("Pea protein + Methylcellulose" in line for line in lines), \
        lines
    assert "Delete limit on the total of each formulation" not in _labels(at)
    assert any(c.value == "Empty the Default batch size box to take "
                          "it off." for c in at.tabs[0].caption), \
        [c.value for c in at.tabs[0].caption]
    # ... and emptying the box is what takes it off.
    _total_box(at).set_value(0.0)
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").formulation_total is None
    assert FoodOptimizer("burger").quantity_constraints == []


def test_the_total_box_is_not_offered_while_the_units_differ(mixed_units):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _total_box(at) is None
    # The box lives in More settings now, and so does the line that stands
    # in its place.
    assert wording.NEEDS_ONE_UNIT in [c.value for c in _more_settings(at).caption]


def test_a_project_that_weighs_nothing_out_is_offered_no_total(ferment):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _total_box(at) is None
    assert wording.NEEDS_ONE_UNIT not in _tab1_captions(at)


def test_a_unit_change_that_splits_the_ingredients_says_the_total_went(burger):
    burger.set_formulation_total(20)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, edited={1: {wording.UNIT_LABEL: "ml"}})
    assert not at.exception
    assert any(w.value == ("The default batch size of 20 g is gone: your ingredients no "
                           "longer share one unit.")
               for w in at.warning), [w.value for w in at.warning]
    assert FoodOptimizer("burger").formulation_total is None
    assert _total_box(at) is None


def test_the_round_tab_opens_its_batch_size_box_at_the_projects_default(
        open_batch):
    """0.5.0: the round screen always asks the size, and a project with a
    default opens the box AT that default rather than showing nothing.

    Until 0.5.0 the box was hidden whenever Set up carried a total, so the
    size of what the bench was about to weigh out was two tabs away and
    making one round bigger meant editing the project."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(n.key == "scale_total" for n in at.number_input)
    # Setting the default discards the open round — its rows were built to
    # the old answer — so a new one is opened under it.
    open_batch.set_formulation_total(20)
    open_batch.set_pending_batch([{"Pea protein": 17.0,
                                   "Methylcellulose": 3.0}])
    again = AppTest.from_file(APP_PATH, default_timeout=180)
    again.run()
    box = next(n for n in again.number_input if n.key == "scale_total")
    assert box.label == "Batch size (g)"
    assert box.value == 20.0
    assert not any(c.value == wording.NEEDS_ONE_UNIT
                   for c in again.tabs[1].caption)
    assert any("20 g" in c.value for c in again.tabs[1].caption), \
        [c.value for c in again.tabs[1].caption]


def test_the_results_amounts_heading_names_the_project_total(burger):
    """The heading names the total the BATCH was made to. A batch recorded
    under the project's total carries it; one recorded before it was typed
    was made as generated and keeps saying so."""
    scored = burger
    scored.set_formulation_total(20)
    scored.set_pending_batch([{"Pea protein": 10.0, "Methylcellulose": 1.0},
                              {"Pea protein": 17.0, "Methylcellulose": 3.0}])
    scored.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 1.0},
                formulation_no=1, batch_no=scored.pending_batch_no,
                note="crumbly")
    scored.tell({"Pea protein": 17.0, "Methylcellulose": 3.0},
                {"Juiciness": 7.0, "Firmness": 8.0},
                formulation_no=2, batch_no=scored.pending_batch_no)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    headings = [m.value for m in at.tabs[2].markdown]
    assert "**Amounts to make it (20 g)**" in headings, headings


def test_the_sample_ships_with_a_hundred_gram_total(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at.main, "Try the sample project").click()
    at.run()
    assert not at.exception
    assert FoodOptimizer(wording.SAMPLE_PROJECT_NAME).formulation_total == 100.0
    assert _total_box(at).value == 100.0
    assert "Default batch size · 100 g" in [t.value for t in at.text]


def test_the_sample_is_written_with_shares_that_add_up_to_a_hundred(
        tmp_path, monkeypatch):
    """Every setter saves, and add_objective deliberately does not
    normalise, so the sample landed on disk holding weights of 1.0 and 1.5
    under CLASS_VERSION 11 — a 0.5.0 file that read as a 0.4.x one, and a
    saved copy of it downloaded that way."""
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at.main, "Try the sample project").click()
    at.run()
    assert not at.exception
    # Read the FILE, not the loaded object: import_json normalises on the
    # way in, so the loaded project has always looked right.
    written = json.loads(
        (tmp_path / f"{wording.SAMPLE_PROJECT_NAME}.pkl").read_text())
    weights = {o['name']: o['weight'] for o in written['objectives']}
    assert weights == {"Juiciness": 40.0, "Firmness": 60.0}, weights
    assert written['CLASS_VERSION'] == FoodOptimizer.CLASS_VERSION


def _warm_burger(opt):
    """Five results on the burger project, Formulation 1 the best of them, so
    the cold start is over and there is something to compare with."""
    for i in range(5):
        opt.tell({"Pea protein": 10.0 + i, "Methylcellulose": 1.0},
                 {"Juiciness": 7.0, "Firmness": 6.0 if i == 0 else 1.0},
                 formulation_no=i + 1, batch_no=1)
    assert opt.best_formulation_no() == 1
    return opt


def test_each_formulation_says_what_it_is_trying(burger):
    """Tab 2's last column, headed with the formulation the changes are
    measured from, and the same line on the sheet directly under the
    formulation's own title — the technician holding the paper is the one who
    asks why this bowl differs from the last."""
    _warm_burger(burger)
    burger.set_pending_batch([{"Pea protein": 11.0, "Methylcellulose": 1.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    assert not at.exception
    table = next(d.value for d in at.dataframe if "Formulation" in d.value.columns)
    assert list(table.columns)[-1] == "Compared with Formulation 1"
    line = "Close to the best · Pea protein +1.00 g"
    assert table["Compared with Formulation 1"].iloc[0] == line
    # On paper there is no column header to read the cell against, so the
    # line carries it.
    sheet = _formulation_sheets(at)[0]
    assert sheet.cell(row=1, column=1).value == \
        "Formulation 6 · Round 2 · burger", sheet.cell(row=1, column=1).value
    assert sheet.cell(row=2, column=1).value == \
        f"Compared with Formulation 1: {line}", sheet.cell(row=2, column=1).value


def test_an_own_formulation_gets_the_same_line(burger):
    """A formulation of the user's own is a row like any other: it says what
    it changes from the best, exactly as a suggested one does."""
    _warm_burger(burger)
    burger.set_pending_batch([{"Pea protein": 11.0, "Methylcellulose": 1.0}])
    burger.add_to_pending_batch({"Pea protein": 20.0, "Methylcellulose": 1.0},
                                wording.OWN_FORMULATION_NOTE)
    table = burger.batch_frame(burger.pending_batch)
    # The changes are a fact about the amounts and are listed for both. The
    # KIND is not: a formulation the user typed is not one of the app's
    # three kinds of suggestion, and calling it one took credit for their
    # own bench standard.
    assert list(table["Compared with Formulation 1"]) == [
        "Close to the best · Pea protein +1.00 g",
        "Your own formulation · Pea protein +10.00 g"]


def test_the_what_it_is_trying_column_is_never_formatted_as_a_number(burger):
    """It is words, not an amount: handing it to a two-decimal format would
    raise the moment the table was drawn."""
    from ui_batch import _amount_format
    _warm_burger(burger)
    burger.set_pending_batch([{"Pea protein": 11.0, "Methylcellulose": 1.0}])
    frame = burger.batch_frame(burger.pending_batch)
    assert "Compared with Formulation 1" not in _amount_format(burger, frame)


# ------------------------------------------------------------------ #
#  Default batch size: fix round 1
# ------------------------------------------------------------------ #

def test_deleting_an_ingredient_says_what_happened_to_the_total(burger):
    """The notice was the only thing that told the user the total went, and
    only an added ingredient produced it — a deleted one left the box holding
    a number the project no longer had, writing it back every run."""
    burger.add_ingredient("Water", 20, 60)
    burger.set_formulation_total(80)          # 25 + 3 + 60 reaches 80 g
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _total_box(at).value == 80.0
    _save_grid(at, ING_GRID, confirm=True, deleted=[2])
    assert not at.exception
    assert any(w.value == ("The default batch size of 80 g is gone: the allowed amounts "
                           "no longer add up to it.")
               for w in at.warning), [w.value for w in at.warning]
    assert FoodOptimizer("burger").formulation_total is None
    assert _total_box(at).value is None
    # ...and it stays gone: nothing writes the old number back on the next run.
    at.run()
    assert not at.error, [e.value for e in at.error]
    assert FoodOptimizer("burger").formulation_total is None


def test_a_recorded_batch_keeps_the_weight_it_was_made_to(scored):
    """Typing a total on tab 1 must not rewrite what the bench already
    weighed out for a batch in the records."""
    scored.set_pending_batch_total(400.0)
    scored._batch_totals()[1] = 400.0
    scored.save()
    scored.set_formulation_total(20)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    headings = [m.value for m in at.tabs[2].markdown]
    assert "**Amounts to make it (400 g)**" in headings, headings


def test_the_limit_picker_asks_for_a_choice_like_every_other_picker(burger):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert at.multiselect(key="qty_pick").proto.placeholder == \
        wording.CHOOSE_MANY_PLACEHOLDER


def test_a_total_at_the_edge_of_what_the_amounts_reach_still_generates(burger):
    """The box accepted 28 g and Generate then failed with a sentence about
    limits: rejection sampling cannot find an equality on a sum."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _total_box(at).set_value(28.0)            # 25 g + 3 g, the very ceiling
    at.run()
    assert not at.error, [e.value for e in at.error]
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    _submit_button(at, "Generate 3 formulations").click()
    at.run()
    assert not at.exception
    assert not at.error, [e.value for e in at.error]
    rows = FoodOptimizer("burger").pending_batch
    assert len(rows) == 3
    for row in rows:
        assert sum(row['recipe'].values()) == pytest.approx(28.0, abs=0.5)


# ------------------------------------------------------------------ #
#  0.4.0 §G: the workbook, from the tab it is downloaded on to the
#  results it brings back.
# ------------------------------------------------------------------ #
def test_the_filled_in_workbook_records_the_results_and_the_ticked_row(
        open_batch):
    """The whole round trip a bench does: download the sheets, write the
    readings in the Measured cells of the summary sheet, tick Not scored on
    the one that never got measured, and upload the file."""
    open_batch.add_to_pending_batch({"Pea protein": 15.0,
                                     "Methylcellulose": 1.5})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    opt = at.session_state["optimizer"]
    book = openpyxl.load_workbook(
        io.BytesIO(opt.workbook_bytes(opt.pending_batch, None)))
    sheet = book[wording.batch_sheet_name(1)]
    labels = [sheet.cell(row=r, column=1).value
              for r in range(1, sheet.max_row + 1)]
    firm = labels.index("Firmness · target 6 N") + 1
    juice = labels.index("Juiciness (/10) · target 7") + 1
    sheet.cell(row=firm, column=2, value=6.0)
    sheet.cell(row=juice, column=2, value=7.0)
    sheet.cell(row=firm, column=4, value=5.0)
    sheet.cell(row=juice, column=4, value=6.5)
    sheet.cell(row=labels.index(wording.NOT_SCORED_CHECKBOX_SHEET) + 1, column=6, value="x")
    sheet.cell(row=labels.index(wording.NOTE) + 1, column=6,
               value="burner failed")
    filled = io.BytesIO()
    book.save(filled)
    filled.seek(0)
    filled.name = "burger · Round 1.xlsx"

    import ui_batch
    at.session_state["_results_upload"] = ui_batch._read_results_file(opt,
                                                                     filled)
    at.run()
    assert not at.exception
    assert not at.warning, [w.value for w in at.warning]
    assert any("Found results for 2 of 3 formulations" in i.value
               and wording.not_scored_counter_suffix(1) in i.value
               for i in at.info), [i.value for i in at.info]
    _submit_button(at, wording.SAVE_UPLOADED_RESULTS).click()
    at.run()
    assert not at.exception
    assert not at.error, [e.value for e in at.error]
    reloaded = FoodOptimizer("burger")
    assert reloaded.results_history == [
        {"Firmness": 6.0, "Juiciness": 7.0},
        {"Firmness": 5.0, "Juiciness": 6.5}], reloaded.results_history
    # The ticked one is in the record with its reason, and with no result.
    assert [(s["formulation"], s["note"]) for s in reloaded.skipped] == \
        [(3, "Not scored · burner failed")], reloaded.skipped
    assert not reloaded.pending_batch, reloaded.pending_batch


def test_the_upload_takes_a_workbook_or_a_comma_separated_file(open_batch):
    """Two shapes, one door. The workbook is what the app hands out; a file
    saved out of a spreadsheet is what a bench that keeps its own sheet
    hands back."""
    import ui_batch
    opt = open_batch
    book = io.BytesIO(opt.workbook_bytes(opt.pending_batch, None))
    book.name = "burger · Round 1.xlsx"
    filled = openpyxl.load_workbook(book)
    sheet = filled[wording.batch_sheet_name(1)]
    labels = [sheet.cell(row=r, column=1).value
              for r in range(1, sheet.max_row + 1)]
    sheet.cell(row=labels.index("Firmness · target 6 N") + 1, column=2,
               value=6.0)
    out = io.BytesIO()
    filled.save(out)
    out.seek(0)
    out.name = "burger · Round 1.xlsx"
    # Both shapes come back as one record: the rows, and beside them what
    # the file said about lots and about what was weighed — nothing, in a
    # file of columns.
    from_workbook = ui_batch._read_results_file(opt, out)
    assert list(from_workbook.frame["Formulation"]) == [1]
    assert from_workbook.frame["Firmness"].iloc[0] == 6.0

    plain = io.BytesIO(b"Formulation,Firmness,Juiciness\n1,6.0,7.0\n")
    plain.name = "results.csv"
    from_csv = ui_batch._read_results_file(opt, plain)
    assert list(from_csv.frame["Formulation"]) == [1]
    assert from_csv.actual == {} and from_csv.lots == {}
    assert opt.parse_batch_results(from_csv.frame, opt.pending_batch) == \
        [(1, {"Firmness": 6.0, "Juiciness": 7.0}, "")]


def test_every_tab_still_has_one_lit_button_with_the_workbook_on_it(
        open_batch):
    """The download is the lit thing on tab 2 until a result is typed, and
    no tab has two."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert not at.warning, [w.value for w in at.warning]
    assert _tab_primaries(at, 1) == [wording.DOWNLOAD_BATCH_SHEETS], \
        _tab_primaries(at, 1)
    for index in (0, 2):
        assert len(_tab_primaries(at, index)) <= 1, _tab_primaries(at, index)


# ------------------------------------------------------------------ #
#  The final fix wave: the total survives everything that is not a
#  change to it, and the screen says what the batch is held to.
# ------------------------------------------------------------------ #

def _sidebar_button(at, label):
    return next(b for b in at.sidebar.button if b.label == label)


@pytest.mark.parametrize("arm,cancel", [
    ("Delete this project", "Cancel"),
    ("Start this project over", "Cancel"),
])
def test_a_sidebar_cancel_never_writes_the_total_away(burger, arm, cancel):
    """Back out, change nothing. Cancel reruns from ABOVE the tabs, so
    Streamlit discards the total box's value — and an emptied box used to be
    read as the user clearing the total every suggestion is held to."""
    burger.set_formulation_total(20.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _total_box(at).value == 20.0
    _sidebar_button(at, arm).click()
    at.run()
    _sidebar_button(at, cancel).click()
    at.run()
    assert not at.exception
    saved = FoodOptimizer("burger")
    assert saved.formulation_total == 20.0
    assert len(saved.quantity_constraints) == 1
    assert _total_box(at).value == 20.0


def test_adding_an_ingredient_keeps_the_total_and_says_so(burger):
    """Adding ingredients is step 2 of the app's own first-run instructions,
    and it used to take the 100 g rule with it, silently."""
    burger.set_formulation_total(20.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, added=[_ingredient_row("Onion powder", high=2.0)])
    assert not at.exception
    saved = FoodOptimizer("burger")
    assert saved.formulation_total == 20.0
    assert saved._formulation_total_index() is not None
    # The limit is over every ingredient, so the new one is in it.
    assert "Onion powder" in saved.quantity_constraints[0]['ingredients']
    assert any(s.value == "Onion powder added. Each formulation still "
                          "totals 20 g." for s in at.success), \
        [s.value for s in at.success]


def test_a_round_with_no_batch_size_of_its_own_says_so_under_the_table(
        open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert any(c.value == "This round has no batch size of its own."
               for c in at.tabs[1].caption), \
        [c.value for c in at.tabs[1].caption]


def test_a_row_that_misses_the_total_says_so_under_the_table(burger):
    """Nothing is rewritten to make a formulation of the user's own add up,
    so this line is the only thing that says it does not."""
    burger.set_formulation_total(20.0)
    burger.set_pending_batch([{"Pea protein": 17.0, "Methylcellulose": 3.0}])
    burger.add_to_pending_batch({"Pea protein": 12.0, "Methylcellulose": 1.0},
                                wording.OWN_FORMULATION_NOTE)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    captions = [c.value for c in at.tabs[1].caption]
    assert ("Formulation 2 adds up to 13.00 g, not the 20 g batch size."
            in captions), captions
    # ... and the row is shown at what it really adds up to.
    frame = next(d.value for d in at.dataframe
                 if "Formulation" in d.value.columns)
    assert list(frame["Total (g)"]) == [20.0, 13.0]


# ------------------------------------------------------------------ #
#  The final fix wave: the words, and the controls that carry them
# ------------------------------------------------------------------ #

def test_the_copy_sentence_points_somewhere_the_reader_can_reach(scored):
    """"Open a saved copy can bring it back" was ungrammatical and named no
    control the reader could find."""
    assert wording.COPY_KEPT == ("A copy is saved first. To bring it back, "
                                 "use Open a saved copy in the sidebar.")


def test_a_copy_the_app_saved_itself_restores_through_the_sidebar(burger,
                                                                  tmp_path):
    """Every destructive confirmation promises a copy. Opening one has to put
    the formulation back — the whole point of the sentence."""
    burger.tell({"Pea protein": 10.0, "Methylcellulose": 1.0},
                {"Juiciness": 7.0, "Firmness": 6.0}, formulation_no=1,
                batch_no=1)
    storage_backend.LocalStorage().archive("burger", "pre_delete", copy=True)
    burger.delete_formulation(1)
    assert FoodOptimizer("burger").formulation_ids == []
    # What the uploader hands on: the archive's own bytes, read as a copy.
    kept = json.loads((tmp_path / "burger_pre_delete.pkl").read_text())
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["_restore_candidate"] = kept
    at.run()
    assert not at.exception
    _submit_button(at.sidebar, wording.YES_REPLACE).click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").formulation_ids == [1]


def test_the_batch_line_says_to_record_on_every_tab(open_batch):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any(c.value == "Round 1 · 2 to record"
               for c in at.tabs[0].caption), [c.value for c in at.tabs[0].caption]


def test_the_correction_boxes_show_the_amounts_the_sheet_printed(scored):
    """A box labelled (g) holding 22.241068936455125 claims the balance read
    that. The sheet said 22.24."""
    scored.edit_amounts(0, {"Pea protein": 22.241068936455125,
                            "Methylcellulose": 1.0})
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    box = at.number_input(key="correct_amount_1_Pea protein")
    assert box.value == 22.24
    assert box.proto.format == "%.2f"


def test_a_scored_formulation_can_have_its_note_corrected(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    at.selectbox(key="correct_formulation").set_value(1)
    at.run()
    note = at.text_input(key="correct_note_1")
    assert note.value == "crumbly"          # what the row was recorded with
    note.set_value("crumbly, second bake")
    at.run()
    _submit_button(at, wording.SAVE_CORRECTION_BUTTON).click()
    at.run()
    assert not at.exception
    assert FoodOptimizer("burger").notes_history[0] == "crumbly, second bake"


def test_the_scoring_row_is_headed_score_and_saves_a_result(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert wording.CORRECT_A_FORMULATION_HEADING in [m.value for m in at.markdown]
    at.selectbox(key="correct_formulation").set_value(3)
    at.run()
    headings = [m.value for m in at.markdown]
    assert wording.SCORE_A_FORMULATION_HEADING in headings, headings
    assert wording.CORRECT_A_FORMULATION_HEADING not in headings
    assert wording.SAVE_RESULT_BUTTON in _labels(at)
    # The caption stays: it is what says a not-scored row can be scored here.
    assert wording.NOT_SCORED_CAN_BE_SCORED_CAPTION in \
        [c.value for c in at.caption]


def test_the_best_so_far_amounts_read_in_set_up_order(scored):
    """The same recipe in two orders on two screens made the reader check it
    line by line."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    table = _amount_tables(at)[0]
    assert list(table[wording.INGREDIENT_OR_SETTING_LABEL]) == [
        "Pea protein", "Methylcellulose"]


def test_no_measurement_cell_reads_nan(scored):
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    table = next(d for d in at.dataframe
                 if "Overall score" in d.value.columns)
    shown = _displayed(table).to_string()
    assert "nan" not in shown.lower(), shown


def test_the_check_step_shows_the_numbers_it_read(open_batch):
    """It counted the formulations it found and showed none of them: a
    firmness of 74 written for 7.4 passed the check."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.session_state["_results_upload"] = pd.DataFrame(
        [{"Formulation": 1, "Firmness": 6.0, "Juiciness": 7.4,
          wording.NOT_SCORED: "", wording.NOTE: "held together"}])
    at.run()
    assert not at.exception
    preview = next(d.value for d in at.dataframe
                   if wording.NOT_SCORED in d.value.columns
                   and wording.NOTE in d.value.columns)
    assert list(preview[wording.FORMULATION_CAP]) == [1]
    assert list(preview["Juiciness (/10)"]) == [7.4]
    assert list(preview[wording.NOTE]) == ["held together"]
    assert any(c.value == wording.UPLOAD_PREVIEW_CAPTION for c in at.caption)


def test_a_refused_row_keeps_what_was_typed_into_it(scored):
    """The refusal used to name a field that was not on the form when it was
    filled in, and what had been typed was gone by the time it arrived. The
    grid answers under the row, and the row itself is untouched: nothing is
    written while there is an error."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    typed = _ingredient_row("Cook time", kind=wording.KIND_SETTING,
                            low=0.0, high=30.0, unit="min")
    _save_grid(at, ING_GRID, added=[typed])
    # Refused for the missing baseline, by row...
    assert [e.value for e in at.error] == [
        wording.row_error(3, wording.ADD_BASELINE_ERROR)]
    # ...and the row is still on the grid, still holding what was typed.
    _grid_edits(at, ING_GRID, added=[typed])
    at.run()
    assert wording.SAVE_CHANGES_BUTTON in _labels(at), _labels(at)
    assert not any(v["name"] == "Cook time"
                   for v in FoodOptimizer("scored").variables)


def test_the_template_is_headers_and_one_example_row(burger):
    """A file arriving with a full ingredient list in it is an export."""
    frame = pd.read_excel(io.BytesIO(
        ingredients_template_workbook(
            os.path.join(os.path.dirname(APP_PATH), "data",
                         "sample_ingredients.csv"))))
    assert list(frame.columns)[:4] == ["Name", "Lowest", "Highest", "Unit"]
    assert len(frame) == 1


def test_the_first_screen_says_where_the_name_box_is(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert any("Name a project** in the sidebar on the left" in m.value
               for m in at.markdown), [m.value for m in at.markdown]


def test_only_an_ingredient_add_says_the_total_still_holds(burger):
    """The total is a sum of AMOUNTS. A process setting is not in it, so
    adding one has nothing to reassure anybody about."""
    burger.set_formulation_total(20.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _save_grid(at, ING_GRID, added=[
        _ingredient_row("Cook temperature", kind=wording.KIND_SETTING,
                        low=0.0, high=220.0, unit="°C")])
    assert not at.exception
    assert [s.value for s in at.success] == ["Cook temperature added."]
    assert FoodOptimizer("burger").formulation_total == 20.0


def test_a_batch_discarded_by_the_total_says_the_total_did_it(burger):
    """'Your set-up changed' is true of the ingredient list, a hold and the
    allowed amounts. When it was the total, the notice names the total."""
    burger.set_formulation_total(20.0)
    burger.set_pending_batch([{"Pea protein": 17.0, "Methylcellulose": 3.0}])
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _total_box(at).set_value(25.0)
    at.run()
    assert not at.exception
    assert any(i.value == ("Round 1 was discarded: the default batch size "
                           "changed after it was made. Generate "
                           "a new one.") for i in at.info), \
        [i.value for i in at.info]
    assert FoodOptimizer("burger").pending_batch is None


def test_the_apps_own_export_imports_back_as_past_formulations(burger):
    """Download all formulations (Excel) opens with a one-cell title row; the
    importer must find the real header under it, or the app's own file is
    refused with a list of missing columns."""
    import io
    burger.tell({"Pea protein": 11.875, "Methylcellulose": 1.0},
                {"Juiciness": 5.0, "Firmness": 6.25}, formulation_no=1,
                batch_no=1)
    import ui_results
    data = io.BytesIO(burger.all_formulations_workbook())
    data.name = "All formulations.xlsx"
    frame = ui_results._read_past_formulations(data)
    for column in ("Pea protein (g)", "Methylcellulose (g)", "Juiciness (/10)",
                   "Firmness (N)"):
        assert any(str(c).startswith(column.split(" (")[0]) for c in frame.columns), (
            column, list(frame.columns))
    assert len(frame) == 1
    col_for, missing = ui_results._import_columns(burger, frame)
    assert missing == [], missing
    assert set(col_for) == {"Pea protein", "Methylcellulose", "Juiciness",
                            "Firmness"}


def test_the_app_says_it_is_starting_before_its_heavy_imports():
    """The desktop wrapper shows the web view on Streamlit's first healthy
    answer, which arrives before app.py has finished importing food_bo (torch
    and friends, seconds even warm). Anything drawn after those imports is a
    blank window until they finish, so the order in the source is the fix and
    is what this test guards: page config, then the placeholder, then the
    heavy imports, then the placeholder is cleared."""
    import ast
    root = pathlib.Path(__file__).resolve().parent.parent
    tree = ast.parse((root / "app.py").read_text())

    def line_of_call(attr, owner=None, arg_is=None):
        """The line of the FIRST matching call. ast.walk visits by node, not
        by line, so the candidates are sorted rather than taken as they come
        — otherwise this guard's own answer depends on the shape of the tree.
        """
        found = []
        for node in ast.walk(tree):
            if not (isinstance(node, ast.Call)
                    and isinstance(node.func, ast.Attribute)
                    and node.func.attr == attr):
                continue
            if owner and getattr(node.func.value, "id", None) != owner:
                continue
            if arg_is:
                if not (node.args and isinstance(node.args[0], ast.Attribute)
                        and node.args[0].attr == arg_is):
                    continue
            found.append(node.lineno)
        return min(found) if found else None

    config_at = line_of_call("set_page_config", owner="st")
    placeholder_at = line_of_call("info", owner="_starting", arg_is="STARTING_APP")
    cleared_at = line_of_call("empty", owner="_starting")
    heavy = []
    for node in ast.walk(tree):   # sorted below, for the same reason
        if isinstance(node, ast.Import):
            names = [a.name for a in node.names]
        elif isinstance(node, ast.ImportFrom):
            names = [node.module or ""]
        else:
            continue
        if any(n == "food_bo" or n.startswith("ui_") for n in names):
            heavy.append(node.lineno)

    heavy.sort()
    assert config_at and placeholder_at and cleared_at and heavy
    assert config_at < placeholder_at < heavy[0], (
        config_at, placeholder_at, heavy[0])
    assert cleared_at > heavy[-1], (cleared_at, heavy[-1])
    # And the sentence itself is the one the window shows while it waits.
    assert wording.STARTING_APP == ("Starting Food Optimizer… loading its "
                                    "components. This takes a few seconds.")


def test_the_starting_line_is_gone_once_the_app_has_drawn(tmp_path, monkeypatch):
    """The placeholder is emptied after the imports, so a finished run never
    shows it — it exists only for the seconds before the first element."""
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    assert not any(wording.STARTING_APP in i.value for i in at.info), \
        [i.value for i in at.info]


def test_an_untouched_older_sample_is_rebuilt_as_the_current_one(
        tmp_path, monkeypatch):
    """A Mac that met an earlier version has that version's sample on disk:
    no formulation total, no targets source. Reopening it unchanged is how a
    sample ends up making batches at different weights, so a sample nobody
    has used yet is rebuilt as the current one."""
    monkeypatch.chdir(tmp_path)
    old = FoodOptimizer(wording.SAMPLE_PROJECT_NAME)
    # Not one of the sample's own eight, so its absence afterwards is proof
    # the project was rebuilt rather than added to.
    old.add_ingredient("Oat flour", 0, 100)
    old.add_objective("Taste", 1.0, goal="max")
    assert old.formulation_total is None and old.targets_source == ""

    at = AppTest.from_file(APP_PATH, default_timeout=300)
    at.run()
    _submit_button(at.sidebar, wording.TRY_SAMPLE_LABEL).click()
    at.run()
    assert not at.exception

    rebuilt = FoodOptimizer(wording.SAMPLE_PROJECT_NAME)
    assert rebuilt.formulation_total == 100
    assert rebuilt.targets_source == wording.SAMPLE_TARGETS_SOURCE
    assert [o["name"] for o in rebuilt.objectives] == ["Juiciness", "Firmness"]
    # The old project's own set-up is gone, not added to.
    assert "Oat flour" not in [v["name"] for v in rebuilt.variables]
    assert len(rebuilt.variables) == 8
    # It opens as the sample it now is, welcome line and all, and it is an
    # opening: the project was already there.
    assert any(s.value == wording.project_opened(wording.SAMPLE_PROJECT_NAME)
               for s in at.success), [s.value for s in at.success]
    assert wording.SAMPLE_TAB1_DESCRIPTION in [c.value for c in at.tabs[0].caption]


def test_a_sample_with_a_result_is_opened_exactly_as_it_stands(
        tmp_path, monkeypatch):
    """The moment the sample holds a result it is the user's own project.
    Rebuilding it would throw away their work."""
    monkeypatch.chdir(tmp_path)
    old = FoodOptimizer(wording.SAMPLE_PROJECT_NAME)
    old.add_ingredient("Water", 0, 100)
    old.add_objective("Taste", 1.0, goal="max")
    old.tell({"Water": 50.0}, {"Taste": 7.0})

    at = AppTest.from_file(APP_PATH, default_timeout=300)
    at.run()
    _submit_button(at.sidebar, wording.TRY_SAMPLE_LABEL).click()
    at.run()
    assert not at.exception

    kept = FoodOptimizer(wording.SAMPLE_PROJECT_NAME)
    assert [v["name"] for v in kept.variables] == ["Water"]
    assert [o["name"] for o in kept.objectives] == ["Taste"]
    assert len(kept.X_history) == 1
    assert kept.formulation_total is None


def test_a_sample_whose_only_batch_was_never_scored_is_left_alone(
        tmp_path, monkeypatch):
    """A row generated and then ticked Not scored leaves X_history empty but
    `skipped` filled. That is a sample someone has used: leave it alone."""
    monkeypatch.chdir(tmp_path)
    old = FoodOptimizer(wording.SAMPLE_PROJECT_NAME)
    old.add_ingredient("Water", 0, 100)
    old.add_objective("Taste", 1.0, goal="max")
    old.record_skipped(1, 1, {"Water": 50.0})
    assert old.X_history == [] and old.skipped

    at = AppTest.from_file(APP_PATH, default_timeout=300)
    at.run()
    _submit_button(at.sidebar, wording.TRY_SAMPLE_LABEL).click()
    at.run()
    assert not at.exception

    kept = FoodOptimizer(wording.SAMPLE_PROJECT_NAME)
    assert [v["name"] for v in kept.variables] == ["Water"]
    assert kept.skipped


def test_a_sample_with_an_open_batch_is_left_exactly_as_it_is(
        tmp_path, monkeypatch):
    """Nothing is scored and nothing was ticked Not scored, but a batch is on
    someone's bench and the formulations in it have been numbered. That is
    work, and rebuilding the sample would take it away."""
    monkeypatch.chdir(tmp_path)
    old = FoodOptimizer(wording.SAMPLE_PROJECT_NAME)
    old.add_ingredient("Oat flour", 0, 100)
    old.add_objective("Taste", 1.0, goal="max")
    old.set_pending_batch([{"Oat flour": 50.0}])
    assert old.X_history == [] and old.skipped == []
    assert old.pending_batch and old.next_formulation_no > 1

    at = AppTest.from_file(APP_PATH, default_timeout=300)
    at.run()
    _submit_button(at.sidebar, wording.TRY_SAMPLE_LABEL).click()
    at.run()
    assert not at.exception

    kept = FoodOptimizer(wording.SAMPLE_PROJECT_NAME)
    assert [v["name"] for v in kept.variables] == ["Oat flour"]
    assert kept.pending_batch == old.pending_batch
    assert kept.formulation_total is None


def test_a_sample_whose_batch_was_made_and_deleted_is_left_alone(
        tmp_path, monkeypatch):
    """A number already issued is the trace a deleted batch leaves: history
    the user made, even though nothing is on file under it."""
    monkeypatch.chdir(tmp_path)
    old = FoodOptimizer(wording.SAMPLE_PROJECT_NAME)
    old.add_ingredient("Oat flour", 0, 100)
    old.add_objective("Taste", 1.0, goal="max")
    old.set_pending_batch([{"Oat flour": 50.0}])
    old.set_pending_batch(None)          # generated, then thrown away
    assert old.pending_batch is None and old.next_formulation_no > 1

    at = AppTest.from_file(APP_PATH, default_timeout=300)
    at.run()
    _submit_button(at.sidebar, wording.TRY_SAMPLE_LABEL).click()
    at.run()
    assert not at.exception
    assert [v["name"] for v in FoodOptimizer(
        wording.SAMPLE_PROJECT_NAME).variables] == ["Oat flour"]


# ------------------------------------------------------------------ #
#  0.5.0 · three tiers on every tab (spec 1.5)
#
#  The order each tab reads in, walked off the element tree rather than
#  taken from a list of labels that happen to be on screen somewhere, and
#  the one lit button in every state the tabs can be in at once.
# ------------------------------------------------------------------ #

def test_the_round_tab_reads_in_one_order(open_batch):
    """Title, Batch size, the round table, the download, Record the results
    — then the optional doors, folded, and the button that throws the round
    away last of all."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert _tab_outline(at, 1) == [
        wording.make_these(1, 2),
        wording.STEP_MAKE_HEADING,
        wording.STEP_PRINT_HEADING,
        wording.STEP_RECORD_HEADING,
        # One heading per row of the round, inside step 3.
        wording.formulation_heading(1),
        wording.formulation_heading(2),
        wording.ADD_OWN_EXPANDER,
        wording.UPLOAD_EXPANDER,
    ], _tab_outline(at, 1)
    # The Batch size box sits between the first step heading and the table
    # it is the size of, and the download under the second.
    tab = at.tabs[1]
    assert tab.number_input[0].label == "Batch size (g)"
    assert [d.label for d in _unknowns(tab, "download_button")] == [
        wording.DOWNLOAD_BATCH_SHEETS]
    # ...and `Generate a different round` is the last thing drawn.
    assert [b.label for b in tab.button][-1] == wording.GENERATE_DIFFERENT_BATCH


def test_the_round_tab_records_then_offers_the_two_folded_doors(open_batch):
    """`Add a formulation of your own` and the results upload are optional
    and collapsed, and both sit below Save results rather than between the
    grid and the button that writes it."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    tab = at.tabs[1]
    folds = [e for e in tab.expander]
    assert [e.label for e in folds] == [wording.ADD_OWN_EXPANDER,
                                        wording.UPLOAD_EXPANDER]
    assert not any(e.proto.expanded for e in folds)
    labels = [b.label for b in tab.button]
    assert labels.index(wording.SAVE_RESULTS) < labels.index(
        wording.ADD_TO_THIS_BATCH), labels


def test_the_advanced_fold_holds_the_model_settings_and_nothing_folded(
        burger):
    """One small fold at the very bottom: the model settings and the two
    explanations, as plain markdown. Nothing inside it is a fold of its own
    — Streamlit cannot nest one expander in another — and nothing inside it
    is coloured."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    fold = _advanced(at)
    assert not fold.proto.expanded
    # Its own label and no other: nothing inside it is a fold.
    assert [e.label for e in fold.expander] == [wording.ADVANCED_EXPANDER]
    assert fold.radio[0].key == "bo_cfg_mode"
    assert not [b.label for b in fold.button if b.proto.type == "primary"]
    # It is the last thing on the tab.
    assert _tab_outline(at, 0)[-1] == wording.ADVANCED_EXPANDER


def test_one_lit_button_per_tab_with_both_grids_mid_edit(burger):
    """Both grids can be mid-edit at once, on a tab whose foot is ready to
    move on. The topmost lit Save wins; the other pair and the foot go
    grey, and the two tabs below keep their own single button."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, edited={0: {wording.HIGHEST_LABEL: 30.0}})
    _grid_edits(at, MEAS_GRID, edited={0: {wording.SHARE_COLUMN: 55.0}})
    at.run()
    assert _tab_primaries(at, 0) == [wording.SAVE_CHANGES_BUTTON], \
        _tab_primaries(at, 0)
    assert len(_tab_primaries(at, 1)) <= 1, _tab_primaries(at, 1)
    assert len(_tab_primaries(at, 2)) <= 1, _tab_primaries(at, 2)


def test_one_lit_button_per_tab_while_a_property_is_being_deleted(burger):
    """A confirmation armed inside More settings, with an edit standing in
    a grid above it: the Yes is the one coloured button, and the grid's Save
    steps aside for it."""
    burger.add_property("Cost")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _grid_edits(at, ING_GRID, edited={0: {wording.HIGHEST_LABEL: 30.0}})
    at.run()
    _submit_button(at, wording.delete_button("Cost")).click()
    _grid_edits(at, ING_GRID, edited={0: {wording.HIGHEST_LABEL: 30.0}})
    at.run()
    assert _tab_primaries(at, 0) == [wording.YES_DELETE], _tab_primaries(at, 0)


def test_one_lit_button_per_tab_while_the_targets_note_is_open(burger):
    """The note's own Save is always grey: opening a one-line box does not
    take the tab's coloured button away from the foot."""
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at, wording.ADD_TARGETS_SOURCE_BUTTON).click()
    at.run()
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON], \
        _tab_primaries(at, 0)
    # ...and with a grid mid-edit under it, the Save above wins and the
    # note's Save is still grey.
    _grid_edits(at, ING_GRID, edited={0: {wording.HIGHEST_LABEL: 30.0}})
    at.run()
    assert _tab_primaries(at, 0) == [wording.SAVE_CHANGES_BUTTON], \
        _tab_primaries(at, 0)


def test_the_tab_says_nothing_twice_with_every_tier_open(burger):
    """Sparse: one line each, folds included. The two tiers hold most of
    tab 1's captions now, so the check that nothing is said twice has to
    reach inside them."""
    burger.add_property("Cost")
    burger.add_constraint("Cost", max_val=2.0)
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    assert not at.exception
    captions = [c.value for c in _tab1(at).caption
                if not c.value.startswith("Overall score = ")]
    assert len(captions) == len(set(captions)), \
        [c for c in captions if captions.count(c) > 1]


def test_moving_the_property_picker_takes_the_question_down(burger):
    """The question belongs to the property it was armed over. Pick another
    and it is taken down rather than left arming every coloured button away
    behind a Yes the reader cannot reach."""
    burger.add_property("Cost")
    burger.add_property("Fat per 100 g")
    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.run()
    _submit_button(at, wording.delete_button("Cost")).click()
    at.run()
    assert any("Delete Cost?" in w.value for w in at.warning)
    at.selectbox(key="prop_delete").select("Fat per 100 g")
    at.run()
    assert not at.warning, [w.value for w in at.warning]
    assert _tab_primaries(at, 0) == [wording.NEXT_MAKE_BATCH_BUTTON], \
        _tab_primaries(at, 0)
    assert FoodOptimizer("burger").properties() == ["Cost", "Fat per 100 g"]


def test_an_uploaded_workbook_records_what_was_weighed_and_keeps_the_lot(
        open_batch):
    """0.5.0 §1.6, end to end: the file the bench hands back carries an
    Actual weight and a lot number, and pressing Save records the amounts
    that were made — not the ones that were printed — with the lot kept
    against the round."""
    import ui_batch
    opt = open_batch
    book = openpyxl.load_workbook(
        io.BytesIO(opt.workbook_bytes(opt.pending_batch, None)))
    summary = book[wording.batch_sheet_name(1)]
    at_row = {summary.cell(row=r, column=1).value: r
              for r in range(1, summary.max_row + 1)}
    summary.cell(row=at_row["Firmness · target 6 N"], column=2, value=6.0)
    summary.cell(row=at_row["Juiciness (/10) · target 7"], column=2, value=7.0)
    # The Lot column sits after the two formulations and their shares.
    lot_column = 2 + 2 * len(opt.pending_batch)
    summary.cell(row=at_row["Pea protein (g)"], column=lot_column,
                 value="PP-42")
    page = book["Formulation 1"]
    rows = {page.cell(row=r, column=2).value: r
            for r in range(1, page.max_row + 1)}
    page.cell(row=rows["Pea protein"], column=4, value=10.4)
    filled = io.BytesIO()
    book.save(filled)
    filled.seek(0)
    filled.name = "burger · Round 1.xlsx"

    at = AppTest.from_file(APP_PATH, default_timeout=180)
    at.session_state["_loaded_project"] = "burger"
    at.session_state["main_tab"] = wording.TAB_BATCH
    at.run()
    at.session_state["_results_upload"] = ui_batch._read_results_file(
        at.session_state["optimizer"], filled)
    at.run()
    assert not at.exception
    _submit_button(at, wording.SAVE_UPLOADED_RESULTS).click()
    at.run()
    assert not at.error, [e.value for e in at.error]
    reloaded = FoodOptimizer("burger")
    assert reloaded.recipe_history[0] == {"Pea protein": 10.4,
                                          "Methylcellulose": 1.0}
    assert reloaded.notes_history[0] == wording.AMOUNTS_AS_WEIGHED
    assert reloaded.lots == {1: {"Pea protein": "PP-42"}}
