"""Continue ux_workflow.py: fill a compact workbook, review, and save a partial round."""
from grid_helpers import *
from openpyxl import load_workbook

with sync_playwright() as pw:
    browser = pw.chromium.launch(channel='chrome', headless=True)
    page = browser.new_page(viewport={'width': 1280, 'height': 860}, accept_downloads=True)
    page.goto(URL); settle(page, 3)
    page.get_by_role('tab').nth(1).click(); settle(page)
    with page.expect_download() as event:
        page.get_by_role('button', name='Download the round sheets (Excel)', exact=True).click()
    path = '/tmp/foodopt-ux-final-workbook.xlsx'
    event.value.save_as(path)
    book = load_workbook(path)
    sheet = book['Results']
    for name, value in [('Firmness', 6), ('Juiciness', 7), ('Cook loss', 12)]:
        row = next(c.row for cells in sheet for c in cells if c.column == 1 and str(c.value).startswith(name))
        sheet.cell(row, 2, value)
    assert any(str(c.value).startswith('Cook loss (%)') for cells in sheet for c in cells)
    book.save(path)
    open_fold(page, 'Or upload results from a file')
    fold(page, 'Or upload results from a file').locator('input[type=file]').set_input_files(path)
    settle(page, 3)
    assert not alerts(page, 'Error'), alerts(page, 'Error')
    assert page.get_by_role('button', name='Save uploaded results', exact=True).is_visible()
    page.screenshot(path='/tmp/foodopt-ux-upload-preview.png', full_page=True)
    click(page, 'Save uploaded results')
    assert not alerts(page, 'Error'), alerts(page, 'Error')
    assert not page.locator('[data-testid="stException"]').count()
    page.reload(); settle(page, 3)
    assert not page.locator('[data-testid="stException"]').count()
    print('PASS compact Excel download → fill → upload → review → save → reload', flush=True)
    browser.close()
