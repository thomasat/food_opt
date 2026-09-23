"""The whole return path in Chrome: download the three-tab workbook, write
measurements into it, upload it, review and save. Runs against a served
copy named in APP_URL; see conftest.py."""
import os

import pytest

pytestmark = pytest.mark.skipif(
    not os.environ.get("APP_URL"),
    reason="set APP_URL to a served copy of the app")
pytest.importorskip("playwright.sync_api")
pytest.importorskip("openpyxl")

from openpyxl import load_workbook        # noqa: E402
from grid_helpers import *                # noqa: E402,F401,F403


def test_a_filled_in_workbook_comes_back_and_saves(playwright_session):
    browser, page = open_sample(playwright_session)
    try:
        ensure_round(page)
        with page.expect_download() as event:
            page.get_by_role(
                'button', name='Download the round sheets (Excel)',
                exact=True).click()
        path = '/tmp/foodopt-ux-final-workbook.xlsx'
        event.value.save_as(path)
        book = load_workbook(path)
        sheet = book['Results']
        for name, value in [('Firmness', 6), ('Juiciness', 7),
                            ('Cook loss', 12)]:
            row = next(c.row for cells in sheet for c in cells
                       if c.column == 1 and str(c.value).startswith(name))
            sheet.cell(row, 2, value)
        # The unit is in the label, because the cell beside it is empty.
        assert any(str(c.value).startswith('Cook loss (%)')
                   for cells in sheet for c in cells)
        book.save(path)
        open_fold(page, 'Or upload results from a file')
        fold(page, 'Or upload results from a file').locator(
            'input[type=file]').set_input_files(path)
        settle(page, 3)
        assert not alerts(page, 'Error'), alerts(page, 'Error')
        assert page.get_by_role('button', name='Save uploaded results',
                                exact=True).is_visible()
        click(page, 'Save uploaded results')
        assert not alerts(page, 'Error'), alerts(page, 'Error')
        assert not page.locator('[data-testid="stException"]').count()
        page.reload()
        settle(page, 3)
        assert not page.locator('[data-testid="stException"]').count()
    finally:
        browser.close()
