"""Run against an isolated empty project directory: APP_URL=... python review_wave.py MODE."""
import io
import sys
from pathlib import Path
from openpyxl import load_workbook
from grid_helpers import *

mode = sys.argv[1]
pw = sync_playwright().start()
browser = pw.chromium.launch(channel='chrome', headless=True)
page = browser.new_context(viewport={'width': 1280, 'height': 860}, accept_downloads=True).new_page()
page.goto(URL, wait_until='domcontentloaded')
settle(page, 5)
click(page, 'Try the sample project', wait=5)
assert not page.locator('[data-testid="stException"]').count()
assert grid_head(page, 0) == ['Name', 'Type', 'Made as', 'Lowest', 'Highest', 'Unit', 'Rule']
assert [r[0] for r in grid_rows(page, 0)] == ['Textured pea protein', 'Dry blend', 'Wheat gluten', 'Fats and oils', 'Seasoning blend', 'Water', 'Mixing time after fat']
print('PASS sample structure and optional defaults', flush=True)
# Read the actual canvas header bounds, not just its accessibility text.
widths = _measure_columns(page, 0) if '_measure_columns' in globals() else None
box = grid(page, 0).bounding_box()
assert column_x(page, 0, 6) < box['width']
page.screenshot(path=f'/tmp/foodopt-{mode}-setup.png', full_page=True)
print('PASS Rule is within default desktop width', flush=True)

if mode == 'grid':
    set_cell(page, 0, row_of(page, 0, 'Wheat gluten'), 4, '7.5')
    click(page, 'Save changes')
    assert grid_rows(page, 0)[row_of(page, 0, 'Wheat gluten')][4] == '7.50'
    print('PASS ingredient edit and save', flush=True)
    open_fold(page, 'Ingredient limits (optional)')
    page.locator('[data-testid="stCheckbox"]').filter(has_text='Vendor').locator('label').click(); settle(page)
    assert 'Vendor' in grid_head(page, 0)
    page.reload(); settle(page, 3)
    assert 'Vendor' in grid_head(page, 0)
    print('PASS optional supplier column persists across reload', flush=True)
    open_fold(page, 'Ingredient limits (optional)')
    page.locator('[data-testid="stCheckbox"]').filter(has_text='Vendor').locator('label').click(); settle(page)
    assert 'Vendor' not in grid_head(page, 0)
    print('PASS optional supplier column can be hidden again', flush=True)
elif mode == 'rule':
    assert any('To write a rule' in c for c in captions(page))
    water = row_of(page, 0, 'Water')
    assert grid_rows(page, 0)[water][6] == '= rest'
    set_cell(page, 0, row_of(page, 0, 'Wheat gluten'), 6, '= Missing ingredient')
    click(page, 'Save changes')
    assert alerts(page, 'Error')
    click(page, 'Discard changes')
    assert not alerts(page, 'Error')
    assert grid_rows(page, 0)[row_of(page, 0, 'Water')][6] == '= rest'
    print('PASS rule discoverability, invalid name refusal and discard', flush=True)
elif mode == 'premix':
    open_fold(page, 'Dry blend · parts')
    parts = grid_named(page, 'Part')
    set_cell(page, parts, 0, 1, '65')
    assert fold_is_open(page, 'Dry blend · parts')
    assert fold(page, 'Dry blend · parts').get_by_role('button', name='Save changes', exact=True).is_visible()
    fold_click(page, 'Dry blend · parts', 'Discard changes')
    print('PASS parts edit keeps its fold and Save visible', flush=True)
    set_select(page, 0, row_of(page, 0, 'Dry blend'), 2, 'Variable-ratio blend')
    click(page, 'Save changes')
    assert grid_rows(page, 0)[1][0] == 'Dry blend'
    open_fold(page, 'Dry blend · parts')
    parts = grid_named(page, 'Part')
    assert all(float(r[1]) > 0 and float(r[2]) > float(r[1]) for r in grid_rows(page, parts))
    set_select(page, 0, 1, 2, 'Fixed-ratio pre-mix')
    click(page, 'Save changes')
    assert grid_rows(page, 0)[1][3:5] == ['16.00', '22.00']
    print('PASS mode round trip preserves nonzero bands and row order', flush=True)
    click(page, 'Generate formulations')
    # The exact button text comes from the rendered generate action.
    generate = next(x for x in labels(page) if x.startswith('Generate') and 'different' not in x)
    click(page, generate, wait=8)
    page.get_by_label("Include individual formulation pages for printing").check(); settle(page)
    with page.expect_download() as event:
        page.get_by_role('button', name='Download the round sheets (Excel)', exact=True).click()
    downloaded = event.value
    path = f'/tmp/foodopt-{mode}-round.xlsx'; downloaded.save_as(path)
    settle(page)
    assert fold_is_open(page, 'Or upload results from a file')
    book = load_workbook(path)
    assert book.sheetnames[:2] == ['Pre-mix · Dry blend', 'Pre-mix · Seasoning blend']
    assert 'Then upload this file' in book.active['A2'].value
    assert any('Actual' in str(c.value) for row in book.active for c in row)
    print('PASS workbook download, preparation amounts and visible upload path', flush=True)
page.screenshot(path=f'/tmp/foodopt-{mode}-finished.png', full_page=True)
assert not page.locator('[data-testid="stException"]').count()
browser.close(); pw.stop()
print(f'PASS {mode} walkthrough complete', flush=True)
