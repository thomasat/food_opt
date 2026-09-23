"""What a reader sees on tab 1 of the served app: the ingredients grid, the
calculations under it, and the pre-mix folds — in Chrome, on the canvas
grids an AppTest cannot touch."""
import io
import os

import pytest

pytestmark = pytest.mark.skipif(
    not os.environ.get("APP_URL"),
    reason="set APP_URL to a served copy of the app")
pytest.importorskip("playwright.sync_api")
pytest.importorskip("openpyxl")

from openpyxl import load_workbook        # noqa: E402
from grid_helpers import *                # noqa: E402,F401,F403


@pytest.fixture
def page(playwright_session):
    browser, page = open_sample(playwright_session)
    yield page
    browser.close()


def test_the_sample_reads_in_the_order_the_method_adds_it(page):
    tab(page, SET_UP_TAB)
    assert grid_head(page, 0) == GRID_HEAD, grid_head(page, 0)
    assert [r[NAME] for r in grid_rows(page, 0)] == SAMPLE_ROWS, \
        [r[NAME] for r in grid_rows(page, 0)]
    # Every column is inside the default desktop width: the last one was
    # off the right edge of a 1280 px window for a whole wave.
    box = grid(page, 0).bounding_box()
    assert column_x(page, 0, len(GRID_HEAD) - 1) < box['width']


def test_the_two_water_rows_say_what_calculates_them(page):
    tab(page, SET_UP_TAB)
    rows = grid_rows(page, 0)
    hydration = rows[row_of(page, 0, 'Hydration water')]
    remaining = rows[row_of(page, 0, 'Remaining water')]
    assert hydration[CALCULATION] == \
        '2.2 * (Textured pea protein + Textured soy protein)', hydration
    assert remaining[CALCULATION] == 'Fill to total', remaining
    # A calculated row has no allowed amounts of its own to show.
    assert (remaining[LOWEST], remaining[HIGHEST]) == ('', ''), remaining
    said = captions(page)
    assert any(c.startswith('Remaining water is calculated to bring the total '
                            'to the 100 g default batch size') for c in said), said
    # The one line that says how to write a calculation is under the grid
    # whether or not a row has one.
    assert any(c.startswith('To calculate a row from the others') for c in said), said


def test_a_calculation_the_project_cannot_take_is_refused_and_discarded(page):
    tab(page, SET_UP_TAB)
    set_cell(page, 0, row_of(page, 0, 'Wheat gluten'), CALCULATION,
             '= Missing ingredient')
    click(page, 'Save changes')
    assert alerts(page, 'Error')
    click(page, 'Discard changes')
    assert not alerts(page, 'Error')
    assert grid_rows(page, 0)[row_of(page, 0, 'Remaining water')][CALCULATION] \
        == 'Fill to total'


def test_a_premix_fold_keeps_its_place_while_its_parts_are_edited(page):
    tab(page, SET_UP_TAB)
    open_fold(page, 'Dry blend · parts')
    parts = grid_named(page, 'Part')
    assert 'Composition (%)' in grid_head(page, parts)
    set_cell(page, parts, 0, 1, '65')
    assert fold_is_open(page, 'Dry blend · parts')
    assert fold(page, 'Dry blend · parts').get_by_role(
        'button', name='Save changes', exact=True).is_visible()
    fold_click(page, 'Dry blend · parts', 'Discard changes')
    assert not alerts(page, 'Error'), alerts(page, 'Error')


def test_a_round_downloads_as_a_print_pack_the_bench_can_carry(page):
    ensure_round(page)
    tick(page, 'Include individual formulation pages for printing')
    assert page.get_by_label(
        'Include individual formulation pages for printing').is_checked()
    with page.expect_download() as event:
        page.get_by_role('button', name='Download the round sheets (Excel)',
                         exact=True).click()
    path = '/tmp/foodopt-review-round.xlsx'
    event.value.save_as(path)
    settle(page)
    assert fold_is_open(page, 'Or upload results from a file')
    book = load_workbook(path)
    assert book.sheetnames[:2] == ['Pre-mix · Dry blend',
                                   'Pre-mix · Seasoning blend']
    assert book['Pre-mix · Dry blend']['A1'].value.startswith(
        'Dry blend · Prepare 150 g'), book['Pre-mix · Dry blend']['A1'].value
    round_sheet = book['Round 1']
    printed = '\n'.join(str(c.value) for row in round_sheet for c in row
                        if c.value)
    for fact in ('180 °C', '74 °C core', 'Contains soy and wheat (gluten).',
                 'Settings', 'Cook loss (%)'):
        assert fact in printed, fact
    assert 'Each formulation totals 100 g' not in printed
    assert not page.locator('[data-testid="stException"]').count()
